"""
Dossier de présentation (le « mémoire technique » de l'acheteur) : du texte rédigé au fichier remis.

Doctrine § 7 (règle 7) : rien d'interne ne part chez l'acheteur.
- les renvois internes aux sources « [S1] », « [S2, S3] » sont retirés du fichier (ils renvoient
  à la base de savoir-faire d'Adjugo, que l'acheteur ne voit pas) ;
- aucun message technique : un dossier dont la rédaction a échoué n'est pas produit du tout ;
- les passages « [À compléter …] » restent VISIBLES (surlignés) et sont comptés : l'utilisateur
  est averti avant le dépôt (le blocage de l'export viendra avec l'atelier, vague 2) ;
- un texte coupé par l'IA (limite de longueur atteinte) n'est jamais livré comme complet : une
  mention « [À compléter : … interrompue …] » est ajoutée à la fin, et l'utilisateur est averti ;
- les tableaux Markdown sont rendus en vrais tableaux (Word et PDF).

Archivage : .docx (modifiable) ET .pdf, sous le nom employé par le règlement de la consultation
quand il est retrouvé mot pour mot dans le dossier (sinon « Mémoire technique »), un nom que la
liste des pièces reconnaît.
"""
import io
import re
import unicodedata

DEFAULT_LABEL = "Mémoire technique"
TRUNCATED_NOTE = "[À compléter : la rédaction a été interrompue ici, la suite du texte manque.]"

# Renvois internes aux sources (même forme que rag.cited_refs : « [S1] », « [S1, S2] », « [S1;S3] »).
_CITE = re.compile(r"[ \t]*\[S\s*\d[\dS,;\s]*\]")
_TODO = re.compile(r"\[\s*[ÀAà]\s+compl[ée]ter[^\]]*\]", re.IGNORECASE)
_FENCE = re.compile(r"^\s*```[a-zA-Z]*\s*\n|\n?\s*```\s*$")
# Noms que les règlements donnent au document de réponse technique.
_FAMILLE = re.compile(r"^(m[ée]moire\b|note\s+(m[ée]thodologique|technique)\b|offre\s+technique\b)",
                      re.IGNORECASE)


def _norm(s) -> str:
    t = unicodedata.normalize("NFKD", str(s or ""))
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    t = t.replace("’", "'")
    return re.sub(r"\s+", " ", t).strip()


# ── Nettoyage du texte livré ───────────────────────────────────────────────────
def clean_markdown(md: str) -> str:
    """Texte prêt pour l'acheteur : sans renvois [S1], sans clôture de code, espaces propres."""
    t = str(md or "").replace("\r\n", "\n")
    t = _FENCE.sub("", t.strip())
    t = _CITE.sub("", t)
    # L'espace AVANT le renvoi part avec lui (« texte [S1]. » → « texte. ») ; celle d'avant « : »
    # reste (typographie française). Reste d'une liste de renvois : « texte, . » → « texte. ».
    t = re.sub(r",[ \t]*([.;])", r"\1", t)
    t = re.sub(r"[ \t]{2,}", " ", t)
    return t.strip()


def placeholders(md: str) -> list:
    """Passages « [À compléter …] » encore présents (à remplir avant dépôt)."""
    return _TODO.findall(str(md or ""))


def finalize(md: str, truncated: bool = False) -> dict:
    """Texte livré + état : {"markdown", "a_completer", "coupe"}."""
    text = clean_markdown(md)
    if truncated:
        text = (text + "\n\n" + TRUNCATED_NOTE).strip()
    return {"markdown": text, "a_completer": len(placeholders(text)), "coupe": bool(truncated)}


def warnings_for(state: dict, label: str) -> list:
    """Avertissements affichés avec le dossier (jamais dans le fichier remis)."""
    out = []
    if state.get("coupe"):
        out.append(f"BLOQUANT — Le dossier de présentation ({label}) a été coupé avant la fin : relancez la "
                   "génération ou terminez le texte vous-même avant de déposer.")
    n = int(state.get("a_completer") or 0) - (1 if state.get("coupe") else 0)
    if n > 0:
        out.append(f"Dossier de présentation : {n} passage{'s' if n > 1 else ''} « À compléter » à remplir "
                   "avant le dépôt (surlignés dans le fichier Word et le PDF).")
    return out


# ── Nom du fichier : celui du règlement, s'il est retrouvé mot pour mot ──────────
def memoire_label(analysis) -> str:
    """Nom que le règlement donne au document (« Mémoire justificatif », « Note méthodologique »…).
    Proposé par l'analyse (liste des pièces), retenu SEULEMENT s'il figure tel quel dans le texte du
    dossier conservé ; sinon « Mémoire technique ». Jamais un nom inventé."""
    a = analysis if isinstance(analysis, dict) else {}
    det = a.get("details") if isinstance(a.get("details"), dict) else {}
    pieces = det.get("pieces_requises") or a.get("pieces_requises") or []
    corpus = _norm(a.get("dce_excerpt") or "")
    if not corpus or not isinstance(pieces, list):
        return DEFAULT_LABEL
    for p in pieces:
        nom = p if isinstance(p, str) else (p.get("nom") if isinstance(p, dict) else "")
        nom = re.sub(r"[\x00-\x1f]", " ", str(nom or "")).strip()
        # « Mémoire justificatif (10 pages max) » / « Mémoire technique — méthodologie » → le nom seul.
        nom = re.split(r"\s+[—–-]\s+|\s*[(:;,]", nom)[0].strip(" .")
        if not (6 <= len(nom) <= 60) or not _FAMILLE.match(nom):
            continue
        if _norm(nom) in corpus:
            return nom[0].upper() + nom[1:]
    return DEFAULT_LABEL


def is_memoire_name(name) -> bool:
    """Le nom d'une pièce désigne-t-il le mémoire (« Mémoire justificatif », « Note méthodologique »…) ?"""
    return bool(_FAMILLE.match(str(name or "").strip()))


def file_stem(label: str) -> str:
    """Nom de fichier sûr, lisible (accents gardés) : « Mémoire technique »."""
    s = re.sub(r"[\x00-\x1f/\\:*?\"<>|]+", " ", str(label or DEFAULT_LABEL))
    s = re.sub(r"\s+", " ", s).strip(" .")[:60]
    return s or DEFAULT_LABEL


# ── Blocs (un seul lecteur pour Word et PDF) ─────────────────────────────────────
def parse_blocks(md: str) -> list:
    """("h1"|"h2"|"h3"|"p"|"bullet"|"num", texte) ou ("table", {"headers", "rows"})."""
    from app.services.md_pdf import is_table_start, split_table_row, _TABLE_ROW
    lines = str(md or "").split("\n")
    out, i = [], 0
    while i < len(lines):
        line = lines[i].rstrip()
        if is_table_start(lines, i):
            headers = split_table_row(line)
            i += 2
            rows = []
            while i < len(lines) and _TABLE_ROW.match(lines[i]):
                rows.append(split_table_row(lines[i]))
                i += 1
            out.append(("table", {"headers": headers, "rows": rows}))
            continue
        i += 1
        if not line.strip():
            continue
        h = re.match(r"^(#{1,6})\s+(.*)", line)
        if h:
            out.append((f"h{min(len(h.group(1)), 3)}", h.group(2).strip()))
            continue
        b = re.match(r"^\s*[-*+]\s+(.*)", line)
        if b:
            out.append(("bullet", b.group(1)))
            continue
        n = re.match(r"^\s*(\d+)[.)]\s+(.*)", line)
        if n:
            out.append(("num", f"{n.group(1)}. {n.group(2)}"))
            continue
        out.append(("p", line.strip()))
    return out


# ── Word (.docx, modifiable) ─────────────────────────────────────────────────────
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|(?<!\*)\*(?!\*)[^*]+?\*(?!\*)|\[\s*[ÀAà]\s+compl[ée]ter[^\]]*\])",
                     re.IGNORECASE)


def _runs(par, text: str) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    text = _CTRL.sub("", str(text or ""))
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**") or tok.startswith("__"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        else:                                   # « [À compléter …] » : surligné, en gras
            r = par.add_run(tok)
            r.bold = True
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def to_docx(md: str, title: str) -> bytes:
    from docx import Document
    doc = Document()
    blocks = parse_blocks(md)
    if blocks and blocks[0][0] == "h1":            # le 1er titre fait office de titre du document
        doc.add_heading(_CTRL.sub("", blocks[0][1]) or title, level=0)
        blocks = blocks[1:]
    else:
        doc.add_heading(title, level=0)
    for kind, val in blocks:
        if kind in ("h1", "h2", "h3"):
            _runs(doc.add_heading("", level={"h1": 1, "h2": 2, "h3": 3}[kind]), val)
        elif kind == "bullet":
            _runs(doc.add_paragraph(style="List Bullet"), val)
        elif kind in ("p", "num"):
            _runs(doc.add_paragraph(), val)
        elif kind == "table":
            headers, rows = val["headers"], val["rows"]
            n = max([len(headers)] + [len(r) for r in rows] + [1])
            t = doc.add_table(rows=1, cols=n)
            t.style = "Table Grid"
            for j in range(n):
                p = t.rows[0].cells[j].paragraphs[0]
                _runs(p, f"**{headers[j]}**" if j < len(headers) and headers[j] else "")
            for row in rows:
                cells = t.add_row().cells
                for j in range(n):
                    _runs(cells[j].paragraphs[0], row[j] if j < len(row) else "")
            doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(md: str, title: str) -> bytes:
    from app.services.md_pdf import markdown_to_pdf
    return markdown_to_pdf(md, title, tables=True)


def render_files(md: str, label: str) -> dict:
    """{"docx": (nom, octets), "pdf": (nom, octets)} pour le dossier livré."""
    stem = file_stem(label)
    return {"docx": (f"{stem}.docx", to_docx(md, label)), "pdf": (f"{stem}.pdf", to_pdf(md, label))}
