"""
Adjugo - Service d'analyse IA des DCE
Prompt expert marches publics francais.
Analyse en deux temps (docs/ANALYSE_PARTAGEE.md) : extraction du dossier SANS donnée d'entreprise
(partageable entre organisations pour un même dossier), puis partie propre à l'entreprise — score
Go/No-Go déterministe (dce_scoring) et appel court qui reçoit l'extraction, le profil et les critères.
"""
import copy
import json
import re
import unicodedata
from app.core.config import get_settings
from app.services.llm import client  # client Anthropic LAZY partagé (pas d'instanciation à l'import)
from app.services.llm import LLMUnavailable

settings = get_settings()

# Termes des clauses sensibles : souvent dans le CCAP, en fin de DCE — on garantit
# qu'elles sont vues par l'analyse même dans un dossier de plusieurs centaines de pages.
# Étendu (co-traitance v2) aux clauses de GROUPEMENT (forme conjoint/solidaire, mandataire,
# appréciation des capacités en groupement, références, effectif) : sans ça, la clause
# « chaque membre doit satisfaire individuellement » page 140 du RC n'entrait jamais dans
# la fenêtre du LLM → exigence invisible → score de groupement faussement optimiste.
_RISK_TERMS = re.compile(
    r"p[ée]nalit|retenue de garantie|garantie [aà] premi[èe]re demande|cautionnement|"
    r"avance|d[ée]lai d'ex[ée]cution|d[ée]lai de remise|crit[èe]res? d['e]|pond[ée]ration|"
    r"notation|valeur technique|prix|reconduction|tranche|variante|sous-trait|"
    r"assurance|qualification|certificat|RGE|capacit[ée]|chiffre d'affaires|"
    r"groupement|conjoint|solidaire|mandataire|co-?traitan|r[ée]f[ée]rences?|effectif",
    re.IGNORECASE,
)


# ── Garde anti-hallucination des CITATIONS du DCE ─────────────────────────────────────
# Le LLM « recopie » des phrases du DCE (CA exigé, forme de groupement imposée…). Une
# phrase reformulée ou inventée serait affichée comme une citation → mensonge au client.
# Règle : un extrait n'est conservé QUE s'il est une sous-chaîne du texte intégral (après
# normalisation tolérante : casse, accents, espaces, apostrophes). Sinon '' (jamais affiché).

_QUOTE_CHARS = "«»\"“”„‘’'"
# Surveillance de la qualité des citations : désactivée pour une extraction resservie (analyse partagée).
import contextvars as _contextvars
_QUALITE = _contextvars.ContextVar("adjugo_analyse_qualite", default=True)


def _norm_citation(s) -> str:
    """Forme normalisée pour la comparaison : minuscules, accents retirés (NFKD), apostrophes
    typographiques → ', suites d'espaces/retours → un espace. Ne lève jamais."""
    try:
        t = unicodedata.normalize("NFKD", str(s or ""))
        t = "".join(c for c in t if not unicodedata.combining(c))
        t = t.replace("’", "'").replace("‘", "'").replace("ʼ", "'")
        t = t.replace("“", '"').replace("”", '"').replace("„", '"')
        t = t.replace(" ", " ").replace(" ", " ")
        t = re.sub(r"\s+", " ", t.lower()).strip()
        return t
    except Exception:
        return ""


def citation_guard(text_integral: str, extrait: str, min_len: int = 12) -> str:
    """Renvoie l'extrait d'ORIGINE (trim, sans guillemets d'encadrement, ≤ 400 caractères)
    s'il est RÉELLEMENT une sous-chaîne du texte intégral du DCE (comparaison normalisée :
    casse, accents, espaces multiples, apostrophes typographiques tolérés) ET s'il fait au
    moins `min_len` caractères utiles ; sinon ''. Jamais d'exception."""
    try:
        ex = str(extrait or "").strip()
        # Tolérance : l'IA encadre souvent sa citation de guillemets « » ou "" → retirés.
        ex = ex.strip(_QUOTE_CHARS + " \t\r\n")
        if not ex:
            return ""
        ex = ex[:400].strip()
        nx = _norm_citation(ex)
        if len(nx) < max(1, int(min_len)):
            return ""
        nt = _norm_cached(text_integral)
        if not nt or nx not in nt:
            return ""
        return ex
    except Exception:
        return ""


# Texte intégral normalisé, mémorisé : un dossier lu en entier fait plusieurs centaines de milliers
# de caractères et chaque citation était comparée après une normalisation complète du dossier.
_NORM_CACHE = {}
_NORM_CACHE_MAX = 8
_NORM_CACHE_MIN_LEN = 20000
_NORM_CACHE_MAX_CHARS = 12_000_000   # ≈ 24 Mo au plus (deux dossiers au plafond de 6 M caractères)
import threading as _threading
_NORM_LOCK = _threading.Lock()


def _norm_cached(text) -> str:
    t = str(text or "")
    if len(t) < _NORM_CACHE_MIN_LEN:
        return _norm_citation(t)
    key = (len(t), hash(t))
    with _NORM_LOCK:
        hit = _NORM_CACHE.get(key)
    if hit is not None:
        return hit
    nt = _norm_citation(t)
    if len(nt) > _NORM_CACHE_MAX_CHARS:
        return nt
    with _NORM_LOCK:
        while _NORM_CACHE and (len(_NORM_CACHE) >= _NORM_CACHE_MAX
                               or sum(len(v) for v in _NORM_CACHE.values()) + len(nt) > _NORM_CACHE_MAX_CHARS):
            _NORM_CACHE.pop(next(iter(_NORM_CACHE)), None)
        _NORM_CACHE[key] = nt
    return nt


def _norm_with_map(text: str):
    """Même normalisation que _norm_citation, caractère par caractère, avec la position
    d'ORIGINE de chaque caractère produit : une citation retrouvée dans le texte normalisé
    est ainsi localisée dans le texte d'origine (page d'une pièce)."""
    out, pos, prev_space = [], [], True
    for i, c in enumerate(str(text or "")):
        d = unicodedata.normalize("NFKD", c)
        d = "".join(x for x in d if not unicodedata.combining(x))
        for x in d:
            x = {"’": "'", "‘": "'", "ʼ": "'", "“": '"', "”": '"', "„": '"'}.get(x, x)
            if x.isspace():
                if prev_space:
                    continue
                x, prev_space = " ", True
            else:
                prev_space = False
            for y in x.lower():
                out.append(y)
                pos.append(i)
    return "".join(out), pos


class Localisateur:
    """Localise des extraits (déjà vérifiés) dans UN texte : recherche exacte d'abord, puis
    tolérante (casse, accents, espaces, apostrophes). La table de correspondance du texte normalisé
    n'est construite qu'une fois, au premier extrait qui en a besoin (plusieurs citations d'une même
    pièce ne refont pas le travail)."""

    def __init__(self, text: str):
        self.t = str(text or "")
        self._map = None

    def offset(self, extrait: str):
        try:
            ex = str(extrait or "").strip().strip(_QUOTE_CHARS + " \t\r\n")[:400].strip()
            if not self.t or not ex:
                return None
            i = self.t.find(ex)
            if i >= 0:
                return i
            if self._map is None:
                self._map = _norm_with_map(self.t)
            nt, pos = self._map
            nx = _norm_citation(ex)
            j = nt.find(nx) if nx else -1
            return pos[j] if 0 <= j < len(pos) else None
        except Exception:
            return None


def citation_offset(text: str, extrait: str):
    """Position (dans `text`) du début d'un extrait déjà vérifié par citation_guard, ou None."""
    return Localisateur(text).offset(extrait)


# Exigences dont le DCE doit être CITÉ (clé → présente dans details.exigences_sources).
_EXIGENCE_KEYS = ("ca_minimum_requis", "qualifications_requises", "references_requises",
                  "effectif_minimum_requis", "capacites_groupement")
_FORMES_GROUPEMENT = ("non_specifie", "libre", "conjoint", "solidaire")
_NON_SPECIFIE = "Non specifie dans le DCE"


def _normalize_exigences(details: dict, text_integral: str) -> dict:
    """Post-traitement DÉTERMINISTE des exigences de groupement extraites par l'IA :
    - chaque citation (exigences_sources.*, forme_groupement.extrait) passe par citation_guard
      sur le texte INTÉGRAL (pas la fenêtre envoyée au LLM) → '' si non vérifiable ;
    - exigences_sources : toujours un dict des 5 clés (str) ;
    - forme_groupement : toujours {"imposee": l'une des 4 valeurs, "extrait": str} ;
    - capacites_groupement / references_requises / effectif_minimum_requis absents → non spécifié ;
    - details["exigences_verifiees"] = {clé: bool(extrait non vide après garde)}.
    Ne lève jamais (données LLM sales → valeurs neutres)."""
    if not isinstance(details, dict):
        return {}
    try:
        raw_src = details.get("exigences_sources")
        raw_src = raw_src if isinstance(raw_src, dict) else {}
        sources = {}
        for k in _EXIGENCE_KEYS:
            v = raw_src.get(k)
            if isinstance(v, (list, tuple)):
                v = " ".join(str(x) for x in v if x)
            elif isinstance(v, dict):
                v = v.get("extrait") or v.get("texte") or ""
            sources[k] = citation_guard(text_integral, v)
        details["exigences_sources"] = sources
        _soumis = [k for k in _EXIGENCE_KEYS if str(raw_src.get(k) or "").strip()]

        fg = details.get("forme_groupement")
        if isinstance(fg, str):                     # l'IA a parfois renvoyé « conjoint » tout court
            fg = {"imposee": fg, "extrait": ""}
        fg = fg if isinstance(fg, dict) else {}
        imposee = _norm_citation(fg.get("imposee")).replace(" ", "_").replace("-", "_")
        if imposee not in _FORMES_GROUPEMENT:
            imposee = "non_specifie"
        fg_extrait = fg.get("extrait")
        if isinstance(fg_extrait, (list, tuple)):
            fg_extrait = " ".join(str(x) for x in fg_extrait if x)
        details["forme_groupement"] = {"imposee": imposee,
                                       "extrait": citation_guard(text_integral, fg_extrait)}
        # Surveillance de la qualité (par fournisseur et par tâche) : citations retenues / rejetées.
        # Une extraction resservie (analyse partagée) n'est pas recomptée : aucune IA n'a répondu ici.
        try:
            if not _QUALITE.get():
                raise LookupError
            from app.services import llm_qualite as _q
            _acc = sum(1 for k in _soumis if sources.get(k))
            _fg_ok = bool(str(fg_extrait or "").strip())
            _q.noter_citations(_acc + (1 if _fg_ok and details["forme_groupement"]["extrait"] else 0),
                               (len(_soumis) - _acc) + (1 if _fg_ok and not details["forme_groupement"]["extrait"] else 0))
        except Exception:
            pass

        for k in ("capacites_groupement", "references_requises", "effectif_minimum_requis"):
            v = details.get(k)
            if isinstance(v, (list, tuple)):
                v = "; ".join(str(x) for x in v if x)
            if v is None or (isinstance(v, str) and not v.strip()) or isinstance(v, (dict, bool)):
                details[k] = _NON_SPECIFIE
            else:
                details[k] = str(v).strip()

        verifiees = {k: bool(sources.get(k)) for k in _EXIGENCE_KEYS}
        verifiees["forme_groupement"] = bool(details["forme_groupement"]["extrait"])
        details["exigences_verifiees"] = verifiees
    except Exception:
        details.setdefault("exigences_sources", {k: "" for k in _EXIGENCE_KEYS})
        details.setdefault("forme_groupement", {"imposee": "non_specifie", "extrait": ""})
        details.setdefault("exigences_verifiees", {k: False for k in _EXIGENCE_KEYS + ("forme_groupement",)})
    return details


# Ce que l'IA reçoit EN ENTIER (règlement + cadre de réponse), et le budget minimal laissé aux
# fenêtres ciblées sur les autres pièces. Borne haute de l'entrée ≈ 84 000 caractères.
SELECT_PRIORITY_MAX = 60000
SELECT_REST_MIN = 24000


def select_dce_content(text: str, head: int = 9000, budget: int = 30000, window: int = 700) -> str:
    """
    Sélection pertinente du DCE plutôt qu'une troncature aveugle des N premiers caractères.

    Dossier assemblé pièce par pièce (« === nom === », cf. extract_dce_text) : le règlement de
    la consultation et le cadre de réponse, reconnus sans IA, sont remis EN ENTIER (dans la
    limite de SELECT_PRIORITY_MAX), puis chaque autre pièce reçoit des fenêtres ciblées
    autour des clauses sensibles, avec un budget partagé équitablement.

    Texte d'un seul tenant sans règlement reconnu : comportement historique — tête du dossier
    (objet, acheteur, contexte) puis fenêtres autour des CLAUSES SENSIBLES (pénalités, retenue
    de garantie, critères de pondération, délais…) où qu'elles se trouvent. Ainsi une clause de
    pénalité située page 180 d'un CCAP est bien analysée. Budget borné pour la fenêtre de contexte.
    """
    text = text or ""
    if len(text) <= budget:
        return text
    try:
        out = _select_by_piece(text, head, budget, window)
        if out:
            return out
    except Exception:
        pass
    return _select_windows(text, head, budget, window)


def _select_by_piece(text: str, head: int, budget: int, window: int) -> str:
    """Sélection pièce par pièce ; '' si aucun règlement ni cadre n'est reconnu (repli)."""
    from app.services import dce_pieces as _dp
    pieces = _dp.typed_pieces(text)
    prio = [p for p in pieces if p["type"] in _dp.PRIORITAIRES]
    if not prio:
        return ""
    rest = [p for p in pieces if p["type"] not in _dp.PRIORITAIRES]
    alloc = _dp.water_fill([len(p["texte"]) for p in prio], SELECT_PRIORITY_MAX)
    out, overflow = [], []
    for p, n in zip(prio, alloc):
        body = p["texte"].strip("\n")
        complet = n >= len(p["texte"])
        out.append(f"=== {p['nom'] or 'dossier'} : {p['libelle'].lower()}, "
                   f"{'lu en entier' if complet else 'début'} ===\n{body[:n]}")
        if not complet:
            overflow.append(dict(p, texte=p["texte"][n:], libelle=p["libelle"] + " (suite)"))
    rest = overflow + rest
    used = sum(len(x) for x in out)
    rest_budget = max(SELECT_REST_MIN, budget - used)
    rest = [p for p in rest if p["texte"].strip()]
    if rest:
        shares = _dp.water_fill([len(p["texte"]) for p in rest], rest_budget)
        picked = []
        for p, share in zip(rest, shares):
            if share <= 0:
                continue
            body = p["texte"]
            if len(body) > share:
                body = _select_windows(body, head=min(head, max(500, share // 4)), budget=share, window=window)
                tag = "extraits ciblés"
            else:
                tag = "lu en entier"
            picked.append(f"=== {p['nom'] or 'dossier'} : {p['libelle'].lower()}, {tag} ===\n{body.strip()}")
        if picked:
            out.append("[… autres pièces du dossier : extraits ciblés des clauses sensibles "
                       "(pénalités, garanties, critères, délais) …]")
            out.extend(picked)
    return "\n\n".join(out)


def _select_windows(text: str, head: int = 9000, budget: int = 30000, window: int = 700) -> str:
    """Tête + fenêtres autour des clauses sensibles (sélection historique, d'un seul tenant)."""
    text = text or ""
    if len(text) <= budget:
        return text
    head_part = text[:head]
    rest = text[head:]
    # Repère les fenêtres autour de chaque occurrence de terme sensible.
    spans = []
    for m in _RISK_TERMS.finditer(rest):
        a = max(0, m.start() - window // 3)
        b = min(len(rest), m.end() + window)
        if spans and a <= spans[-1][1] + 120:   # fusionne les fenêtres proches
            spans[-1] = (spans[-1][0], max(spans[-1][1], b))
        else:
            spans.append((a, b))
    picked, used = [], 0
    room = budget - head
    for a, b in spans:
        if used >= room:
            break
        seg = rest[a:b]
        if used + len(seg) > room:
            seg = seg[: room - used]
        picked.append(seg)
        used += len(seg)
    if not picked:
        # Aucune clause sensible repérée : on reprend une coupe simple jusqu'au budget.
        return text[:budget]
    return head_part + "\n\n[… extraits ciblés des clauses sensibles, ailleurs dans le DCE …]\n\n" + \
        "\n[…]\n".join(picked)

SYSTEM_PROMPT = """Tu es un expert en marches publics francais, specialise dans l'analyse des Dossiers de Consultation des Entreprises (DCE). Tu travailles pour un logiciel appele Adjugo qui aide les entreprises du BTP et des services a repondre aux appels d'offres.

Ton role est d'analyser le DCE fourni et de produire une analyse structuree et actionnable.

Tu dois TOUJOURS repondre en JSON valide, sans texte avant ou apres. Voici le format exact attendu :

{
  "match_score": 75,
  "go_decision": "go",
  "summary": "Resume en 3-4 phrases de l'appel d'offres...",
  "details": {
    "intitule_marche": "Nom exact du marche",
    "acheteur": "Nom de l'acheteur / pouvoir adjudicateur",
    "contact": {"nom": "Nom de la personne a contacter pour les questions (si indique)", "fonction": "Fonction / service", "email": "Email de contact", "telephone": "Telephone"},
    "type_marche": "Travaux / Services / Fournitures",
    "nature_marche": "Marche a procedure adaptee (MAPA) / Appel d'offres ouvert / etc.",
    "forme_prix": "Prix global et forfaitaire / Prix unitaires / Mixte",
    "budget_estime": "150 000 EUR HT",
    "date_limite": "15/06/2026 a 12h00",
    "delai_execution": "6 mois a compter de l'ordre de service",
    "lieu_execution": "Quimper (29)",
    "allotissement": "Lot unique / Lot 1 : Gros oeuvre, Lot 2 : Electricite...",
    "criteres_attribution": [
      {"critere": "Prix", "ponderation": "60 %", "extrait": "ligne ou phrase litterale du reglement qui cite ce critere et son poids"},
      {"critere": "Valeur technique", "ponderation": "40 %", "extrait": "ligne ou phrase litterale du reglement"}
    ],
    "garanties_exigees": "Garantie a premiere demande 5% / Retenue de garantie 5%",
    "penalites": "1/1000 du montant HT par jour de retard, plafonnees a 10%",
    "avance": "Avance de 10% si marche > 50 000 EUR",
    "sous_traitance": "Autorisee / Interdite / Limitee a 30%",
    "variantes": "Autorisees / Non autorisees",
    "critere_rse": "Oui - clause d'insertion sociale / Non",
    "ca_minimum_requis": "CA > 300 000 EUR sur 3 ans / Non specifie",
    "qualifications_requises": ["Qualibat 1312", "RGE"],
    "references_requises": "ex. 3 references de travaux similaires sur 5 ans / Non specifie dans le DCE",
    "effectif_minimum_requis": "ex. 10 salaries / Non specifie dans le DCE",
    "capacites_groupement": "cumul admis / chaque membre doit satisfaire individuellement / mandataire seul / Non specifie dans le DCE",
    "forme_groupement": {"imposee": "non_specifie | libre | conjoint | solidaire", "extrait": "phrase litterale du DCE ou vide"},
    "exigences_sources": {"ca_minimum_requis": "phrase litterale du DCE ou vide", "qualifications_requises": "", "references_requises": "", "effectif_minimum_requis": "", "capacites_groupement": ""},
    "assurances_requises": ["RC Professionnelle", "Decennale"],
    "visite_obligatoire": "Oui, le 01/06/2026 a 14h / Non",
    "pieces_requises": [
      {"nom": "DC1 - Lettre de candidature", "disponible": true},
      {"nom": "DC2 - Declaration du candidat", "disponible": true},
      {"nom": "Kbis de moins de 3 mois", "disponible": false},
      {"nom": "Attestation URSSAF", "disponible": false},
      {"nom": "Attestation fiscale", "disponible": false},
      {"nom": "Attestation assurance RC Pro", "disponible": false},
      {"nom": "Attestation assurance decennale", "disponible": false},
      {"nom": "Memoire technique", "disponible": false},
      {"nom": "DPGF / Bordereau des prix", "disponible": false},
      {"nom": "Planning previsionnel", "disponible": false},
      {"nom": "Certificat Qualibat", "disponible": false},
      {"nom": "References de travaux similaires", "disponible": false}
    ],
    "points_attention": [
      "Delai tres court - 6 mois pour un chantier complexe",
      "Penalites elevees - 1/1000 par jour",
      "Visite obligatoire le 01/06"
    ],
    "recommandation": "Analyse redigee en 3 a 4 paragraphes (180 a 350 mots), separes par une ligne vide : (1) de quoi parle le marche et ce qu'il exige, (2) les attentes de l'acheteur et la mecanique d'attribution, (3) les points techniques et administratifs a maitriser et les clauses a negocier, (4) la strategie et la recommandation finale",
    "clauses_risque": [
      {"clause": "Penalites de retard", "niveau": "eleve", "pourquoi": "1/1000 du HT par jour plafonnees a 10%, au-dela du seuil tolere par l'entreprise", "levier_negociation": "Demander un plafond a 5% et une franchise de 5 jours ouvres"},
      {"clause": "Retenue de garantie", "niveau": "moyen", "pourquoi": "Retenue de 5% sur 1 an immobilise la tresorerie", "levier_negociation": "Proposer une caution bancaire de substitution (art. R2191-33)"}
    ]
  }
}

Regles importantes :
- "match_score" est un entier entre 0 et 100
- "go_decision" est "go", "no_go" ou "a_etudier"
- Pour les pieces_requises, "disponible" est true pour les CERFA (DC1, DC2, DC4, ATTRI1) car le logiciel les genere automatiquement, et false pour le reste
- Les "points_attention" sont les elements de risque ou les contraintes particulieres
- La "recommandation" est une ANALYSE REDIGEE et substantielle, PAS un resume d'une phrase : 3 a 4 paragraphes, entre 180 et 350 mots (jamais moins de 180), separes par une ligne vide (double saut de ligne \n\n). Structure attendue : (1) de quoi parle le marche et ce qu'il demande concretement ; (2) les attentes de l'acheteur et la mecanique d'attribution (poids des criteres, ce qui fait gagner) ; (3) les points techniques et administratifs a maitriser + les clauses a negocier ; (4) la strategie et la recommandation finale (y aller ou non, comment, seul ou en groupement). Reste factuel et fonde UNIQUEMENT sur le DCE fourni — n'invente jamais un chiffre, une exigence ou une echeance. Sur une simple annonce (pas de DCE complet), reste plus prudent mais redige quand meme une analyse structuree de ce qui est connu.
- "clauses_risque" liste les clauses contractuelles a surveiller : penalites, retenue de garantie, garantie a premiere demande, avance, delais d'execution intenables, revision de prix figee, assurances exigees, criteres financiers (CA mini). Pour CHAQUE clause notable :
  * "niveau" vaut "faible", "moyen" ou "eleve"
  * Tu DOIS comparer aux seuils de l'entreprise fournis dans les criteres (penalites max, garantie max, retenue max, avance min, delai de reponse min) : marque "eleve" des qu'un seuil est depasse, et dis-le explicitement dans "pourquoi"
  * "levier_negociation" propose une contre-proposition concrete et actionnable (article du CCP si pertinent)
  * Renvoie une liste vide [] si aucune clause notable, n'invente jamais
- Pour "contact" : n'extrais le nom/email/telephone QUE s'ils figurent dans le DCE (souvent dans le reglement de consultation, rubrique "renseignements"). Si rien n'est indique, laisse les champs vides "" — n'invente JAMAIS de contact.
- Exigences de candidature et de GROUPEMENT ("ca_minimum_requis", "qualifications_requises", "references_requises", "effectif_minimum_requis", "capacites_groupement", "forme_groupement") : n'extrais QUE si c'est ecrit dans le DCE, sinon "Non specifie dans le DCE". "capacites_groupement" dit comment les capacites sont appreciees quand plusieurs entreprises repondent ensemble : "cumul admis" (les capacites des membres s'additionnent), "chaque membre doit satisfaire individuellement", "mandataire seul", ou "Non specifie dans le DCE". "forme_groupement.imposee" vaut "conjoint" ou "solidaire" UNIQUEMENT si le reglement impose cette forme, "libre" s'il laisse le choix, "non_specifie" sinon.
- Pour CHAQUE extrait ("exigences_sources".*, "forme_groupement.extrait") : recopie LITTERALEMENT la phrase du DCE (300 caracteres maximum), sans reformuler, sans resumer, sans corriger. Laisse "" si l'information n'est pas dans le DCE. Un extrait qui n'est pas une copie exacte du DCE sera rejete.
- "criteres_attribution" : les criteres de jugement des OFFRES ecrits dans le reglement de la consultation (criteres principaux seulement, pas les sous-criteres), dans l'ORDRE ou le reglement les cite. "ponderation" : le poids TEL QU'ECRIT dans le reglement ("60 %", "40 points"...), ou "" si le reglement ne donne pas de poids. N'invente JAMAIS de poids ni d'ordre d'importance. "extrait" : recopie LITTERALEMENT la ligne ou la phrase du reglement qui cite ce critere (avec son poids s'il existe), 300 caracteres maximum, sans reformuler ; "" si tu ne la trouves pas. Un extrait qui n'est pas une copie exacte du DCE sera rejete.
- Si une information n'est pas dans le DCE, indique "Non specifie dans le DCE"
- Extrais TOUTES les pieces demandees dans le reglement de consultation"""


# ── Analyse PARTAGÉE (docs/ANALYSE_PARTAGEE.md) : l'extraction du dossier ne reçoit plus le profil ──
# L'analyse d'un dossier se fait en deux temps :
#   (a) EXTRACTION, qui ne dépend QUE du dossier : la consigne d'origine (SYSTEM_PROMPT), dont seules
#       les phrases qui parlaient de l'entreprise candidate sont retirées ou ramenées au dossier (liste
#       exacte ci-dessous, vérifiée à l'import du module) ; l'IA ne reçoit plus ni le profil ni les
#       critères. Cette extraction peut donc être resservie, telle quelle, à toute organisation qui
#       analyse EXACTEMENT le même dossier (app.services.analyse_partagee) ;
#   (b) PARTIE PROPRE À L'ENTREPRISE : score et verdict (déjà calculés sans IA : dce_scoring), puis un
#       appel IA court (SYSTEM_ENTREPRISE) qui reçoit l'extraction, le profil, les critères et la
#       compatibilité calculée — jamais le dossier entier — et rédige ce que l'IA personnalisait
#       jusqu'ici : la stratégie (4e paragraphe de la recommandation), les points d'attention propres à
#       l'entreprise, le niveau des clauses au regard de SES seuils (pénalités, garanties, retenue,
#       avance, délai de remise) et une clause du dossier qui dépasse un seuil et manquerait à la liste.
# ANALYSE_PARTAGEE=0 : retour à l'analyse d'origine en un seul appel (consigne SYSTEM_PROMPT, profil et
# critères dans le prompt), à l'identique.
# Version des consignes de l'EXTRACTION (clé de l'analyse partagée ; le texte de SYSTEM_EXTRACTION y entre
# aussi : le modifier sans changer ce numéro invalide quand même). La consigne de la partie propre à
# l'entreprise n'y entre pas : cette partie n'est jamais gardée, elle est refaite à chaque analyse.
VERSION_CONSIGNES = "1"

# (texte d'origine, texte de l'extraction) : TOUT le reste de SYSTEM_PROMPT est repris mot pour mot.
_EXTRACTION_REMPLACEMENTS = (
    # Score et verdict : propres à l'entreprise, calculés sans IA (dce_scoring) — ceux de l'IA n'étaient
    # déjà jamais gardés (finalize_analysis les remplace).
    ('{\n  "match_score": 75,\n  "go_decision": "go",\n  "summary"', '{\n  "summary"'),
    ('"recommandation": "Analyse redigee en 3 a 4 paragraphes (180 a 350 mots), separes par une ligne vide : '
     '(1) de quoi parle le marche et ce qu\'il exige, (2) les attentes de l\'acheteur et la mecanique '
     'd\'attribution, (3) les points techniques et administratifs a maitriser et les clauses a negocier, '
     '(4) la strategie et la recommandation finale"',
     '"recommandation": "Analyse du marche redigee en 3 paragraphes (140 a 280 mots), separes par une ligne '
     'vide : (1) de quoi parle le marche et ce qu\'il exige, (2) les attentes de l\'acheteur et la mecanique '
     'd\'attribution, (3) les points techniques et administratifs a maitriser et les clauses a negocier"'),
    ('"pourquoi": "1/1000 du HT par jour plafonnees a 10%, au-dela du seuil tolere par l\'entreprise"',
     '"pourquoi": "1/1000 du HT par jour, plafonnees a 10% du montant du marche"'),
    ('- "match_score" est un entier entre 0 et 100\n- "go_decision" est "go", "no_go" ou "a_etudier"\n', ''),
    ('3 a 4 paragraphes, entre 180 et 350 mots (jamais moins de 180), separes par une ligne vide (double saut '
     'de ligne \n\n). Structure attendue : (1) de quoi parle le marche et ce qu\'il demande concretement ; (2) '
     'les attentes de l\'acheteur et la mecanique d\'attribution (poids des criteres, ce qui fait gagner) ; (3) '
     'les points techniques et administratifs a maitriser + les clauses a negocier ; (4) la strategie et la '
     'recommandation finale (y aller ou non, comment, seul ou en groupement). Reste factuel',
     '3 paragraphes, entre 140 et 280 mots (jamais moins de 140), separes par une ligne vide (double saut de '
     'ligne \n\n). Structure attendue : (1) de quoi parle le marche et ce qu\'il demande concretement ; (2) '
     'les attentes de l\'acheteur et la mecanique d\'attribution (poids des criteres, ce qui fait gagner) ; (3) '
     'les points techniques et administratifs a maitriser + les clauses a negocier. La strategie propre a '
     'l\'entreprise (y aller ou non, seul ou en groupement) est redigee ensuite, a part. Reste factuel'),
    ('  * "niveau" vaut "faible", "moyen" ou "eleve"\n'
     '  * Tu DOIS comparer aux seuils de l\'entreprise fournis dans les criteres (penalites max, garantie max, '
     'retenue max, avance min, delai de reponse min) : marque "eleve" des qu\'un seuil est depasse, et dis-le '
     'explicitement dans "pourquoi"\n',
     '  * "niveau" vaut "faible", "moyen" ou "eleve", au regard des usages des marches publics (les seuils de '
     'chaque entreprise sont compares ensuite)\n'
     '  * Dans "pourquoi", donne les valeurs EXACTES du DCE (taux, plafond, duree, montant)\n'),
    ('- Extrais TOUTES les pieces demandees dans le reglement de consultation',
     '- Extrais TOUTES les pieces demandees dans le reglement de consultation\n'
     '- Tu ne connais pas l\'entreprise candidate : n\'ecris rien qui suppose son profil (taille, CA, '
     'qualifications, lieu)'),
)


def _consigne_extraction() -> str:
    out = SYSTEM_PROMPT
    for avant, apres in _EXTRACTION_REMPLACEMENTS:
        if out.count(avant) != 1:       # la consigne d'origine a changé : à revoir ici, jamais en silence
            raise RuntimeError("consigne d'extraction : passage d'origine introuvable : " + avant[:60])
        out = out.replace(avant, apres)
    return out


SYSTEM_EXTRACTION = _consigne_extraction()


SYSTEM_ENTREPRISE = """Tu es un expert en marches publics francais. Tu travailles pour un logiciel appele Adjugo qui aide les entreprises du BTP et des services a repondre aux appels d'offres.

On te donne : (1) l'analyse OBJECTIVE d'un appel d'offres, deja extraite du dossier de consultation complet (faits, criteres et poids, exigences, clauses a risque, points d'attention, analyse du marche en 3 paragraphes) ; (2) la compatibilite de l'entreprise avec ce marche, calculee par un bareme ouvert et deterministe (score en %, detail par critere), et la detection d'une opportunite de groupement ; (3) le profil de l'entreprise candidate (identite, chiffres d'affaires, effectif, activite : specialites, zone d'intervention, qualifications) et ses criteres.

Ton role : rediger la partie PROPRE A CETTE ENTREPRISE de l'analyse. Tu dois TOUJOURS repondre en JSON valide, sans texte avant ou apres. Voici le format exact attendu :

{
  "strategie": "Un paragraphe de 60 a 130 mots : la strategie et la recommandation finale pour cette entreprise",
  "points_attention_entreprise": ["CA exige (1 500 000 EUR sur 3 ans) superieur a votre CA N-1 (900 000 EUR)"],
  "clauses": [{"index": 0, "niveau": "eleve", "pourquoi_entreprise": "Au-dela de votre seuil de penalites de retard (5%)."}],
  "clauses_ajoutees": []
}

Regles importantes :
- "strategie" est le 4e et dernier paragraphe de l'analyse (les 3 premiers, sur le marche lui-meme, te sont fournis : ne les repete pas) : y aller ou non, comment, seul ou en groupement, et ce qui fera la difference pour CETTE entreprise. Il doit etre COHERENT avec la compatibilite calculee, que l'entreprise voit a l'ecran : ne la contredis jamais. Compatibilite « a etudier » : dis ce qu'il faut verifier ; « peu adapte » : explique pourquoi ne pas y aller, ou a quelle condition (par exemple en groupement si une opportunite est detectee). Parle de compatibilite en %, jamais de « Go » ni de « No-Go ».
- "points_attention_entreprise" : 0 a 4 points PROPRES a cette entreprise (seuil de l'entreprise depasse, chiffre d'affaires exige superieur au sien, qualification exigee absente de son profil, delai de remise inferieur a son minimum, lieu hors de sa zone...), jamais un point deja present dans les points d'attention fournis.
- "clauses" : pour CHAQUE clause de la liste numerotee fournie qui depasse un seuil de l'entreprise (penalites de retard max, garantie a premiere demande max, retenue de garantie max, avance minimale souhaitee, delai de remise minimal) : "index" est son numero, "niveau" vaut "eleve", et "pourquoi_entreprise" le dit explicitement, avec le seuil de l'entreprise. N'inclus aucune autre clause. Liste vide [] si aucun seuil n'est fourni ou depasse.
- "clauses_ajoutees" : seulement une clause du dossier (presente dans les faits fournis) qui depasse un seuil de l'entreprise et qui manque a la liste numerotee, au format {"clause": "...", "niveau": "eleve", "pourquoi": "...", "levier_negociation": "contre-proposition concrete et actionnable (article du CCP si pertinent)"} ; sinon [].
- Fonde-toi UNIQUEMENT sur les elements fournis : n'invente jamais un chiffre, une exigence, une echeance, ni une capacite de l'entreprise. Si une donnee du profil manque, dis-le plutot que de la supposer."""


def _bloc_profil(company) -> str:
    """Profil de l'entreprise tel que l'analyse d'origine le remettait à l'IA (mêmes champs)."""
    prompt = ""
    if company:
        prompt += "=== PROFIL DE L'ENTREPRISE CANDIDATE ===\n"
        if company.get("name"):
            prompt += f"Denomination : {company['name']}\n"
        if company.get("siret"):
            prompt += f"SIRET : {company['siret']}\n"
        if company.get("code_ape"):
            prompt += f"Code APE : {company['code_ape']}\n"
        if company.get("forme_juridique"):
            prompt += f"Forme juridique : {company['forme_juridique']}\n"
        if company.get("city"):
            prompt += f"Ville : {company['city']}\n"
        if company.get("ca_n1"):
            prompt += f"CA N-1 : {company['ca_n1']:,.0f} EUR\n"
        if company.get("ca_n2"):
            prompt += f"CA N-2 : {company['ca_n2']:,.0f} EUR\n"
        if company.get("ca_n3"):
            prompt += f"CA N-3 : {company['ca_n3']:,.0f} EUR\n"
        if company.get("effectif"):
            prompt += f"Effectif : {company['effectif']} salaries\n"
        prompt += "\n"
    return prompt


def _bloc_criteres(criteria) -> str:
    """Critères Go/No-Go de l'entreprise tels que l'analyse d'origine les remettait à l'IA (sans la
    phrase finale « Utilise ces critères… », propre à l'analyse d'origine)."""
    prompt = ""
    if criteria:
        prompt += "=== CRITERES GO/NO-GO DE L'ENTREPRISE ===\n"
        if criteria.get("budget_min"):
            prompt += f"Budget minimum acceptable : {criteria['budget_min']:,.0f} EUR\n"
        if criteria.get("budget_max"):
            prompt += f"Budget maximum acceptable : {criteria['budget_max']:,.0f} EUR\n"
        if criteria.get("max_distance_km"):
            prompt += f"Distance max : {criteria['max_distance_km']} km\n"
        if criteria.get("penalty_max"):
            prompt += f"Penalites de retard max acceptables : {criteria['penalty_max']}%\n"
        if criteria.get("garantie_max"):
            prompt += f"Garantie a premiere demande max acceptable : {criteria['garantie_max']}%\n"
        if criteria.get("retenue_garantie_max"):
            prompt += f"Retenue de garantie max acceptable : {criteria['retenue_garantie_max']}%\n"
        if criteria.get("avance_min"):
            prompt += f"Avance minimale souhaitee : {criteria['avance_min']}%\n"
        if criteria.get("delai_reponse_min"):
            prompt += f"Delai de remise minimal acceptable : {criteria['delai_reponse_min']} jours\n"
        if criteria.get("go_threshold"):
            prompt += f"Seuil Go : score >= {criteria['go_threshold']}%\n"
        if criteria.get("nogo_threshold"):
            prompt += f"Seuil No-Go : score < {criteria['nogo_threshold']}%\n"
        if criteria.get("excluded_keywords"):
            prompt += f"Mots-cles a exclure : {criteria['excluded_keywords']}\n"
        prompt += "\n"
    return prompt


def build_user_prompt(text, company=None, criteria=None, lang_name=None, integral: bool = False,
                      note_lecture: str = ""):
    """Construit le prompt utilisateur avec le contexte entreprise. integral=True : le contenu
    du dossier est remis tel quel (lecture intégrale), précédé de `note_lecture`."""
    prompt = "Analyse le DCE suivant et produis l'analyse JSON.\n\n"
    prompt += _bloc_profil(company)
    if criteria:
        prompt += _bloc_criteres(criteria)
        prompt += "Utilise ces criteres pour ajuster le match_score et la go_decision.\n"
        prompt += "Si le budget du marche est hors des limites, penalise le score.\n"
        prompt += "Si le lieu d'execution est trop loin, penalise le score.\n"
        prompt += "Compare IMPERATIVEMENT penalites/retenue de garantie/garantie/avance/delai du DCE a ces "
        prompt += "seuils et renseigne le tableau clauses_risque (niveau eleve si un seuil est depasse).\n\n"

    if lang_name and lang_name.lower() != "français":
        prompt += (f"=== LANGUE DE SORTIE ===\nRédige TOUTES les valeurs textuelles "
                   f"(summary, recommandation, points_attention, details.*, clauses_risque : "
                   f"clause/pourquoi/levier_negociation) en {lang_name}. Conserve les CLÉS JSON "
                   f"et les énumérations en l'état (go_decision: go/no_go/a_etudier ; "
                   f"niveau: faible/moyen/eleve).\n\n")

    prompt += "=== CONTENU DU DCE ===\n"
    if integral:
        if note_lecture:
            prompt += note_lecture.strip() + "\n\n"
        return prompt + (text or "")
    selected = select_dce_content(text)
    prompt += selected
    if len(selected) < len(text):
        prompt += ("\n\n[... DCE volumineux : règlement de la consultation et cadre de réponse "
                   "remis en entier quand ils sont reconnus (sinon tête du dossier), puis extraits "
                   "ciblés des clauses sensibles (pénalités, garanties, critères, délais) dans les "
                   "autres pièces. Analyse sur la base de ces éléments.]")

    return prompt


# ── Extraction de texte BORNÉE (anti-bombe de décompression / PDF géant) ─────────
# Process unique (1 worker sans Redis) : un seul fichier piégé ne doit jamais pouvoir
# épuiser la mémoire ou le CPU pour tous les clients. Tout est donc plafonné AVANT lecture :
#   • ZIP : chaque fichier lu en flux borné (jamais zf.read sans limite), budget cumulé
#     sur l'archive entière, ZIP imbriqué accepté sur 1 seul niveau, nombre d'entrées borné ;
#   • DOCX (archive elle aussi) : tailles décompressées vérifiées avant python-docx ;
#   • PDF : nombre de pages lues borné + arrêt dès que le volume de texte utile est atteint ;
#   • OCR : seules les premières pages sont RENDUES (pas toutes puis tronquées), page par
#     page, résolution réduite pour les formats géants, budget de pages partagé dans un ZIP.
# Un vrai DCE (RC + CCAP + CCTP + DPGF + plans, ~30 Mo zippés) reste très en dessous.
_MB = 1024 * 1024
ZIP_MEMBER_MAX_BYTES = 40 * _MB      # un fichier de l'archive, une fois décompressé
ZIP_TOTAL_MAX_BYTES = 160 * _MB      # cumul décompressé lu dans toute l'archive
ZIP_MAX_ENTRIES = 2000               # entrées déclarées (un DCE en compte quelques dizaines)
ZIP_MAX_DEPTH = 1                    # ZIP dans le ZIP : oui ; ZIP dans un ZIP dans un ZIP : non
ZIP_MAX_FILES = 25                   # fichiers réellement exploités (inchangé)
DOCX_PART_MAX_BYTES = 40 * _MB       # une partie d'un .docx (image, XML…) décompressée
DOCX_XML_MAX_BYTES = 25 * _MB        # une partie XML (document.xml…) décompressée
DOCX_XML_TOTAL_MAX_BYTES = 40 * _MB  # tout le XML parsé par python-docx
DOCX_TOTAL_MAX_BYTES = 120 * _MB     # cumul décompressé d'un .docx
PDF_MAX_PAGES = 600                  # pages dont on extrait la couche texte
PDF_TIME_BUDGET_S = 25.0             # temps CPU max d'extraction texte par dossier (archive comprise)
PDF_PAGE_TIME_S = 6.0                # tranche de temps max pour UNE page (interrompue en cours)
PDF_PAGE_CONTENT_MAX_BYTES = 2 * _MB # contenu décompressé d'une page (flux + formulaires) ;
                                     # une page de texte pèse quelques Ko, l'analyse ≈ 1 s/Mo
PDF_DOC_CONTENT_MAX_BYTES = 12 * _MB # cumul décompressé sur un dossier (archive comprise)
PDF_PAGE_MAX_STREAMS = 500           # flux concaténés dans /Contents d'une seule page
PDF_PAGE_MAX_FORMS = 500             # formulaires (XObject /Form) distincts sous une page
OCR_MAX_PAGES = 15                   # pages scannées OCR-isées par PDF (inchangé)
OCR_TOTAL_MAX_PAGES = 30             # pages OCR-isées au total sur une archive
OCR_DPI = 200
OCR_MAX_SIDE_PX = 4000               # plus grand côté d'une page rendue (A4 @200 dpi ≈ 2339)
OCR_TIMEOUT_S = 60                   # rendu d'UNE page par poppler


# Seules compressions dont zipfile borne la sortie à la taille demandée.
_ZIP_SAFE_COMPRESSION = (0, 8)       # ZIP_STORED, ZIP_DEFLATED


class _Budget:
    """Budget partagé sur une extraction (archive imbriquée comprise). Sans argument : les
    plafonds historiques ci-dessus (extract_dce_text). La lecture intégrale
    (app.services.dce_lecture) passe ses propres plafonds, plus larges et réglables."""
    def __init__(self, total_bytes: int = None, ocr_pages: int = None, time_s: float = None,
                 pdf_content: int = None):
        self.bytes_left = ZIP_TOTAL_MAX_BYTES if total_bytes is None else int(total_bytes)
        self.ocr_pages_left = OCR_TOTAL_MAX_PAGES if ocr_pages is None else int(ocr_pages)
        self.skipped_big = 0        # fichiers ignorés car trop volumineux décompressés
        self.skipped_nested = 0     # archives ignorées car imbriquées trop profondément
        import time as _t
        # partagé par tous les PDF de l'archive (et par l'OCR de la lecture intégrale)
        self.deadline = _t.monotonic() + (PDF_TIME_BUDGET_S if time_s is None else float(time_s))
        # idem, en octets de contenu décompressé
        self.pdf_content_left = PDF_DOC_CONTENT_MAX_BYTES if pdf_content is None else int(pdf_content)
        self.pdf_heavy_pages = 0    # pages sautées car trop lourdes (plans vectoriels, bombes)


def _zip_entry_count_ok(content: bytes) -> bool:
    """Lit le nombre d'entrées DÉCLARÉ dans l'enregistrement de fin d'archive AVANT de
    construire ZipFile (qui matérialise un objet par entrée)."""
    import io, zipfile
    try:
        endrec = zipfile._EndRecData(io.BytesIO(content))   # API interne stable (3.8 → 3.14)
        if endrec is None:
            return True                                     # ZipFile lèvera son erreur claire
        return int(endrec[zipfile._ECD_ENTRIES_TOTAL]) <= ZIP_MAX_ENTRIES
    except Exception:
        return True


def _read_member_limited(zf, info, limit: int) -> bytes:
    """Lit un membre d'archive en flux, jamais plus de `limit` + 1 octets en mémoire.
    Retourne None si le membre dépasse la limite (taille déclarée OU réelle) ou s'il
    utilise une compression autre que « stockée »/deflate : bzip2/LZMA décompressent un
    bloc entier d'un coup, sans plafond de sortie (contournement de la lecture bornée)."""
    if info.file_size > limit or info.compress_type not in _ZIP_SAFE_COMPRESSION:
        return None
    with zf.open(info) as fh:
        # zipfile tronque la sortie à la taille DÉCLARÉE ; on ne demande donc jamais plus que
        # celle-ci (+1) : un en-tête qui ment « 1 Ko » ne fait décompresser que ~1 Ko.
        data = fh.read(min(info.file_size, limit) + 1)
    if len(data) > limit:
        return None
    return data


def _docx_sizes_ok(data: bytes) -> bool:
    """Un .docx est une archive : on vérifie ses tailles décompressées AVANT python-docx
    (qui charge toutes les parties en mémoire et parse le XML)."""
    import io, zipfile
    try:
        if not _zip_entry_count_ok(data):
            return False
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            total = xml_total = 0
            for i in z.infolist():
                if i.compress_type not in _ZIP_SAFE_COMPRESSION:
                    return False
                is_xml = i.filename.lower().endswith((".xml", ".rels"))
                if i.file_size > (DOCX_XML_MAX_BYTES if is_xml else DOCX_PART_MAX_BYTES):
                    return False
                total += i.file_size
                xml_total += i.file_size if is_xml else 0
                if total > DOCX_TOTAL_MAX_BYTES or xml_total > DOCX_XML_TOTAL_MAX_BYTES:
                    return False
        return True
    except Exception:
        return False   # archive illisible → python-docx échouerait de toute façon


def _ref_key(raw):
    """Identité d'un objet PDF référencé (numéro d'objet), None pour un objet direct."""
    return (raw.idnum, raw.generation) if hasattr(raw, "idnum") else None


def _res_signature(obj, budget: list, depth: int = 0):
    """Signature d'un dictionnaire de ressources DIRECT (inscrit dans la page) : sa structure, les
    objets référencés étant identifiés par leur numéro (sans les suivre). Deux pages aux ressources
    écrites à l'identique (même police, même formulaire référencé) ont la même signature. None si
    la structure est trop grande ou contient un flux direct (pas de signature sûre)."""
    budget[0] -= 1
    if budget[0] < 0 or depth > 6:
        return None
    if hasattr(obj, "idnum"):
        return ("ref", obj.idnum, obj.generation)
    if hasattr(obj, "get_data"):
        return None                               # flux direct
    if hasattr(obj, "items"):
        out = []
        for k in sorted(obj.keys(), key=str):
            v = _res_signature(obj.raw_get(k) if hasattr(obj, "raw_get") else obj[k], budget, depth + 1)
            if v is None:
                return None
            out.append((str(k), v))
        return ("dict", tuple(out))
    if isinstance(obj, (list, tuple)):
        out = []
        for x in obj:
            v = _res_signature(x, budget, depth + 1)
            if v is None:
                return None
            out.append(v)
        return ("list", tuple(out))
    return ("val", str(obj)[:200])


def _page_content_key(page):
    """Identité de ce qui fait le texte d'une page : ses flux de contenu ET ses ressources
    (polices, formulaires, images). Deux pages qui pointent vers les MÊMES flux et les MÊMES
    ressources ont le même texte : inutile de le redécompresser. Un flux partagé par des pages
    aux ressources différentes (PDF optimisé : « /Fm0 Do » sur chaque page, chacune avec SON
    formulaire ou SON image) n'est PAS un doublon. None si l'identité n'est pas sûre (objet
    direct) : la page est alors lue normalement."""
    try:
        raw = page.raw_get("/Contents") if hasattr(page, "raw_get") else page.get("/Contents")
        if hasattr(raw, "idnum"):
            contents = (_ref_key(raw),)
        else:
            refs = []
            for x in (raw or []):
                if not hasattr(x, "idnum"):
                    return None                   # flux direct : pas de clé fiable
                refs.append(_ref_key(x))
            contents = tuple(refs)
        if not contents:
            return None
        # Ressources : propres à la page (objet référencé) ou héritées d'un nœud parent.
        node, hops, res = page, 0, None
        while node is not None and hops < 32:
            if "/Resources" in node:
                r = node.raw_get("/Resources") if hasattr(node, "raw_get") else node.get("/Resources")
                if hasattr(r, "idnum"):
                    res = ("objet",) + _ref_key(r)
                elif node is not page and hasattr(getattr(node, "indirect_reference", None), "idnum"):
                    res = ("noeud",) + _ref_key(node.indirect_reference)
                else:
                    sig = _res_signature(r, [400])   # ressources écrites dans la page
                    if sig is None:
                        return None
                    res = ("direct", sig)
                break
            parent = node.raw_get("/Parent") if hasattr(node, "raw_get") else node.get("/Parent")
            node = parent.get_object() if parent is not None else None
            hops += 1
        return (contents, res or ("aucune",))
    except Exception:
        return None


class _PdfAbort(BaseException):
    """Interruption d'une page trop coûteuse. BaseException EXPRÈS : pypdf avale `Exception`
    autour des formulaires (XObject), une Exception ordinaire ne remonterait pas."""


def _stream_size(s, limit: int) -> int:
    """Taille décompressée d'un flux, mesurée sans jamais produire plus de `limit` + 1 octets
    pour le cas courant (Flate). Les autres filtres passent par pypdf (plafonné par flux)."""
    import zlib
    try:
        flt = s.get("/Filter")
        flt = flt.get_object() if flt is not None else None
        if isinstance(flt, list):
            flt = flt[0] if len(flt) == 1 else tuple(flt)
        enc = getattr(s, "_data", None)
        if flt is None and isinstance(enc, (bytes, bytearray)):
            return len(enc)
        if flt in ("/FlateDecode", "/Fl") and isinstance(enc, (bytes, bytearray)):
            try:
                return len(zlib.decompressobj().decompress(bytes(enc), limit + 1))
            except zlib.error:
                pass                               # flux abîmé : pypdf sait parfois le réparer
        return len(s.get_data())
    except Exception:
        return 0


def _page_content_bytes(page, limit: int) -> int:
    """Ce que extract_text va réellement analyser : les flux de /Contents ET les formulaires
    (XObject /Form, récursivement, chacun compté une fois). Renvoie dès que `limit` est
    dépassé (bombe « une page = des milliers de flux » ou « formulaire géant »)."""
    total = 0
    try:
        contents = page.get("/Contents")
        contents = contents.get_object() if contents is not None else None
        streams = []
        if isinstance(contents, list):
            if len(contents) > PDF_PAGE_MAX_STREAMS:
                return limit + 1
            streams = [c.get_object() for c in contents]
        elif contents is not None:
            streams = [contents]
        for st in streams:
            total += _stream_size(st, limit - total)
            if total > limit:
                return total
        # Formulaires référencés par les ressources (héritées des nœuds parents au besoin).
        res, node, hops = None, page, 0
        while node is not None and hops < 32:
            if "/Resources" in node:
                res = node["/Resources"].get_object()
                break
            node = node.get("/Parent")
            node = node.get_object() if node is not None else None
            hops += 1
        seen, todo = set(), [res]
        while todo:
            r = todo.pop()
            if not hasattr(r, "get"):
                continue
            xo = r.get("/XObject")
            xo = xo.get_object() if xo is not None else None
            if not hasattr(xo, "items"):
                continue
            for _name, ref in xo.items():
                key = (ref.idnum, ref.generation) if hasattr(ref, "idnum") else id(ref)
                if key in seen:
                    continue
                seen.add(key)
                if len(seen) > PDF_PAGE_MAX_FORMS:
                    return limit + 1
                f = ref.get_object()
                if not hasattr(f, "get") or f.get("/Subtype") != "/Form":
                    continue                       # image : extract_text ne la décode pas
                total += _stream_size(f, limit - total)
                if total > limit:
                    return total
                sub = f.get("/Resources")
                if sub is not None:
                    todo.append(sub.get_object())
        return total
    except Exception:
        return total


def _page_has_images(page) -> bool:
    """La page porte-t-elle une image (XObject /Image), directement ou dans un formulaire ?
    Distingue une page BLANCHE (rien à lire) d'une page SCANNÉE insérée dans un PDF texte
    (annexe signée, page tamponnée) : celle-ci n'a pas de texte mais un contenu à lire par
    reconnaissance de caractères. Ne lève jamais."""
    try:
        res, node, hops = None, page, 0
        while node is not None and hops < 32:
            if "/Resources" in node:
                res = node["/Resources"].get_object()
                break
            node = node.get("/Parent")
            node = node.get_object() if node is not None else None
            hops += 1
        todo, seen = [res], 0
        while todo and seen < 200:
            r = todo.pop()
            if not hasattr(r, "get"):
                continue
            xo = r.get("/XObject")
            xo = xo.get_object() if xo is not None else None
            if not hasattr(xo, "items"):
                continue
            for _name, ref in xo.items():
                seen += 1
                f = ref.get_object()
                if not hasattr(f, "get"):
                    continue
                if f.get("/Subtype") == "/Image":
                    return True
                if f.get("/Subtype") == "/Form" and f.get("/Resources") is not None:
                    todo.append(f.get("/Resources").get_object())
        return False
    except Exception:
        return False


# Page « scannée » : une image couvre l'essentiel de la page et le texte tapé est court (bandeau
# « Annexe 2 – document numérisé », tampon de la plateforme de téléchargement) — elle passe par la
# reconnaissance de caractères comme une page sans texte.
SCAN_COUVERTURE_MIN = 0.5            # part de la surface de la page couverte par une image
SCAN_TEXTE_MAX = 200                 # caractères tapés au plus sur une page scannée


class _ImageCover:
    """Suit, pendant l'extraction du texte d'une page, la plus grande surface couverte par une
    image (XObject /Image), y compris dans un formulaire : surface = |déterminant| de la matrice
    courante × celles des formulaires englobants. Ne lève jamais."""

    def __init__(self, page):
        self.max_area = 0.0
        try:
            res = page.get("/Resources")
            res = res.get_object() if res is not None else None
        except Exception:
            res = None
        self.stack = [(res, 1.0)]
        try:
            mb = page.mediabox
            self.page_area = abs(float(mb.width) * float(mb.height))
        except Exception:
            self.page_area = 0.0

    @staticmethod
    def _det(m) -> float:
        try:
            return abs(float(m[0]) * float(m[3]) - float(m[1]) * float(m[2]))
        except Exception:
            return 0.0

    def before(self, op, args, cm):
        if op != b"Do":
            return
        res, scale = self.stack[-1]
        push = (res, scale)
        try:
            xo = res.get("/XObject") if hasattr(res, "get") else None
            xo = xo.get_object() if xo is not None else None
            obj = xo[args[0]].get_object() if (xo is not None and args and args[0] in xo) else None
            if obj is not None:
                area = self._det(cm) * scale
                sub = obj.get("/Subtype")
                if sub == "/Image":
                    self.max_area = max(self.max_area, area)
                elif sub == "/Form":
                    fres = obj.get("/Resources")
                    fres = fres.get_object() if fres is not None else res
                    mtx = obj.get("/Matrix")
                    mdet = self._det([float(x) for x in mtx]) if mtx is not None else 1.0
                    push = (fres, area * mdet)
        except Exception:
            pass
        self.stack.append(push)

    def after(self, op):
        if op == b"Do" and len(self.stack) > 1:
            self.stack.pop()

    def couverture(self) -> float:
        return (self.max_area / self.page_area) if self.page_area > 0 else 0.0


def extract_text_from_pdf(file_bytes, max_pages: int = PDF_MAX_PAGES, max_chars: int = None,
                          deadline: float = None, budget: "_Budget" = None, info: dict = None,
                          pages_out: list = None):
    """Extrait le texte d'un PDF — au plus `max_pages` pages, arrêt dès `max_chars`
    caractères (le texte au-delà serait de toute façon tronqué par l'appelant).
    Garde-fous CPU et mémoire (process unique) :
      • une page qui réutilise les flux ET les ressources d'une page déjà lue reprend son texte
        sans le redécompresser (« toutes les pages pointent vers le même flux de 74 Mo ») ;
      • une page dont le contenu décompressé (flux + formulaires) dépasse
        PDF_PAGE_CONTENT_MAX_BYTES est sautée, et le cumul est plafonné par dossier ;
      • chaque page dispose d'une tranche de temps, contrôlée PENDANT l'analyse (à chaque
        opérateur, y compris dans les formulaires appelés en boucle) ;
      • l'extraction s'arrête au-delà du budget de temps global (partagé par une archive).
    info (dict facultatif) : reçoit "complet" = False si des pages n'ont pas été lues (plafond de
    pages, page trop lourde sautée, temps écoulé, page interrompue) — jamais « lu en entier » à
    tort ; "pages" = nombre de pages du PDF ; "protege" = True pour un PDF chiffré illisible.
    pages_out (liste facultative, lecture intégrale) : reçoit UNE entrée par page du PDF,
    {"num", "texte", "etat"}, etat ∈ lu | vide (page blanche) | image (page scannée : sans
    texte, ou image couvrant la page avec un simple bandeau tapé) | lourde (plan ou dessin trop
    lourd) | coupee (lecture interrompue) | temps (temps écoulé) | plafond (au-delà de max_pages) |
    plafond_caracteres (au-delà de max_chars) : chaque page non lue est nommée, avec sa raison.
    Une page identique à une page déjà vue (mêmes flux, mêmes ressources) reprend son état et son
    texte."""
    import time
    if info is not None:
        info["complet"] = True
    if budget is not None:
        deadline = budget.deadline
    if deadline is None:
        deadline = time.monotonic() + PDF_TIME_BUDGET_S
    content_left = budget.pdf_content_left if budget is not None else PDF_DOC_CONTENT_MAX_BYTES
    track = pages_out is not None

    def _mark(i, etat, texte=""):
        if track:
            pages_out.append({"num": i + 1, "texte": texte, "etat": etat})

    try:
        from pypdf import PdfReader
        import io
        reader = PdfReader(io.BytesIO(file_bytes))
        if getattr(reader, "is_encrypted", False):
            # Chiffré sans mot de passe d'ouverture (droits restreints seulement) : lisible.
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                if info is not None:
                    info["complet"] = False
                    info["protege"] = True
                return "Erreur extraction PDF : fichier protégé par un mot de passe"
        total_pages = len(reader.pages)
        if info is not None:
            info["pages"] = total_pages
        parts, total, seen = [], 0, {}
        n = min(total_pages, max_pages)
        if info is not None and total_pages > n:
            info["complet"] = False                # plafond de pages atteint
        stopped = None
        for i in range(n):
            if stopped:
                _mark(i, stopped)
                continue
            if time.monotonic() > deadline:
                if info is not None:
                    info["complet"] = False
                if not track:
                    break
                stopped = "temps"
                _mark(i, stopped)
                continue
            page = reader.pages[i]
            key = _page_content_key(page)
            if key is not None and key in seen:
                # Même contenu ET mêmes ressources qu'une page déjà vue : même texte, même état.
                etat0, txt0 = seen[key]
                _mark(i, etat0, txt0)
                if txt0:
                    parts.append(txt0)
                    total += len(txt0) + 2
                    if max_chars and total >= max_chars:
                        if i + 1 < total_pages and info is not None:
                            info["complet"] = False
                        if not track:
                            break
                        stopped = "plafond_caracteres"
                continue
            size = _page_content_bytes(page, min(PDF_PAGE_CONTENT_MAX_BYTES, content_left))
            if size > PDF_PAGE_CONTENT_MAX_BYTES or size > content_left:
                # Page trop lourde (plan vectoriel, bombe) ou budget presque épuisé : on la
                # saute, les pages de texte (quelques Ko) suivantes passent encore.
                if budget is not None:
                    budget.pdf_heavy_pages += 1
                if info is not None:
                    info["complet"] = False
                _mark(i, "lourde")
                if key is not None:
                    seen[key] = ("lourde", "")
                continue
            content_left -= size
            page_deadline = min(deadline, time.monotonic() + PDF_PAGE_TIME_S)
            cap = (max_chars - total + 2000) if max_chars else 2_000_000
            frags, got = [], [0]
            cover = _ImageCover(page) if track else None

            def _before(_op, _args, _cm, _tm):
                if time.monotonic() > page_deadline:
                    raise _PdfAbort()
                if cover is not None:
                    cover.before(_op, _args, _cm)

            def _after(_op, _args, _cm, _tm):
                if cover is not None:
                    cover.after(_op)

            def _text(t, _cm, _tm, _font, _size):
                if t:
                    room = cap - got[0]
                    got[0] += len(t)
                    if got[0] <= cap:
                        frags.append(t)
                    else:
                        if room > 0:
                            frags.append(t[:room])
                        raise _PdfAbort()

            etat = "lu"
            try:
                page_text = page.extract_text(visitor_operand_before=_before, visitor_operand_after=_after,
                                              visitor_text=_text)
            except _PdfAbort:
                page_text = "".join(frags)         # ce qui a été lu avant l'interruption
                etat = "coupee"
                if info is not None:
                    info["complet"] = False
            if track and etat == "lu":
                court = len((page_text or "").strip())
                if court < 20:
                    etat = "image" if _page_has_images(page) else "vide"
                elif court < SCAN_TEXTE_MAX and cover is not None and cover.couverture() >= SCAN_COUVERTURE_MIN:
                    etat = "image"          # page scannée portant un bandeau tapé
            _mark(i, etat, page_text or "")
            if key is not None:
                seen[key] = (etat, page_text or "")
            if page_text:
                parts.append(page_text)
                total += len(page_text) + 2
                if max_chars and total >= max_chars:
                    if i + 1 < total_pages and info is not None:
                        info["complet"] = False
                    if not track:
                        break
                    stopped = "plafond_caracteres"
        for i in range(n, total_pages if track else n):
            _mark(i, "plafond")
        return "\n\n".join(parts).strip()
    except Exception as e:
        if info is not None:
            info["complet"] = False
            if type(e).__name__ in ("FileNotDecryptedError", "DependencyError"):
                info["protege"] = True
        return f"Erreur extraction PDF : {str(e)}"
    finally:
        if budget is not None:
            budget.pdf_content_left = content_left


OCR_MIN_DPI = 10                     # en dessous, on borne directement en pixels (size=)


def _ocr_side_pt(page) -> float:
    """Plus grand côté de la page en points, tel que poppler le rendra : valeurs ABSOLUES
    (une MediaBox inversée [9500 9500 0 0] donne chez pypdf une largeur NÉGATIVE, alors que
    poppler remet les coins dans l'ordre et rend une page de 132 pouces), × /UserUnit s'il
    agrandit la page. 0.0 si la taille est illisible ou absurde (→ rendu borné en pixels)."""
    try:
        import math
        b = page.mediabox
        side = max(abs(float(b.right) - float(b.left)), abs(float(b.top) - float(b.bottom)))
        try:
            uu = float(page.get("/UserUnit", 1) or 1)
        except Exception:
            uu = 1.0
        if math.isfinite(uu) and uu > 1:
            side *= uu
        return side if math.isfinite(side) and side > 0 else 0.0
    except Exception:
        return 0.0


def _ocr_dpi_for(page) -> int:
    """Résolution de rendu : 200 dpi, réduite pour qu'une page géante (affiche, plan A0,
    format piégé 200×200 pouces) ne produise jamais une image de plusieurs Go.
    0 si la taille de la page est inconnue : l'appelant borne alors en pixels (size=)."""
    side_pt = _ocr_side_pt(page)
    if side_pt <= 0:
        return 0
    return min(OCR_DPI, int(OCR_MAX_SIDE_PX * 72 / side_pt))


def _ocr_render_kw(page) -> dict:
    """Paramètres de rendu poppler pour UNE page. Cas normal : dpi calculé. Si pypdf ne lit
    pas la page (page=None), si sa taille est inconnue ou si elle est si grande que même
    OCR_MIN_DPI dépasserait le plafond : size=OCR_MAX_SIDE_PX (pdf2image → pdftoppm
    -scale-to), borne en pixels du plus grand côté, quelle que soit la boîte déclarée."""
    dpi = _ocr_dpi_for(page) if page is not None else 0
    if dpi >= OCR_MIN_DPI:
        return {"dpi": dpi}
    return {"dpi": OCR_DPI // 2, "size": OCR_MAX_SIDE_PX}


def _ocr_pdf(content: bytes, max_pages: int = OCR_MAX_PAGES, budget: "_Budget" = None,
             info: dict = None) -> str:
    """OCR de repli pour les DCE SCANNÉS (PDF image, fréquents chez les petites
    collectivités). Gracieux : retourne '' si l'OCR n'est pas disponible (binaire
    tesseract/poppler absent) — le code appelant retombe alors sur l'erreur claire.
    BORNÉ : poppler ne rend QUE les premières pages (first_page/last_page), une à la fois,
    au lieu de rendre tout le PDF en images puis d'en garder 15.
    info (dict facultatif) : reçoit "complet" = True seulement si TOUTES les pages ont été lues
    (un scan de 40 pages lu sur 15 n'est jamais dit « lu en entier »)."""
    if info is not None:
        info["complet"] = False
    try:
        import pytesseract
        from pdf2image import convert_from_bytes
    except Exception:
        return ""
    if budget is not None:
        max_pages = min(max_pages, budget.ocr_pages_left)
    if max_pages <= 0:
        return ""
    pages, total_pages = None, None
    try:
        import io
        from pypdf import PdfReader
        pages = PdfReader(io.BytesIO(content)).pages
        total_pages = len(pages)
        max_pages = min(max_pages, len(pages))
    except Exception:
        pages = None   # PDF que pypdf ne lit pas : poppler borne quand même via last_page
    out, rendered, reached_end = [], 0, False
    for num in range(1, max_pages + 1):
        try:
            page = pages[num - 1] if pages is not None else None
        except Exception:
            page = None
        kw = dict(_ocr_render_kw(page), first_page=num, last_page=num)
        try:
            try:
                images = convert_from_bytes(content, timeout=OCR_TIMEOUT_S, **kw)
            except TypeError:      # pdf2image ancien sans `timeout` : on garde la borne de pages
                images = convert_from_bytes(content, **kw)
        except Exception:
            break
        if not images:
            reached_end = True
            break          # au-delà de la dernière page
        rendered += 1
        if budget is not None:
            budget.ocr_pages_left -= 1
        for img in images[:1]:
            try:
                out.append(pytesseract.image_to_string(img, lang="fra"))
            except Exception:
                try:
                    out.append(pytesseract.image_to_string(img))   # repli langue par défaut
                except Exception:
                    pass
        del images
    if info is not None:
        info["complet"] = bool(reached_end or (total_pages is not None and rendered >= total_pages))
    return "\n\n".join(t for t in out if t).strip()


def ocr_available() -> bool:
    """La reconnaissance de caractères est-elle installée (modules + binaires) ? Ne lève jamais."""
    try:
        import pytesseract
        from pdf2image import convert_from_bytes  # noqa: F401
    except Exception:
        return False
    try:
        import shutil
        cmd = getattr(getattr(pytesseract, "pytesseract", None), "tesseract_cmd", "tesseract") or "tesseract"
        return bool(shutil.which(cmd)) or not hasattr(pytesseract, "pytesseract")
    except Exception:
        return True


def _tesseract(img) -> str:
    import pytesseract
    try:
        return pytesseract.image_to_string(img, lang="fra") or ""
    except Exception:
        try:
            return pytesseract.image_to_string(img) or ""   # repli langue par défaut
        except Exception:
            return ""


def ocr_pdf_pages(content: bytes, page_nums, budget: "_Budget", results: dict,
                  stop=None) -> dict:
    """Lecture intégrale : OCR des SEULES pages demandées (pages scannées d'un PDF, ou toutes les
    pages d'un scan), une à la fois, sous le budget partagé de la lecture (pages et temps).
    results[num] reçoit le texte reconnu de chaque page rendue. stop() (facultatif) : vrai pour
    interrompre (temps écoulé). Renvoie {"disponible", "non_lues": {num: raison}} — raison ∈
    plafond_ocr | temps | rendu (page que poppler n'a pas pu rendre). Ne lève jamais."""
    import time
    out = {"disponible": True, "non_lues": {}}
    nums = [int(n) for n in (page_nums or [])]
    try:
        from pdf2image import convert_from_bytes
        import pytesseract  # noqa: F401
    except Exception:
        out["disponible"] = False
        for n in nums:
            out["non_lues"][n] = "ocr_indisponible"
        return out
    pages = None
    try:
        import io
        from pypdf import PdfReader
        pages = PdfReader(io.BytesIO(content)).pages
    except Exception:
        pages = None
    for idx, num in enumerate(nums):
        if (stop is not None and stop()) or time.monotonic() > budget.deadline:
            for n in nums[idx:]:
                out["non_lues"][n] = "temps"
            break
        if budget.ocr_pages_left <= 0:
            for n in nums[idx:]:
                out["non_lues"][n] = "plafond_ocr"
            break
        try:
            page = pages[num - 1] if pages is not None else None
        except Exception:
            page = None
        kw = dict(_ocr_render_kw(page), first_page=num, last_page=num)
        try:
            try:
                images = convert_from_bytes(content, timeout=OCR_TIMEOUT_S, **kw)
            except TypeError:
                images = convert_from_bytes(content, **kw)
        except Exception:
            images = None
        budget.ocr_pages_left -= 1
        if not images:
            out["non_lues"][num] = "rendu"
            continue
        results[num] = _tesseract(images[0])
        del images
    return out


OCR_IMAGE_MAX_PIXELS = 40_000_000    # image envoyée par l'acheteur (scan d'une page) : ~A0 à 150 dpi


def ocr_image(data: bytes, budget: "_Budget") -> tuple:
    """Lecture intégrale : OCR d'une IMAGE (.png, .jpg, .tif : page scannée envoyée telle quelle).
    BORNÉ : nombre de pixels plafonné AVANT décodage (bombe d'image), réduite au-delà de
    OCR_MAX_SIDE_PX, une page du budget OCR par image (les .tif multipages : première page).
    Renvoie (texte, raison) ; raison '' si lue, sinon ocr_indisponible | plafond_ocr | temps |
    illisible."""
    import time
    if budget.ocr_pages_left <= 0:
        return "", "plafond_ocr"
    if time.monotonic() > budget.deadline:
        return "", "temps"
    try:
        import pytesseract  # noqa: F401
        from PIL import Image
    except Exception:
        return "", "ocr_indisponible"
    try:
        import io
        img = Image.open(io.BytesIO(data))
        w, h = img.size
        if w * h > OCR_IMAGE_MAX_PIXELS or w <= 0 or h <= 0:
            return "", "illisible"
        img.load()
        side = max(w, h)
        if side > OCR_MAX_SIDE_PX:
            r = OCR_MAX_SIDE_PX / float(side)
            img = img.resize((max(1, int(w * r)), max(1, int(h * r))))
        budget.ocr_pages_left -= 1
        return _tesseract(img), ""
    except Exception:
        return "", "illisible"


# ── Tableurs et documents OpenDocument (sans dépendance : zipfile + lxml déjà présent) ──
_XLSX_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_XLSX_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
SHEET_MAX_CELLS = 400_000            # cellules lues dans un tableur (un BPU en compte quelques milliers)


def _xml_root(data: bytes):
    """Analyse XML SÛRE (pas d'entités externes, pas de réseau, pas d'arbre géant)."""
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False,
                             remove_comments=True)
    return etree.fromstring(data, parser=parser)


def _office_parts_ok(data: bytes) -> bool:
    """Un .xlsx/.ods/.odt est une archive : tailles décompressées vérifiées avant lecture
    (mêmes plafonds qu'un .docx)."""
    return _docx_sizes_ok(data)


def _col_index(ref: str) -> int:
    n = 0
    for ch in ref:
        if "A" <= ch <= "Z":
            n = n * 26 + (ord(ch) - 64)
        else:
            break
    return max(0, n - 1)


def extract_xlsx(data: bytes, info: dict = None) -> str:
    """Texte d'un classeur .xlsx, feuille par feuille (« === Feuille : nom === »), une ligne de
    texte par ligne du tableau (« cellule | cellule »). Valeurs affichées des formules lues
    telles qu'enregistrées. info["complet"] = False si le plafond de cellules est atteint.
    '' si illisible."""
    import io
    import zipfile
    if info is not None:
        info["complet"] = True
    if not _office_parts_ok(data):
        return ""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = set(z.namelist())
            shared = []
            if "xl/sharedStrings.xml" in names:
                root = _xml_root(z.read("xl/sharedStrings.xml"))
                for si in root.iter(_XLSX_NS + "si"):
                    shared.append("".join(t.text or "" for t in si.iter(_XLSX_NS + "t")))
            sheets = []
            try:
                wb = _xml_root(z.read("xl/workbook.xml"))
                rels = {}
                if "xl/_rels/workbook.xml.rels" in names:
                    rr = _xml_root(z.read("xl/_rels/workbook.xml.rels"))
                    for r in rr:
                        rels[r.get("Id")] = r.get("Target") or ""
                for sh in wb.iter(_XLSX_NS + "sheet"):
                    target = rels.get(sh.get(_XLSX_REL_NS + "id"), "")
                    target = target.lstrip("/")
                    path = target if target.startswith("xl/") else "xl/" + target
                    sheets.append((sh.get("name") or "", path))
            except Exception:
                sheets = []
            if not sheets:
                sheets = [("", n) for n in sorted(names) if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
            out, cells = [], 0
            for title, path in sheets:
                if path not in names:
                    continue
                root = _xml_root(z.read(path))
                rows = []
                for row in root.iter(_XLSX_NS + "row"):
                    vals = {}
                    for c in row.iter(_XLSX_NS + "c"):
                        cells += 1
                        if cells > SHEET_MAX_CELLS:
                            break
                        t = c.get("t") or ""
                        v = c.find(_XLSX_NS + "v")
                        if t == "s" and v is not None and (v.text or "").strip().isdigit():
                            i = int(v.text)
                            txt = shared[i] if 0 <= i < len(shared) else ""
                        elif t == "inlineStr":
                            txt = "".join(x.text or "" for x in c.iter(_XLSX_NS + "t"))
                        else:
                            txt = (v.text or "") if v is not None else ""
                        txt = txt.strip()
                        if txt:
                            vals[_col_index(c.get("r") or "")] = txt
                    if vals:
                        rows.append(" | ".join(vals[k] for k in sorted(vals)))
                    if cells > SHEET_MAX_CELLS:
                        if info is not None:
                            info["complet"] = False
                        break
                if rows:
                    out.append((f"=== Feuille : {title} ===\n" if title else "") + "\n".join(rows))
                if cells > SHEET_MAX_CELLS:
                    break
            return "\n\n".join(out).strip()
    except Exception:
        if info is not None:
            info["complet"] = False
        return ""


_ODF_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_ODF_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
ODF_MAX_REPEAT = 50                  # répétitions de lignes/cellules vides recopiées (tableurs ODF)


def _odf_text(el) -> str:
    """Texte d'un élément ODF (espaces <text:s/>, tabulations, sauts de ligne compris)."""
    parts = []
    if el.text:
        parts.append(el.text)
    for ch in el:
        tag = ch.tag if isinstance(ch.tag, str) else ""
        if tag == _ODF_TEXT + "s":
            parts.append(" " * min(int(ch.get(_ODF_TEXT + "c") or 1), 50))
        elif tag in (_ODF_TEXT + "tab",):
            parts.append("\t")
        elif tag == _ODF_TEXT + "line-break":
            parts.append("\n")
        else:
            parts.append(_odf_text(ch))
        if ch.tail:
            parts.append(ch.tail)
    return "".join(parts)


def extract_odf(data: bytes, info: dict = None) -> str:
    """Texte d'un document OpenDocument (.odt traitement de texte, .ods tableur) : paragraphes et
    titres dans l'ordre, tableaux ligne par ligne (« cellule | cellule »). '' si illisible."""
    import io
    import zipfile
    if info is not None:
        info["complet"] = True
    if not _office_parts_ok(data):
        return ""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if "content.xml" not in z.namelist():
                return ""
            root = _xml_root(z.read("content.xml"))
    except Exception:
        if info is not None:
            info["complet"] = False
        return ""
    out, cells = [], [0]

    def _walk(el):
        for ch in el:
            tag = ch.tag if isinstance(ch.tag, str) else ""
            if tag in (_ODF_TEXT + "p", _ODF_TEXT + "h"):
                t = _odf_text(ch).strip()
                if t:
                    out.append(t)
            elif tag == _ODF_TABLE + "table":
                name = ch.get(_ODF_TABLE + "name") or ""
                rows = []
                for row in ch.iter(_ODF_TABLE + "table-row"):
                    vals = []
                    for c in row:
                        if not isinstance(c.tag, str) or not c.tag.endswith("table-cell"):
                            continue
                        cells[0] += 1
                        if cells[0] > SHEET_MAX_CELLS:
                            break
                        t = " ".join(_odf_text(p).strip() for p in c.iter(_ODF_TEXT + "p")).strip()
                        if t:
                            rep = min(int(c.get(_ODF_TABLE + "number-columns-repeated") or 1), ODF_MAX_REPEAT)
                            vals.extend([t] * rep)
                    if vals:
                        rows.append(" | ".join(vals))
                    if cells[0] > SHEET_MAX_CELLS:
                        if info is not None:
                            info["complet"] = False
                        break
                if rows:
                    out.append((f"=== Feuille : {name} ===\n" if name else "") + "\n".join(rows))
            else:
                _walk(ch)

    try:
        _walk(root)
    except Exception:
        if info is not None:
            info["complet"] = False
    return "\n".join(out).strip()


_W_NS ="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# Blocs parcourus dans un .docx (paragraphes, où qu'ils soient, et lignes de tableau) : plafond
# TECHNIQUE (temps processeur), bien au-delà d'un vrai document — le XML lu est déjà borné
# (DOCX_XML_MAX_BYTES). Atteint : la pièce est dite « lue en partie ».
_DOCX_MAX_BLOCKS = 200000
_CT_DOCM = "application/vnd.ms-word.document.macroEnabled.main+xml"


def _roman(n: int) -> str:
    vals = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
            (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))
    out, n = "", max(1, min(int(n), 3999))
    for v, r in vals:
        while n >= v:
            out += r
            n -= v
    return out


def _fmt_list_number(n: int, fmt: str) -> str:
    """Numéro d'un niveau de liste Word selon son format (w:numFmt)."""
    fmt = fmt or "decimal"
    if fmt == "bullet":
        return "-"
    if fmt == "none":
        return ""
    if fmt == "decimalZero":
        return f"{n:02d}"
    if fmt in ("lowerLetter", "upperLetter"):
        k = max(1, n)
        s = chr(ord("a") + (k - 1) % 26) * ((k - 1) // 26 + 1)
        return s.upper() if fmt == "upperLetter" else s
    if fmt in ("lowerRoman", "upperRoman"):
        r = _roman(n)
        return r.upper() if fmt == "upperRoman" else r
    return str(n)


class _DocxNumbering:
    """Rejoue la numérotation AUTOMATIQUE des listes Word (« Question 1 », « 2.1 », « a) »),
    absente de paragraph.text : sans elle, un cadre de réponse perdait ses numéros de
    questions. Gère les formats courants, les niveaux, les styles numérotés (titres) et les
    redémarrages (startOverride). Tout écart de format → numéro décimal (jamais d'exception)."""

    def __init__(self, doc):
        self.nums, self.abstracts, self.style_num = {}, {}, {}
        self.counters, self.started = {}, set()
        try:
            root = doc.part.numbering_part.element
            for an in root.findall(_W_NS + "abstractNum"):
                lvls = {}
                for lvl in an.findall(_W_NS + "lvl"):
                    il = int(lvl.get(_W_NS + "ilvl") or 0)
                    f, t, st, ps = (lvl.find(_W_NS + x) for x in ("numFmt", "lvlText", "start", "pStyle"))
                    lvls[il] = {"fmt": f.get(_W_NS + "val") if f is not None else "decimal",
                                "text": t.get(_W_NS + "val") if t is not None else f"%{il + 1}.",
                                "start": int(st.get(_W_NS + "val")) if st is not None else 1,
                                "pstyle": ps.get(_W_NS + "val") if ps is not None else None}
                self.abstracts[an.get(_W_NS + "abstractNumId")] = lvls
            for num in root.findall(_W_NS + "num"):
                a = num.find(_W_NS + "abstractNumId")
                ov = {}
                for o in num.findall(_W_NS + "lvlOverride"):
                    so = o.find(_W_NS + "startOverride")
                    if so is not None:
                        ov[int(o.get(_W_NS + "ilvl") or 0)] = int(so.get(_W_NS + "val"))
                self.nums[num.get(_W_NS + "numId")] = (a.get(_W_NS + "val") if a is not None else None, ov)
        except Exception:
            self.nums, self.abstracts = {}, {}
        try:
            styles = {}
            for st in doc.styles.element.findall(_W_NS + "style"):
                sid = st.get(_W_NS + "styleId")
                ppr = st.find(_W_NS + "pPr")
                npr = ppr.find(_W_NS + "numPr") if ppr is not None else None
                based = st.find(_W_NS + "basedOn")
                styles[sid] = (self._numpr(npr), based.get(_W_NS + "val") if based is not None else None)
            for sid in styles:
                cur, hops = sid, 0
                while cur in styles and hops < 10:
                    npr, based = styles[cur]
                    if npr and npr[0]:
                        self.style_num[sid] = npr
                        break
                    cur, hops = based, hops + 1
        except Exception:
            self.style_num = {}

    @staticmethod
    def _numpr(npr):
        if npr is None:
            return None
        nid, il = npr.find(_W_NS + "numId"), npr.find(_W_NS + "ilvl")
        return (nid.get(_W_NS + "val") if nid is not None else None,
                int(il.get(_W_NS + "val")) if il is not None else None)

    def prefix(self, p_el) -> str:
        try:
            ppr = p_el.find(_W_NS + "pPr")
            direct = self._numpr(ppr.find(_W_NS + "numPr")) if ppr is not None else None
            sid = None
            if ppr is not None:
                ps = ppr.find(_W_NS + "pStyle")
                sid = ps.get(_W_NS + "val") if ps is not None else None
            from_style = self.style_num.get(sid) if sid else None
            num_id = (direct[0] if direct and direct[0] else None) or (from_style[0] if from_style else None)
            if not num_id or num_id == "0" or num_id not in self.nums:
                return ""
            abs_id, overrides = self.nums[num_id]
            lvls = self.abstracts.get(abs_id) or {}
            il = direct[1] if direct and direct[1] is not None else (from_style[1] if from_style else None)
            if il is None:
                il = next((k for k, v in lvls.items() if sid and v.get("pstyle") == sid), 0)
            il = max(0, min(int(il), 8))
            st = self.counters.setdefault(abs_id, {})
            if num_id not in self.started:
                self.started.add(num_id)
                for k, v in overrides.items():
                    st[k] = v - 1
            lv = lvls.get(il) or {"fmt": "decimal", "text": f"%{il + 1}.", "start": 1}
            st[il] = st.get(il, lv["start"] - 1) + 1
            for k in [k for k in st if k > il]:
                del st[k]
            if lv["fmt"] == "bullet":
                return "- "

            def _sub(m):
                k = int(m.group(1)) - 1
                lk = lvls.get(k) or {"fmt": "decimal", "start": 1}
                return _fmt_list_number(st.get(k, lk.get("start", 1)), lk.get("fmt"))
            txt = re.sub(r"%([1-9])", _sub, lv.get("text") or "")
            return (txt.strip() + " ") if txt.strip() else ""
        except Exception:
            return ""


_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_DML_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
# Sous-arbres d'un paragraphe dont le texte n'est PAS celui du document : texte supprimé par une
# modification suivie (w:del, w:moveFrom), double de compatibilité (mc:Fallback), encadrés (lus à
# part, à leur place), propriétés.
_DOCX_SKIP = frozenset({_W_NS + "del", _W_NS + "moveFrom", _W_NS + "txbxContent", _MC_FALLBACK,
                        _W_NS + "pPr", _W_NS + "rPr", _W_NS + "sdtPr", _W_NS + "sdtEndPr",
                        _W_NS + "instrText", _W_NS + "delText"})
_DOCX_WRAPPERS = (_W_NS + "sdt", _W_NS + "customXml")


def _docx_inline_text(el, parts: list, notes=None) -> None:
    """Texte d'un paragraphe Word à partir de TOUS ses descendants (w:t, tabulations, sauts de
    ligne) : zones de saisie en ligne (w:sdt), insertions suivies (w:ins, w:moveTo), champs
    (w:fldSimple, résultat des champs complexes), liens, balises personnalisées compris — que
    python-docx (Paragraph.text) ne lit pas. notes : appels de notes « (note n) » à leur place."""
    for ch in el:
        tag = ch.tag
        if not isinstance(tag, str) or tag in _DOCX_SKIP:
            continue
        if tag == _W_NS + "t":
            parts.append(ch.text or "")
        elif tag in (_W_NS + "tab", _W_NS + "ptab"):
            parts.append("\t")
        elif tag == _W_NS + "br":
            if (ch.get(_W_NS + "type") or "textWrapping") == "textWrapping":
                parts.append("\n")
        elif tag == _W_NS + "cr":
            parts.append("\n")
        elif tag == _W_NS + "noBreakHyphen":
            parts.append("-")
        elif tag in (_W_NS + "footnoteReference", _W_NS + "endnoteReference") and notes is not None:
            parts.append(notes.ref(tag, ch.get(_W_NS + "id")))
        else:
            _docx_inline_text(ch, parts, notes)


class _DocxNotes:
    """Appels de notes de bas de page et de fin, numérotés dans l'ordre où le texte les appelle."""

    def __init__(self):
        self.num = {}

    def ref(self, tag, nid) -> str:
        kind = "fin" if tag.endswith("endnoteReference") else "bas"
        key = (kind, str(nid))
        if key not in self.num:
            self.num[key] = sum(1 for k in self.num if k[0] == kind) + 1
        n = self.num[key]
        return f" (note {n})" if kind == "bas" else f" (note de fin {n})"


def _docx_unwrap(container, tag):
    """Enfants `tag` d'un élément, y compris ceux rangés dans une zone de saisie (w:sdt) ou une
    balise personnalisée (w:customXml) : section répétitive de lignes, cellule dans un contrôle."""
    for ch in container.iterchildren():
        if ch.tag == tag:
            yield ch
        elif ch.tag == _W_NS + "sdt":
            c = ch.find(_W_NS + "sdtContent")
            if c is not None:
                yield from _docx_unwrap(c, tag)
        elif ch.tag == _W_NS + "customXml":
            yield from _docx_unwrap(ch, tag)


def _docx_blocks(doc, info: dict = None) -> list:
    """Texte d'un .docx DANS L'ORDRE du document : paragraphes (avec leur numéro de liste et tout
    leur contenu : zones de saisie, insertions suivies, champs), tableaux ligne par ligne
    (« cellule | cellule », fusion verticale lue une fois, tableaux imbriqués, lignes et cellules
    rangées dans des zones de saisie), contrôles de contenu (w:sdt) et zones de texte (encadrés) —
    un cadre de réponse range souvent ses questions dans des tableaux, des zones de saisie ou des
    encadrés. Puis les notes de bas de page et de fin, les en-têtes et pieds de page et le texte des
    schémas (SmartArt), chacun dans sa rubrique. info["complet"] = False si le plafond de blocs
    est atteint (jamais « lu en entier » à tort)."""
    numbering = _DocxNumbering(doc)
    notes = _DocxNotes()
    out, seen = [], [0]

    def _para(el) -> str:
        parts = []
        _docx_inline_text(el, parts, notes)
        t = "".join(parts)
        return (numbering.prefix(el) + t) if t.strip() else ""

    def _over() -> bool:
        if seen[0] > _DOCX_MAX_BLOCKS:
            if info is not None:
                info["complet"] = False
            return True
        return False

    def _textboxes(el, dest):
        # Encadrés (wps/v:textbox) : le contenu est rangé deux fois (mc:Choice et mc:Fallback),
        # on ne lit que le premier.
        for tb in el.iter(_W_NS + "txbxContent"):
            if any(a.tag == _MC_FALLBACK for a in tb.iterancestors()):
                continue
            _walk(tb, dest)

    def _cell_text(tc) -> str:
        sub = []
        _walk(tc, sub, sep=" / ")
        return "\n".join(sub)

    def _walk(container, dest, sep=" | "):
        for el in container.iterchildren():
            if _over():
                return
            tag = el.tag
            if tag == _W_NS + "p":
                seen[0] += 1
                t = _para(el)
                if t:
                    dest.append(t)
                if el.find(".//" + _W_NS + "txbxContent") is not None:
                    _textboxes(el, dest)
            elif tag == _W_NS + "tbl":
                for tr in _docx_unwrap(el, _W_NS + "tr"):
                    seen[0] += 1
                    if _over():
                        return
                    cells = []
                    for tc in _docx_unwrap(tr, _W_NS + "tc"):
                        tcpr = tc.find(_W_NS + "tcPr")
                        vm = tcpr.find(_W_NS + "vMerge") if tcpr is not None else None
                        if vm is not None and (vm.get(_W_NS + "val") or "continue") == "continue":
                            continue        # suite d'une fusion verticale : déjà lue au-dessus
                        txt = _cell_text(tc)
                        if txt:
                            cells.append(txt)
                    if cells:
                        dest.append(sep.join(cells))
            elif tag == _W_NS + "sdt":
                content = el.find(_W_NS + "sdtContent")
                if content is not None:
                    _walk(content, dest, sep)
            elif tag == _W_NS + "customXml":
                _walk(el, dest, sep)

    _walk(doc.element.body, out)
    _docx_other_parts(doc, _walk, notes, out)
    return out


def _docx_part_root(part):
    el = getattr(part, "element", None) if hasattr(type(part), "element") else None
    if el is not None:
        return el
    from docx.oxml.parser import parse_xml
    return parse_xml(part.blob)


def _docx_other_parts(doc, walk, notes, out) -> None:
    """Notes de bas de page et de fin, en-têtes et pieds de page, texte des schémas (SmartArt) :
    rangés ailleurs que dans le corps du document, lus ici, chacun dans sa rubrique."""
    try:
        rels = list(doc.part.rels.values())
    except Exception:
        return
    groups = {"bas": [], "fin": [], "entete": [], "schema": []}
    seen_parts, seen_lines = set(), set()
    for rel in rels:
        try:
            if getattr(rel, "is_external", False):
                continue
            kind = str(rel.reltype or "").rsplit("/", 1)[-1]
            if kind not in ("footnotes", "endnotes", "header", "footer", "diagramData"):
                continue
            part = rel.target_part
            if id(part) in seen_parts:
                continue
            seen_parts.add(id(part))
            root = _docx_part_root(part)
            if kind in ("footnotes", "endnotes"):
                tag = _W_NS + ("footnote" if kind == "footnotes" else "endnote")
                nk = "bas" if kind == "footnotes" else "fin"
                for fn in root.iterchildren(tag):
                    if (fn.get(_W_NS + "type") or "normal") != "normal":
                        continue          # séparateurs
                    lines = []
                    walk(fn, lines)
                    txt = " ".join(x.strip() for x in lines if x.strip())
                    if not txt:
                        continue
                    n = notes.num.get((nk, str(fn.get(_W_NS + "id"))))
                    lab = (f"(note {n})" if nk == "bas" else f"(note de fin {n})") if n else "(note)"
                    groups[nk].append(f"{lab} {txt}")
            elif kind in ("header", "footer"):
                lines = []
                walk(root, lines)
                for x in lines:
                    k = " ".join(x.split())
                    if k and k not in seen_lines:
                        seen_lines.add(k)
                        groups["entete"].append(x)
            else:
                txt = " ".join((t.text or "").strip() for t in root.iter(_DML_A_NS + "t") if (t.text or "").strip())
                if txt:
                    groups["schema"].append(txt)
        except Exception:
            continue
    for key, title in (("bas", "Notes de bas de page :"), ("fin", "Notes de fin :"),
                       ("entete", "En-têtes et pieds de page :"), ("schema", "Texte des schémas :")):
        if groups[key]:
            out.append("")
            out.append(title)
            out.extend(groups[key])


def _docx_open(data: bytes):
    """Ouvre un .docx (ou un .docm, document à macros : même contenu, autre type déclaré, que
    docx.Document refuse) avec python-docx."""
    import io
    import docx
    try:
        return docx.Document(io.BytesIO(data))
    except ValueError:
        from docx.opc.constants import CONTENT_TYPE as CT
        from docx.opc.part import PartFactory
        from docx.package import Package
        from docx.parts.document import DocumentPart
        PartFactory.part_type_for.setdefault(_CT_DOCM, DocumentPart)
        part = Package.open(io.BytesIO(data)).main_document_part
        if part.content_type not in (CT.WML_DOCUMENT_MAIN, _CT_DOCM):
            raise
        return part.document


def _extract_docx(data: bytes, info: dict = None) -> str:
    """Extrait le texte d'un .docx (Word), tableaux compris, DANS L'ORDRE du document et avec
    la numérotation automatique des listes (numéros des questions d'un cadre de réponse).
    Gracieux si python-docx absent. Tailles décompressées vérifiées AVANT ouverture (un .docx
    est une archive ZIP). info["complet"] = False si le document n'a pas été lu en entier."""
    if info is not None:
        info["complet"] = True
    if not _docx_sizes_ok(data):
        if info is not None:
            info["complet"] = False
        return ""
    try:
        d = _docx_open(data)
    except Exception:
        if info is not None:
            info["complet"] = False
        return ""
    try:
        return "\n".join(_docx_blocks(d, info=info)).strip()
    except Exception:
        # Repli historique : paragraphes puis tableaux (sans numérotation).
        try:
            parts = [p.text for p in d.paragraphs if p.text]
            for tbl in d.tables:
                for row in tbl.rows:
                    cells = [c.text for c in row.cells if c.text]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts).strip()
        except Exception:
            return ""


_ZIP_OKEXT = (".pdf", ".docx", ".docm", ".txt", ".md", ".csv", ".xlsx", ".xlsm", ".ods", ".odt")


def _extract_zip(content: bytes, max_chars: int, budget: "_Budget", depth: int, parts: list) -> None:
    """Parcourt une archive (et au plus 1 niveau d'archive imbriquée) en ajoutant à
    `parts` une pièce {nom, texte, type, complet} par fichier exploitable. Toute lecture est
    bornée par `budget`.

    Lecture PIÈCE PAR PIÈCE (rédaction, vague 1) : le règlement de la consultation et le
    cadre de réponse sont lus EN PREMIER (reconnus par leur nom, sans IA), puis les pièces au
    nom inconnu, l'acte d'engagement, le CCAP, les bordereaux, et le CCTP en dernier. Chaque
    pièce a son propre plafond (`max_chars`) et la lecture totale reste bornée à
    2 × `max_chars` : un gros CCTP ne fait plus disparaître le règlement ni le cadre."""
    import io, os, time, zipfile
    from app.services import dce_pieces as _dp
    if not _zip_entry_count_ok(content):
        raise ValueError("Archive ZIP trop complexe (trop de fichiers). "
                         "Importez directement les pièces utiles (RC, CCAP, CCTP, DPGF).")
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except Exception:
        raise ValueError("Archive ZIP corrompue ou illisible.")
    read_cap = 2 * max(1, int(max_chars))
    with zf:
        # Règlement et cadre d'abord, archives imbriquées en dernier ; à rang égal, PDF puis
        # Word/texte, puis par nom (ordre stable, reproductible d'un import à l'autre).
        def _rank(i):
            ln = i.filename.lower()
            if ln.endswith(".zip"):
                return (9, 2, ln)
            kind, _sure = _dp.classify_name(os.path.basename(i.filename))
            return (_dp.read_rank(kind), 0 if ln.endswith(".pdf") else 1, ln)
        for info in sorted(zf.infolist(), key=_rank):
            if _chars(parts) > read_cap or len(parts) >= ZIP_MAX_FILES or budget.bytes_left <= 0:
                break
            n = info.filename
            ln = n.lower()
            if info.is_dir() or os.path.basename(n).startswith("."):
                continue
            is_zip = ln.endswith(".zip")
            if not (ln.endswith(_ZIP_OKEXT) or is_zip):
                continue
            if is_zip and depth >= ZIP_MAX_DEPTH:
                budget.skipped_nested += 1
                continue
            if info.compress_type not in _ZIP_SAFE_COMPRESSION:
                continue    # bzip2/LZMA : sortie non bornable → ignoré (format non géré)
            try:
                data = _read_member_limited(zf, info, min(ZIP_MEMBER_MAX_BYTES, budget.bytes_left))
                if data is None:
                    budget.skipped_big += 1
                    continue
                budget.bytes_left -= len(data)
                if is_zip:
                    _extract_zip(data, max_chars, budget, depth + 1, parts)
                    continue
                # Plafond PAR PIÈCE (et jamais au-delà du plafond de lecture total).
                remaining = max(0, min(max_chars, read_cap - _chars(parts))) + 1
                complet = True
                if ln.endswith(".pdf"):
                    heavy_before = budget.pdf_heavy_pages
                    pinfo = {}
                    t = extract_text_from_pdf(data, max_chars=remaining, budget=budget, info=pinfo)
                    # Pages non lues (plafond, page trop lourde sautée, temps écoulé) : la pièce
                    # n'est pas « lue en entier ».
                    if time.monotonic() > budget.deadline or pinfo.get("complet") is False:
                        complet = False
                    # PDF scanné dans l'archive → OCR de repli. Pas pour un PDF dont on a sauté
                    # des pages trop lourdes (plans vectoriels) : l'OCR y serait long et inutile.
                    if (not (t and not t.startswith("Erreur") and len(t) > 30)
                            and budget.pdf_heavy_pages == heavy_before):
                        oinfo = {}
                        t = _ocr_pdf(data, budget=budget, info=oinfo)
                        complet = bool(oinfo.get("complet"))     # scan lu sur 15 pages au plus
                elif ln.endswith((".docx", ".docm")):
                    t = _extract_docx(data)
                elif ln.endswith((".xlsx", ".xlsm")):
                    t = extract_xlsx(data)
                elif ln.endswith((".ods", ".odt")):
                    t = extract_odf(data)
                else:
                    t = data[:remaining * 4].decode("utf-8", errors="replace")
                del data
                if t and not t.startswith("Erreur") and len(t) > 30:
                    base = os.path.basename(n)
                    body = t[:remaining - 1] if len(t) >= remaining else t
                    parts.append({"nom": base, "texte": body,
                                  "type": _dp.classify_piece(base, body),
                                  "complet": complet and len(t) < remaining})
            except ValueError:
                continue    # archive imbriquée illisible : on passe aux autres pièces
            except Exception:
                continue


def _chars(parts: list) -> int:
    return sum(len(p["texte"]) if isinstance(p, dict) else len(p) for p in parts)


def _single_piece(filename: str, text: str, max_chars: int, complet: bool = True) -> list:
    from app.services import dce_pieces as _dp
    import os
    base = os.path.basename((filename or "").replace("\\", "/"))
    body = text[:max_chars]
    return [{"nom": base, "texte": body, "type": _dp.classify_piece(base, body),
             "complet": bool(complet) and len(text) < max_chars}]


def extract_dce_pieces(filename: str, content: bytes, max_chars: int = 40000) -> tuple:
    """Lecture PIÈCE PAR PIÈCE d'un dossier : renvoie (pieces, archive) où pieces =
    [{nom, texte, type, complet}] (type reconnu sans IA, cf. app.services.dce_pieces) et
    archive = True pour un ZIP. Lève ValueError avec un message clair si illisible/vide/non
    supporté. Mémoire et CPU BORNÉS quel que soit le fichier (cf. plafonds ci-dessus)."""
    name = (filename or "").lower()

    if name.endswith(".pdf") or content[:5] == b"%PDF-":
        pinfo = {}
        txt = extract_text_from_pdf(content, max_chars=max_chars, info=pinfo)
        if not txt or txt.startswith("Erreur") or len(txt) < 30:
            # PDF sans couche texte → tentative OCR (document scanné).
            oinfo = {}
            ocr = _ocr_pdf(content, info=oinfo)
            if ocr and len(ocr) >= 30:
                return _single_piece(filename, ocr, max_chars, complet=oinfo.get("complet", False)), False
            raise ValueError("PDF illisible (document scanné/protégé ?) ou vide. "
                             "Si c'est un scan, l'OCR n'a rien pu extraire.")
        return _single_piece(filename, txt, max_chars, complet=pinfo.get("complet", True)), False

    if name.endswith((".docx", ".docm")):   # AVANT le ZIP : un .docx est aussi une archive PK
        if not _docx_sizes_ok(content):
            raise ValueError("Document Word trop volumineux ou illisible. Exportez-le en PDF puis réessayez.")
        t = _extract_docx(content)
        if t and len(t) >= 20:
            return _single_piece(filename, t, max_chars), False
        raise ValueError("Document Word illisible ou vide. Exportez-le en PDF si le souci persiste.")

    if name.endswith((".xlsx", ".xlsm", ".ods", ".odt")):   # AVANT le ZIP : archives PK elles aussi
        t = extract_xlsx(content) if name.endswith((".xlsx", ".xlsm")) else extract_odf(content)
        if t and len(t) >= 20:
            return _single_piece(filename, t, max_chars), False
        raise ValueError("Tableur ou document illisible ou vide. Exportez-le en PDF si le souci persiste.")

    if name.endswith((".txt", ".md", ".csv")):
        t = content[:max_chars * 4].decode("utf-8", errors="replace").strip()
        if t:
            return _single_piece(filename, t, max_chars), False
        raise ValueError("Fichier texte vide ou illisible.")

    if name.endswith(".zip") or content[:2] == b"PK":
        budget = _Budget()
        parts = []
        _extract_zip(content, max_chars, budget, 0, parts)
        if not parts:
            if budget.skipped_big:
                raise ValueError("Un fichier de l'archive est trop volumineux une fois décompressé "
                                 f"({ZIP_MEMBER_MAX_BYTES // _MB} Mo max par fichier). "
                                 "Importez directement les pièces utiles (RC, CCAP, CCTP, DPGF).")
            if budget.skipped_nested:
                raise ValueError("Archive trop imbriquée (un ZIP dans un ZIP dans un ZIP). "
                                 "Décompressez-la puis importez les pièces ou un seul ZIP.")
            raise ValueError("Aucun fichier exploitable dans l'archive (PDF scanné ou format non géré ?).")
        return parts, True

    raise ValueError("Format non supporté. Importez le dossier en PDF, Word (.docx), texte ou archive ZIP.")


def extract_dce_text(filename: str, content: bytes, max_chars: int = 40000, with_pieces: bool = False):
    """
    Extrait le texte d'un DCE réel : PDF unique OU archive ZIP de plusieurs pièces.
    Lève ValueError avec un message clair si illisible/vide/non supporté.
    Mémoire et CPU BORNÉS quel que soit le fichier (cf. plafonds ci-dessus).

    Archive : une pièce par fichier, précédée de « === nom === ». Le règlement de la
    consultation et le cadre de réponse viennent EN PREMIER et EN ENTIER ; les autres pièces
    se partagent équitablement le reste du plafond (app.services.dce_pieces.compose).
    with_pieces=True → (texte, meta) avec meta = [{nom, type, libelle, caracteres,
    lu_en_entier}] (ce qui a été lu, pièce par pièce, pour l'écran).
    """
    from app.services import dce_pieces as _dp
    pieces, archive = extract_dce_pieces(filename, content, max_chars=max_chars)
    text, meta = _dp.compose(pieces, max_chars, headers=archive)
    return (text, meta) if with_pieces else text


def _detect_groupement(details: dict, company: dict = None) -> dict:
    """Détecte si le marché est une OPPORTUNITÉ DE GROUPEMENT (répondre à plusieurs),
    de façon DÉTERMINISTE à partir des faits extraits — jamais inventé :
    (a) marché alloti, (b) qualification exigée que l'entreprise n'affiche pas,
    (c) CA minimum requis supérieur au CA de l'entreprise."""
    import re
    company = company or {}
    raisons = []
    allot = str(details.get("allotissement") or "").lower()
    # Détecte l'allotissement sous ses formulations courantes : énumération « Lot 1… Lot 2… »
    # MAIS aussi « 3 lots », « marché alloti en 4 lots », etc. (les sorties LLM réelles).
    n = re.search(r"\b(\d+)\s+lots?\b", allot)
    multi = (n and int(n.group(1)) >= 2) or allot.count("lot ") >= 2 or "lot 2" in allot
    if "lot unique" in allot and allot.count("lot ") < 2:
        multi = False
    if multi:
        raisons.append("Marché alloti : vous pouvez répondre à plusieurs, un partenaire par lot.")
    reqs = details.get("qualifications_requises") or []
    own_q = str(company.get("qualifications") or "").lower()
    if isinstance(reqs, list) and own_q:
        manquantes = [q for q in reqs if isinstance(q, str) and q.strip()
                      and not any(t for t in re.split(r"[\s/]+", q.lower()) if len(t) > 2 and t in own_q)]
        if manquantes:
            raisons.append("Qualification(s) exigée(s) que vous n'affichez pas ("
                           + ", ".join(manquantes[:3]) + ") — un partenaire peut les apporter.")
    cam = str(details.get("ca_minimum_requis") or "")
    m = re.search(r"(\d[\d\s  .]{3,})", cam)
    req_ca = int(re.sub(r"\D", "", m.group(1))) if m else 0
    ca = company.get("ca_n1") or 0
    if req_ca and ca and req_ca > ca:
        raisons.append("CA minimum exigé supérieur au vôtre — un groupement additionne les capacités.")
    # Exigences de groupement (copie brute des faits extraits) : c'est le routeur qui calcule
    # le score de groupement (app.services.groupement_scoring) — ici on ne fait que transmettre.
    fg = details.get("forme_groupement")
    if not isinstance(fg, dict):
        fg = {"imposee": str(fg or "non_specifie"), "extrait": ""} if fg else \
             {"imposee": "non_specifie", "extrait": ""}
    exigences = {
        "ca_minimum_requis": details.get("ca_minimum_requis") or "",
        "qualifications_requises": reqs if isinstance(reqs, list) else ([reqs] if reqs else []),
        "references_requises": details.get("references_requises") or "Non specifie dans le DCE",
        "effectif_minimum_requis": details.get("effectif_minimum_requis") or "Non specifie dans le DCE",
        "capacites_groupement": details.get("capacites_groupement") or "Non specifie dans le DCE",
        "forme_groupement": fg,
    }
    return {"recommande": bool(raisons), "raisons": raisons[:3], "exigences": exigences}


ANALYSE_MODEL = "claude-sonnet-4-6"   # modèle de l'ANCIEN routage ; le modèle réel = llm.route("analyse")
ANALYSE_TACHE = "analyse"              # Mistral Large 4, secours Claude Sonnet 5.5 (app.services.llm)
ANALYSE_MAX_TOKENS = 6000              # sortie JSON (~30 champs) ; payée à l'usage, pas au plafond
                                       # (×1,3 pour le découpage des nouveaux modèles, dans llm.py)


class AnalyseCoupee(RuntimeError):
    """La réponse de l'IA a été coupée (plafond de sortie) et n'est pas exploitable : traitée
    comme une panne (le quota est remboursé, l'utilisateur relance)."""


def _system_blocks(consigne: str = None):
    """Consigne de l'analyse (d'origine, d'extraction ou de la partie propre à l'entreprise), identique
    pour tous les clients : marquée pour le cache de prompt (relue à 10 % du prix par l'appel suivant,
    dans les 5 minutes)."""
    from app.services.llm import cached_text_block
    return [cached_text_block(SYSTEM_PROMPT if consigne is None else consigne)]


def _parse_analysis_json(response_text: str) -> dict:
    t = (response_text or "").strip()
    if t.startswith("```json"):
        t = t[7:]
    if t.startswith("```"):
        t = t[3:]
    if t.endswith("```"):
        t = t[:-3]
    t = t.strip()
    try:
        return json.loads(t)
    except json.JSONDecodeError:
        # Réponse « bavarde » (texte autour du JSON) : on repêche le premier objet JSON
        # complet plutôt que de jeter une analyse déjà payée par le client.
        start = t.find("{")
        if start < 0:
            raise
        result, _ = json.JSONDecoder().raw_decode(t[start:])
        return result


def finalize_analysis(result: dict, text_integral: str, company=None, criteria=None, pieces=None) -> dict:
    """Post-traitement DÉTERMINISTE d'une réponse d'analyse (sans IA) : champs obligatoires,
    citations vérifiées contre le texte INTÉGRAL, grille « Comment vous serez noté », score
    Go/No-Go par barème ouvert, opportunité de groupement."""
    if not isinstance(result, dict):
        result = {}
    if "match_score" not in result:
        result["match_score"] = 50
    if "go_decision" not in result:
        result["go_decision"] = "a_etudier"
    if "summary" not in result:
        result["summary"] = "Analyse terminee."
    if "details" not in result or not isinstance(result.get("details"), dict):
        result["details"] = {}

    # GARDE ANTI-HALLUCINATION des citations : chaque extrait « recopié » par l'IA est
    # vérifié contre le texte INTÉGRAL du DCE (pas la fenêtre envoyée au LLM) ; un extrait
    # non retrouvé devient '' (jamais affiché). Exigences de groupement normalisées.
    _normalize_exigences(result["details"], text_integral)

    # « COMMENT VOUS SEREZ NOTÉ » (rédaction, vague 1) : critères retrouvés dans le texte
    # avec leur phrase (citation_guard), poids TEL QU'ÉCRIT, contrôle du total et ce que le
    # dossier impose (cadre de réponse, limites, visite) — tout est vérifié SANS IA. Une
    # erreur ici ne casse jamais l'analyse déjà payée : la grille est simplement absente.
    try:
        from app.services.grille_notation import build_grille
        result["details"]["grille_notation"] = build_grille(result["details"], text_integral, pieces=pieces)
    except Exception as e:      # noqa: BLE001
        import logging
        logging.getLogger("adjugo").warning("grille de notation non construite : %s", type(e).__name__)

    # SCORE GO/NO-GO DÉTERMINISTE : l'IA a EXTRAIT les faits ci-dessus ; le score est
    # recalculé par un barème ouvert et reproductible (jamais laissé au LLM).
    from app.services.dce_scoring import score_dce
    det = score_dce(result.get("details") or {}, company, criteria)
    result["match_score"] = det["score"]
    result["go_decision"] = det["go_decision"]
    result["score_breakdown"] = det["breakdown"]
    result["score_deterministe"] = True
    # Auto-flag « opportunité de groupement » (déterministe, n'altère pas le score)
    result["groupement"] = _detect_groupement(result.get("details") or {}, company)
    return result


def analyze_dce_text(text, company=None, criteria=None, lang_name=None, pieces=None,
                     integral: bool = False, note_lecture: str = "", text_integral: str = None,
                     annonce_seule: bool = False, extraction: dict = None, apres_extraction=None,
                     portee_extraction=None):
    """Analyse un DCE depuis du texte brut (notice BOAMP, DCE déjà extrait...).
    lang_name : langue de rédaction des valeurs textuelles (adaptation par pays).
    pieces : ce qui a été lu, pièce par pièce (meta de la lecture), repris tel quel dans la
    grille « Comment vous serez noté ».
    integral=True (lecture intégrale, app.services.analyse_dossier) : `text` est remis TEL QUEL
    à l'IA (pas de sélection de passages) ; `note_lecture` la prévient des pièces non lues ;
    `text_integral` (défaut : `text`) est le texte complet du dossier contre lequel toutes les
    citations sont vérifiées (en lecture en deux temps, `text` est la synthèse des pièces).
    Analyse en deux temps (défaut, docs/ANALYSE_PARTAGEE.md) : (a) `extraire` — le dossier SEUL, aucune
    donnée d'entreprise —, puis (b) `completer` — post-traitement déterministe (citations vérifiées,
    grille, score, groupement) et partie propre à l'entreprise (appel court : l'extraction et le profil,
    jamais le dossier entier). Partage (app.services.analyse_partagee) : `extraction` = extraction (a)
    déjà faite pour CE dossier, exactement (aucun appel pour (a)) ; sinon `portee_extraction` (portée de
    coût de (a) : plateforme ou essai) et `apres_extraction(extraction, coût €, durée s)`, appelé dès
    que (a) est prête, avant (b) (une panne de (b) n'oblige pas à refaire (a)).
    `annonce_seule` : avis sans dossier complet (aucune compatibilité affichée).
    ANALYSE_PARTAGEE=0 : analyse d'origine en un seul appel, à l'identique (partage ignoré)."""
    if not text or len(text) < 50:
        return {
            "match_score": 0,
            "go_decision": "a_etudier",
            "summary": "Contenu insuffisant pour une analyse fiable.",
            "details": {}
        }
    ti = text_integral if text_integral is not None else text
    from app.services import analyse_partagee as _ap
    if not _ap.actif():
        return _analyse_un_appel(text, company, criteria, lang_name, pieces, integral, note_lecture, ti)
    # Toute PANNE de l'IA (extraction ou partie propre à l'entreprise : timeout, 5xx/429, disjoncteur,
    # réponse coupée…) est PROPAGÉE, jamais transformée en verdict : l'appelant rembourse le quota et
    # propose de relancer (voir _analyse_un_appel).
    servie = isinstance(extraction, dict)
    if not servie:
        import contextlib
        import time as _time
        from app.services import llm as _llm
        t0 = _time.monotonic()
        with (portee_extraction if portee_extraction is not None else contextlib.nullcontext()), \
                _llm.cost_meter("extraction du dossier", _llm._CURRENT_TENANT.get(), log=False) as cm:
            ext = extraire(text, lang_name, integral=integral, note_lecture=note_lecture)
        if ext.donnees is None:
            # Le LLM a bien RÉPONDU mais dans un format illisible : verdict NEUTRE honnête
            # (« à étudier »), jamais un faux NO-GO — comme l'analyse d'origine.
            return _resultat_illisible(ext.reponse)
        extraction = ext.donnees
        if apres_extraction is not None:
            try:
                apres_extraction(copy.deepcopy(extraction), cm.cost_eur(), _time.monotonic() - t0)
            except Exception as e:      # noqa: BLE001 — le rangement ne casse jamais l'analyse
                import logging
                logging.getLogger("adjugo").warning("extraction non rangée : %s", type(e).__name__)
    return completer(extraction, ti, company, criteria, pieces, lang_name, annonce_seule=annonce_seule,
                     noter_qualite=not servie)


class Extraction:
    """Résultat de (a) : `donnees` (réponse décodée de l'IA, None si illisible) et `reponse` (texte brut)."""
    __slots__ = ("donnees", "reponse")

    def __init__(self, donnees, reponse: str = ""):
        self.donnees = donnees if isinstance(donnees, dict) else None
        self.reponse = reponse or ""


def build_extraction_prompt(text, lang_name=None, integral: bool = False, note_lecture: str = "") -> str:
    """Prompt de l'EXTRACTION : celui de l'analyse d'origine SANS profil ni critères (même dossier, même
    sélection, même note de lecture, même consigne de langue)."""
    return build_user_prompt(text, None, None, lang_name, integral=integral, note_lecture=note_lecture)


def extraire(text, lang_name=None, integral: bool = False, note_lecture: str = "") -> "Extraction":
    """(a) EXTRACTION du dossier : la consigne SYSTEM_EXTRACTION et le dossier, AUCUNE donnée
    d'entreprise. Même tâche, même modèle, mêmes réglages, même plafond de sortie que l'analyse
    d'origine ; le dossier remis est exactement le même (rien n'est retiré de ce qui est lu). Propage
    une panne et une réponse coupée (AnalyseCoupee). L'espace facturé (tenant_scope) et, pour
    l'extraction partagée, la portée « mutualisé » sont posés par l'appelant."""
    user_prompt = build_extraction_prompt(text, lang_name, integral=integral, note_lecture=note_lecture)
    response_text, stop = _appel_analyse(SYSTEM_EXTRACTION, user_prompt, ANALYSE_MAX_TOKENS)
    try:
        data = _parse_analysis_json(response_text)
    except json.JSONDecodeError:
        if stop == "max_tokens":
            raise AnalyseCoupee("Réponse de l'IA coupée avant la fin.")
        return Extraction(None, response_text)
    return Extraction(data, response_text)


def completer(extraction: dict, text_integral: str, company=None, criteria=None, pieces=None, lang_name=None,
              annonce_seule: bool = False, noter_qualite: bool = True) -> dict:
    """(b) À partir de l'extraction (neuve ou resservie, jamais modifiée) : post-traitement DÉTERMINISTE
    (finalize_analysis : citations vérifiées contre le texte intégral de CE dossier, grille « Comment
    vous serez noté » d'après les pièces de CE dossier, score et verdict de CETTE entreprise,
    opportunité de groupement), puis partie propre à l'entreprise (personnaliser). `noter_qualite` :
    False pour une extraction resservie (ses citations ont déjà été comptées quand elle a été faite)."""
    tok = _QUALITE.set(bool(noter_qualite))
    try:
        result = finalize_analysis(copy.deepcopy(extraction), text_integral, company, criteria, pieces)
    finally:
        _QUALITE.reset(tok)
    return personnaliser(result, company, criteria, lang_name, annonce_seule=annonce_seule)


# ── (b) Partie propre à l'entreprise ──────────────────────────────────────────────────────────

ENTREPRISE_MAX_TOKENS = 2000           # un paragraphe et quelques points ; payée à l'usage, pas au plafond
ENTREPRISE_POINTS_MAX = 4
ENTREPRISE_CLAUSES_AJOUTEES_MAX = 3
_NIVEAUX = {"faible": 0, "moyen": 1, "eleve": 2}
_COMPATIBILITE = {"go": "bon potentiel", "a_etudier": "a etudier", "no_go": "peu adapte"}
_HORS_FAITS = ("grille_notation", "clauses_risque")   # grille : dérivée des critères ; clauses : numérotées à part
# Seuils de l'entreprise auxquels les clauses sont comparées (comme dans l'analyse d'origine).
SEUILS_CLAUSES = ("penalty_max", "garantie_max", "retenue_garantie_max", "avance_min", "delai_reponse_min")


def _seuils_fournis(criteria) -> bool:
    return any((criteria or {}).get(k) for k in SEUILS_CLAUSES)


def _bloc_activite(company, criteria) -> str:
    """Activité de l'entreprise (spécialités, zone, qualifications) : ce que le score déterministe
    compare déjà, remis aussi à la partie propre à l'entreprise pour qu'elle en parle juste. Aucun nom
    de personne (l'équipe n'est jamais remise ; ses noms seraient masqués de toute façon)."""
    company, criteria = company or {}, criteria or {}
    lignes = []
    spe = str(criteria.get("specialites") or "").strip()
    if spe:
        lignes.append(f"Specialites : {spe[:400]}")
    deps = criteria.get("departements")
    if isinstance(deps, (list, tuple)):
        deps = ", ".join(str(d).strip() for d in deps if str(d).strip())
    if str(deps or "").strip():
        lignes.append(f"Departements d'intervention : {str(deps)[:300]}")
    quals = company.get("qualifications") or []
    noms = []
    if isinstance(quals, (list, tuple)):
        for q in quals[:20]:
            n = q.get("name") if isinstance(q, dict) else q
            if str(n or "").strip():
                noms.append(str(n).strip()[:120])
    elif str(quals).strip():
        noms.append(str(quals).strip()[:400])
    if noms:
        lignes.append("Qualifications : " + ", ".join(noms))
    if not lignes:
        return ""
    return "=== ACTIVITE DE L'ENTREPRISE ===\n" + "\n".join(lignes) + "\n\n"


def build_entreprise_prompt(result: dict, company=None, criteria=None, lang_name=None,
                            annonce_seule: bool = False) -> str:
    """Prompt de la partie propre à l'entreprise : l'analyse objective (faits vérifiés, jamais le
    dossier entier), la compatibilité calculée, le profil et les critères (les champs que l'analyse
    d'origine remettait, plus l'activité que le score compare)."""
    det = result.get("details") if isinstance(result.get("details"), dict) else {}
    faits = {k: v for k, v in det.items() if k not in _HORS_FAITS}
    p = "Redige la partie propre a l'entreprise candidate et produis le JSON demande.\n\n"
    p += "=== ANALYSE OBJECTIVE DU MARCHE (extraite du dossier complet) ===\n"
    # JSON compact : toute l'analyse objective, sans les blancs de mise en forme (moins de tokens, rien de retiré).
    p += json.dumps({"summary": result.get("summary") or "", "details": faits}, ensure_ascii=False,
                    separators=(",", ":"))
    p += "\n\n=== CLAUSES A RISQUE (liste numerotee) ===\n"
    clauses = det.get("clauses_risque") if isinstance(det.get("clauses_risque"), list) else []
    lignes = []
    for i, c in enumerate(clauses):
        if isinstance(c, dict):
            lignes.append(f"{i}. {c.get('clause') or ''} (niveau {c.get('niveau') or 'non precise'}) : "
                          f"{c.get('pourquoi') or ''} — levier : {c.get('levier_negociation') or ''}")
    p += ("\n".join(lignes) if lignes else "(aucune)") + "\n\n"
    if annonce_seule:
        p += ("=== ANNONCE SEULE ===\nLe dossier complet du marche n'a pas encore ete importe : aucune "
              "compatibilite n'est affichee. N'annonce aucune decision ferme ; dis ce qu'il faudra verifier "
              "dans le dossier complet.\n\n")
    else:
        dec = _COMPATIBILITE.get(str(result.get("go_decision") or ""), "a etudier")
        p += "=== COMPATIBILITE CALCULEE (bareme deterministe, affichee a l'entreprise) ===\n"
        p += f"Compatibilite : {result.get('match_score')} % ({dec})\n"
        for c in result.get("score_breakdown") or []:
            if isinstance(c, dict):
                p += f"- {c.get('label')} : {c.get('points')}/{c.get('max_points')} — {c.get('detail')}\n"
        p += "\n"
    grp = result.get("groupement") if isinstance(result.get("groupement"), dict) else {}
    if grp.get("recommande"):
        p += ("=== OPPORTUNITE DE GROUPEMENT DETECTEE ===\n"
              + "\n".join("- " + str(r) for r in (grp.get("raisons") or [])) + "\n\n")
    p += _bloc_profil(company) + _bloc_activite(company, criteria) + _bloc_criteres(criteria)
    if not _seuils_fournis(criteria):
        p += "(Aucun seuil de clause n'est fourni par l'entreprise : \"clauses\" et \"clauses_ajoutees\" sont vides.)\n\n"
    if lang_name and lang_name.lower() != "français":
        p += (f"=== LANGUE DE SORTIE ===\nRédige TOUTES les valeurs textuelles (strategie, "
              f"points_attention_entreprise, pourquoi_entreprise, clauses_ajoutees) en {lang_name}. Conserve "
              f"les CLÉS JSON et les énumérations en l'état (niveau: faible/moyen/eleve).\n\n")
    return p.rstrip() + "\n"


class PartieEntrepriseIllisible(LLMUnavailable):
    """La partie propre à l'entreprise est restée illisible (deux réponses de suite) : traitée comme une
    panne — jamais une analyse en apparence complète mais privée de ses alertes propres à l'entreprise.
    L'appelant rembourse le quota ; l'utilisateur relance (l'extraction du dossier, déjà rangée dans le
    partage, n'est alors pas refaite)."""


MESSAGE_ENTREPRISE_ILLISIBLE = ("Les conseils propres à votre entreprise n'ont pas pu être rédigés par l'IA. "
                                "Votre analyse n'a pas été débitée : relancez-la.")


def personnaliser(result: dict, company=None, criteria=None, lang_name=None, annonce_seule: bool = False) -> dict:
    """(b) Partie propre à l'entreprise, après le calcul déterministe du score : un appel court (même
    tâche, même modèle, mêmes réglages que l'analyse d'origine), dans l'espace facturé du client, qui
    reçoit l'extraction, la compatibilité calculée, le profil et les critères — jamais le dossier entier.
    Noms des salariés masqués (pii_scope). Une réponse illisible est relancée une fois ; illisible
    encore : PartieEntrepriseIllisible (panne : l'appelant rembourse, l'utilisateur relance). Propage
    une panne de l'IA (l'appelant rembourse)."""
    import logging
    from app.services.llm import pii_scope
    from app.services.pii import from_team
    user = build_entreprise_prompt(result, company, criteria, lang_name, annonce_seule=annonce_seule)
    perso = None
    with pii_scope(from_team((company or {}).get("team") or [])):
        for _essai in range(2):
            txt, _stop = _appel_analyse(SYSTEM_ENTREPRISE, user, ENTREPRISE_MAX_TOKENS)
            try:
                perso = _parse_analysis_json(txt)
            except json.JSONDecodeError:
                perso = None
            if isinstance(perso, dict):
                break
    if isinstance(perso, dict):
        if not _seuils_fournis(criteria):
            # Comme l'analyse d'origine : une clause n'est « au-delà du seuil de l'entreprise » que si
            # l'entreprise a donné ce seuil (jamais un seuil supposé).
            perso = dict(perso, clauses=[], clauses_ajoutees=[])
        fusionner_entreprise(result, perso)
    else:
        logging.getLogger("adjugo").warning("partie propre à l'entreprise illisible deux fois : traitée comme une panne")
        raise PartieEntrepriseIllisible(MESSAGE_ENTREPRISE_ILLISIBLE)
    return result


def _liste(v) -> list:
    if isinstance(v, list):
        return v
    return [v] if isinstance(v, str) and v.strip() else []


def fusionner_entreprise(result: dict, perso: dict) -> dict:
    """Ajoute la partie propre à l'entreprise à l'analyse objective, sans rien en retirer : la
    stratégie devient le dernier paragraphe de la recommandation, les points propres à l'entreprise
    s'ajoutent à ceux du dossier (sans doublon, 4 au plus), une clause qui dépasse un seuil de
    l'entreprise passe au niveau « élevé » (jamais abaissé) avec la raison, une clause manquante qui
    dépasse un seuil s'ajoute (3 au plus)."""
    det = result.get("details")
    if not isinstance(det, dict):
        det = result["details"] = {}
    strat = perso.get("strategie")
    if isinstance(strat, str) and strat.strip():
        base = str(det.get("recommandation") or "").strip()
        det["recommandation"] = (base + "\n\n" + strat.strip()) if base else strat.strip()
    pts = perso.get("points_attention_entreprise")
    if isinstance(pts, list) and pts:
        cur = [x for x in _liste(det.get("points_attention"))]
        vus = {_norm_citation(x) for x in cur if isinstance(x, str)}
        ajoutes = 0
        for x in pts:
            if ajoutes >= ENTREPRISE_POINTS_MAX:
                break
            if isinstance(x, str) and x.strip() and _norm_citation(x) not in vus:
                cur.append(x.strip())
                vus.add(_norm_citation(x))
                ajoutes += 1
        det["points_attention"] = cur
    clauses = det.get("clauses_risque")
    if not isinstance(clauses, list):
        clauses = None
    for adj in perso.get("clauses") or []:
        if clauses is None or not isinstance(adj, dict):
            continue
        try:
            i = int(adj.get("index"))
        except (TypeError, ValueError):
            continue
        if not (0 <= i < len(clauses)) or not isinstance(clauses[i], dict):
            continue
        cl = clauses[i]
        niv = str(adj.get("niveau") or "").strip().lower()
        if _NIVEAUX.get(niv, -1) > _NIVEAUX.get(str(cl.get("niveau") or "").strip().lower(), -1):
            cl["niveau"] = niv
        why = str(adj.get("pourquoi_entreprise") or "").strip()
        if why and _norm_citation(why) not in _norm_citation(cl.get("pourquoi") or ""):
            cl["pourquoi"] = (str(cl.get("pourquoi") or "").strip() + " " + why).strip()
    ajouts = [c for c in (perso.get("clauses_ajoutees") or []) if isinstance(c, dict)
              and isinstance(c.get("clause"), str) and c["clause"].strip()][:ENTREPRISE_CLAUSES_AJOUTEES_MAX]
    if ajouts:
        if clauses is None:
            clauses = det["clauses_risque"] = []
        noms = {_norm_citation(c.get("clause") or "") for c in clauses if isinstance(c, dict)}
        for c in ajouts:
            if _norm_citation(c["clause"]) in noms:
                continue
            niv = str(c.get("niveau") or "").strip().lower()
            clauses.append({"clause": c["clause"].strip(), "niveau": niv if niv in _NIVEAUX else "eleve",
                            "pourquoi": str(c.get("pourquoi") or "").strip(),
                            "levier_negociation": str(c.get("levier_negociation") or "").strip()})
            noms.add(_norm_citation(c["clause"]))
    return result


def _appel_analyse(consigne: str, user_prompt: str, max_tokens: int):
    """Un appel de la tâche « analyse » (même route, mêmes réglages que l'analyse d'origine) :
    (texte de la réponse, raison de l'arrêt). Une bascule de secours est signalée à l'analyse partagée
    (une extraction produite par le secours n'est jamais gardée pour être resservie)."""
    from app.services.llm import messages_create
    response = messages_create(
        tache=ANALYSE_TACHE,
        model=ANALYSE_MODEL,
        max_tokens=max_tokens,
        temperature=0,  # extraction reproductible (jamais envoyée à Sonnet 5.5, qui la refuse)
        system=_system_blocks(consigne),
        messages=[{"role": "user", "content": user_prompt}]
    )
    from app.services import analyse_partagee as _ap
    _ap.noter_reponse(response)
    text = "".join(getattr(b, "text", "") or "" for b in (response.content or [])
                   if getattr(b, "type", "text") == "text").strip()
    return text, (getattr(response, "stop_reason", "") or "")


def _resultat_illisible(response_text: str) -> dict:
    return {
        "match_score": 50,
        "go_decision": "a_etudier",
        "summary": response_text[:500] if response_text else "Erreur de format dans la reponse IA.",
        "details": {}
    }


def _analyse_un_appel(text, company, criteria, lang_name, pieces, integral, note_lecture, text_integral):
    """Analyse d'ORIGINE en un seul appel (ANALYSE_PARTAGEE=0) : consigne SYSTEM_PROMPT, profil et
    critères dans le prompt — à l'identique de l'avant-partage."""
    user_prompt = build_user_prompt(text, company, criteria, lang_name, integral=integral,
                                    note_lecture=note_lecture)
    response_text = ""
    try:
        response_text, stop = _appel_analyse(SYSTEM_PROMPT, user_prompt, ANALYSE_MAX_TOKENS)
        try:
            result = _parse_analysis_json(response_text)
        except json.JSONDecodeError:
            if stop == "max_tokens":
                # Réponse COUPÉE : rien d'exploitable. Jamais un faux verdict facturé : on
                # propage (l'appelant rembourse le quota et propose de relancer).
                raise AnalyseCoupee("Réponse de l'IA coupée avant la fin.")
            raise
        return finalize_analysis(result, text_integral, company, criteria, pieces)

    except json.JSONDecodeError:
        # Le LLM a bien RÉPONDU mais dans un format illisible : verdict NEUTRE honnête
        # (« à étudier »), jamais un faux NO-GO.
        return _resultat_illisible(response_text)
    except Exception:
        # PANNE de l'IA (timeout, coupure réseau, 5xx/429 du fournisseur, plafond de tokens,
        # disjoncteur ouvert, réponse coupée…) : on NE fabrique JAMAIS un verdict — un faux
        # « NO-GO » serait mensonger ET débiterait le client pour rien. On PROPAGE : les
        # appelants (sourcing) remboursent le quota (refund_analysis) et renvoient un 503
        # « réessayez ». Avant, seul LLMUnavailable remontait ; une vraie erreur SDK (avant
        # l'ouverture du disjoncteur à 5 échecs) retournait un faux NO-GO facturé — c'est ce
        # trou qu'on ferme ici.
        raise
