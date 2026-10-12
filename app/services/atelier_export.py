"""
Export du dossier de présentation (rédaction, vague 2 — l'atelier) : Word (.docx) et PDF, et le
compteur (mots, caractères, pages estimées) comparé aux limites du dossier.

Doctrine (étude §7, règle 7) : rien d'interne ne part chez l'acheteur.
  • ni renvoi aux sources (« [S1] »), ni marque d'Adjugo, ni nom d'outil dans les propriétés ;
  • page de garde sobre (logo de l'entreprise s'il existe, titre du document tel que le règlement
    le nomme, marché, acheteur, candidat, mois), sommaire, titres, listes, tableaux simples ;
  • quand le cadre de réponse est imposé : ses numéros et ses intitulés, dans son ordre ;
  • propriétés du fichier : « Préparé avec assistance IA » SEULEMENT si l'IA a réellement servi
    (décision 9a), sinon rien.

Honnêteté du compteur : une page Word n'est pas calculable exactement (le rendu dépend du
logiciel et des polices installées). On compte les pages du PDF produit ICI, avec la police
demandée par le règlement quand elle est connue (à défaut, Arial 11), et on dit « environ ».
"""
import io
import re
from datetime import datetime
from typing import Optional

from app.services import memoire_export as me

IA_MENTION = "Préparé avec assistance IA"
POLICE_DEFAUT = "Arial"
TAILLE_DEFAUT = 11
MARGE_CM = 2.5

# Familles de police → police intégrée du moteur PDF aux mêmes largeurs (ou proches).
_SERIF = re.compile(r"times|garamond|georgia|cambria|book antiqua|palatino|serif", re.IGNORECASE)
_MONO = re.compile(r"courier|consolas|mono", re.IGNORECASE)
_MOIS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
         "novembre", "décembre")
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _s(v) -> str:
    return _CTRL.sub("", str(v if v is not None else "")).strip()


def police(limites: Optional[dict]) -> dict:
    """{famille, taille, connue, phrase, pdf_regular, pdf_bold, pdf_italic} : la police demandée
    par le règlement si elle a été repérée, sinon Arial 11 (dit comme tel)."""
    p = (limites or {}).get("police") if isinstance(limites, dict) else None
    fam = _s((p or {}).get("police")) or POLICE_DEFAUT
    fam = fam[0].upper() + fam[1:] if fam else POLICE_DEFAUT
    try:
        taille = float((p or {}).get("taille") or TAILLE_DEFAUT)
    except (TypeError, ValueError):
        taille = TAILLE_DEFAUT
    if not (6 <= taille <= 24):
        taille = TAILLE_DEFAUT
    if _SERIF.search(fam):
        base = ("Times-Roman", "Times-Bold", "Times-Italic")
    elif _MONO.search(fam):
        base = ("Courier", "Courier-Bold", "Courier-Oblique")
    else:
        base = ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")
    connue = bool(p and (p.get("police") or p.get("taille")))
    return {"famille": fam, "taille": int(taille) if float(taille).is_integer() else taille,
            "connue": connue, "phrase": _s((p or {}).get("phrase"))[:400],
            "pdf_regular": base[0], "pdf_bold": base[1], "pdf_italic": base[2]}


def mois_annee(dt: Optional[datetime] = None) -> str:
    try:
        from zoneinfo import ZoneInfo
        d = dt or datetime.now(ZoneInfo("Europe/Paris"))
    except Exception:
        d = dt or datetime.now()
    return f"{_MOIS[d.month - 1].capitalize()} {d.year}"


def titre_section(sec: dict) -> str:
    num = _s(sec.get("numero"))
    intit = _s(sec.get("intitule"))
    if not num:
        return intit
    return f"{num} {intit}" if "." in num else f"{num}. {intit}"


def _parents(sec: dict) -> list:
    return [p for p in (sec.get("parents") or []) if isinstance(p, dict) and (p.get("numero") or p.get("intitule"))]


def plan_rendu(sections: list) -> list:
    """Suite ordonnée de titres et de sections : [("titre", niveau, texte) | ("section", niveau, sec)].
    Les titres du cadre au-dessus d'une section (« 1. Moyens humains ») sont écrits une fois."""
    out, last = [], []
    for sec in sections:
        par = _parents(sec)
        keys = [(p.get("numero"), p.get("intitule")) for p in par]
        common = 0
        while common < min(len(keys), len(last)) and keys[common] == last[common]:
            common += 1
        for i in range(common, len(keys)):
            out.append(("titre", i + 1, titre_section(par[i])))
        last = keys
        out.append(("section", len(keys) + 1, sec))
    return out


def texte_export(md: str) -> str:
    """Texte remis à l'acheteur : sans renvois [S1], sans clôture de code."""
    return me.clean_markdown(md or "")


# ── Comptage ─────────────────────────────────────────────────────────────────────────

_WORD = re.compile(r"[0-9A-Za-zÀ-ÖØ-öø-ÿŒœ]+(?:['’\-][0-9A-Za-zÀ-ÖØ-öø-ÿŒœ]+)*")


def texte_brut(md: str) -> str:
    """Texte lisible (sans syntaxe Markdown) pour compter mots et caractères."""
    lines = []
    for raw in str(md or "").split("\n"):
        line = raw.strip()
        if re.match(r"^\|?\s*:?-{2,}", line):
            continue
        line = re.sub(r"^#{1,6}\s+", "", line)
        line = re.sub(r"^\s*[-*+]\s+", "", line)
        if line.startswith("|"):
            line = " ".join(c.strip() for c in line.strip("|").split("|"))
        line = re.sub(r"\*\*|__|`", "", line)
        line = re.sub(r"(?<!\w)\*(?=\S)|(?<=\S)\*(?!\w)", "", line)
        lines.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def compter_texte(md: str) -> dict:
    t = texte_brut(texte_export(md))
    flat = re.sub(r"\s+", " ", t).strip()
    return {"mots": len(_WORD.findall(t)), "caracteres": len(flat),
            "caracteres_sans_espaces": len(re.sub(r"\s", "", t))}


# ── Rendu PDF (reportlab) ────────────────────────────────────────────────────────────

def _styles(pol: dict):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    size = float(pol["taille"])
    ink = colors.HexColor("#111827")
    body = ParagraphStyle("body", fontName=pol["pdf_regular"], fontSize=size, leading=round(size * 1.2, 2),
                          textColor=ink, spaceAfter=size * 0.5)
    return {
        "body": body,
        "bullet": ParagraphStyle("bul", parent=body, leftIndent=size * 1.4, bulletIndent=size * 0.4,
                                 spaceAfter=size * 0.2),
        "cell": ParagraphStyle("cell", parent=body, fontSize=max(6, size - 1.5), leading=round(max(6, size - 1.5) * 1.2, 2),
                               spaceAfter=0),
        "cellh": ParagraphStyle("cellh", parent=body, fontName=pol["pdf_bold"], fontSize=max(6, size - 1.5),
                                leading=round(max(6, size - 1.5) * 1.2, 2), spaceAfter=0),
        1: ParagraphStyle("h1", parent=body, fontName=pol["pdf_bold"], fontSize=size + 4, leading=(size + 4) * 1.2,
                          spaceBefore=size * 1.2, spaceAfter=size * 0.5),
        2: ParagraphStyle("h2", parent=body, fontName=pol["pdf_bold"], fontSize=size + 2, leading=(size + 2) * 1.2,
                          spaceBefore=size, spaceAfter=size * 0.4),
        3: ParagraphStyle("h3", parent=body, fontName=pol["pdf_bold"], fontSize=size + 1, leading=(size + 1) * 1.2,
                          spaceBefore=size * 0.8, spaceAfter=size * 0.3),
        4: ParagraphStyle("h4", parent=body, fontName=pol["pdf_bold"], fontSize=size, leading=size * 1.2,
                          spaceBefore=size * 0.6, spaceAfter=size * 0.2),
        "cover_title": ParagraphStyle("ct", parent=body, fontName=pol["pdf_bold"], fontSize=size + 12,
                                      leading=(size + 12) * 1.25, spaceAfter=size * 1.5),
        "cover": ParagraphStyle("cv", parent=body, fontSize=size + 1, leading=(size + 1) * 1.4, spaceAfter=size * 0.6),
        "toc": ParagraphStyle("toc", parent=body, spaceAfter=size * 0.25),
    }


def _inline(s: str, pol: dict) -> str:
    """Gras, italique et passages « [À compléter …] » surlignés (comme dans le fichier Word : un
    export avec dérogation les montre, pour qu'on les complète avant de déposer)."""
    from app.services.md_pdf import _inline as md_inline, _mark_todo
    return _mark_todo(md_inline(_s(s))).replace("<b>", f"<font name='{pol['pdf_bold']}'>") \
        .replace("</b>", "</font>").replace("<i>", f"<font name='{pol['pdf_italic']}'>").replace("</i>", "</font>")


def _section_flowables(md: str, level: int, st: dict, pol: dict, width: float) -> list:
    """Flowables du texte d'une section ; les titres internes sont décalés sous le titre de section."""
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    out = []
    for kind, val in me.parse_blocks(texte_export(md)):
        if kind in ("h1", "h2", "h3"):
            lvl = min(4, level + {"h1": 1, "h2": 1, "h3": 2}[kind])
            out.append(Paragraph(_inline(val, pol), st[lvl]))
        elif kind == "bullet":
            out.append(Paragraph(_inline(val, pol), st["bullet"], bulletText="•"))
        elif kind in ("p", "num"):
            out.append(Paragraph(_inline(val, pol), st["body"]))
        elif kind == "table":
            headers, rows = val["headers"], val["rows"]
            n = max([len(headers)] + [len(r) for r in rows] + [1])

            def _row(vals, style):
                vals = list(vals) + [""] * (n - len(vals))
                return [Paragraph(_inline(v, pol), style) for v in vals[:n]]
            t = Table([_row(headers, st["cellh"])] + [_row(r, st["cell"]) for r in rows],
                      colWidths=[width / n] * n, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9CA3AF")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
            out.append(t)
            out.append(Spacer(1, float(pol["taille"]) * 0.5))
    return out


def _logo_flowable(logo: Optional[bytes], max_w: float, max_h: float):
    if not logo:
        return None
    try:
        from reportlab.lib.utils import ImageReader
        from reportlab.platypus import Image
        ir = ImageReader(io.BytesIO(logo))
        w, h = ir.getSize()
        if not w or not h:
            return None
        k = min(max_w / w, max_h / h)
        img = Image(io.BytesIO(logo), width=w * k, height=h * k)
        img.hAlign = "LEFT"
        return img
    except Exception:
        return None


def _cover(ctx: dict, st: dict, pol: dict, width: float) -> list:
    from reportlab.platypus import Paragraph, Spacer
    from reportlab.lib.units import cm
    out = [Spacer(1, 2 * cm)]
    img = _logo_flowable(ctx.get("logo"), 6 * cm, 3 * cm)
    if img is not None:
        out += [img, Spacer(1, 1.2 * cm)]
    out.append(Paragraph(_inline(ctx.get("titre") or me.DEFAULT_LABEL, pol), st["cover_title"]))
    for label, key in (("Marché", "marche"), ("Acheteur", "acheteur"), ("Candidat", "entreprise")):
        v = _s(ctx.get(key))
        if v:
            out.append(Paragraph(f"{label} : " + _inline(v, pol), st["cover"]))
    if _s(ctx.get("date")):
        out.append(Paragraph(_inline(ctx["date"], pol), st["cover"]))
    return out


def _toc(ctx: dict, st: dict, pol: dict) -> list:
    from reportlab.platypus import Paragraph
    out = [Paragraph(_inline("Sommaire", pol), st[1])]
    size = float(pol["taille"])
    for kind, lvl, val in plan_rendu(ctx.get("sections") or []):
        label = val if kind == "titre" else titre_section(val)
        from reportlab.lib.styles import ParagraphStyle
        style = ParagraphStyle(f"toc{lvl}", parent=st["toc"], leftIndent=size * 1.2 * (lvl - 1))
        out.append(Paragraph(_inline(label, pol), style))
    return out


def _body(ctx: dict, st: dict, pol: dict, width: float, per_section: Optional[dict] = None) -> list:
    from reportlab.platypus import Paragraph
    out = []
    for kind, lvl, val in plan_rendu(ctx.get("sections") or []):
        if kind == "titre":
            out.append(Paragraph(_inline(val, pol), st[min(lvl, 4)]))
            continue
        head = Paragraph(_inline(titre_section(val), pol), st[min(lvl, 4)])
        flows = [head] + _section_flowables(val.get("texte") or "", lvl, st, pol, width)
        if per_section is not None:
            per_section[val.get("id")] = flows
        out += flows
    return out


def _doc(buf, ctx: dict):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate
    ia = bool(ctx.get("ia"))
    return SimpleDocTemplate(buf, pagesize=A4, leftMargin=MARGE_CM * cm, rightMargin=MARGE_CM * cm,
                             topMargin=MARGE_CM * cm, bottomMargin=MARGE_CM * cm,
                             title=_s(ctx.get("titre") or me.DEFAULT_LABEL)[:200], author="", creator="",
                             producer="", subject=IA_MENTION if ia else "", keywords="assistance IA" if ia else "",
                             lang="fr-FR")


def _page_number(canvas, doc):
    canvas.saveState()
    pol = getattr(doc, "_adj_pol", None) or {"pdf_regular": "Helvetica", "taille": 10}
    canvas.setFont(pol["pdf_regular"], max(7, float(pol["taille"]) - 2))
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, doc.bottomMargin / 2, str(doc.page))
    canvas.restoreState()


def to_pdf(ctx: dict) -> tuple:
    """(octets du PDF, nombre de pages)."""
    from reportlab.platypus import PageBreak
    pol = ctx.get("police_rendu") or police(None)
    buf = io.BytesIO()
    doc = _doc(buf, ctx)
    doc._adj_pol = pol
    st = _styles(pol)
    flow = []
    if ctx.get("page_garde", True):
        flow += _cover(ctx, st, pol, doc.width) + [PageBreak()]
    if ctx.get("sommaire", True):
        flow += _toc(ctx, st, pol) + [PageBreak()]
    flow += _body(ctx, st, pol, doc.width)
    if not flow:
        from reportlab.platypus import Paragraph
        flow = [Paragraph(" ", st["body"])]
    # Page de garde sobre : pas de numéro de page sur elle.
    first = (lambda _c, _d: None) if ctx.get("page_garde", True) else _page_number
    doc.build(flow, onFirstPage=first, onLaterPages=_page_number)
    return buf.getvalue(), int(doc.page)


def compteur(ctx: dict) -> dict:
    """Mots, caractères, pages et lignes estimées, par section et au total. Les pages du total sont
    celles du PDF produit avec la police du règlement (page de garde et sommaire selon les options)."""
    pol = ctx.get("police_rendu") or police(None)
    st = _styles(pol)
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    width = A4[0] - 2 * MARGE_CM * cm
    height = A4[1] - 2 * MARGE_CM * cm
    per = {}
    _body(ctx, st, pol, width, per_section=per)
    sections = []
    for sec in ctx.get("sections") or []:
        c = compter_texte(sec.get("texte") or "")
        h, lines = 0.0, 0
        for f in per.get(sec.get("id"), [])[1:]:          # sans le titre de la section
            try:
                _w, fh = f.wrap(width, height)
                h += fh + float(getattr(getattr(f, "style", None), "spaceAfter", 0) or 0) \
                    + float(getattr(getattr(f, "style", None), "spaceBefore", 0) or 0)
                bl = getattr(f, "blPara", None)
                if bl is not None and getattr(bl, "lines", None) is not None:
                    lines += len(bl.lines)
                elif hasattr(f, "_cellvalues"):
                    lines += len(f._cellvalues)
            except Exception:
                continue
        c["pages_estimees"] = round(h / height, 1) if h else 0.0
        c["lignes_estimees"] = lines
        sections.append(dict(c, id=sec.get("id")))
    _pdf, pages = to_pdf(ctx)
    n_garde = 1 if ctx.get("page_garde", True) else 0
    n_toc = 0
    if ctx.get("sommaire", True):
        toc_ctx = dict(ctx, page_garde=False, sections=ctx.get("sections") or [])
        n_toc = _toc_pages(toc_ctx, pol)
    total = {"mots": sum(x["mots"] for x in sections), "caracteres": sum(x["caracteres"] for x in sections),
             "caracteres_sans_espaces": sum(x["caracteres_sans_espaces"] for x in sections),
             "lignes_estimees": sum(x["lignes_estimees"] for x in sections),
             "pages": pages, "pages_contenu": max(0, pages - n_garde - n_toc),
             "page_garde": bool(n_garde), "pages_sommaire": n_toc}
    return {"sections": sections, "total": total, "police": {k: pol[k] for k in ("famille", "taille", "connue", "phrase")}}


def _toc_pages(ctx: dict, pol: dict) -> int:
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=MARGE_CM * cm, rightMargin=MARGE_CM * cm,
                            topMargin=MARGE_CM * cm, bottomMargin=MARGE_CM * cm)
    st = _styles(pol)
    doc.build(_toc(ctx, st, pol))
    return max(1, int(doc.page))


# ── Rendu Word (python-docx) ─────────────────────────────────────────────────────────

def _docx_fonts(doc, pol: dict) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    fam, size = pol["famille"], float(pol["taille"])
    names = ["Normal", "List Bullet", "Table Grid"] + [f"Heading {i}" for i in range(1, 5)] + ["Title"]
    for name in names:
        try:
            stl = doc.styles[name]
        except KeyError:
            continue
        f = stl.font
        f.name = fam
        rpr = stl.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(att), fam)
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(att), None)
        if name.startswith("Heading"):
            lvl = int(name.split()[-1])
            f.size = Pt(size + {1: 4, 2: 2, 3: 1, 4: 0}[lvl])
            f.bold = True
            f.italic = False
            f.color.rgb = RGBColor(0x11, 0x18, 0x27)
        elif name == "Title":
            f.size = Pt(size + 12)
            f.color.rgb = RGBColor(0x11, 0x18, 0x27)
        else:
            f.size = Pt(size)


def _docx_page(doc) -> None:
    from docx.shared import Cm
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(MARGE_CM)


def _docx_properties(doc, ctx: dict) -> None:
    """Aucune trace de l'outil ; titre du document ; mention IA seulement si l'IA a servi."""
    me._neutral_properties(doc, _s(ctx.get("titre") or me.DEFAULT_LABEL))
    if ctx.get("ia"):
        cp = doc.core_properties
        cp.subject = IA_MENTION
        cp.comments = IA_MENTION
        cp.keywords = "assistance IA"


def _docx_text(doc, md: str, level: int) -> None:
    for kind, val in me.parse_blocks(texte_export(md)):
        if kind in ("h1", "h2", "h3"):
            lvl = min(4, level + {"h1": 1, "h2": 1, "h3": 2}[kind])
            me._runs(doc.add_heading("", level=lvl), val)
        elif kind == "bullet":
            me._runs(doc.add_paragraph(style="List Bullet"), val)
        elif kind in ("p", "num"):
            me._runs(doc.add_paragraph(), val)
        elif kind == "table":
            headers, rows = val["headers"], val["rows"]
            n = max([len(headers)] + [len(r) for r in rows] + [1])
            t = doc.add_table(rows=1, cols=n)
            t.style = "Table Grid"
            for j in range(n):
                me._runs(t.rows[0].cells[j].paragraphs[0], f"**{headers[j]}**" if j < len(headers) and headers[j] else "")
            for row in rows:
                cells = t.add_row().cells
                for j in range(n):
                    me._runs(cells[j].paragraphs[0], row[j] if j < len(row) else "")
            doc.add_paragraph()


def to_docx(ctx: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm, Pt
    pol = ctx.get("police_rendu") or police(None)
    doc = Document()
    _docx_page(doc)
    _docx_fonts(doc, pol)
    _docx_properties(doc, ctx)
    size = float(pol["taille"])
    if ctx.get("page_garde", True):
        for _ in range(3):
            doc.add_paragraph()
        if ctx.get("logo"):
            try:
                doc.add_picture(io.BytesIO(ctx["logo"]), width=Cm(5))
                doc.add_paragraph()
            except Exception:
                pass
        p = doc.add_paragraph()
        r = p.add_run(_s(ctx.get("titre") or me.DEFAULT_LABEL))
        r.bold = True
        r.font.size = Pt(size + 12)
        for label, key in (("Marché", "marche"), ("Acheteur", "acheteur"), ("Candidat", "entreprise")):
            v = _s(ctx.get(key))
            if v:
                q = doc.add_paragraph()
                q.add_run(f"{label} : ").bold = True
                q.add_run(v)
        if _s(ctx.get("date")):
            doc.add_paragraph(_s(ctx["date"]))
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    if ctx.get("sommaire", True):
        doc.add_heading("Sommaire", level=1)
        for kind, lvl, val in plan_rendu(ctx.get("sections") or []):
            label = val if kind == "titre" else titre_section(val)
            q = doc.add_paragraph()
            q.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
            q.add_run(_s(label))
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for kind, lvl, val in plan_rendu(ctx.get("sections") or []):
        if kind == "titre":
            doc.add_heading(_s(val), level=min(lvl, 4))
            continue
        doc.add_heading(_s(titre_section(val)), level=min(lvl, 4))
        _docx_text(doc, val.get("texte") or "", lvl)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
