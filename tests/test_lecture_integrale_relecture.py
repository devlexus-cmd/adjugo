"""
Lecture INTÉGRALE — relecture croisée (trois relecteurs) : chaque défaut confirmé a son test.

 • questions : TOUS les règlements remis (RC modificatif en premier), recherche sans accents ni
   pluriels, « pas trouvé dans les extraits consultés » (jamais « absent du dossier ») ;
 • Word lu en entier (zones de saisie, modifications suivies, champs, notes, en-têtes), .docm, .xlsm ;
 • PDF : pages qui partagent un flux mais pas leurs ressources lues, page scannée à bandeau tapé
   envoyée à l'OCR (bandeau gardé), numéros de page justes après des pages blanches, raison
   « texte trop long » au plafond de caractères ;
 • pièces nommées par leur lot (« LOT 02/CCTP.pdf »), fichiers Office protégés signalés comme tels ;
 • lecture en deux temps : petites pièces regroupées, réponse coupée relue en deux, panne relancée
   puis nommée, étiquettes honnêtes, reprise des parties déjà lues, échec massif remboursé ;
 • places de lecture attribuées dans l'ordre d'arrivée ;
 • débit hors de la boucle d'événements, unité Stripe annulable après un redémarrage ;
 • lecture en tâche de fond : une par organisation, échec visible et définitif, « terminé » avec le
   résultat, ré-analyse qui complète une lecture inachevée ;
 • écrans : parties non exploitées listées, lecture liée à son marché.
Aucun appel réseau ni IA réel.
"""
import asyncio
import io
import json
import re
import threading
import time
import types
import zipfile

import pytest

from app.services import llm
from tests.test_lecture_integrale import (FAIT, RC_TETE, _answer, _blank_pdf, _db, _pdf, _project,
                                          _project_row, _static, _texte, _used, _wait, _zip, fake_ia)  # noqa: F401

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


# ═══ 1. Questions : tous les règlements, recherche normalisée, « pas trouvé » honnête ═══════════

def _dossier_deux_rc():
    from app.services import dce_lecture as lec
    rc = RC_TETE + _texte("RC", 6000, 51) + "\nLa date limite de remise des offres est fixee au 3 novembre 2026 a 12h00."
    modif = ("REGLEMENT DE LA CONSULTATION MODIFIE N 1\n" + _texte("MODIF", 1500, 52)
             + "\nLa date limite de remise des offres est reportee au 17 novembre 2026 a 12h00.")
    ccap = "CAHIER DES CLAUSES ADMINISTRATIVES PARTICULIERES\n" + _texte("CCAP", 200000, 53)
    lignes = ccap.split("\n")
    lignes.insert(len(lignes) // 2, "Les pénalités de retard sont fixées à 1/1000e du montant HT par jour calendaire.")
    return lec.lire_dossier("DCE.zip", _zip({"RC.txt": rc.encode(), "RC_modificatif_1.txt": modif.encode(),
                                              "CCAP.txt": "\n".join(lignes).encode()}))


def test_questions_tous_les_reglements_recherche_normalisee_et_pas_trouve_honnete(fake_ia):
    from app.routers import sourcing as S
    from app.services import dce_texte
    lecture = _dossier_deux_rc()
    assert sorted(p["type"] for p in lecture["pieces"])[-2:] == ["reglement", "reglement"]
    proj = types.SimpleNamespace(dce_file_key=None, name="Voirie", client="Ville", id=1)

    def pas_trouve(kw):
        return json.dumps({"trouve": False, "reponse": "Cette information ne figure pas dans le dossier.",
                           "citations": []}), "end_turn"
    fake = fake_ia(pas_trouve)
    out = S._ask_integral(proj, lecture, "La date limite a-t-elle été reportée ?", "français")
    regs = fake.calls[-1]["messages"][0]["content"][0]["text"]
    # Les DEUX règlements sont remis, le modificatif en premier.
    assert "17 novembre 2026" in regs and "3 novembre 2026" in regs
    assert regs.index("17 novembre 2026") < regs.index("3 novembre 2026") and "[R1]" in regs and "[R2]" in regs
    # Non trouvé : ce qui a été consulté, jamais « absent du dossier ».
    assert out["trouve"] is False and out["answer"].startswith("Pas trouvé dans les extraits consultés")
    assert "ne figure pas" not in out["answer"] and "2 règlements de la consultation" in out["answer"]
    # Accès commun : tous les règlements, chacun sous son nom, le modificatif en premier.
    ordre = [p["nom"] for p in dce_texte.reglements(lecture["pieces"])]
    assert ordre == ["RC_modificatif_1.txt", "RC.txt"]

    # Recherche sans accents, sans pluriels, sans casse : le passage du CCAP est retrouvé.
    for q in ("Quelles sont les pénalités de retard ?", "Quelles sont les penalites de retard ?",
              "Quelle est la pénalité de retard ?", "PENALITES DE RETARD ?"):
        hits = dce_texte.passages(proj, q, k=8, payload=lecture)
        assert any("1/1000e" in h["texte"] for h in hits), q

    # Une citation du règlement modificatif est vérifiée et sourcée.
    def cite(kw):
        return json.dumps({"trouve": True, "reponse": "Le 17 novembre 2026.", "citations": [
            {"ref": "R1", "phrase": "La date limite de remise des offres est reportee au 17 novembre 2026 a 12h00."}]}), "end_turn"
    fake._answer = cite
    out = S._ask_integral(proj, lecture, "Date limite ?", "français")
    assert out["verifiee"] is True and out["sources"][0]["piece"] == "RC_modificatif_1.txt"


def test_dce_reglement_rend_tous_les_reglements(monkeypatch):
    from app.services import dce_texte
    lecture = _dossier_deux_rc()
    monkeypatch.setattr(dce_texte, "charger", lambda project: lecture)
    t = dce_texte.dce_reglement(object())
    assert t.startswith("=== RC_modificatif_1.txt ===") and "=== RC.txt ===" in t and "17 novembre" in t


# ═══ 2. Word lu en entier ════════════════════════════════════════════════════════════════════════

def _docx_complet() -> bytes:
    import docx
    from docx.oxml import parse_xml
    d = docx.Document()
    d.add_paragraph("CADRE DE REPONSE TECHNIQUE")
    sect = d.element.body[-1]

    def ins(xml):
        sect.addprevious(parse_xml(xml))
    ins(f'<w:p {W}><w:r><w:t xml:space="preserve">Question 1 : </w:t></w:r><w:sdt><w:sdtPr/><w:sdtContent>'
        '<w:r><w:t>FAIT-SDT-ENLIGNE delai de 8 semaines</w:t></w:r></w:sdtContent></w:sdt></w:p>')
    ins(f'<w:p {W}><w:r><w:t xml:space="preserve">Delai : </w:t></w:r><w:del w:id="2" w:author="a" '
        'w:date="2026-01-01T00:00:00Z"><w:r><w:delText>TEXTE-SUPPRIME 6 mois</w:delText></w:r></w:del>'
        '<w:ins w:id="3" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:t>FAIT-INS 9 mois</w:t></w:r></w:ins></w:p>')
    ins(f'<w:tbl {W}><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>Critere'
        '</w:t></w:r></w:p></w:tc><w:tc><w:sdt><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>FAIT-SDT-CELLULE valeur '
        'technique 60 %</w:t></w:r></w:p></w:sdtContent></w:sdt></w:tc></w:tr></w:tbl>')
    ins(f'<w:tbl {W}><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>Entete</w:t></w:r>'
        '</w:p></w:tc></w:tr><w:sdt><w:sdtPr/><w:sdtContent><w:tr><w:tc><w:p><w:r><w:t>FAIT-SDT-LIGNES prix 40 %'
        '</w:t></w:r></w:p></w:tc></w:tr></w:sdtContent></w:sdt></w:tbl>')
    ins(f'<w:p {W}><w:fldSimple w:instr="DOCPROPERTY Montant"><w:r><w:t>FAIT-CHAMP 120 000 euros</w:t></w:r>'
        '</w:fldSimple></w:p>')
    p = d.add_paragraph("Le titulaire est penalise en cas de retard")
    p._p.append(parse_xml(f'<w:r {W}><w:footnoteReference w:id="2"/></w:r>'))
    d.sections[0].header.paragraphs[0].text = "FAIT-ENTETE Marche n 2026-045 lot 3"
    b = io.BytesIO()
    d.save(b)
    zin = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    out = io.BytesIO()
    fn = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes ' + W + '>'
          '<w:footnote w:id="0" w:type="separator"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
          '<w:footnote w:id="2"><w:p><w:r><w:t>FAIT-NOTE Penalite de 300 euros par jour.</w:t></w:r></w:p>'
          '</w:footnote></w:footnotes>')
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in zin.infolist():
            data = zin.read(i.filename)
            if i.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            if i.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdFn" Type="http://schemas.openxml'
                                    b'formats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
                                    b'</Relationships>')
            z.writestr(i, data)
        z.writestr("word/footnotes.xml", fn)
    return out.getvalue()


def _macro(data: bytes, de: bytes, vers: bytes) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in zin.infolist():
            d = zin.read(i.filename)
            if i.filename == "[Content_Types].xml":
                d = d.replace(de, vers)
            z.writestr(i, d)
    return out.getvalue()


def test_word_lu_en_entier_et_formats_a_macros():
    from app.services import dce_lecture as lec
    from tests.test_lecture_integrale import _xlsx
    data = _docx_complet()
    lu = lec.lire_dossier("Cadre_de_reponse.docx", data)
    t = lu["pieces"][0]["texte"]
    for fait in ("FAIT-SDT-ENLIGNE", "FAIT-INS", "FAIT-SDT-CELLULE", "FAIT-SDT-LIGNES", "FAIT-CHAMP",
                 "FAIT-NOTE", "FAIT-ENTETE"):
        assert fait in t, fait
    assert "TEXTE-SUPPRIME" not in t                     # texte supprimé par une modification suivie
    assert "retard (note 1)" in t and "Notes de bas de page :" in t
    # .docm et .xlsm : mêmes formats, autres types déclarés — lus, dans une archive comme seuls.
    docm = _macro(data, b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                  b"application/vnd.ms-word.document.macroEnabled.main+xml")
    lu = lec.lire_dossier("DCE.zip", _zip({"RC.docm": docm, "DPGF.xlsm": _xlsx([["Poste", "Bordure", "45 EUR"]])}))
    noms = {p["nom"]: p for p in lu["pieces"]}
    assert noms["RC.docm"]["statut"] == "lu_en_entier" and "FAIT-INS" in noms["RC.docm"]["texte"]
    assert noms["DPGF.xlsm"]["statut"] == "lu_en_entier" and "Bordure | 45 EUR" in noms["DPGF.xlsm"]["texte"]
    assert lec.lire_dossier("RC.docm", docm)["pieces"][0]["statut"] == "lu_en_entier"


def test_fichiers_office_proteges_signales_comme_tels():
    from app.services import dce_lecture as lec
    chiffre = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + "EncryptedPackage".encode("utf-16-le") + b"\0" * 600
    lu = lec.lire_dossier("DCE.zip", _zip({"RC.txt": RC_TETE.encode(), "Memoire.docx": chiffre, "BPU.xlsx": chiffre}))
    noms = {p["nom"]: p for p in lu["pieces"]}
    for n in ("Memoire.docx", "BPU.xlsx"):
        assert noms[n]["raison"] == "protege" and "mot de passe" in noms[n]["detail"], n
    with pytest.raises(ValueError) as e:
        lec.lire_dossier("RC.docx", chiffre)
    assert "protégé par un mot de passe" in str(e.value) and "volumineux" not in str(e.value)


# ═══ 3. PDF : flux partagé, scan à bandeau, pages après des pages blanches ══════════════════════

def _pdf_flux_partage() -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, StreamObject
    w = PdfWriter()
    font = w._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                           NameObject("/Subtype"): NameObject("/Type1"),
                                           NameObject("/BaseFont"): NameObject("/Helvetica")}))
    content = StreamObject()
    content._data = b"q /Fm0 Do Q"
    cref = w._add_object(content)
    for i in range(3):
        page = w.add_blank_page(595, 842)
        form = StreamObject()
        form._data = (f"BT /F1 12 Tf 50 700 Td (Page {i + 1} : penalite de {100 * (i + 1)} euros par jour.) Tj ET").encode()
        form.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Form"),
                     NameObject("/BBox"): ArrayObject([FloatObject(0), FloatObject(0), FloatObject(595), FloatObject(842)]),
                     NameObject("/Resources"): DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject(
            {NameObject("/Fm0"): w._add_object(form)})})
        page[NameObject("/Contents")] = cref
    b = io.BytesIO()
    w.write(b)
    return b.getvalue()


def _png(seed: int) -> bytes:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (400, 560), "white")
    ImageDraw.Draw(im).text((10, 10 + seed), f"SCAN {seed}", fill="black")
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def _pdf_scan_a_bandeau() -> bytes:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(30, 800, "REGLEMENT DE LA CONSULTATION - Article 1 objet du marche de voirie communale")
    c.showPage()
    for i in (2, 3):
        c.drawImage(ImageReader(io.BytesIO(_png(i))), 0, 0, 595, 842)
        c.drawString(30, 820, f"Annexe {i} au reglement - document numerise")
        c.showPage()
    c.save()
    return buf.getvalue()


def test_pdf_flux_partage_scan_a_bandeau_et_pages_apres_pages_blanches(monkeypatch):
    from app.services import analysis as an, dce_lecture as lec
    # Flux de contenu partagé, ressources différentes : les trois pages sont lues.
    lu = lec.lire_dossier("CCAP.pdf", _pdf_flux_partage())
    t = lu["pieces"][0]["texte"]
    assert "100 euros" in t and "200 euros" in t and "300 euros" in t and lu["pieces"][0]["pages_lues"] == 3

    # Page scannée portant un bandeau tapé : à lire par OCR (en tâche de fond), bandeau gardé.
    data = _pdf_scan_a_bandeau()
    direct = lec.lire_dossier("RC.pdf", data, mode="direct")
    assert direct["a_poursuivre"] is True and direct["pieces"][0]["pages_non_lues"] == [2, 3]
    monkeypatch.setattr(an, "ocr_available", lambda: True)
    envoyees = []

    def fake_ocr(content, nums, budget, results, stop=None):
        envoyees.extend(nums)
        for n in nums:
            results[n] = f"Texte reconnu de la page scannee {n} : visite obligatoire le 21 octobre."
        return {"disponible": True, "non_lues": {}}
    monkeypatch.setattr(an, "ocr_pdf_pages", fake_ocr)
    fond = lec.lire_dossier("RC.pdf", data, mode="fond", reprise=direct)
    p = fond["pieces"][0]
    assert envoyees == [2, 3] and p["statut"] == "lu_en_entier"
    assert "page scannee 2" in p["texte"] and "Annexe 2 au reglement - document numerise" in p["texte"]
    assert lec.page_of(p, p["texte"].index("page scannee 3")) == 3

    # PDF qui commence par une page blanche : les numéros de page restent justes.
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    w.add_blank_page(595, 842)
    for pg in PdfReader(io.BytesIO(_pdf("Article 1. Objet du marche : travaux de voirie.\n" * 50
                                        + "Article 2. Le delai d'execution est de 4 mois.", 50))).pages:
        w.add_page(pg)
    b = io.BytesIO()
    w.write(b)
    p = lec.lire_dossier("CCTP.pdf", b.getvalue())["pieces"][0]
    assert lec.page_of(p, p["texte"].index("Article 1.")) == 2
    assert lec.page_of(p, p["texte"].index("Article 2.")) == 3


def test_plafond_de_caracteres_nomme_comme_tel(monkeypatch):
    from app.services import dce_lecture as lec
    monkeypatch.setattr(lec, "MAX_CARACTERES_PIECE", 3000)
    p = lec.lire_dossier("CCTP.pdf", _pdf(_texte("CCTP", 30000, 61)))["pieces"][0]
    assert p["statut"] == "lu_en_partie" and "texte trop long" in p["detail"] and "trop de pages" not in p["detail"]


# ═══ 4. Pièces nommées par leur lot ══════════════════════════════════════════════════════════════

def test_pieces_nommees_par_leur_lot(fake_ia):
    from app.routers import sourcing as S
    from app.services import dce_lecture as lec
    z = _zip({"DCE_2026-045/RC.txt": RC_TETE.encode(),
              "DCE_2026-045/LOT 01 - VRD/CCTP.txt": ("CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n"
                                                    "Le delai d'execution est de 4 mois.").encode(),
              "DCE_2026-045/LOT 02 - Eclairage/CCTP.txt": ("CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n"
                                                          "Le delai d'execution est de 9 mois.").encode()})
    lu = lec.lire_dossier("DCE.zip", z)
    noms = sorted(p["nom"] for p in lu["pieces"])
    assert noms == ["LOT 01 - VRD/CCTP.txt", "LOT 02 - Eclairage/CCTP.txt", "RC.txt"]   # dossier commun retiré
    assert all(p["type"] == "cctp" for p in lu["pieces"] if p["nom"].endswith("CCTP.txt"))
    texte, _ = lec.texte_integral(lu)
    assert "=== LOT 01 - VRD/CCTP.txt ===" in texte and "=== LOT 02 - Eclairage/CCTP.txt ===" in texte
    fake = fake_ia(lambda kw: (json.dumps({"trouve": True, "reponse": "4 mois", "citations": [
        {"ref": "P1", "phrase": "Le delai d'execution est de 4 mois."}]}), "end_turn"))
    proj = types.SimpleNamespace(dce_file_key=None, name="Voirie", client="Ville", id=1)
    out = S._ask_integral(proj, lu, "Quel est le délai d'exécution ?", "français")
    assert "LOT 02 - Eclairage/CCTP.txt" in llm._flatten_content(fake.calls[-1]["messages"][0]["content"])
    assert out["sources"][0]["piece"] == "LOT 01 - VRD/CCTP.txt"


# ═══ 5. Lecture en deux temps : regroupement, réponse coupée, panne, reprise ═════════════════════

def _lecture_deux_temps():
    from app.services import dce_lecture as lec
    files = {"RC.txt": (RC_TETE + _texte("RC", 3000, 71)).encode()}
    for i in range(1, 5):
        files[f"CCTP_lot{i}.txt"] = ("CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n"
                                     + _texte(f"CCTP{i}", 70000, 71 + i,
                                              fin=f"Le lot {i} exige la reference FIN-LOT-{i} pour le compacteur.")).encode()
    for i in range(30):
        files[f"PLAN_{i:02d}.txt"] = (f"Plan de detail {i}\n" + _texte(f"PLAN{i}", 600, 100 + i,
                                      fin=f"Cartouche du plan numero FIN-PLAN-{i:02d} valide.")).encode()
    return lec.lire_dossier("DCE.zip", _zip(files))


def _lecteur(coupe=(), panne_toujours=(), panne_une_fois=(), vus=None):
    """IA simulée : relève la dernière ligne « FIN- » de chaque pièce (avec son nom) ; coupe la
    réponse d'un appel long qui contient un marqueur ; tombe en panne sur d'autres."""
    deja = set()

    def ans(kw):
        system = llm._flatten_content(kw.get("system"))
        user = llm._flatten_content(kw["messages"][0]["content"])
        if "Tu lis UNE pièce" not in system:
            from tests.test_lecture_integrale import ANALYSE
            return json.dumps(ANALYSE), "end_turn"
        if vus is not None:
            vus.append(user)
        for m in panne_toujours:
            if m in user:
                raise RuntimeError("panne simulée")
        for m in panne_une_fois:
            if m in user and m not in deja:
                deja.add(m)
                raise RuntimeError("panne passagère")
        faits = []
        for nom, seg in re.findall(r"^Pièce : (.+?) \(.*?\n\n<<<\n(.*?)\n>>>", user, flags=re.S | re.M):
            for ln in seg.split("\n"):
                if "FIN-" in ln and len(ln) > 20:
                    faits.append({"piece": nom, "rubrique": "technique", "fait": "Exigence.", "citation": ln})
        txt = json.dumps({"faits": faits})
        if any(m in user for m in coupe) and len(user) > 30000:
            return txt[: len(txt) // 2], "max_tokens"
        return txt, "end_turn"
    return ans


def test_deux_temps_regroupe_relit_les_reponses_coupees_et_nomme_les_pannes(fake_ia, monkeypatch):
    from app.services import analyse_dossier as ad
    monkeypatch.setattr(ad, "BUDGET_TOKENS", 12_000)          # force la lecture en deux temps
    monkeypatch.setattr(ad, "PAUSE_RELANCE_S", 0.0)
    lecture = _lecture_deux_temps()
    vus, ranges = [], {}
    fake = fake_ia(_lecteur(coupe=("DEBUT-PIECE-CCTP1",), panne_toujours=("DEBUT-PIECE-CCTP3",),
                            panne_une_fois=("DEBUT-PIECE-CCTP4",), vus=vus))
    with llm.tenant_scope(515151):
        a = ad.analyser(lecture, {"name": "Test"}, {}, lang_name="français", on_parties=ranges.update)
    assert a["mode_lecture"] == "deux_temps"
    synth = llm._flatten_content(fake.calls[-1]["messages"][0]["content"])
    # 30 plans d'une page regroupés : bien moins d'appels que de pièces, chaque plan lu.
    assert len([u for u in vus if u.count("Pièce : ") >= 2]) >= 2
    for i in range(30):
        assert f"FIN-PLAN-{i:02d}" in synth, i
    # Réponse coupée : relue en deux moitiés, rien de perdu.
    assert "FIN-LOT-1" in synth
    # Panne passagère : relancée ; panne persistante : nommée, jamais « lue en entier ».
    assert "FIN-LOT-4" in synth and "FIN-LOT-3" not in synth
    head3 = re.search(r"=== CCTP_lot3\.txt[^\n]*", synth).group(0)
    assert "non exploitée par l'IA (IA indisponible)" in head3 and "lue en entier" not in head3
    assert "ATTENTION" in synth and "jamais qu'elle n'existe pas" in synth
    cov = a["couverture"]
    assert cov["complet"] is False and any(m["nom"] == "CCTP_lot3.txt" for m in cov["analyse"]["manques"])
    pc = next(p for p in a["lecture"]["pieces"] if p["nom"] == "CCTP_lot3.txt")
    assert pc["statut"] == "lu_en_entier" and "non exploitée par l'IA" in pc["ia"]
    grille = a["details"]["grille_notation"]["lecture"]["pieces"]
    assert "non exploitée par l'IA" in next(p for p in grille if p["nom"] == "CCTP_lot3.txt")["ia"]

    # Reprise : les parties déjà lues sont rangées ; une nouvelle analyse ne relit que la partie manquante.
    assert ranges and len(ranges) >= 30
    vus.clear()
    fake._answer = _lecteur(vus=vus)
    lecture2 = dict(lecture, analyse_parties=ranges)
    with llm.tenant_scope(515151):
        a2 = ad.analyser(lecture2, {"name": "Test"}, {}, lang_name="français")
    assert len(vus) == 1 and "DEBUT-PIECE-CCTP3" in vus[0]
    assert a2["couverture"]["complet"] is True and a2["couverture"]["analyse"]["parties_reprises"] >= 30
    assert "FIN-LOT-3" in llm._flatten_content(fake.calls[-1]["messages"][0]["content"])


def test_deux_temps_echec_massif_leve(fake_ia, monkeypatch):
    from app.services import analyse_dossier as ad
    monkeypatch.setattr(ad, "BUDGET_TOKENS", 12_000)
    monkeypatch.setattr(ad, "PAUSE_RELANCE_S", 0.0)
    fake_ia(_lecteur(panne_toujours=("DEBUT-PIECE-",)))
    with llm.tenant_scope(525252), pytest.raises(llm.LLMUnavailable):
        ad.analyser(_lecture_deux_temps(), {"name": "Test"}, {}, lang_name="français")


def test_places_de_lecture_dans_l_ordre_d_arrivee():
    from app.services.analyse_dossier import _FairSlots
    slots = _FairSlots(1)
    assert slots.acquire(1)
    ordre = []

    def attend():
        assert slots.acquire(5)
        ordre.append("attente")
        slots.release()
    t = threading.Thread(target=attend)
    t.start()
    time.sleep(0.1)
    slots.release()
    # Le fil qui vient de rendre sa place ne la reprend pas avant celui qui attendait.
    assert slots.acquire(5)
    ordre.append("reprise")
    slots.release()
    t.join(2)
    assert ordre == ["attente", "reprise"]


# ═══ 6. Débit hors de la boucle, unité Stripe annulable après redémarrage ════════════════════════

def test_debit_hors_boucle_et_unite_stripe_annulable_apres_redemarrage(client, auth, fake_ia, monkeypatch):
    import stripe
    import app.services.billing_usage as bu
    from app.core.quota import _limit, _period, refund_analysis
    from app.models import Job, PlanType, User
    from tests.test_lecture_integrale import _petit_dossier
    fake_ia()
    events, cancels, boucle = [], [], []

    def create(**kw):
        try:
            asyncio.get_running_loop()
            boucle.append(True)
        except RuntimeError:
            boucle.append(False)
        events.append(kw)
        return kw
    monkeypatch.setattr(bu.settings, "STRIPE_PRICE_OVERAGE", "price_overage_test")
    monkeypatch.setattr(stripe.billing.MeterEvent, "create", staticmethod(create))
    monkeypatch.setattr(stripe.billing.MeterEventAdjustment, "create", staticmethod(lambda **kw: cancels.append(kw)))
    db = _db()
    try:
        u = db.query(User).filter(User.email == auth["email"]).first()
        u.plan, u.stripe_customer_id, u.overage_enabled = PlanType.pro, "cus_boucle", True
        u.analyses_period, u.analyses_used_this_month = _period(), _limit(u)
        db.commit()
    finally:
        db.close()
    h = auth["headers"]
    pid = _project(client, h, "Hors forfait")
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE.zip", _petit_dossier(), "application/zip")})
    assert r.status_code == 200, r.text
    assert boucle == [False]                  # l'appel à Stripe ne fige pas la boucle d'événements
    ident = events[0]["identifier"]
    # Redémarrage : la mémoire du process est perdue, l'identifiant est en base.
    bu._SENT.clear()
    db = _db()
    try:
        u = db.query(User).filter(User.email == auth["email"]).first()
        row = db.query(Job).filter(Job.user_id == u.id, Job.kind == bu.UNITE_KIND).one()
        assert row.label == ident and row.status == "envoyee"
        refund_analysis(u, db)
        assert cancels == [{"event_name": "adjugo_analyse_overage", "type": "cancel", "cancel": {"identifier": ident}}]
        db.refresh(row)
        assert row.status == "annulee"
    finally:
        db.close()


# ═══ 7. Tâche de fond : une par organisation, échec définitif, reprise d'une lecture inachevée ══

def test_lecture_de_fond_bornee_echec_definitif_et_reprise(client, auth, fake_ia, monkeypatch):
    from app.models import Job, Project
    from app.services import analysis as an, lecture_taches
    fake = fake_ia()
    h = auth["headers"]
    monkeypatch.setattr(an, "ocr_available", lambda: False)
    z = _zip({"RC.txt": (RC_TETE + _texte("RC", 3000, 81)).encode(), "Annexe_scan.pdf": _blank_pdf(2)})

    # Une lecture déjà en cours pour l'organisation : refus AVANT tout débit.
    from tests.test_redaction_socle0 import _user_id
    uid = _user_id(auth["email"])
    db = _db()
    try:
        j = Job(user_id=uid, kind=lecture_taches.KIND, status="running", label="autre lecture")
        db.add(j)
        db.commit()
        jid = j.id
    finally:
        db.close()
    pid = _project(client, h, "Lecture bornée")
    before = _used(auth["email"])
    monkeypatch.setattr(an, "ocr_available", lambda: True)          # des pages à lire : tâche de fond
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE.zip", z, "application/zip")})
    assert r.status_code == 429 and "Rien n'a été débité" in r.json()["detail"]
    assert _used(auth["email"]) == before
    db = _db()
    try:
        db.get(Job, jid).status = "done"
        db.commit()
    finally:
        db.close()

    # Lecture de fond sans OCR installé : la pièce scannée est signalée ; le marché se termine.
    monkeypatch.setattr(an, "ocr_available", lambda: False)
    fake._answer = _answer
    from app.services import dce_lecture as lec
    real = lec.lire_dossier

    def direct_puis_fond(*a, **k):
        out = real(*a, **k)
        if k.get("mode") == "direct":
            out["a_poursuivre"] = True                               # force la tâche de fond
        return out
    monkeypatch.setattr(lec, "lire_dossier", direct_puis_fond)
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE.zip", z, "application/zip")})
    assert r.status_code == 202, r.text
    assert _wait(client, h, pid)["etat"] == "termine"
    a = client.get(f"/api/projects/{pid}", headers=h).json()["ai_analysis"]
    scan = next(p for p in a["lecture"]["pieces"] if p["nom"] == "Annexe_scan.pdf")
    assert scan["raison"] == "ocr_indisponible"

    # La lecture des images est maintenant disponible : « Ré-analyser » complète la lecture en
    # tâche de fond (et ne se contente pas de la lecture rangée).
    monkeypatch.setattr(lec, "lire_dossier", real)
    monkeypatch.setattr(an, "ocr_available", lambda: True)

    def fake_ocr(content, nums, budget, results, stop=None):
        for n in nums:
            results[n] = f"Plan signe page {n} : visite obligatoire du site le 21 octobre 2026."
        return {"disponible": True, "non_lues": {}}
    monkeypatch.setattr(an, "ocr_pdf_pages", fake_ocr)
    rr = client.post("/api/sourcing/reanalyze", headers=h, data={"project_id": str(pid)})
    assert rr.status_code == 202, rr.text
    assert _wait(client, h, pid)["etat"] == "termine"
    a = client.get(f"/api/projects/{pid}", headers=h).json()["ai_analysis"]
    scan = next(p for p in a["lecture"]["pieces"] if p["nom"] == "Annexe_scan.pdf")
    assert scan["statut"] == "lu_en_entier" and "21 octobre 2026" in fake.sent()

    # Échec de la tâche : le marché passe « erreur » pour de bon (plus de suivi à chaque ouverture),
    # l'analyse est remboursée.
    def panne(kw):
        raise RuntimeError("fournisseur en panne")
    fake._answer = panne
    monkeypatch.setattr(lec, "lire_dossier", direct_puis_fond)
    before = _used(auth["email"])
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE2.zip", z, "application/zip")})
    assert r.status_code == 202, r.text
    st = _wait(client, h, pid)
    assert st["etat"] == "erreur" and st["erreur"]
    time.sleep(0.3)
    assert _project_row(pid).workflow["lecture"]["etat"] == "erreur"
    assert _used(auth["email"]) == before
    llm._CB.update(fails=0, open_until=0.0)
    llm._TENANT_CB.clear()

    # « termine » n'est dit qu'avec le résultat du job enregistré.
    db = _db()
    try:
        j = Job(user_id=uid, kind=lecture_taches.KIND, status="running", label="x")
        db.add(j)
        db.commit()
        p = db.get(Project, pid)
        p.workflow = dict(p.workflow or {}, lecture={"etat": "termine", "job_id": j.id})
        db.commit()
        jid = j.id
    finally:
        db.close()
    assert client.get(f"/api/sourcing/lecture/{pid}", headers=h).json()["etat"] == "en_cours"
    db = _db()
    try:
        db.get(Job, jid).status = "done"
        db.commit()
    finally:
        db.close()


# ═══ 8. Écrans ═══════════════════════════════════════════════════════════════════════════════════

def test_ecrans_parties_non_exploitees_et_lecture_liee_au_marche():
    js, html = _static("app.js"), _static("app.html")
    assert "c.analyse && c.analyse.manques" in js                        # parties non exploitées listées
    assert "p.ia" in js and "!p.ia" in js                                # « Pièces lues » le dit aussi
    assert "aoBusy()" in html and 'v-if="ao.uploading"' not in html      # voile lié au marché suivi
    assert "lecturePid" in js and "aoLectureFin(pid)" in js
    assert html.count(".docm,.xlsx,.xlsm") == 2                          # formats à macros acceptés


def test_cache_compte_l_index_au_reel():
    from app.services import dce_texte
    payload = {"pieces": [{"texte": "x" * 1000}]}
    assert dce_texte._taille(payload) == 1000 * dce_texte.TEXTE_OCTETS_PAR_CAR
    assert dce_texte._taille(payload, index={}) == 1000 * (dce_texte.TEXTE_OCTETS_PAR_CAR + dce_texte.INDEX_OCTETS_PAR_CAR)
