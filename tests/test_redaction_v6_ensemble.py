"""
Rédaction, vague 6 — ENSEMBLE (F + H + J réunis ; contrat docs/REDACTION_V6.md, section « Ensemble »).

Le parcours de bout en bout d'un marché alloti, sans IA ni réseau :
  marché réparti en 3 lots, réponse aux lots 2 et 3 (J) → parties communes et propres → un export par lot, avec la
  page « Comment ce dossier a été préparé » cochée (H : la page du fichier de CE lot, ses parties seulement) puis
  décochée → marché passé à « Perdu » → « Demandez vos notes » (F : courrier pré-rempli, article exact, lots cités,
  Adjugo n'envoie rien) → courrier de l'acheteur lu sans IA LOT PAR LOT → notes validées (un relevé par lot,
  « lue dans le courrier » seulement si la ligne est dans la partie de SON lot) → « Où vous perdez des points » (le
  marché compte une fois) → nouveau dossier : « La dernière fois, l'acheteur a écrit » (avec le lot) ;
  plus : la lecture avec l'IA d'un courrier à plusieurs lots (contrôlée lot par lot), un groupement (le partenaire
  du lot 3 voit les notes communes, SES parties du lot 3 signalées, jamais le courrier), un marché simple inchangé,
  la démo (la page de chaque lot du marché de Trémorvan, rien d'écrit, rien qui coûte).
IA simulée, aucun réseau.
"""
import io
import re

import pytest

from app.services import atelier_preparation as prep
from app.services import carnet_notes as cn
from tests.test_redaction_socle0 import _company, _db, _user_id, fake_llm  # noqa: F401
from tests.test_redaction_v2_atelier import _creer, _etat
from tests.test_redaction_v6_carnet import ARTICLES_PERMIS, _marche, _questions_du_jour, _reponse_ia
from tests.test_redaction_v6_lots import RC, _creer_lots, _lot, _projet, _tout_ecrire

TITRE = prep.TITRE

# Le règlement du marché alloti (J), avec sa procédure écrite (F la lit sans IA).
RC_PROC = RC.replace("Article 5 - Contenu de l'offre",
                     "Article 2 - Procédure\nLa présente consultation est passée en procédure adaptée.\n"
                     "Article 5 - Contenu de l'offre")

# Le courrier de l'acheteur : une partie par lot (une ligne qui cite un seul lot ouvre sa partie).
COURRIER_LOTS = "\n".join([
    "Objet : restructuration de l'école maternelle Les Hirondelles, motifs du rejet de vos offres",
    "Madame, Monsieur,",
    "Vous trouverez ci-dessous vos notes pour les lots 2 et 3.",
    "Lot 2 - Couverture et zinguerie",
    "Valeur technique : 41 / 60 ; offre retenue : 52 / 60",
    "1.1 Méthodologie d'exécution : 20 sur 30",
    "Appréciation : phasage précis, mais la protection des abords de l'école reste générale.",
    "1.2 Moyens humains et matériels affectés au lot : 14/20",
    "1.3 Gestion environnementale du chantier : 7/10",
    "Prix : 36 / 40",
    "Lot 3 - Étanchéité des toitures-terrasses",
    "Valeur technique : 47 / 60",
    "1.1 Méthodologie d'exécution : 25 sur 30",
    "1.2 Moyens humains et matériels affectés au lot : 15/20",
    "1.3 Gestion environnementale du chantier : 7/10",
    "Prix : 33 / 40",
    "Contact : service des marchés, marches@tremorvan.exemple, 02 98 00 00 00.",
])


# ── Aides ───────────────────────────────────────────────────────────────────────────

def _docx_texte(contenu: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(contenu))
    lignes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lignes.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lignes)


def _docx_tables(contenu: bytes) -> list:
    import docx
    return docx.Document(io.BytesIO(contenu)).tables


def _pdf_pages(contenu: bytes) -> list:
    from pypdf import PdfReader
    return [re.sub(r"\s+", " ", p.extract_text() or "") for p in PdfReader(io.BytesIO(contenu)).pages]


def _export(client, h, pid, fmt, **corps):
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json=dict({"format": fmt}, **corps))
    assert r.status_code == 200, r.text
    return r


def _form_avec(etat, props):
    """Le formulaire tel que l'écran l'envoie après « Lire le courrier » (cnFusion) : les valeurs proposées,
    chacune avec sa ligne du courrier, marquées « lue dans le courrier »."""
    out = []
    for l in etat["criteres"]["lignes"]:
        x = props.get(l["cle"]) or {}
        ln = dict(cle=l["cle"], critere=l["critere"], parent=l["parent"], poids=l["poids"], note=l["note"],
                  note_max=l["note_max"], note_gagnant=l["note_gagnant"], commentaire=l["commentaire"],
                  famille=l["famille"], origine="saisie", lot=l.get("lot") or "")
        if x:
            ln["origine"] = x.get("origine") or "courrier"
            if x.get("note"):
                ln.update(note=x["note"], note_max=x.get("note_max") or ln["note_max"],
                          extrait_note=x.get("extrait_note") or "")
            if x.get("note_gagnant"):
                ln.update(note_gagnant=x["note_gagnant"], extrait_gagnant=x.get("extrait_gagnant") or "")
            if x.get("commentaire"):
                ln.update(commentaire=x["commentaire"], commentaire_du_courrier=True)
        out.append(ln)
    return out


def _cle(etat, lot, critere):
    return next(l["cle"] for l in etat["criteres"]["lignes"] if l.get("lot") == lot and l["critere"] == critere)


def _articles(texte: str) -> set:
    return set(re.findall(r"R\s?\d{4}-\d+", texte))


@pytest.fixture
def mails(monkeypatch):
    envois = []
    import app.services.email as em
    monkeypatch.setattr(em, "send_email", lambda *a, **k: envois.append(a) or True)
    return envois


@pytest.fixture
def ia_interdite(fake_llm):
    def ans(_kw):
        raise AssertionError("aucun appel d'IA n'est permis dans ce parcours")
    return fake_llm(ans)


# ═══ 1. Le parcours de bout en bout d'un marché alloti ═══════════════════════════════════════════

def test_parcours_complet_lots_page_preparation_carnet_et_derniere_fois(client, auth, mails, ia_interdite):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h, rc=RC_PROC)

    # (J) Deux lots : parties communes (une fois) et parties propres à chaque lot.
    st = _creer_lots(client, h, pid, lots=("2", "3"))
    assert st["lots"]["actif"] is True and [x["numero"] for x in st["lots"]["liste"]] == ["2", "3"]
    communes = [s for s in st["sections"] if s["commune"]]
    propres = [s for s in st["sections"] if not s["commune"]]
    assert len(communes) == 1 and {s["lot"] for s in propres} == {"2", "3"}
    st = _tout_ecrire(client, h, pid, st)
    l2, l3 = _lot(st, "2"), _lot(st, "3")
    assert len(l2["parties"]) == 3 and len(l3["parties"]) == 3          # 1 commune + 2 propres chacun

    # (H × J) La page « Comment ce dossier a été préparé » : celle du fichier d'UN lot.
    r = client.get(f"/api/atelier/{pid}/preparation", headers=h)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "lot_requis"
    ap2 = client.get(f"/api/atelier/{pid}/preparation?lot=2", headers=h).json()
    assert ap2["lot"] == "2" and ap2["page"]["lot"]["numero"] == "2"
    assert ap2["page"]["resume"][0] == ("Ce fichier répond au lot 2 — Couverture et zinguerie. Les parties communes "
                                        "à plusieurs lots y sont reprises telles quelles.")
    assert ap2["page"]["parties"] == 3 and ap2["page"]["parties_validees"] == 3
    numeros = {s["id"]: s for s in st["sections"]}
    assert {x["section_id"] for x in ap2["page"]["lignes"]} == {numeros[i]["id"] for i in
                                                                 (s["id"] for s in st["sections"]
                                                                  if s["lot"] in ("", "2"))}
    assert client.get(f"/api/atelier/{pid}/preparation?lot=1", headers=h).json()["detail"]["code"] == "lot_inconnu"

    # Export du lot 2 avec la page (Word), du lot 3 avec la page (PDF), puis du lot 2 sans la page.
    r = _export(client, h, pid, "docx", lot="2", page_preparation=True)
    assert "Lot%202" in r.headers["content-disposition"] or "Lot 2" in r.headers["content-disposition"]
    t = _docx_texte(r.content)
    assert TITRE in t and "Ce fichier répond au lot 2 — Couverture et zinguerie." in t
    assert "lot 3" not in t.split(TITRE, 1)[1].lower()
    tab = _docx_tables(r.content)[-1]
    assert len(tab.rows) == 1 + 3                                          # en-tête + les 3 parties du lot 2
    pages = _pdf_pages(_export(client, h, pid, "pdf", lot="3", page_preparation=True).content)
    assert TITRE in pages[-1] and "Ce fichier répond au lot 3 — Étanchéité des toitures-terrasses." in pages[-1]
    sans = _docx_texte(_export(client, h, pid, "docx", lot="2", page_preparation=False).content)
    assert TITRE not in sans
    from app.models import AuditLog, DossierPresentation
    db = _db()
    try:
        d = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).one()
        assert d.dernier_export["lots"]["2"]["page_preparation"] is False
        assert d.dernier_export["lots"]["3"]["page_preparation"] is True
        assert "page_preparation" not in (d.options_export or {})        # jamais enregistrée
        metas = [(e.meta or {}) for e in db.query(AuditLog).filter(AuditLog.action == "atelier.export",
                                                                    AuditLog.project_id == pid)
                 .order_by(AuditLog.id).all()]
        assert [(m.get("lot"), m.get("page_preparation")) for m in metas] == [("2", True), ("3", True), ("2", False)]
    finally:
        db.close()

    # (F) Le marché est perdu : « Demandez vos notes », courrier pré-rempli, article exact, lots cités.
    assert client.put(f"/api/projects/{pid}", headers=h, json={"status": "perdu"}).status_code == 200
    e = client.get(f"/api/carnet/{pid}", headers=h).json()
    assert e["statut"] == "perdu" and e["demande"]["proposee"] is True
    assert e["demande"]["procedure"]["code"] == cn.ADAPTEE
    assert e["demande"]["procedure"]["phrase"] == "La présente consultation est passée en procédure adaptée."
    assert [x["numero"] for x in e["lots"]] == ["2", "3"]
    lots_form = {l.get("lot") for l in e["criteres"]["lignes"]}
    assert lots_form == {"2", "3"} and e["criteres"]["message"].startswith(cn.MSG_LOTS)
    assert all(l["cle"].startswith(("L2-", "L3-")) for l in e["criteres"]["lignes"])
    lettre = client.get(f"/api/carnet/{pid}/courrier", headers=h).json()
    assert lettre["procedure"] == cn.ADAPTEE and lettre["article"] == "R2181-2"
    assert "article R2181-2 du code de la commande publique" in lettre["corps"]
    assert _articles(lettre["corps"] + lettre["message_delai"]) <= ARTICLES_PERMIS
    assert ("au titre des lots auxquels nous avons répondu : lot 2 — Couverture et zinguerie ; lot 3 — Étanchéité "
            "des toitures-terrasses.") in lettre["corps"]
    assert "pour chacun de ces lots et pour chaque critère et sous-critère" in lettre["corps"]
    assert "15 jours" in lettre["message_delai"] and "pas toujours obtenues" in lettre["message_notes"]
    assert lettre["message_lots"] == cn.MSG_COURRIER_LOTS and lettre["message_envoi"] == cn.MSG_PAS_ENVOYE
    assert mails == []                                                      # Adjugo n'envoie rien

    # Le courrier de l'acheteur, lu SANS IA lot par lot : aucune note d'un lot rangée sous l'autre.
    r = client.post(f"/api/carnet/{pid}/lire", headers=h, json={"courrier": COURRIER_LOTS})
    assert r.status_code == 200, r.text
    props = r.json()["propositions"]
    m2, m3 = _cle(e, "2", "Méthodologie d'exécution"), _cle(e, "3", "Méthodologie d'exécution")
    v2, v3 = _cle(e, "2", "Valeur technique"), _cle(e, "3", "Valeur technique")
    assert props[m2]["note"] == "20" and props[m2]["note_max"] == "30"
    assert props[m3]["note"] == "25" and props[m3]["extrait_note"] == "1.1 Méthodologie d'exécution : 25 sur 30"
    assert props[v2]["note"] == "41" and props[v2]["note_gagnant"] == "52"
    assert props[v3]["note"] == "47" and "note_gagnant" not in props[v3]
    assert props[m2]["commentaire"].startswith("phasage précis, mais la protection") and "commentaire" not in props[m3]
    assert r.json()["trouvees"] == 10
    assert cn.reponse(_projet_row(pid)) is None                             # rien d'enregistré avant « Valider »

    # « Valider les notes » : un relevé par lot ; une ligne d'un autre lot n'est jamais « lue dans le courrier ».
    lignes = _form_avec(e, props)
    triche = next(l for l in lignes if l["cle"] == _cle(e, "2", "Moyens humains et matériels affectés au lot"))
    triche.update(note="15", note_max="20", origine="courrier",
                  extrait_note="1.2 Moyens humains et matériels affectés au lot : 15/20")    # la ligne du lot 3
    r = client.put(f"/api/carnet/{pid}/notes", headers=h,
                   json={"lignes": lignes, "rev": e["rev"], "courrier": COURRIER_LOTS})
    assert r.status_code == 200, r.text
    e2 = r.json()
    par_cle = {x["cle"]: x for x in e2["notes"]["lignes"]}
    assert par_cle[m2]["origine"] == "courrier" and par_cle[m2]["lot"] == "2"
    assert par_cle[m2]["lot_libelle"] == "Lot 2 — Couverture et zinguerie"
    assert par_cle[triche["cle"]]["origine"] == "saisie"                  # extrait hors de la partie du lot 2
    assert [x["lot"] for x in e2["notes"]["lignes"]] == sorted(x["lot"] for x in e2["notes"]["lignes"])
    assert e2["notes"]["total"] is None                                    # jamais une somme de deux lots
    # Note d'ensemble lot par lot : valeur technique + prix (41 + 36 ; 47 + 33).
    assert [(t["lot"], t["note"], t["note_max"], t["pct"]) for t in e2["notes"]["totaux"]] == [
        ("2", "77", "100", 77), ("3", "80", "100", 80)]
    assert "(lot " in e2["notes"]["constat"]["message"]

    # « Où vous perdez des points » : le marché alloti compte UNE fois ; moyenne à partir de 3 marchés.
    autres = []
    for nom, statut, notes in (("Toiture du gymnase", "perdu", ("30", "28")), ("Toiture de la mairie", "gagne", ("36", "32"))):
        q = _marche(client, h, name=nom, acheteur="Ville de Brest", statut=statut)
        eq = client.get(f"/api/carnet/{q}", headers=h).json()
        assert eq["lots"] == [] and all(not l.get("lot") for l in eq["criteres"]["lignes"])
        ls = [{"cle": "c1.1", "critere": "Méthodologie d'exécution", "note": notes[0], "note_max": "40"},
              {"cle": "c1.2", "critere": "Gestion environnementale du chantier", "note": "10", "note_max": "20"},
              {"cle": "c2", "critere": "Prix des prestations", "note": notes[1], "note_max": "40"}]
        assert client.put(f"/api/carnet/{q}/notes", headers=h, json={"lignes": ls, "rev": eq["rev"]}).status_code == 200
        autres.append(q)
    tb = client.get("/api/carnet/tableau", headers=h).json()
    assert tb["marches"] == 3
    tech = next(f for f in tb["familles"] if f["code"] == "technique")
    assert tech["marches"] == 3 and tech["moyenne"] is not None
    nous = (20 + 14 + 25 + 15) / (30 + 20 + 30 + 20)                       # les lots 2 et 3 réunis : 74 %
    attendu = round(100 * (nous + 30 / 40 + 36 / 40) / 3)
    assert abs(tech["moyenne"] - attendu) <= 1

    # Un nouveau dossier, même acheteur : « La dernière fois, l'acheteur a écrit », avec le lot.
    nouveau = _projet(client, h, rc=RC, name="Extension de l'école élémentaire")
    st_n = client.post(f"/api/atelier/{nouveau}", headers=h, json={}).json()
    assert st_n["lots"]["actif"] is False
    df = st_n["derniere_fois"]
    meth = next(s for s in st_n["sections"] if s["intitule"].startswith("Méthodologie"))
    x = df[str(meth["id"])]
    assert x["commentaire"].startswith("phasage précis, mais la protection") and x["meme_acheteur"] is True
    assert x["lot"] == "Lot 2 — Couverture et zinguerie" and x["marche"].startswith("Restructuration")
    assert mails == []


def _projet_row(pid):
    from app.models import Project
    db = _db()
    try:
        p = db.get(Project, pid)
        db.expunge(p)
        return p
    finally:
        db.close()


# ═══ 2. La lecture avec l'IA d'un courrier à plusieurs lots ══════════════════════════════════════

def test_lecture_ia_controlee_lot_par_lot(client, auth, fake_llm):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h, rc=RC_PROC)
    _tout_ecrire(client, h, pid, _creer_lots(client, h, pid))
    assert client.put(f"/api/projects/{pid}", headers=h, json={"status": "perdu"}).status_code == 200
    e = client.get(f"/api/carnet/{pid}", headers=h).json()
    m2, mo2 = _cle(e, "2", "Méthodologie d'exécution"), _cle(e, "2", "Moyens humains et matériels affectés au lot")
    m3 = _cle(e, "3", "Méthodologie d'exécution")
    fake = fake_llm(_reponse_ia([
        {"id": m2, "note": "20", "note_max": "30", "extrait": "1.1 Méthodologie d'exécution : 20 sur 30"},
        {"id": mo2, "note": "15", "note_max": "20",                       # la ligne du lot 3 : écartée
         "extrait": "1.2 Moyens humains et matériels affectés au lot : 15/20"},
        {"id": m3, "note": "20", "note_max": "30",                        # la ligne du lot 2 : écartée
         "extrait": "1.1 Méthodologie d'exécution : 20 sur 30"}]))
    q0 = _questions_du_jour(auth["email"])
    r = client.post(f"/api/carnet/{pid}/lire-ia", headers=h, json={"courrier": COURRIER_LOTS})
    assert r.status_code == 200, r.text
    assert set(r.json()["propositions"]) == {m2} and r.json()["refusees"] == 2
    assert _questions_du_jour(auth["email"]) == q0 + 1
    envoye = fake.sent()
    assert "Ne prends la note d'un critère que dans la partie du courrier qui parle de CE lot." in envoye
    assert f"[{m2}] (Lot 2 — Couverture et zinguerie) — Méthodologie d'exécution (30 %)" in envoye
    assert "marches@tremorvan" not in envoye and "02 98 00 00 00" not in envoye
    # Un courrier qui ne dit pas à quel lot se rapporte chaque note : aucun appel, rien de décompté.
    n = len(fake.calls)
    sans_lot = "\n".join(x for x in COURRIER_LOTS.split("\n") if not x.startswith("Lot "))
    r = client.post(f"/api/carnet/{pid}/lire-ia", headers=h, json={"courrier": sans_lot})
    assert r.status_code == 200 and r.json()["propositions"] == {} and r.json()["message"] == cn.MSG_LOT_INCONNU
    assert len(fake.calls) == n and _questions_du_jour(auth["email"]) == q0 + 1
    r = client.post(f"/api/carnet/{pid}/lire", headers=h, json={"courrier": sans_lot})
    assert r.json()["propositions"] == {} and r.json()["message"] == cn.MSG_LOT_INCONNU
    # Un courrier propre à un seul lot (il ne cite que le lot 3) : tout le courrier en parle.
    seul3 = "\n".join(["Objet : lot n° 3, motifs du rejet de votre offre", "Madame, Monsieur,",
                       "Valeur technique : 47 / 60", "1.1 Méthodologie d'exécution : 25 sur 30"])
    p = client.post(f"/api/carnet/{pid}/lire", headers=h, json={"courrier": seul3}).json()["propositions"]
    assert set(p) == {_cle(e, "3", "Valeur technique"), m3}
    # Un critère ajouté doit dire son lot.
    r = client.put(f"/api/carnet/{pid}/notes", headers=h, json={"rev": e["rev"], "lignes": [
        {"cle": "u1", "critere": "Visite de site", "note": "3", "note_max": "5"}]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "lot_requis"
    r = client.put(f"/api/carnet/{pid}/notes", headers=h, json={"rev": e["rev"], "lignes": [
        {"cle": "u1", "critere": "Visite de site", "note": "3", "note_max": "5", "lot": "3"}]})
    assert r.status_code == 200, r.text
    u1 = next(x for x in r.json()["notes"]["lignes"] if x["cle"] == "u1")
    assert u1["lot"] == "3" and u1["lot_libelle"] == "Lot 3 — Étanchéité des toitures-terrasses"


# ═══ 3. Groupement : le partenaire du lot 3 ══════════════════════════════════════════════════════

def test_groupement_partenaire_du_lot_3_voit_les_notes_communes_ses_parties_signalees(client, auth, ia_interdite):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h, rc=RC_PROC)
    r = client.post(f"/api/projects/{pid}/invites", headers=h, json={"recipient": "", "company_name": "Étanche Ouest"})
    tok = r.json()["token"]
    v = client.get(f"/api/invite/{tok}/contribution").json()["version"]
    assert client.put(f"/api/invite/{tok}/contribution", json={
        "version": v, "company_name": "Étanche Ouest", "lot": "Lot 3 : étanchéité des toitures-terrasses",
        "siret": "55566677700012"}).status_code == 200
    _creer_lots(client, h, pid)
    assert client.post(f"/api/atelier/{pid}/repartition", headers=h, json={}).status_code == 200
    assert client.put(f"/api/projects/{pid}", headers=h, json={"status": "perdu"}).status_code == 200
    e = client.get(f"/api/carnet/{pid}", headers=h).json()
    props = client.post(f"/api/carnet/{pid}/lire", headers=h, json={"courrier": COURRIER_LOTS}).json()["propositions"]
    # Le chef de file voit quelle partie du lot 3 a été rédigée par le partenaire (jamais celle du lot 2).
    form = {l["cle"]: l for l in e["criteres"]["lignes"]}
    m2, m3 = _cle(e, "2", "Méthodologie d'exécution"), _cle(e, "3", "Méthodologie d'exécution")
    assert (form[m3]["auteur"] or {}).get("nom") == "Étanche Ouest"
    assert (form[m2]["auteur"] or {}).get("type") != "invitation"
    r = client.put(f"/api/carnet/{pid}/notes", headers=h,
                   json={"lignes": _form_avec(e, props), "rev": e["rev"], "courrier": COURRIER_LOTS})
    assert r.status_code == 200, r.text
    # L'espace du partenaire : toutes les notes communes, SES parties (lot 3) signalées, rien du courrier.
    cle = client.post(f"/api/invite/{tok}/redaction/espace", json={}).json()["cle"]
    vue = client.get(f"/api/invite/{tok}/redaction", headers={"X-Adjugo-Espace": cle}).json()
    c = vue["carnet"]
    assert c and {x["lot"] for x in c["lignes"]} == {"2", "3"}
    vous = {x["cle"] for x in c["vos_parties"]}
    assert m3 in vous and m2 not in vous and all(x["lot"] == "3" for x in c["vos_parties"])
    assert [t["lot"] for t in c["totaux"]] == ["2", "3"] and c["total"] is None
    brut = str(c)
    assert "extrait" not in brut and "marches@tremorvan" not in brut and "Contact :" not in brut
    assert all("origine" not in x for x in c["lignes"])


# ═══ 4. Un marché simple : rien ne change ═════════════════════════════════════════════════════════

def test_marche_simple_inchange_page_carnet_et_courrier(client, auth, mails, ia_interdite):
    from tests.test_redaction_v2_atelier import _projet as _projet_un_lot
    h = auth["headers"]
    _company(client, h)
    pid = _projet_un_lot(client, h)
    _creer(client, h, pid)
    ap = client.get(f"/api/atelier/{pid}/preparation", headers=h).json()
    assert ap["lot"] == "" and ap["page"]["lot"] is None
    assert not any(x.startswith("Ce fichier répond") for x in ap["page"]["resume"])
    r = client.get(f"/api/atelier/{pid}/preparation?lot=2", headers=h)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "lot_inconnu"
    st = _etat(client, h, pid)
    assert st["lots"]["actif"] is False
    # Le carnet d'un marché simple : clés de la grille, aucun lot, courrier sans lots.
    q = _marche(client, h)
    e = client.get(f"/api/carnet/{q}", headers=h).json()
    assert e["lots"] == [] and [l["cle"] for l in e["criteres"]["lignes"]] == ["c1", "c1.1", "c1.2", "c2"]
    assert all("lot" not in l for l in e["criteres"]["lignes"])
    assert e["criteres"]["message"] == cn.MSG_GRILLE
    lettre = client.get(f"/api/carnet/{q}/courrier", headers=h).json()
    assert "lots" not in lettre["corps"] and lettre["message_lots"] == "" and lettre["lots"] == []
    assert "pour chaque critère et sous-critère de jugement" in lettre["corps"]
    ls = [{"cle": "c1", "critere": "Valeur technique", "note": "42", "note_max": "60"},
          {"cle": "c2", "critere": "Prix des prestations", "note": "36", "note_max": "40"}]
    e2 = client.put(f"/api/carnet/{q}/notes", headers=h, json={"lignes": ls, "rev": e["rev"]}).json()
    assert e2["notes"]["total"] == {"note": "78", "note_max": "100", "pct": 78} and e2["notes"]["totaux"] == []
    assert all(x["lot"] == "" for x in e2["notes"]["lignes"])
    assert "(lot" not in e2["notes"]["constat"]["message"]
    assert mails == []


def test_fonctions_pures_lots_du_courrier():
    L = ["Objet : vos notes", "Pour les lots 2 et 3 :", "Lot 2 : Couverture", "Prix : 36 / 40", "LOT N° 03 - Étanchéité",
         "Prix : 33 / 40", "Moyens affectés au lot : 14/20"]
    assert cn.lots_des_lignes(L, ["2", "3"]) == [None, None, "2", "2", "3", "3", "3"]
    assert cn.lots_des_lignes(["Objet : lot 3", "Prix : 33 / 40"], ["2", "3"]) == ["3", "3"]
    assert cn.lots_des_lignes(["Pour les lots 2 et 3", "Prix : 33 / 40"], ["2", "3"]) is None
    assert cn.lots_des_lignes(["Lot 1 : Désamiantage", "Prix : 1 / 2"], ["2", "3"]) is None   # pas un de nos lots
    assert cn.cle_lot("2", "c1.1") == "L2-c1.1" and cn.sans_lot("L2-c1.1") == "c1.1" and cn.sans_lot("c1") == "c1"
    assert cn._SYS_IA_LOTS != cn._SYS_IA and cn._SYS_IA_LOTS.endswith(cn._SYS_IA.split("Réponds UNIQUEMENT")[1])


def test_ecrans_un_releve_par_lot_et_page_du_fichier_de_chaque_lot():
    import json
    import os
    from tests.test_redaction_v6_carnet import _node
    lignes = [{"cle": "L2-c1", "lot": "2", "lot_libelle": "Lot 2 — Couverture <b>", "critere": "Valeur technique"},
              {"cle": "L2-c2", "lot": "2", "lot_libelle": "Lot 2 — Couverture <b>", "critere": "Prix"},
              {"cle": "L3-c1", "lot": "3", "lot_libelle": "Lot 3 — Étanchéité", "critere": "Valeur technique"},
              {"cle": "u1", "lot": "", "critere": "Visite", "note": "3", "note_max": "5", "ajoutee": True}]
    part = {"message": "Notes", "lignes": [dict(x, note="4", note_max="10", pct=40, vous=x["lot"] == "3")
                                           for x in lignes[:3]],
            "vos_parties": [dict(lignes[2], note="4", note_max="10", pct=40)],
            "totaux": [{"lot": "2", "lot_libelle": "Lot 2 — Couverture", "note": "8", "note_max": "20", "pct": 40}]}
    r = _node(f"const L = {json.dumps(lignes)}, P = {json.dumps(part)};\n"
              "console.log(JSON.stringify({ t: R.cnTetes(L, true), t0: R.cnTetes(L.slice(0, 3), false),"
              " un: R.cnTetes([{cle: 'c1'}, {cle: 'c2'}], true), pl: R.cnPayload(L).map(x => x.lot),"
              " c: R.cnControle(L, [{numero: '2'}, {numero: '3'}]), c1: R.cnControle(L, []),"
              " h: R.cnPartenaireHtml(P) }));")
    assert r["t"] == ["Lot 2 — Couverture <b>", "", "Lot 3 — Étanchéité", "Critères ajoutés"]
    assert r["t0"] == ["Lot 2 — Couverture <b>", "", "Lot 3 — Étanchéité"] and r["un"] == ["", ""]
    assert r["pl"] == ["2", "2", "3", ""]
    assert "choisissez le lot" in r["c"] and r["c1"] == ""
    assert "<b>" not in r["h"] and "&lt;b&gt;" in r["h"] and "cn-lot-h" in r["h"]
    assert "Valeur technique (lot 3 — Étanchéité)" in r["h"]
    assert "Note d'ensemble du groupement, lot 2 — Couverture" in r["h"]
    racine = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    html = open(os.path.join(racine, "app", "static", "app.html"), encoding="utf-8").read()
    js = open(os.path.join(racine, "app", "static", "app.js"), encoding="utf-8").read()
    assert 'aria-label="Page du fichier de quel lot"' in html and "atlPrepLot(l.numero)" in html
    assert "atlDerniereFois(atlSec()).lot" in html
    assert '"/preparation" + (voulu ? "?lot=" + encodeURIComponent(voulu) : "")' in js
    assert "lot: l.numero, page_preparation: !!a.prepJoindre" in js


# ═══ 5. La démo : la page de chaque lot du marché de Trémorvan, rien d'écrit ═════════════════════

def test_demo_page_de_chaque_lot_sans_rien_ecrire(client, ia_interdite):
    from app.services.demo_seed import ensure_demo
    db = _db()
    try:
        ensure_demo(db, force=True)
    finally:
        db.close()
    r = client.post("/api/auth/demo")
    assert r.status_code == 200, r.text
    h = {"Authorization": f"Bearer {r.json()['access_token']}"}
    pid = next(x for x in client.get("/api/projects/", headers=h).json() if "Hirondelles" in x["name"])["id"]
    avant = _empreinte_demo(pid)
    for n in ("2", "3"):
        r = client.get(f"/api/atelier/{pid}/preparation?lot={n}", headers=h)
        assert r.status_code == 200, r.text
        assert r.json()["page"]["resume"][0].startswith(f"Ce fichier répond au lot {n} — ")
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx", "lot": "2",
                                                                   "page_preparation": True})
    assert r.status_code == 403
    assert _empreinte_demo(pid) == avant


def _empreinte_demo(pid):
    import json
    from app.models import DossierPresentation, Project
    db = _db()
    try:
        d = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).one()
        p = db.get(Project, pid)
        return json.dumps([d.dernier_export, d.options_export, d.derogation, p.carnet_notes], sort_keys=True,
                          default=str)
    finally:
        db.close()
