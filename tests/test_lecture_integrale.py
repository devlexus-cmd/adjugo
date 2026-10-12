"""
Lecture INTÉGRALE du dossier du marché (« on ne peut pas avoir que la moitié de l'appel d'offres
analysé »).

 • un dossier synthétique volumineux (≥ 10 pièces, ≥ 600 000 caractères, un fait unique à la fin du
   dernier CCTP et dans une annexe) : l'IA simulée reçoit TOUTES les pièces, en entier ; une seule
   analyse est débitée ; le coût est journalisé ;
 • un petit dossier est lu d'un seul tenant, règlement en tête, pendant la requête ;
 • les pièces non lues sont signalées avec leur raison (scan illisible, format, fichier protégé,
   fichier abîmé) ; les pages scannées sont lues en tâche de fond ;
 • une question sur le fait final retrouve la bonne pièce (et sa page), avec la phrase vérifiée ;
   une réponse absente du dossier est dite telle ; les questions sont bornées par jour ;
 • le texte intégral est rangé avec le dossier, listé dans l'export RGPD et supprimé avec le marché ;
 • la facturation au-delà du forfait part bien chez Stripe (simulé), et s'annule si l'IA tombe.
Aucun appel réseau ni IA réel : le fournisseur IA et Stripe sont simulés.
"""
import io
import json
import logging
import random
import re
import time
import zipfile

import pytest

from app.services import llm
from tests.test_redaction_socle0 import FakeLLM, _db, _user_id

FAIT = "Le titulaire du lot 6 fournit la carte grise du compacteur vibrant XQZ-4471 avant le chantier."
LIGNES_PAR_PAGE = 45
MOTS = ("chantier ouvrage entreprise titulaire prestation execution travaux reseau voirie chaussee "
        "terrassement beton enrobe bordure trottoir signalisation eclairage assainissement collecteur "
        "regard tranchee remblai canalisation fourniture pose controle reception essai etude plan "
        "dossier maitre oeuvre coordination securite environnement dechet planning phase").split()


def _phrase(rnd: random.Random, n=14) -> str:
    return " ".join(rnd.choice(MOTS) for _ in range(n)).capitalize() + "."


def _texte(code: str, chars: int, seed: int, fin: str = "") -> str:
    rnd = random.Random(seed)
    lignes, total = [f"DEBUT-PIECE-{code}"], 0
    i = 0
    while total < chars:
        i += 1
        ln = f"Article {i}. " + _phrase(rnd)
        lignes.append(ln)
        total += len(ln) + 1
        if i % 6 == 0:
            lignes.append("")
    if fin:
        lignes.append(fin)
    lignes.append(f"FIN-PIECE-{code}")
    return "\n".join(lignes)


def _pdf(texte: str, lignes_par_page: int = LIGNES_PAR_PAGE) -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    lignes = texte.split("\n")
    for k in range(0, len(lignes), lignes_par_page):
        y = 800
        for ln in lignes[k:k + lignes_par_page]:
            c.drawString(30, y, ln[:110])
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def _docx(texte: str) -> bytes:
    import docx
    d = docx.Document()
    for para in texte.split("\n"):
        d.add_paragraph(para)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx(rows) -> bytes:
    """Classeur .xlsx minimal (chaînes partagées), sans bibliothèque."""
    strings, idx = [], {}

    def si(v):
        if v not in idx:
            idx[v] = len(strings)
            strings.append(v)
        return idx[v]
    xml_rows = []
    for r, row in enumerate(rows, 1):
        cells = "".join(f'<c r="{chr(65 + c)}{r}" t="s"><v>{si(str(v))}</v></c>' for c, v in enumerate(row))
        xml_rows.append(f'<row r="{r}">{cells}</row>')
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rns = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    from xml.sax.saxutils import escape
    files = {
        "[Content_Types].xml": '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        "xl/workbook.xml": f'<?xml version="1.0"?><workbook {ns} {rns}><sheets><sheet name="BPU" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/sharedStrings.xml": f'<?xml version="1.0"?><sst {ns}>' + "".join(f"<si><t>{escape(x)}</t></si>" for x in strings) + "</sst>",
        "xl/worksheets/sheet1.xml": f'<?xml version="1.0"?><worksheet {ns}><sheetData>' + "".join(xml_rows) + "</sheetData></worksheet>",
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in files.items():
            z.writestr(n, d)
    return buf.getvalue()


def _blank_pdf(pages: int) -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _protected_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(io.BytesIO(_pdf("Document confidentiel")))
    w = PdfWriter()
    for p in r.pages:
        w.add_page(p)
    w.encrypt("motdepasse")
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _zip(members: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buf.getvalue()


RC_TETE = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Criteres de jugement des offres\n"
           "Valeur technique : 60 %\nPrix : 40 %\nLe memoire technique ne devra pas exceder 15 pages.\n")


def _gros_dossier():
    """≥ 10 pièces lisibles, ≥ 600 000 caractères ; le fait unique à la fin du dernier CCTP (PDF,
    dernière page) et dans une annexe ; plus des pièces illisibles à signaler."""
    textes = {
        "01_RC_Reglement_de_la_consultation.txt": RC_TETE + _texte("RC", 18000, 1),
        "02_Cadre_de_reponse.docx": "CADRE DE REPONSE TECHNIQUE\n" + _texte("CADRE", 4000, 2),
        "03_Acte_d_engagement.txt": "ACTE D'ENGAGEMENT\n" + _texte("AE", 3000, 3),
        "04_CCAP.txt": "CAHIER DES CLAUSES ADMINISTRATIVES PARTICULIERES\n" + _texte("CCAP", 60000, 4),
    }
    for lot in range(1, 6):
        textes[f"05_CCTP_lot{lot}.txt"] = ("CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n"
                                          + _texte(f"CCTP{lot}", 92000, 10 + lot))
    cctp6 = "CAHIER DES CLAUSES TECHNIQUES PARTICULIERES LOT 6\n" + _texte("CCTP6", 60000, 16, fin=FAIT)
    annexe = "ANNEXE 3 - PRESCRIPTIONS PARTICULIERES\n" + _texte("ANNEXE", 8000, 17, fin=FAIT)
    members = {}
    for n, t in textes.items():
        members[n] = _docx(t) if n.endswith(".docx") else t.encode("utf-8")
    members["06_CCTP_lot6.pdf"] = _pdf(cctp6)
    members["07_BPU.xlsx"] = _xlsx([["DEBUT-PIECE-BPU", "", ""], ["Poste", "Designation", "Prix unitaire"]]
                                   + [[str(i), f"Fourniture et pose de bordure type {i}", f"{i * 3} EUR"] for i in range(1, 120)]
                                   + [["FIN-PIECE-BPU", "", ""]])
    members["08_Annexe_3_prescriptions.txt"] = annexe.encode("utf-8")
    # Pièces qui ne peuvent pas être lues : elles doivent être SIGNALÉES, jamais ignorées.
    members["09_plan_signe_scan.pdf"] = _blank_pdf(2)
    members["10_ancien_format.doc"] = b"\xd0\xcf\x11\xe0 ancien format"
    members["11_fichier_abime.pdf"] = b"%PDF-1.4 " + bytes(random.Random(5).getrandbits(8) for _ in range(3000))
    members["12_protege.pdf"] = _protected_pdf()
    members["__MACOSX/._01_RC.txt"] = b"x"
    codes = ["RC", "CADRE", "AE", "CCAP"] + [f"CCTP{lot}" for lot in range(1, 7)] + ["BPU", "ANNEXE"]
    total = sum(len(t) for t in textes.values()) + len(cctp6) + len(annexe)
    page_du_fait = cctp6.split("\n").index(FAIT) // LIGNES_PAR_PAGE + 1
    return _zip(members), codes, total, page_du_fait


# ── IA simulée : enregistre tout ce qu'elle reçoit, répond selon la consigne ─────────

ANALYSE = {"match_score": 55, "go_decision": "a_etudier", "summary": "Analyse simulée du dossier complet.",
           "details": {"intitule_marche": "Aménagement de voirie en 6 lots", "acheteur": "Commune test",
                       "criteres_attribution": [
                           {"critere": "Valeur technique", "ponderation": "60 %", "extrait": "Valeur technique : 60 %"},
                           {"critere": "Prix", "ponderation": "40 %", "extrait": "Prix : 40 %"}]}}


def _segment(user: str) -> str:
    m = re.search(r"<<<\n(.*)\n>>>", user, flags=re.S)
    return m.group(1) if m else ""


def _qa_answer(user: str, phrase=None, trouve=True):
    blocks = re.findall(r"\[(P\d+)\] pièce « (.+?) » \(.*?\) :\n<<<\n(.*?)\n>>>", user, flags=re.S)
    cites = [{"ref": ref, "phrase": phrase or FAIT} for ref, _nom, txt in blocks if "XQZ-4471" in txt]
    if not trouve:
        return json.dumps({"trouve": False, "reponse": "Cette information ne figure pas dans le dossier.",
                           "citations": []})
    return json.dumps({"trouve": True, "reponse": "La référence exigée est XQZ-4471.", "citations": cites})


def _answer(kw):
    system = llm._flatten_content(kw.get("system"))
    content = kw["messages"][0]["content"]
    user = llm._flatten_content(content)
    if "Tu lis UNE pièce" in system:
        seg = _segment(user)
        lignes = [ln for ln in seg.split("\n") if len(ln) > 25]
        faits = [{"rubrique": "technique", "fait": "Exigence relevée.", "citation": lignes[0]}] if lignes else []
        if "XQZ-4471" in seg:
            faits.append({"rubrique": "technique", "fait": "Carte grise du compacteur exigée.", "citation": FAIT})
        # Une citation INVENTÉE : doit être écartée (introuvable dans la pièce).
        faits.append({"rubrique": "penalites", "fait": "Pénalité inventée.",
                      "citation": "Une penalite de 9999 euros par jour sera appliquee sans exception."})
        return json.dumps({"faits": faits}), "end_turn"
    if "réponds à une question sur un appel d'offres" in system:
        return _qa_answer(user), "end_turn"
    return json.dumps(ANALYSE), "end_turn"


@pytest.fixture
def fake_ia(monkeypatch):
    def install(answer=_answer):
        fake = FakeLLM(answer)
        monkeypatch.setattr(llm, "client", lambda: fake)
        monkeypatch.setattr(llm, "_use_mistral", lambda: False)
        llm._CB.update(fails=0, open_until=0.0)
        llm._TENANT_CB.clear()
        return fake
    return install


def _project(client, h, name="Voirie 6 lots"):
    r = client.post("/api/projects/", headers=h, json={"name": name, "client": "Commune test"})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _used(email):
    from app.models import User
    db = _db()
    try:
        return db.query(User).filter(User.email == email).first().analyses_used_this_month or 0
    finally:
        db.close()


def _wait(client, h, pid, timeout=120):
    t0 = time.time()
    while time.time() - t0 < timeout:
        st = client.get(f"/api/sourcing/lecture/{pid}", headers=h).json()
        if st["etat"] in ("termine", "erreur"):
            return st
        time.sleep(0.2)
    raise AssertionError("lecture en tâche de fond trop longue")


def _project_row(pid):
    from app.models import Project
    db = _db()
    try:
        p = db.get(Project, pid)
        db.expunge(p)
        return p
    finally:
        db.close()


# ═══ 1. Gros dossier : toutes les pièces passent par l'IA, en entier ════════════════════════════

def test_gros_dossier_toutes_les_pieces_lues_par_l_ia(client, auth, fake_ia, monkeypatch, caplog):
    from app.services import analysis as an
    monkeypatch.setattr(an, "ocr_available", lambda: False)      # pas d'OCR installé : signalé
    fake = fake_ia()
    h = auth["headers"]
    pid = _project(client, h)
    z, codes, total, page_du_fait = _gros_dossier()
    assert total >= 600_000 and len(codes) >= 10
    before = _used(auth["email"])
    caplog.set_level(logging.INFO, logger="adjugo")

    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE_voirie.zip", z, "application/zip")})
    assert r.status_code == 202, r.text                        # gros dossier : tâche de fond
    assert r.json()["en_cours"] is True and r.json()["job_id"]
    st = _wait(client, h, pid)
    assert st["etat"] == "termine", st

    # TOUTES les pièces lisibles ont été remises à l'IA, du premier au dernier mot.
    sent = fake.sent()
    for code in codes:
        assert f"DEBUT-PIECE-{code}" in sent, code
        assert f"FIN-PIECE-{code}" in sent, code
    assert sent.count("XQZ-4471") >= 2                          # fin du dernier CCTP ET annexe
    piece_calls = [k for k in fake.calls if "Tu lis UNE pièce" in llm._flatten_content(k.get("system"))]
    synth = [k for k in fake.calls if "Tu lis UNE pièce" not in llm._flatten_content(k.get("system"))]
    assert len(piece_calls) >= 9 and len(synth) == 1            # lecture en deux temps + synthèse
    synth_user = llm._flatten_content(synth[0]["messages"][0]["content"])
    # Synthèse : règlement en entier, en tête ; les faits relevés citent la phrase exacte, page comprise.
    assert synth_user.index("FIN-PIECE-RC") < synth_user.index("05_CCTP_lot1.txt")
    assert "FIN-PIECE-RC" in synth_user
    assert "XQZ-4471" in synth_user and re.search(r"« Le titulaire du lot 6 .*?» \(p\. \d+\)", synth_user)
    assert "9999 euros" not in synth_user                        # citation inventée : écartée
    # Consigne de l'analyse mise en cache de prompt.
    assert synth[0]["system"][0].get("cache_control") == {"type": "ephemeral"}

    # Une seule analyse débitée pour tout l'import ; coût journalisé (cost_meter).
    assert _used(auth["email"]) == before + 1
    assert any("Coût IA « analyse du dossier (deux_temps" in rec.getMessage() for rec in caplog.records)

    # Couverture honnête : pièces illisibles nommées avec leur raison.
    p = client.get(f"/api/projects/{pid}", headers=h).json()
    a = p["ai_analysis"]
    cov = a["couverture"]
    assert cov["mode"] == "deux_temps" and cov["complet"] is False
    raisons = {x["nom"]: x["raison"] for x in cov["non_lues"]}
    assert "pages en image" in raisons["09_plan_signe_scan.pdf"]
    assert "format non pris en charge" in raisons["10_ancien_format.doc"]
    assert "abîmé ou illisible" in raisons["11_fichier_abime.pdf"]
    assert "protégé par un mot de passe" in raisons["12_protege.pdf"]
    assert not any("MACOSX" in x["nom"] for x in a["lecture"]["pieces"])
    assert cov["totaux"]["lues_en_entier"] == len(codes)
    assert a["mode_lecture"] == "deux_temps"
    # La grille « Pièces lues » dit la même chose, pièce par pièce.
    lus = a["details"]["grille_notation"]["lecture"]["pieces"]
    assert {x["nom"] for x in lus if x.get("statut") == "non_lu"} >= {"09_plan_signe_scan.pdf", "12_protege.pdf"}
    cctp6 = next(x for x in a["lecture"]["pieces"] if x["nom"] == "06_CCTP_lot6.pdf")
    assert cctp6["statut"] == "lu_en_entier" and cctp6["pages"] >= 10 and cctp6["pages_lues"] == cctp6["pages"]

    # ── Question sur le fait final : la bonne pièce, la bonne page, la phrase vérifiée ──
    q = client.post("/api/sourcing/ask", headers=h,
                    json={"project_id": pid, "question": "Quelle référence de compacteur vibrant est exigée au lot 6 ?"})
    assert q.status_code == 200, q.text
    rep = q.json()
    assert rep["integral"] is True and rep["verifiee"] is True and "XQZ-4471" in rep["answer"]
    pieces_src = {s["piece"]: s for s in rep["sources"]}
    assert "06_CCTP_lot6.pdf" in pieces_src
    assert pieces_src["06_CCTP_lot6.pdf"]["page"] == page_du_fait == cctp6["pages"]   # dernière page
    assert FAIT in pieces_src["06_CCTP_lot6.pdf"]["citation"]
    qa_call = fake.calls[-1]
    blocks = qa_call["messages"][0]["content"]
    assert blocks[0].get("cache_control") == {"type": "ephemeral"}          # règlement en cache
    assert "Règlement de la consultation" in blocks[0]["text"] and "FIN-PIECE-RC" in blocks[0]["text"]
    assert "08_Annexe_3_prescriptions.txt" in blocks[1]["text"]             # l'annexe aussi retrouvée


# ═══ 2. Petit dossier : d'un seul tenant, pendant la requête ═════════════════════════════════════

def _petit_dossier():
    return _zip({
        "RC.txt": (RC_TETE + _texte("RC", 3000, 21)).encode(),
        "CCTP.docx": _docx("CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n" + _texte("CCTP", 15000, 22, fin=FAIT)),
        "BPU.xlsx": _xlsx([["DEBUT-PIECE-BPU"], ["Poste 1", "Bordure T2", "45 EUR"], ["FIN-PIECE-BPU"]]),
        "Annexe.odt": b"PK",                                  # OpenDocument abîmé : signalé
        "vieux.xls": b"\xd0\xcf\x11\xe0",                     # ancien Excel : signalé avec conseil
    })


def test_petit_dossier_lu_d_un_seul_tenant(client, auth, fake_ia):
    fake = fake_ia()
    h = auth["headers"]
    pid = _project(client, h, "Petit dossier")
    before = _used(auth["email"])
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE.zip", _petit_dossier(), "application/zip")})
    assert r.status_code == 200, r.text
    assert len(fake.calls) == 1                               # tout tient : un seul appel
    user = fake.calls[0]["messages"][0]["content"]
    assert isinstance(user, str)
    for code in ("RC", "CCTP", "BPU"):
        assert f"DEBUT-PIECE-{code}" in user and f"FIN-PIECE-{code}" in user
    assert user.index("DEBUT-PIECE-RC") < user.index("DEBUT-PIECE-CCTP")   # règlement en tête
    assert "Attention : certaines pièces n'ont pas pu être lues" in user and "vieux.xls" in user
    cov = r.json()["couverture"]
    assert cov["mode"] == "integral" and cov["complet"] is False
    raisons = {x["nom"]: x["raison"] for x in cov["non_lues"]}
    assert "ancien format Excel" in raisons["vieux.xls"]
    assert "Annexe.odt" in raisons
    assert _used(auth["email"]) == before + 1

    # Le texte intégral est rangé avec le dossier et sert à tous les usages (accès commun).
    from app.services import dce_texte
    proj = _project_row(pid)
    assert proj.dce_file_key and "/lecture/" in proj.dce_file_key
    full = dce_texte.dce_fulltext(proj)
    assert full.index("DEBUT-PIECE-RC") < full.index("DEBUT-PIECE-CCTP") and FAIT in full
    pieces = dce_texte.dce_pieces_text(proj)
    assert {p["nom"] for p in pieces if p["statut"] == "lu_en_entier"} == {"RC.txt", "CCTP.docx", "BPU.xlsx"}
    assert "Bordure T2 | 45 EUR" in next(p["texte"] for p in pieces if p["nom"] == "BPU.xlsx")
    assert dce_texte.dce_reglement(proj).startswith("REGLEMENT DE LA CONSULTATION")
    assert dce_texte.dce_texte_ou_extrait(proj)[1] == "integral"

    # Ré-analyse : le texte déjà lu est réutilisé tel quel (aucune relecture du fichier).
    import app.services.dce_lecture as lec

    def _interdit(*a, **k):
        raise AssertionError("relecture inutile")
    from pytest import MonkeyPatch
    mp = MonkeyPatch()
    mp.setattr(lec, "lire_dossier", _interdit)
    try:
        rr = client.post("/api/sourcing/reanalyze", headers=h, data={"project_id": str(pid)})
    finally:
        mp.undo()
    assert rr.status_code == 200, rr.text
    assert len(fake.calls) == 2 and FAIT in fake.calls[1]["messages"][0]["content"]
    assert _project_row(pid).dce_file_key == proj.dce_file_key


# ═══ 3. Pages scannées : lues en tâche de fond, ou signalées ═════════════════════════════════════

def test_scan_lu_en_tache_de_fond_et_page_illisible_signalee(client, auth, fake_ia, monkeypatch):
    from app.services import analysis as an
    fake = fake_ia()
    monkeypatch.setattr(an, "ocr_available", lambda: True)

    def fake_ocr(content, page_nums, budget, results, stop=None):
        for n in page_nums:
            results[n] = ("" if n == 2 else
                          "REGLEMENT DE LA CONSULTATION scanne. Les offres sont remises avant le 15 decembre 2026 a midi.")
        return {"disponible": True, "non_lues": {}}
    monkeypatch.setattr(an, "ocr_pdf_pages", fake_ocr)
    h = auth["headers"]
    pid = _project(client, h, "Dossier scanné")
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("RC_scan.pdf", _blank_pdf(3), "application/pdf")})
    assert r.status_code == 202, r.text                      # OCR : jamais pendant la requête
    st = _wait(client, h, pid)
    assert st["etat"] == "termine", st
    a = client.get(f"/api/projects/{pid}", headers=h).json()["ai_analysis"]
    piece = a["lecture"]["pieces"][0]
    assert piece["statut"] == "lu_en_partie" and piece["pages"] == 3 and piece["pages_lues"] == 2
    assert "1 page sur 3 non lue" in piece["detail"] and "aucun texte reconnu" in piece["detail"]
    assert "15 decembre 2026" in fake.sent()
    assert a["couverture"]["complet"] is False


def test_scan_sans_ocr_signale_et_rembourse_si_rien_n_est_lisible(client, auth, fake_ia, monkeypatch):
    from app.services import analysis as an
    fake_ia()
    monkeypatch.setattr(an, "ocr_available", lambda: False)
    h = auth["headers"]
    pid = _project(client, h, "Scan seul")
    before = _used(auth["email"])
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("scan.pdf", _blank_pdf(2), "application/pdf")})
    assert r.status_code == 202, r.text
    st = _wait(client, h, pid)
    assert st["etat"] == "erreur" and "scannées" in st["erreur"]
    assert _used(auth["email"]) == before                     # rien d'analysé : remboursé


def test_formats_non_lisibles_refuses_clairement(client, auth, fake_ia):
    fake_ia()
    h = auth["headers"]
    pid = _project(client, h, "Formats")
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("dossier.zip", _zip({"plan.dwg": b"x", "vieux.doc": b"y"}), "application/zip")})
    assert r.status_code == 422 and "formats non pris en charge" in r.json()["detail"]


# ═══ 4. Questions : réponse absente, citation non vérifiée, ancienne analyse, borne du jour ══════

def test_questions_absent_non_verifie_ancienne_analyse_et_borne(client, auth, fake_ia, monkeypatch):
    h = auth["headers"]
    fake = fake_ia()
    pid = _project(client, h, "Questions")
    assert client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                       files={"file": ("DCE.zip", _petit_dossier(), "application/zip")}).status_code == 200

    # Pas dans le dossier : l'IA le dit, Adjugo nomme les pièces non lues.
    def absent(kw):
        if "réponds à une question" in llm._flatten_content(kw.get("system")):
            return _qa_answer("", trouve=False), "end_turn"
        return _answer(kw)
    fake._answer = absent
    r = client.post("/api/sourcing/ask", headers=h, json={"project_id": pid, "question": "Date de la visite ?"})
    assert r.status_code == 200, r.text
    assert r.json()["trouve"] is False and r.json()["sources"] == []
    assert "ne figure pas" in r.json()["answer"] and "vieux.xls" in r.json()["answer"]

    # Phrase inventée par l'IA : jamais présentée comme une citation.
    def invente(kw):
        if "réponds à une question" in llm._flatten_content(kw.get("system")):
            user = llm._flatten_content(kw["messages"][0]["content"])
            return _qa_answer(user, phrase="Le compacteur doit etre peint en rouge vif."), "end_turn"
        return _answer(kw)
    fake._answer = invente
    r = client.post("/api/sourcing/ask", headers=h, json={"project_id": pid, "question": "Compacteur vibrant ?"})
    assert r.status_code == 200 and r.json()["sources"] == [] and r.json()["verifiee"] is False
    assert "non retrouvée telle quelle" in r.json()["message"]

    # Borne hors quota : N questions par organisation et par jour, l'échec de l'IA ne compte pas.
    from app.core import quota
    monkeypatch.setattr(quota.settings, "QUESTIONS_PAR_JOUR", 4)
    fake._answer = _answer
    ok = client.post("/api/sourcing/ask", headers=h, json={"project_id": pid, "question": "Compacteur vibrant ?"})
    assert ok.status_code == 200

    def panne(kw):
        raise RuntimeError("fournisseur en panne")
    fake._answer = panne
    assert client.post("/api/sourcing/ask", headers=h, json={"project_id": pid, "question": "Prix ?"}).status_code == 502
    fake._answer = _answer
    llm._CB.update(fails=0, open_until=0.0)
    llm._TENANT_CB.clear()
    assert client.post("/api/sourcing/ask", headers=h, json={"project_id": pid, "question": "Prix ?"}).status_code == 200
    stop = client.post("/api/sourcing/ask", headers=h, json={"project_id": pid, "question": "Lots ?"})
    assert stop.status_code == 429 and "par organisation et par jour" in stop.json()["detail"]

    # Analyse antérieure à la lecture intégrale : comportement d'origine + invitation à ré-analyser.
    from app.models import Project
    db = _db()
    try:
        old = Project(user_id=_user_id(auth["email"]), name="Ancien",
                      ai_analysis={"dce_excerpt": "REGLEMENT. Date limite le 3 mars.", "summary": "s", "details": {}})
        db.add(old)
        db.commit()
        oid = old.id
    finally:
        db.close()
    monkeypatch.setattr(quota.settings, "QUESTIONS_PAR_JOUR", 50)
    fake._answer = lambda kw: ("Le 3 mars.", "end_turn")
    r = client.post("/api/sourcing/ask", headers=h, json={"project_id": oid, "question": "Date limite ?"})
    assert r.status_code == 200 and r.json()["integral"] is False
    assert "Ré-analysez ou ré-importez le dossier" in r.json()["message"]


# ═══ 5. Stockage : rangé avec le marché, dans l'export RGPD, supprimé avec lui ═══════════════════

def test_texte_integral_supprime_avec_le_marche_et_dans_l_export(client, auth, fake_ia):
    from app.services.storage import get_storage
    fake_ia()
    h = auth["headers"]
    pid = _project(client, h, "À supprimer")
    assert client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                       files={"file": ("DCE.zip", _petit_dossier(), "application/zip")}).status_code == 200
    key = _project_row(pid).dce_file_key
    assert get_storage().load(key)
    exp = client.get("/api/rgpd/export", headers=h)
    assert exp.status_code == 200
    lus = [x for x in exp.json()["texte_lu_des_dossiers"] if x["appel_offres_id"] == pid]
    assert lus and {p["nom"] for p in lus[0]["pieces"]} >= {"RC.txt", "CCTP.docx", "vieux.xls"}
    assert lus[0]["taille"] and "DEBUT-PIECE" not in json.dumps(lus)        # métadonnées, pas le texte

    assert client.delete(f"/api/projects/{pid}", headers=h).status_code == 204
    from app.routers.rgpd import purge_trash
    db = _db()
    try:
        purge_trash(db, older_than_days=0)
    finally:
        db.close()
    assert _project_row(pid).dce_file_key is None
    with pytest.raises(Exception):
        get_storage().load(key)


# ═══ 6. Facturation au-delà du forfait (Stripe simulé) ═══════════════════════════════════════════

def test_stripe_price_overage_lu_depuis_l_environnement(monkeypatch):
    from app.core.config import Settings
    monkeypatch.setenv("STRIPE_PRICE_OVERAGE", "price_overage_test")
    assert Settings().STRIPE_PRICE_OVERAGE == "price_overage_test"
    assert Settings().STRIPE_METER_EVENT_OVERAGE == "adjugo_analyse_overage"


def test_depassement_envoye_a_stripe_puis_annule_si_rembourse(client, auth, monkeypatch):
    import stripe
    from app.core.quota import consume_analysis, refund_analysis, _limit, _period
    from app.models import User, PlanType
    import app.services.billing_usage as bu
    events, cancels = [], []
    monkeypatch.setattr(bu.settings, "STRIPE_PRICE_OVERAGE", "price_overage_test")
    monkeypatch.setattr(stripe.billing.MeterEvent, "create", staticmethod(lambda **kw: events.append(kw) or kw))
    monkeypatch.setattr(stripe.billing.MeterEventAdjustment, "create",
                        staticmethod(lambda **kw: cancels.append(kw) or kw))
    db = _db()
    try:
        u = db.query(User).filter(User.email == auth["email"]).first()
        u.plan, u.stripe_customer_id, u.overage_enabled = PlanType.pro, "cus_overage_test", True
        u.analyses_period, u.analyses_used_this_month = _period(), _limit(u)
        db.commit()
        consume_analysis(u, db)                                  # au-delà du forfait
        db.refresh(u)
        assert u.overage_count == 1 and u.analyses_used_this_month == _limit(u) + 1
        assert len(events) == 1
        ev = events[0]
        assert ev["event_name"] == "adjugo_analyse_overage"
        assert ev["payload"] == {"value": "1", "stripe_customer_id": "cus_overage_test"}
        assert ev["identifier"].startswith(f"adjugo-ovr-{u.id}-")
        refund_analysis(u, db)                                   # l'IA est tombée : on rembourse
        db.refresh(u)
        assert u.overage_count == 0
        assert cancels == [{"event_name": "adjugo_analyse_overage", "type": "cancel",
                            "cancel": {"identifier": ev["identifier"]}}]
    finally:
        db.close()


def test_abonnement_porte_le_prix_a_l_usage(client, auth, monkeypatch):
    import stripe
    from types import SimpleNamespace
    import app.routers.stripe_pay as sp
    import app.services.billing_usage as bu
    from app.models import User, PlanType
    monkeypatch.setattr(sp.settings, "STRIPE_PRICE_OVERAGE", "price_overage_test")
    monkeypatch.setattr(sp.settings, "STRIPE_PRICE_PRO", "price_pro_test")
    seen = {}
    monkeypatch.setattr(stripe.Customer, "create", staticmethod(lambda **kw: SimpleNamespace(id="cus_new")))
    monkeypatch.setattr(stripe.checkout.Session, "create",
                        staticmethod(lambda **kw: seen.update(kw) or SimpleNamespace(url="https://x", id="cs_1")))
    h = auth["headers"]
    r = client.post("/api/stripe/create-checkout?plan=pro", headers=h)
    assert r.status_code == 200, r.text
    assert seen["line_items"] == [{"price": "price_pro_test", "quantity": 1}, {"price": "price_overage_test"}]

    # Le forfait est lu même quand le prix à l'usage est le premier élément de l'abonnement.
    items = [{"price": {"id": "price_overage_test"}}, {"price": {"id": "price_pro_test"}}]
    assert bu.plan_price_id(items) == "price_pro_test"

    # Abonnement plus ancien : activer le paiement à l'usage y ajoute le prix à l'usage.
    added = []
    sub = {"id": "sub_1", "items": {"data": [{"price": {"id": "price_pro_test"}}]}}
    monkeypatch.setattr(stripe.Subscription, "list", staticmethod(lambda **kw: SimpleNamespace(data=[sub])))
    monkeypatch.setattr(stripe.SubscriptionItem, "create", staticmethod(lambda **kw: added.append(kw)))
    db = _db()
    try:
        u = db.query(User).filter(User.email == auth["email"]).first()
        u.plan, u.stripe_customer_id = PlanType.pro, "cus_new"
        db.commit()
    finally:
        db.close()
    ok = client.post("/api/stripe/overage", params={"enabled": "true"}, headers=h)
    assert ok.status_code == 200, ok.text
    assert added == [{"subscription": "sub_1", "price": "price_overage_test"}]


# ═══ 7. Écrans : couverture honnête, sources des réponses, progression ═══════════════════════════

def _static(name):
    import os
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    return open(os.path.join(root, "app", "static", name), encoding="utf-8").read()


def test_ecrans_couverture_sources_et_progression():
    js, html = _static("app.js"), _static("app.html")
    # « Pièces lues » : lue en entier / en partie / non lue, avec la raison.
    assert "piecesResume" in js and "· non lue" in js and "· lue en partie" in js and "p.raison" in js
    # Carte d'analyse : la couverture, et l'invitation à ré-analyser une ancienne analyse.
    assert "aoCouverture().message" in html and "aoAncienneLecture()" in html
    assert "Ré-analysez ou ré-importez le dossier pour qu'Adjugo lise toutes les pièces" in html
    # Réponse d'une question : la phrase, la pièce, la page.
    assert 'v-for="(sc,si) in m.sources"' in html and "sc.piece" in html and "page {{ sc.page }}" in html
    # Gros dossier : progression de la lecture affichée pendant la tâche de fond.
    assert "/api/sourcing/lecture/" in js and "ao.lectureMsg" in html and "Lecture du dossier : " in js
    # Sobre : aucun émoji dans ce qui a été ajouté.
    zone = html[html.index("Couverture de la lecture"):html.index("Couverture de la lecture") + 2500]
    assert not re.search(r"[\U0001F300-\U0001FAFF☀-➿]", zone)


# ═══ 8. Lecture pièce par pièce : cas particuliers (sans IA) ═════════════════════════════════════

def test_lecture_cas_particuliers(monkeypatch):
    from app.services import dce_lecture as lec, analysis as an
    # Plafond technique : pièce lue en partie, raison dite (jamais « lue en entier » à tort).
    monkeypatch.setattr(lec, "MAX_CARACTERES_PIECE", 5000)
    lu = lec.lire_dossier("DCE.zip", _zip({"RC.txt": (RC_TETE + _texte("RC", 20000, 31)).encode(),
                                           "notes.csv": "Désignation;Prix\nBordure;45 €\n".encode("cp1252")}))
    rc = next(p for p in lu["pieces"] if p["nom"] == "RC.txt")
    assert rc["statut"] == "lu_en_partie" and "plafond technique" in rc["detail"] and rc["caracteres"] == 5000
    csv = next(p for p in lu["pieces"] if p["nom"] == "notes.csv")
    assert "Désignation" in csv["texte"] and "45 €" in csv["texte"]          # export Excel (Windows-1252)
    monkeypatch.undo()

    # PDF sans extension reconnu, PDF chiffré sans mot de passe d'ouverture lisible.
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for pg in PdfReader(io.BytesIO(_pdf("REGLEMENT DE LA CONSULTATION lisible malgre le chiffrement"))).pages:
        w.add_page(pg)
    w.encrypt(user_password="", owner_password="proprietaire")
    buf = io.BytesIO()
    w.write(buf)
    lu = lec.lire_dossier("DCE.zip", _zip({"reglement": _pdf("REGLEMENT DE LA CONSULTATION sans extension"),
                                           "droits_restreints.pdf": buf.getvalue()}))
    noms = {p["nom"]: p for p in lu["pieces"]}
    assert noms["reglement"]["statut"] == "lu_en_entier" and noms["reglement"]["format"] == "pdf"
    assert noms["droits_restreints.pdf"]["statut"] == "lu_en_entier"

    # Word : tableaux imbriqués lus.
    import docx
    d = docx.Document()
    d.add_paragraph("CADRE DE REPONSE TECHNIQUE")
    t = d.add_table(rows=1, cols=2)
    inner = t.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Question 4 : moyens humains affectes"
    t.cell(0, 0).text = "Rubrique"
    b = io.BytesIO()
    d.save(b)
    lu = lec.lire_dossier("Cadre.docx", b.getvalue())
    assert "Question 4 : moyens humains affectes" in lu["pieces"][0]["texte"]

    # Citation retrouvée malgré accents, casse et espaces : position dans le texte d'origine.
    texte = "Préambule.\n\nLe  Titulaire   fournit\nla carte grise du compacteur."
    loc = an.Localisateur(texte)
    assert loc.offset("le titulaire fournit la carte grise") == texte.index("Le  Titulaire")
    assert an.citation_offset(texte, "phrase absente du texte") is None


def test_ancienne_extraction_lit_aussi_les_tableurs():
    from app.services.analysis import extract_dce_text
    z = _zip({"RC.txt": RC_TETE.encode(), "BPU.xlsx": _xlsx([["Poste", "Bordure T2", "45 EUR"]])})
    t = extract_dce_text("DCE.zip", z, max_chars=50000)
    assert "Bordure T2 | 45 EUR" in t
