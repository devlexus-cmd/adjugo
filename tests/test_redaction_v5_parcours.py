"""
Rédaction, vague 5 — le parcours complet, de bout en bout (contrat : docs/REDACTION_V5.md, section « I »).

Un seul dossier traverse les trois chantiers de la vague (C2 complet, D, E2), comme un client le vivrait :
 1. le dossier de l'acheteur est IMPORTÉ (route réelle, IA de l'analyse simulée) ;
 2. « Comment vous serez noté » : critères ET sous-critères tels qu'écrits, chacun avec sa phrase retrouvée ;
    la liste des exigences (sans IA, plus une proposée par l'IA et contrôlée) ; l'environnement ; R2142-21 ;
    la phrase du règlement sur l'IA ; les questions à poser à l'acheteur et leur date limite ;
 3. « Mes preuves » : un marché public PROPOSÉ depuis DECP (données d'exemple, aucun réseau), coché par
    l'utilisateur, enregistré « marché attribué, source publique DECP », daté ;
 4. l'entretien chantier : des réponses (« déclaré par vous »), une question PASSÉE, une question adressée à un
    partenaire qui répond DANS SON ESPACE (lien + clé) ;
 5. l'atelier (IA simulée) : les faits « déclaré par vous » [R#] et « source publique DECP » [F#] cités et reliés par
    le contrôle des faits, la réponse du partenaire « déclaré par » lui, « n exigences traitées sur m », le
    « [À compléter] » de la question passée ;
 6. le mode « sans IA » que le règlement demande : plus aucun appel d'IA (rédaction, raccourci, questions de
    l'entretien ; l'aide du partenaire : tests/test_redaction_v5_c2.py), le partenaire le lit chez lui ; le contrôle
    des faits, la liste des exigences et les questions de l'entretien continuent ;
 7. l'export Word : aucune marque interne (ni renvoi, ni « [À compléter] », ni libellé de source, ni Adjugo).
Aucun appel réseau : les sources publiques et l'IA sont simulées.
"""
import io
import json
import zipfile

import pytest

from app.services import llm
from app.services import preuves as pv
from app.services import sources_publiques as sp
from app.services.analysis import citation_guard
from tests.test_redaction_socle0 import _db, fake_llm  # noqa: F401
from tests.test_redaction_v5_e2 import FauxSources

SIRET_CHEF = "12345678900011"
SIRET_PART = "55566677700012"
ACH = "21290232900016"                        # acheteur fictif (nom retrouvé dans l'annuaire simulé)
NOM_ACH = "COMMUNE DE KERBRAT (FICTIVE)"

RC = ("REGLEMENT DE LA CONSULTATION\n"
      "Article 1 - Objet\nRéfection de la toiture du groupe scolaire Paul Bert.\n"
      "Article 2 - Présentation des candidatures\n"
      "Un même opérateur ne peut pas présenter plusieurs offres en agissant à la fois en qualité de candidat "
      "individuel et de membre d'un groupement.\n"
      "Article 5 - Visite\nLa visite du site est obligatoire avant la remise des offres.\n"
      "Article 7 - Critères de jugement des offres\n"
      "1. Valeur technique : 70 %\n"
      "1.1 Méthodologie d'exécution en site occupé : 40 %\n"
      "1.2 Moyens humains et encadrement : 20 %\n"
      "1.3 Démarche environnementale : 10 %\n"
      "2. Prix : 30 %\n"
      "Article 8 - Contenu du mémoire technique\n"
      "Le mémoire technique comprendra :\n"
      "- la méthodologie d'exécution en site occupé ;\n"
      "- l'encadrement du chantier ;\n"
      "- la gestion des déchets du chantier ;\n"
      "- un planning par phase, calé sur les vacances scolaires.\n"
      "Le mémoire technique est rédigé sans recours à un système d'intelligence artificielle.\n"
      "Le mémoire technique comportera au maximum 10 pages.\n"
      "Article 9 - Questions des candidats\n"
      "Les candidats peuvent poser leurs questions sur la plateforme au plus tard le 05/12/2026.\n"
      "Article 10 - Forme de l'offre\n"
      "Le mémoire technique ne dépassera pas 15 pages.\n")
CCTP = ("CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n"
        "Article 1 - Contraintes du site\n"
        "L'école reste en activité pendant toute la durée des travaux.\n"
        "Les livraisons sont interdites aux heures d'entrée et de sortie des élèves.\n"
        "Article 3 - Couverture\nLa couverture existante en zinc est déposée par zones successives.\n"
        "Article 4 - Déchets\nLes déchets sont triés sur le chantier et évacués vers des filières agréées.\n"
        "Article 6 - Propreté\nLa réduction des trajets des véhicules est à exposer dans l'offre.\n")

ANALYSE = {
    "match_score": 64, "go_decision": "a_etudier", "summary": "Toiture d'école en site occupé.",
    "details": {
        "intitule_marche": "Réfection de la toiture du groupe scolaire Paul Bert", "acheteur": "Ville de Quimperlé",
        "criteres_attribution": [
            {"critere": "Valeur technique", "ponderation": "70 %", "extrait": "1. Valeur technique : 70 %",
             "sous_criteres": [
                 {"critere": "Méthodologie d'exécution en site occupé", "ponderation": "40 %",
                  "extrait": "1.1 Méthodologie d'exécution en site occupé : 40 %"},
                 {"critere": "Moyens humains et encadrement", "ponderation": "20 %",
                  "extrait": "1.2 Moyens humains et encadrement : 20 %"},
                 {"critere": "Démarche environnementale", "ponderation": "10 %",
                  "extrait": "1.3 Démarche environnementale : 10 %"}]},
            {"critere": "Prix", "ponderation": "30 %", "extrait": "2. Prix : 30 %", "sous_criteres": []}],
        # Proposées par l'IA : gardées seulement si leur phrase est retrouvée mot pour mot (et sans nombre inventé).
        "exigences_memoire": ["La réduction des trajets des véhicules est à exposer dans l'offre.",
                              "Le candidat garantira un taux de valorisation des déchets de 95 %."],
        "questions_acheteur": [{"question": "Le planning doit-il tenir en 3 phases ?",
                                "extraits": ["- un planning par phase, calé sur les vacances scolaires."]}]}}


def sp_jour() -> str:
    return pv.aujourd_hui().isoformat()


def _zip(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, text in files.items():
            zf.writestr(name, text)
    return buf.getvalue()


def _docx_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return "\n".join(zf.read(n).decode("utf-8", "replace") for n in zf.namelist()
                         if n.endswith(".xml") or n.endswith(".rels"))


@pytest.fixture
def sources(monkeypatch):
    """DECP, liste RGE et annuaire simulés ; le transport HTTP réel est coupé (un appel ferait échouer le test)."""
    import httpx
    fake = FauxSources()
    base = {"nature": "Marché", "titulaire_id_2": "CDL", "titulaire_id_3": "CDL", "lieuexecution_code": "29"}
    fake.decp = [
        dict(base, id="M-2024-01", objet="Réfection de la toiture de l'école Jules-Ferry", montant=85000.0,
             datenotification="2024-05-10", acheteur_id=ACH, titulaire_id_1=SIRET_CHEF,
             typegroupementoperateurs="Pas de groupement"),
        dict(base, id="M-2022-03", objet="CDL", montant=12000.0, datenotification="2022-03-01", acheteur_id=ACH,
             titulaire_id_1=SIRET_CHEF, typegroupementoperateurs="Pas de groupement"),
        dict(base, id="M-2023-07", objet="Couverture du gymnase municipal", montant=142000.0,
             datenotification="2023-09-03", acheteur_id=ACH, titulaire_id_1=SIRET_PART,
             typegroupementoperateurs="Pas de groupement")]
    fake.rge = []
    monkeypatch.setattr(sp, "_http_get", fake)
    reels = []

    def coupe(self, request):
        reels.append(str(request.url))
        raise httpx.ConnectError("réseau coupé (test)", request=request)
    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", coupe)
    sp.vider_cache()
    yield fake
    sp.vider_cache()
    assert reels == [], f"appel réseau réel : {reels}"


def _attendre(client, h, pid, jid):
    import time
    r = None
    for _ in range(300):
        r = client.get(f"/api/atelier/{pid}/taches/{jid}", headers=h)
        assert r.status_code == 200, r.text
        if r.json()["status"] in ("done", "error"):
            return r.json()
        time.sleep(0.05)
    raise AssertionError(f"tâche {jid} jamais terminée : {r.json() if r is not None else None}")


def _etat(client, h, pid):
    r = client.get(f"/api/atelier/{pid}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _rediger(client, h, pid, sid):
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
    assert r.status_code == 202, r.text
    return _attendre(client, h, pid, r.json()["tache"]["id"])


def _ecrire(client, h, pid, sid, texte):
    sec = next(s for s in _etat(client, h, pid)["sections"] if s["id"] == sid)
    r = client.put(f"/api/atelier/{pid}/sections/{sid}/texte", headers=h,
                   json={"texte": texte, "base_version_id": (sec.get("version_courante") or {}).get("id")})
    assert r.status_code == 200, r.text
    return r.json()


def _faits(client, h, pid, sid):
    r = client.get(f"/api/atelier/{pid}/sections/{sid}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()["version_courante"]


def _mesure():
    from app.services import atelier as at
    db = _db()
    try:
        return at.mesure_fichier_impose(db)
    finally:
        db.close()


def test_parcours_complet_vague_5(client, auth, fake_llm, sources, monkeypatch):
    from tests.test_convention_groupement import _invite, _put
    from tests.test_redaction_v4_groupement import P
    monkeypatch.setenv("ATELIER_LIBELLES_SOURCES", "1")       # les libellés de source montrés à l'écran
    h = auth["headers"]
    r = client.put("/api/company/", headers=h, json={
        "name": "Toitures du Léon", "siret": SIRET_CHEF, "representant_legal": "Anne Le Goff",
        "address": "3 quai de l'Odet", "city": "Quimper", "postal_code": "29000", "effectif": 14})
    assert r.status_code == 200, r.text

    # ── 1. Le dossier de l'acheteur, importé par la route réelle (IA de l'analyse simulée) ────────────────
    def analyse(kw):
        if "analyse des Dossiers de Consultation" in llm._flatten_content(kw.get("system")):
            return json.dumps(ANALYSE, ensure_ascii=False), "end_turn"
        raise AssertionError("seule l'analyse appelle l'IA à cette étape")
    fake_llm(analyse)
    m0 = _mesure()
    pid = client.post("/api/projects/", headers=h, json={"name": "Toiture du groupe scolaire Paul Bert",
                                                          "client": "Ville de Quimperlé"}).json()["id"]
    dce = {"01_Reglement_de_la_consultation.txt": RC, "03_CCTP.txt": CCTP}
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE.zip", _zip(dce), "application/zip")})
    assert r.status_code == 200, r.text
    g = r.json()["details"]["grille_notation"]
    texte_dce = RC + "\n" + CCTP
    # Décision 11 (c) : l'import est COMPTÉ (agrégat, aucun contenu) ; ce dossier n'impose pas son fichier.
    m1 = _mesure()
    assert m1["dossiers_importes"] == m0["dossiers_importes"] + 1 and m1["fichier_impose"] == m0["fichier_impose"]

    # ── 2. « Comment vous serez noté » : critères, sous-critères, exigences, avertissements, questions ─────
    vt = next(c for c in g["criteres"] if c["critere"] == "Valeur technique")
    assert [(x["critere"], x["poids"]) for x in vt["sous_criteres"]] == [
        ("Méthodologie d'exécution en site occupé", "40 %"), ("Moyens humains et encadrement", "20 %"),
        ("Démarche environnementale", "10 %")]
    assert vt["controle_sous"]["statut"] == "ok" and g["controle_total"]["statut"] == "ok"
    exig = g["exigences"]["liste"]
    assert [x["texte"] for x in exig] == [
        "la méthodologie d'exécution en site occupé", "l'encadrement du chantier",
        "la gestion des déchets du chantier", "un planning par phase, calé sur les vacances scolaires",
        "La réduction des trajets des véhicules est à exposer dans l'offre"]
    assert exig[-1]["origine"] == "ia"                              # proposée par l'IA, phrase retrouvée
    assert "95 %" not in json.dumps(g, ensure_ascii=False)          # l'exigence inventée par l'IA est écartée
    assert g["environnement"]["dans_la_note"] is True
    assert [m["code"] for m in g["candidatures_multiples"]["messages"]] == ["seul_et_groupement"]
    assert g["ia"]["nature"] == "interdit" and g["ia"]["propose_sans_ia"] is True
    qa = g["questions_acheteur"]
    assert [q["code"] for q in qa["liste"]] == ["limites"]          # la question de l'IA (« 3 phases ») : rejetée
    assert qa["date_limite"]["date"] == "2026-12-05" and "Adjugo n'envoie rien" in qa["message"]
    # Chaque phrase affichée est retrouvée mot pour mot dans le dossier (citation_guard).
    phrases = ([x["phrase"] for x in vt["sous_criteres"]] + [x["phrase"] for x in exig]
               + [x["phrase"] for x in g["environnement"]["lieux"]] + [x["phrase"] for x in g["ia"]["phrases"]]
               + [x["phrase"] for q in qa["liste"] for x in q["phrases"]] + [qa["date_limite"]["phrase"]])
    for ph in phrases:
        assert citation_guard(texte_dce, ph, min_len=4), ph

    # ── 3. Atelier : une partie par sous-critère, sous le titre du critère ; « 0 exigence traitée sur 5 » ──
    r = client.post(f"/api/atelier/{pid}", headers=h, json={"base": "auto"})
    assert r.status_code == 201, r.text
    st = r.json()
    assert [(s["numero_affiche"], s["intitule"]) for s in st["sections"]] == [
        ("1.1", "Méthodologie d'exécution en site occupé"), ("1.2", "Moyens humains et encadrement"),
        ("1.3", "Démarche environnementale")]
    assert all(s["parents"] == [{"numero": "1", "intitule": "Valeur technique"}] for s in st["sections"])
    assert st["exigences"]["message"] == "0 exigence traitée sur 5"
    assert st["sans_ia"]["actif"] is False and st["sans_ia"]["reglement"]["nature"] == "interdit"
    sid = {s["intitule"]: s["id"] for s in st["sections"]}
    s_met, s_moy, s_env = (sid["Méthodologie d'exécution en site occupé"], sid["Moyens humains et encadrement"],
                           sid["Démarche environnementale"])

    # ── 4. Un partenaire du groupement (lien, sans compte), « Qui écrit quoi » préparé ──────────────────────
    a = _invite(client, h, pid, "Kerbrat Zinguerie")
    assert _put(client, a["token"], company_name="Kerbrat Zinguerie", lot="Couverture et zinguerie",
                siret=SIRET_PART).status_code == 200
    assert client.post(f"/api/atelier/{pid}/repartition", headers=h, json={}).status_code == 200

    # ── 5. « Mes preuves » : un marché public PROPOSÉ (DECP), coché, « marché attribué, source publique DECP » ──
    d = client.get("/api/preuves/publiques/marches", headers=h).json()
    assert d["statut"] == "ok" and d["a_cocher"] == 1 and len(d["marches"]) == 2     # la ligne « CDL » : montrée
    m_ok = next(m for m in d["marches"] if m["decp_id"] == "M-2024-01")
    assert not m_ok["deja"] and m_ok["source_libelle"] == "marché attribué, source publique DECP"
    assert m_ok["affiche"]["acheteur"] == NOM_ACH and "réalisé" not in m_ok["texte"] and "attribué" in m_ok["texte"]
    assert next(m for m in d["marches"] if m["decp_id"] == "M-2022-03")["manque"]
    r = client.post("/api/preuves/publiques", headers=h, json={"cles": [m_ok["cle"]]})
    assert r.status_code == 200, r.text
    preuve = r.json()["creees"][0]
    assert preuve["statut"] == "acceptee" and preuve["origine"] == "import_public"
    assert preuve["date_source"] == preuve["source"]["consulte_le"] == sp_jour()
    f_decp = preuve["ref"]

    # Le partenaire reprend SES marchés publics dans SA part (SIRET de sa part ; rien du chef de file).
    dp_ = client.get(f"/api/invite/{a['token']}/contribution/publiques/marches").json()
    assert [m["decp_id"] for m in dp_["marches"]] == ["M-2023-07"]
    v = client.get(f"/api/invite/{a['token']}/contribution").json()["version"]
    r = client.post(f"/api/invite/{a['token']}/contribution/reprendre",
                    json={"publiques": [dp_["marches"][0]["cle"]], "version": v})
    assert r.status_code == 200, r.text
    refs = r.json()["contribution"]["references"]
    assert [(x["intitule"], x["provenance"]) for x in refs] == [("Couverture du gymnase municipal", "decp")]

    # ── 6. L'entretien chantier : une réponse, une question passée, une question posée au partenaire ───────
    r = client.post(f"/api/atelier/{pid}/entretien", headers=h)
    assert r.status_code == 200, r.text
    e = r.json()
    q = {x["code"]: x for x in e["questions"]}
    assert e["message"] == f"{len(q)} questions sur ce chantier" and len(q) <= 8
    # C2 × D : une question se rattache au SOUS-CRITÈRE qu'elle concerne (son poids tel qu'écrit, sa partie) ;
    # les questions se classent par ce poids (du plus lourd au plus léger), jamais par un poids déduit.
    def crit(code):
        c = q[code]["critere"]
        return (c["critere"], c["poids"], q[code]["partie"]["titre"])
    assert crit("site_occupe") == ("Méthodologie d'exécution en site occupé", "40 %",
                                   "Méthodologie d'exécution en site occupé")
    assert crit("chef_chantier") == ("Moyens humains et encadrement", "20 %", "Moyens humains et encadrement")
    env = next(c for c in ("dechets", "environnement") if c in q)        # l'environnement est toujours posé
    assert crit(env) == ("Démarche environnementale", "10 %", "Démarche environnementale")
    poids = [float(x["critere"]["poids"].rstrip(" %")) for x in e["questions"] if x["critere"]]
    assert poids == sorted(poids, reverse=True)
    assert q["site_occupe"]["critere"]["sous"] is True and q["chef_chantier"]["critere"]["sous"] is True
    for x in e["questions"]:                 # chaque question porte sa phrase du dossier, retrouvée mot pour mot
        if x["motif"]:
            assert citation_guard(texte_dce, x["motif"]["phrase"], min_len=4), x["motif"]
    assert q["site_occupe"]["motif"]["phrase"] == "L'école reste en activité pendant toute la durée des travaux."
    ad = q["acces_toiture"]["adresse"]                       # proposée sans IA d'après le lot du partenaire
    assert (ad["type"], ad["invite_id"], ad["nom"]) == ("invitation", a["id"], "Kerbrat Zinguerie")
    rep = "Livraisons entre 9 h et 11 h 30, hors des entrées et des sorties des élèves ; accès par le portail nord."
    r = client.put(f"/api/atelier/{pid}/entretien/{q['site_occupe']['id']}", headers=h,
                   json={"statut": "repondue", "reponse": rep, "lock_version": q["site_occupe"]["lock_version"]})
    assert r.status_code == 200, r.text
    r = client.put(f"/api/atelier/{pid}/entretien/{q['chef_chantier']['id']}", headers=h, json={"statut": "sautee"})
    assert r.status_code == 200, r.text
    # Le chef de file ne répond pas à la place du partenaire.
    r = client.put(f"/api/atelier/{pid}/entretien/{q['acces_toiture']['id']}", headers=h,
                   json={"statut": "repondue", "reponse": "Par l'escalier."})
    assert r.status_code == 403 and r.json()["detail"]["code"] == "question_d_un_partenaire"
    # Le partenaire, dans SON espace (lien + clé), ne voit que SA question, jamais une réponse du chef de file.
    pa = P(client, a["token"])
    vue = pa.get().json()
    assert [x["id"] for x in vue["entretien"]["questions"]] == [q["acces_toiture"]["id"]]
    assert "portail nord" not in json.dumps(vue, ensure_ascii=False)
    r = pa.put(f"/entretien/{q['acces_toiture']['id']}", {"statut": "repondue", "avertissement_lu": True,
                                                         "reponse": "Échafaudage de pied côté cour et monte-matériaux."})
    assert r.status_code == 200, r.text
    e = client.get(f"/api/atelier/{pid}/entretien", headers=h).json()
    assert {(x["ref"], x["libelle"]) for x in e["fiche"]} == {
        (q["site_occupe"]["ref"], "déclaré par vous"), (q["acces_toiture"]["ref"], "déclaré par Kerbrat Zinguerie")}

    # ── 7. Atelier (IA simulée) : les faits cités sont reliés à LEUR source ─────────────────────────────────
    r_site, r_toit = q["site_occupe"]["ref"], q["acces_toiture"]["ref"]
    texte_met = (f"Notre méthodologie d'exécution en site occupé protège les élèves : les livraisons ont lieu entre "
                 f"9 h et 11 h 30 [{r_site}].\n\nKerbrat Zinguerie accède à la toiture par un échafaudage de pied côté "
                 f"cour [{r_toit}].\n\nNous avons été retenus en 2024 pour la réfection de la toiture de l'école "
                 f"Jules-Ferry, un marché de 85 000 € HT [{f_decp}].")
    fake = fake_llm(lambda kw: (texte_met, "end_turn"))
    assert _rediger(client, h, pid, s_met)["status"] == "done"
    envoye = fake.sent()
    assert f"[{r_site}] (déclaré par vous) Entretien chantier, réponse à « " in envoye
    assert f"[{r_toit}] (déclaré par votre partenaire) Kerbrat Zinguerie — entretien chantier" in envoye
    assert f"[{f_decp}] (marché attribué, source publique DECP) Référence : " in envoye
    assert "SOUS-CRITÈRE DE L'ACHETEUR : Méthodologie d'exécution en site occupé" in envoye
    assert "Anne Le Goff" not in envoye                             # aucun nom en clair à l'IA
    vc = _faits(client, h, pid, s_met)
    src = {f["fait"]: f["source"] for f in vc["faits_sources"]}
    assert src["11 h"]["ref"] == r_site and src["11 h"]["libelle"] == "déclaré par vous"
    assert src["échafaudage"]["ref"] == r_toit and src["échafaudage"]["partenaire"] == "Kerbrat Zinguerie"
    assert src["85 000 €"]["ref"] == f_decp and src["85 000 €"]["libelle"] == "marché attribué, source publique DECP"
    assert src["85 000 €"]["date"] == sp_jour()                      # une source publique est datée
    assert vc["controle"]["sans_source"] == 0
    # La question PASSÉE laisse son « [À compléter] » dans la partie concernée, jamais un texte de remplissage.
    fake_llm(lambda kw: ("Un encadrement de proximité suit le chantier au quotidien.", "end_turn"))
    assert _rediger(client, h, pid, s_moy)["status"] == "done"
    st = _etat(client, h, pid)
    t_moy = next(s for s in st["sections"] if s["id"] == s_moy)["version_courante"]["texte"]
    assert t_moy.count("[À compléter : chef de chantier]") == 1
    # « n exigences traitées sur m », avec le nom de celles qui manquent (repérage de mots, sans IA).
    assert st["exigences"]["message"] == "2 exigences traitées sur 5"
    assert st["exigences"]["message_manque"] == (
        "Il manque : « la gestion des déchets du chantier », « un planning par phase, calé sur les vacances "
        "scolaires », « La réduction des trajets des véhicules est à exposer dans l'offre ».")
    assert st["entretien"]["repondues"] == 2 and st["entretien"]["passees"] == 1

    # ── 8. Le règlement écarte l'IA : mode « sans IA » ; plus aucun appel d'IA, le contrôle continue ────────
    fake = fake_llm(lambda kw: (_ for _ in ()).throw(AssertionError("aucun appel d'IA en mode sans IA")))
    r = client.put(f"/api/atelier/{pid}/options", headers=h, json={"sans_ia": True})
    assert r.status_code == 200, r.text
    sia = r.json()["sans_ia"]
    assert sia["actif"] is True and sia["parties_ia"] == 2          # deux parties à réécrire soi-même
    for chemin in (f"/sections/{s_env}/rediger", "/rediger", "/entretien/ia", f"/sections/{s_met}/raccourcir"):
        r = client.post(f"/api/atelier/{pid}{chemin}", headers=h, json={"cible_mots": 30})
        assert r.status_code == 409 and r.json()["detail"]["code"] == "mode_sans_ia", (chemin, r.text)
    e = client.get(f"/api/atelier/{pid}/entretien", headers=h).json()
    assert e["ia"]["permise"] is False                              # l'écran ne propose plus la recherche IA
    assert pa.get().json()["sans_ia"]["actif"] is True             # le partenaire le lit dans son espace
    # Les questions de l'entretien (sans IA) continuent : le chef de chantier est donné, puis tout est réécrit.
    e = client.get(f"/api/atelier/{pid}/entretien", headers=h).json()
    qc = next(x for x in e["questions"] if x["code"] == "chef_chantier")
    r = client.put(f"/api/atelier/{pid}/entretien/{qc['id']}", headers=h,
                   json={"statut": "repondue", "reponse": "Anne Le Goff, conductrice de travaux, présente chaque jour.",
                         "lock_version": qc["lock_version"]})
    assert r.status_code == 200, r.text
    r_chef = qc["ref"]
    # Réécrit SANS reprendre de phrase de l'IA (sinon la partie reste « préparée avec l'IA », décision 9).
    _ecrire(client, h, pid, s_met,
            f"Notre méthodologie d'exécution en site occupé tient en trois règles. L'école reste ouverte : nous livrons "
            f"de 9 h à 11 h 30 seulement [{r_site}].\n\nNotre partenaire Kerbrat Zinguerie monte sur la toiture par un "
            f"échafaudage de pied installé dans la cour [{r_toit}].\n\nLa commune nous a attribué en 2024 les travaux "
            f"de couverture de l'école Jules-Ferry, pour 85 000 € HT [{f_decp}].")
    _ecrire(client, h, pid, s_moy, f"L'encadrement du chantier est confié à Anne Le Goff, conductrice de travaux, "
                                   f"présente chaque jour [{r_chef}].")
    _ecrire(client, h, pid, s_env, "La gestion des déchets du chantier : tri sur place et filières agréées. Notre "
                                   "planning par phase est calé sur les vacances scolaires. La réduction des trajets "
                                   "des véhicules passe par des livraisons groupées.")
    st = _etat(client, h, pid)
    assert st["sans_ia"]["parties_ia"] == 0
    assert st["exigences"]["message"] == "Les 5 exigences repérées sont traitées dans le dossier"
    vc = _faits(client, h, pid, s_met)                              # le contrôle des faits continue, sans IA
    src = {f["fait"]: f["source"] for f in vc["faits_sources"]}
    assert src["85 000 €"]["ref"] == f_decp and src["11 h"]["ref"] == r_site
    assert fake.calls == []

    # ── 9. Questions à l'acheteur : au chef de file seul (Adjugo n'envoie rien) ─────────────────────────────
    # Elles restent lisibles en mode « sans IA » (repérées sans IA), avec leur phrase et leur date limite ; les
    # retirer ou les remettre n'appelle pas l'IA.
    gp = client.get(f"/api/projects/{pid}", headers=h).json()["ai_analysis"]["details"]["grille_notation"]
    qa2 = gp["questions_acheteur"]
    assert [q["code"] for q in qa2["liste"]] == ["limites"] and qa2["date_limite"]["date"] == "2026-12-05"
    assert all(citation_guard(texte_dce, x["phrase"], min_len=4) for q in qa2["liste"] for x in q["phrases"])
    qid = qa2["liste"][0]["id"]
    for choix in ("ecarte", "annule"):
        r = client.post("/api/sourcing/grille/element", headers=h,
                        json={"project_id": pid, "type": "question", "element_id": qid, "choix": choix})
        assert r.status_code == 200, r.text
    assert fake.calls == []
    pub = client.get(f"/api/invite/{a['token']}").json()
    assert "Laquelle s'applique" not in json.dumps(pub, ensure_ascii=False)
    assert "questions_acheteur" not in json.dumps(pub, ensure_ascii=False)

    # ── 10. Export Word : aucune marque interne ─────────────────────────────────────────────────────────────
    # Les parties propres au groupement (tirées de la convention) sont retirées de ce dossier : la convention et
    # le DC1 les portent ; on exporte les trois parties notées.
    for s in st["sections"]:
        if s["id"] not in (s_met, s_moy, s_env):
            assert client.delete(f"/api/atelier/{pid}/sections/{s['id']}", headers=h).status_code == 200
    for s in _etat(client, h, pid)["sections"]:
        r = client.post(f"/api/atelier/{pid}/sections/{s['id']}/valider", headers=h,
                        json={"version_id": s["version_courante"]["id"]})
        assert r.status_code == 200, r.text
    st = _etat(client, h, pid)
    assert st["blocage"]["peut_exporter"] is True, st["blocage"]
    assert st["export"]["ia_utilisee"] is False and st["export"]["mention_ia"] == ""
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx",
                                                                   "etat_hash": st["blocage"]["etat_hash"]})
    assert r.status_code == 200, r.text
    import docx
    doc = docx.Document(io.BytesIO(r.content))
    assert doc.core_properties.subject == ""                        # D9 : aucune mention, aucune IA dans le texte
    plein = "\n".join(p.text for p in doc.paragraphs)
    assert "85 000 € HT" in plein and "Anne Le Goff" in plein and "Kerbrat Zinguerie" in plein
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "1. Valeur technique" in heads and "1.1 Méthodologie d'exécution en site occupé" in heads
    xml = _docx_xml(r.content)
    for interdit in (f"[{r_site}]", f"[{r_toit}]", f"[{f_decp}]", "[R", "[F", "À compléter", "Adjugo", "adjugo",
                     "déclaré par", "source publique", "DECP", "Entretien chantier", "python-docx"):
        assert interdit not in xml, interdit
    assert sources.calls.count("decp") >= 1 and "rge" not in sources.calls


# ═══ La vitrine de la démo : les trois nouveautés ensemble, sans IA ni réseau, rien d'écrit ═══════════════

def test_vitrine_de_la_demo_montre_les_trois_nouveautes_sans_ia_ni_reseau(client, sources, monkeypatch):
    from app.models import DossierPresentation, EntretienQuestion, Preuve, Project, User
    from app.services.demo_seed import DEMO_EMAIL, ensure_demo
    db = _db()
    try:
        ensure_demo(db)
        u = db.query(User).filter(User.email == DEMO_EMAIL).one()
        toiture = next(p for p in db.query(Project).filter(Project.user_id == u.id) if p.name.startswith("Réfection toiture"))
        pid, uid = toiture.id, u.id
        did = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).one().id
    finally:
        db.close()

    def comptes():
        db = _db()
        try:
            return (db.query(EntretienQuestion).filter(EntretienQuestion.dossier_id == did).count(),
                    db.query(Preuve).filter(Preuve.owner_id == uid).count())
        finally:
            db.close()
    avant = comptes()
    appels = []
    monkeypatch.setattr(llm, "messages_create", lambda *a, **k: appels.append(1) or (_ for _ in ()).throw(
        llm.LLMUnavailable("IA coupée")))
    tok = client.post("/api/auth/demo").json()["access_token"]
    hd = {"Authorization": f"Bearer {tok}"}
    st = client.get(f"/api/atelier/{pid}", headers=hd).json()
    # C2 : « 5 exigences traitées sur 7 », la phrase du règlement sur l'IA, le mode « sans IA » proposé.
    assert st["exigences"]["message"] == "5 exigences traitées sur 7"
    assert st["sans_ia"]["actif"] is False and st["sans_ia"]["reglement"]["mentionne"] is True
    # D : l'entretien chantier du marché toiture, avec quelques réponses et une question passée.
    ent = st["entretien"]
    assert ent["prepare"] is True and ent["repondues"] >= 3 and ent["passees"] == 1
    e = client.get(f"/api/atelier/{pid}/entretien", headers=hd).json()
    assert e["ia"]["permise"] is True and e["fiche"]
    # E2 : des références proposées depuis les marchés publics, à partir de données CLAIREMENT fictives.
    dm = client.get("/api/preuves/publiques/marches", headers=hd).json()
    assert dm["exemple_fictif"] is True and all("fictif" in m["affiche"]["objet"].lower() for m in dm["marches"])
    assert all(m["source_libelle"] == "marché attribué, source publique DECP" for m in dm["marches"] if not m["manque"])
    # Toute écriture et tout appel d'IA sont refusés à la démo (lecture guidée).
    for m, url, body in (("post", f"/api/atelier/{pid}/entretien/ia", {}),
                         ("put", f"/api/atelier/{pid}/options", {"sans_ia": True}),
                         ("post", "/api/sourcing/grille/relire", {"project_id": pid}),
                         ("post", "/api/preuves/publiques", {"cles": [dm["marches"][0]["cle"]]})):
        r = getattr(client, m)(url, headers=hd, json=body)
        assert r.status_code == 403 and "lecture guidée" in r.text, (url, r.status_code, r.text)
    assert comptes() == avant and appels == [] and sources.calls == []


# ═══ Les jointures entre chantiers, côté écran ═════════════════════════════════════════════════════════

def test_ecrans_sans_ia_entretien_et_source_publique_datee():
    import os
    import shutil
    import subprocess
    from app.services import atelier as at
    static = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app", "static")
    html = open(os.path.join(static, "app.html"), encoding="utf-8").read()
    app_js = open(os.path.join(static, "app.js"), encoding="utf-8").read()
    rp_js = open(os.path.join(static, "redaction-partenaire.js"), encoding="utf-8").read()
    # Dossier « sans IA » : « Chercher d'autres questions avec Adjugo » n'est pas proposé.
    i = html.index("Chercher d'autres questions avec Adjugo")
    assert "atl.ent.d.prepare && !atlSansIa()" in html[html.rfind("<template", 0, i):i]
    assert "ni question proposée par l'IA" in at.MSG_SANS_IA_ACTIF and "questions sur ce chantier" in at.MSG_SANS_IA_ACTIF
    # Une question rattachée à un sous-critère le dit (« sous-critère : … (10 %) »), sans v-html.
    i = html.index("Partie concernée : {{ q.partie.titre }}")
    meta = html[i:html.index("</div>", i)]
    assert "q.critere.sous ? 'sous-critère' : 'critère'" in meta and "v-html" not in meta
    # Une source publique est datée dans la liste des faits (atelier et espace du partenaire).
    assert '", consultée le "' in app_js[app_js.index("function atlSrcLib"):][:400]
    node = shutil.which("node")
    if not node:
        pytest.skip("node indisponible")
    prog = ("const window = {};\n" + rp_js + "\nconst R = window.AdjugoRedactionPartenaire._pur;\n"
            "console.log(JSON.stringify([R.rpSrcLib({type: 'public', libelle: 'marché attribué, source publique DECP', "
            "date: '2026-10-19'}), R.rpSrcLib({type: 'declare', libelle: 'déclaré par vous', date: '2026-10-19'}), "
            "R.rpSrcLib(null)]));")
    out = subprocess.run([node, "-e", prog], capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stderr
    assert json.loads(out.stdout) == ["marché attribué, source publique DECP, consultée le 19/10/2026",
                                      "déclaré par vous", ""]


def test_entretien_rattache_une_question_a_son_sous_critere_sans_poids_deduit(monkeypatch):
    """C2 × D, sans IA ni base : un critère découpé en sous-critères (deux au moins, comme le plan) donne une ligne
    par sous-critère avec SON poids tel qu'écrit ; un sous-critère sans poids écrit n'a aucun poids (rien n'est
    déduit) et se classe sur le poids ÉCRIT de son critère (contre-relecture de la vague 5 : il passait après un
    critère de 10 %) ; un sous-critère retrouvé l'emporte sur son critère ; la partie concernée est celle du
    sous-critère ; un critère à un seul sous-critère reste une seule ligne."""
    from types import SimpleNamespace as NS
    from app.services import entretien_chantier as ec
    from app.services import grille_notation as gn

    def sc(i, nom, poids, statut=gn.RETROUVE):
        return {"id": f"s{i}", "critere": nom, "poids": poids, "statut": statut, "phrase": f"{nom} : {poids}"}
    grille = {"criteres": [
        {"id": "c1", "critere": "Valeur technique", "poids_retenu": "60 %", "poids_source": "reglement",
         "dans_le_plan": True, "phrase": "Valeur technique : 60 %",
         "sous_criteres": [sc(1, "Organisation du chantier", "40 %"),
                           sc(2, "Démarche environnementale", "20 %", gn.A_VERIFIER)]},
        {"id": "c2", "critere": "Délai d'exécution", "poids_retenu": "10 %", "poids_source": "reglement",
         "dans_le_plan": True, "phrase": "Délai : 10 %", "sous_criteres": [sc(3, "Planning détaillé", "10 %")]},
        {"id": "c3", "critere": "Prix", "poids_retenu": "30 %", "poids_source": "reglement", "dans_le_plan": True,
         "phrase": "Prix : 30 %", "sous_criteres": []}]}
    projet = NS(ai_analysis={"details": {"grille_notation": grille}})
    lg = ec.lignes_grille(projet)
    assert [(x["critere"], x["poids"], x["sous"], x["_w"]) for x in lg] == [
        ("Valeur technique", "60 %", "", 60.0), ("Organisation du chantier", "40 %", "Organisation du chantier", 40.0),
        ("Démarche environnementale", "", "Démarche environnementale", 60.0), ("Délai d'exécution", "10 %", "", 10.0)]
    assert lg[2]["parent_poids"] == "60 %" and lg[1]["parent_poids"] == ""      # dit, jamais attribué au sous-critère
    assert ec._critere_de(["organis"], lg)["critere"] == "Organisation du chantier"
    assert ec._critere_de(["environ"], lg)["critere"] == "Démarche environnementale"
    assert ec._critere_de(["technique"], lg)["critere"] == "Valeur technique"     # le critère, faute de mieux
    assert ec._critere_de(["inconnu"], lg)["critere"] == "Valeur technique"      # repli : le seul critère technique
    # La partie concernée : celle du sous-critère, même quand aucun début de mot ne la désigne.
    parts = [NS(id=1, intitule="Organisation du chantier", critere="Valeur technique", critere_id="c1", question="",
                parents=[], generee=False),
             NS(id=2, intitule="Démarche environnementale", critere="Valeur technique", critere_id="c1", question="",
                parents=[], generee=False)]
    monkeypatch.setattr(ec.at, "est_generee", lambda s: False)
    assert ec.partie_concernee(["zzz"], parts, ec._critere_de(["environ"], lg)).id == 2
    assert ec.partie_concernee(["zzz"], parts, ec._critere_de(["inconnu"], lg)).id == 1
