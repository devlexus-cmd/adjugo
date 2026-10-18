"""
Rédaction, vague 1 — socle 0 « réparations et promesses honnêtes » (étude rédaction § 6.2).

Chaque test échoue sur le code d'avant (commit f91b9c1). Depuis la vague 2, le dossier de
présentation est rédigé dans l'ATELIER (le moteur qui écrivait tout le mémoire d'un seul appel est
retiré) : les mêmes garanties y sont vérifiées.
 1. l'équipe arrive au dossier de présentation, ses NOMS ne partent jamais à l'IA (faux client IA
    qui reçoit le prompt), y compris quand un document du savoir-faire les contient ;
 2. la liste des pièces lit les pièces exigées là où l'analyse les range (details), sans casser
    les anciennes analyses (racine) ;
 4. export propre : ni « [S1] », ni message technique ; « [À compléter] » signalé ; tableaux ;
 5. archivage Word + PDF sous un nom que la liste des pièces reconnaît (celui du règlement) ;
 6. texte coupé par l'IA : dit à l'utilisateur, jamais livré comme complet ;
 7. promesses : plus d'« Aucune invention », plus de faux « sourcé », aide de la liste corrigée ;
 8. compte-rendu de groupement : « document interne au chef de file » ;
 9. un seul vocabulaire à l'écran (« dossier de présentation ») ;
10. limite globale d'appels IA simultanés (sans interblocage) + coût IA d'un dossier journalisé.
Aucun appel réseau : le client IA est remplacé par un faux.
"""
import base64
import io
import json
import logging
import os
import threading
import time
import zipfile

import pytest

from app.services import llm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEAM = [{"nom": "Yann Tanguy", "fonction": "Chef de chantier", "qualifications": "BTS Bâtiment, 15 ans",
         "references": "12 toitures scolaires"},
        {"nom": "Hélène Le Gall", "fonction": "Conductrice de travaux", "qualifications": "Ingénieure",
         "references": ""}]
NOMS = ("Tanguy", "Yann", "Hélène", "Helene", "Le Gall", "LE GALL")


def _static(name):
    with open(os.path.join(ROOT, "app", "static", name), encoding="utf-8") as f:
        return f.read()


# ── Faux fournisseur IA : enregistre chaque requête, répond selon la consigne ──────────────────
class _Block:
    def __init__(self, text):
        self.type = "text"
        self.text = text


class _Usage:
    def __init__(self, i=100, o=50):
        self.input_tokens = i
        self.output_tokens = o


class _Resp:
    def __init__(self, text, stop="end_turn", usage=None):
        self.content = [_Block(text)]
        self.stop_reason = stop
        self.usage = usage or _Usage()


class FakeLLM:
    """Remplace llm.client() : `answer(kwargs) -> (texte, motif d'arrêt)`."""

    def __init__(self, answer):
        self.calls = []
        self._answer = answer
        self.messages = self

    def create(self, **kwargs):
        self.calls.append(kwargs)
        text, stop = self._answer(kwargs)
        return _Resp(text, stop)

    def sent(self) -> str:
        out = []
        for k in self.calls:
            out.append(llm._flatten_content(k.get("system")))
            for m in k.get("messages") or []:
                out.append(llm._flatten_content(m.get("content")))
        return "\n".join(out)


@pytest.fixture
def fake_llm(monkeypatch):
    def install(answer):
        fake = FakeLLM(answer)
        monkeypatch.setattr(llm, "client", lambda: fake)
        monkeypatch.setattr(llm, "_use_mistral", lambda: False)
        llm._CB.update(fails=0, open_until=0.0)
        llm._TENANT_CB.clear()
        return fake
    return install


def _memoire_answer(text, stop="end_turn"):
    """Réponse du rédacteur ; la fusion des apports (si appelée) reçoit un texte neutre."""
    def ans(kwargs):
        if "architecte d'un groupement" in llm._flatten_content(kwargs.get("system")):
            return "Synthèse des moyens du groupement.", "end_turn"
        return text, stop
    return ans


def _db():
    from app.core.database import SessionLocal
    return SessionLocal()


def _user_id(email):
    from app.models import User
    db = _db()
    try:
        return db.query(User).filter(User.email == email).first().id
    finally:
        db.close()


def _set_analysis(pid, analysis):
    from app.models import Project
    db = _db()
    try:
        p = db.get(Project, pid)
        p.ai_analysis = analysis
        db.commit()
    finally:
        db.close()


def _project(client, h, name="Réfection de la toiture de l'école"):
    r = client.post("/api/projects/", headers=h, json={"name": name, "client": "Ville de Lorient"})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _company(client, h, team=TEAM):
    r = client.put("/api/company/", headers=h, json={
        "name": "Breizh Couverture", "siret": "12345678900011", "representant_legal": "Eliot Viegas",
        "address": "1 rue du Port", "city": "Lorient", "postal_code": "56100", "effectif": 12,
        "team": team})
    assert r.status_code == 200, r.text


ANALYSIS = {"summary": "Toiture", "details": {
    "intitule_marche": "Réfection de la toiture de l'école", "acheteur": "Ville de Lorient",
    "pieces_requises": [{"nom": "Attestation de visite obligatoire", "disponible": False},
                        {"nom": "Mémoire justificatif (10 pages)", "disponible": False}]},
    "dce_excerpt": "Le candidat remettra un mémoire justificatif des dispositions qu'il propose."}


def _zip(dossier):
    return zipfile.ZipFile(io.BytesIO(base64.b64decode(dossier["zip_b64"])))


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    return " ".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)


def _docx(data: bytes):
    import docx
    return docx.Document(io.BytesIO(data))


# ═══ 1. L'équipe arrive au dossier, ses noms ne partent pas à l'IA ═══════════════════════════════

def test_company_dict_reprend_l_equipe_et_le_groupement_ne_la_recopie_pas():
    from app.models import Company
    from app.services.groupement_verdict import lead_member
    from app.services.profile import company_dict
    c = Company(name="Breizh Couverture", team=TEAM)
    d = company_dict(c)
    assert d["team"] == TEAM
    # Le membre « chef de file » d'un groupement (stocké, montré aux partenaires) n'a pas l'équipe.
    assert "team" not in lead_member(d, {})


def test_pseudonymiseur_masque_et_remet_les_noms_sans_ia():
    from app.services.pii import Pseudonymizer
    pz = Pseudonymizer(TEAM + [{"nom": "Yann Martin", "fonction": "Chef de chantier"}])
    txt = ("Yann Tanguy et TANGUY Yann ; M. Tanguy ; Mme Le Gall ; Helene LE GALL ; CV_Yann_Tanguy.pdf ; "
           "Yann seul ; Saint-Martin ; un petit martin ; Martin")
    m = pz.mask(txt)
    for nom in ("Tanguy", "TANGUY", "Gall", "GALL", "Helene", "Yann", "; Martin"):
        assert nom not in m, (nom, m)
    assert "[Chef de chantier 1]" in m and "[Conductrice de travaux 1]" in m
    assert "Saint-Martin" in m and "un petit martin" in m          # pas de faux positif
    back = pz.unmask(m)
    # Nom complet → nom complet ; morceau seul → rendu tel quel (jamais attribué à quelqu'un).
    assert back.startswith("Yann Tanguy et Yann Tanguy ; M. Tanguy ; Mme Le Gall ; Hélène Le Gall")
    assert "CV_Yann Tanguy.pdf" in back and "Yann seul" in back and back.endswith("; Martin")


def test_pseudonymiseur_ni_mauvaise_attribution_ni_mot_courant_abime():
    from app.services.pii import Pseudonymizer
    pz = Pseudonymizer([{"nom": "Jean Petit", "fonction": "Couvreur"}, {"nom": "Paul Martin", "fonction": "Couvreur"}])
    txt = "Petit chantier en site occupé. M. Petit encadre. Contact mairie : M. Martin, service technique."
    m = pz.mask(txt)
    assert m.startswith("Petit chantier") and "M. Petit" not in m and "Martin" not in m
    assert pz.unmask(m) == txt                     # « M. Martin » de la mairie reste « M. Martin »


def _atelier(client, h, pid):
    r = client.post(f"/api/atelier/{pid}", headers=h, json={})
    assert r.status_code == 201, r.text
    return r.json()


def _attendre_tache(client, h, pid, jid):
    for _ in range(200):
        j = client.get(f"/api/atelier/{pid}/taches/{jid}", headers=h).json()
        if j["status"] in ("done", "error"):
            return j
        time.sleep(0.05)
    raise AssertionError("tâche jamais terminée")


def _rediger_tout(client, h, pid):
    r = client.post(f"/api/atelier/{pid}/rediger", headers=h, json={})
    assert r.status_code == 202, r.text
    return _attendre_tache(client, h, pid, r.json()["tache"]["id"])


def _ecrire_tout(client, h, pid, texte, valider=False):
    st = client.get(f"/api/atelier/{pid}", headers=h).json()
    for sec in st["sections"]:
        base = (sec.get("version_courante") or {}).get("id")
        r = client.put(f"/api/atelier/{pid}/sections/{sec['id']}/texte", headers=h,
                       json={"texte": texte, "base_version_id": base})
        assert r.status_code == 200, r.text
        if valider:
            r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/valider", headers=h,
                            json={"version_id": r.json()["version_courante"]["id"]})
            assert r.status_code == 200, r.text
    return client.get(f"/api/atelier/{pid}", headers=h).json()


def test_dossier_route_equipe_presentee_et_aucun_nom_a_l_ia(client, auth, fake_llm):
    """Bout en bout : profil avec équipe + CV dans le savoir-faire → l'atelier rédige sans qu'aucun
    nom ne parte à l'IA ; le dossier livré (Word) les contient, remis sans IA à la place des
    étiquettes."""
    h = auth["headers"]
    _company(client, h)
    r = client.post("/api/knowledge/text", headers=h, json={
        "name": "CV Yann Tanguy", "kind": "reference",
        "text": "Yann Tanguy, chef de chantier couverture : 12 toitures scolaires, sécurité, méthodologie, "
                "références de chantiers en site occupé. Hélène Le Gall conduit les travaux."})
    assert r.status_code == 200, r.text
    pid = _project(client, h)
    _set_analysis(pid, ANALYSIS)
    fake = fake_llm(_memoire_answer(
        "[Chef de chantier 1] encadre le chantier [S1]. [Conductrice de travaux 1] assure la conduite des travaux."))
    _atelier(client, h, pid)
    assert _rediger_tout(client, h, pid)["status"] == "done"
    sent = fake.sent()
    assert fake.calls, "les parties doivent avoir été demandées à l'IA (simulée)"
    for nom in NOMS:
        assert nom not in sent, f"« {nom} » est parti à l'IA"
    assert "[Chef de chantier 1]" in sent and "12 toitures scolaires" in sent   # l'équipe est bien là
    assert "recopie ces étiquettes" in sent.lower() or "Recopie ces étiquettes" in sent
    r = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid, "cotraitants": []})
    assert r.status_code == 200, r.text
    d = r.json()["dossier"]
    assert "Yann Tanguy encadre le chantier" in d["memoire_markdown"]
    assert "[Chef de chantier 1]" not in d["memoire_markdown"]
    zf = _zip(d)
    word = " ".join(p.text for p in _docx(zf.read("Mémoire justificatif.docx")).paragraphs)
    assert "Yann Tanguy encadre le chantier." in word and "Hélène Le Gall assure" in word


def test_entreprise_seule_pas_de_groupement_invente():
    from app.services import atelier_plan as ap
    seule = [s["intitule"] for s in ap.build_plan({"details": {}}, en_groupement=False)["sections"]]
    assert seule[0] == "Présentation de l'entreprise" and "Présentation du groupement" not in seule
    grp = [s["intitule"] for s in ap.build_plan({"details": {}}, en_groupement=True)["sections"]]
    assert grp[0] == "Présentation du groupement"


def test_note_d_information_aux_salaries_dans_le_profil():
    html = _static("app.html")
    assert "Note d'information à remettre à vos salariés" in html
    assert "ne lui est pas envoyé" in html and "cnil.fr" in html
    assert "copyRefText('teamNote'" in html and "copyRefText(ref, ok)" in _static("app.js")


# ═══ 2. Pièces exigées lues au bon niveau ═════════════════════════════════════════════════════════

def test_pieces_exigees_lues_dans_details_et_ancien_format(client, auth):
    h = auth["headers"]
    pid = _project(client, h, "Marché pièces")
    _set_analysis(pid, {"details": {"intitule_marche": "Marché pièces",
                                    "pieces_requises": [{"nom": "Attestation de visite obligatoire"},
                                                        "Certificat de capacité amiante", None, 3]}})
    names = [i["name"] for i in client.get(f"/api/checklist/{pid}", headers=h).json()["checklist"]]
    assert "Attestation de visite obligatoire" in names and "Certificat de capacité amiante" in names
    # Ancienne analyse : pièces à la racine → toujours lues.
    _set_analysis(pid, {"details": {}, "pieces_requises": [{"nom": "Plan de prévention signé"}]})
    names = [i["name"] for i in client.get(f"/api/checklist/{pid}", headers=h).json()["checklist"]]
    assert "Plan de prévention signé" in names


# ═══ 4 et 6. Export propre, texte coupé ═══════════════════════════════════════════════════════════

MD_SALE = """```markdown
# Mémoire technique

## Méthodologie
Nous intervenons en site occupé [S1]. Équipe de 6 compagnons [S1, S2].
- Point **important** [S3]
- [À compléter par l'entreprise] : références

| Phase | Durée |
|---|---|
| Dépose | 5 j |
| Pose | [À compléter] |
```"""


def _ctx_rendu(texte):
    from app.services import atelier_export as ax
    return {"titre": "Mémoire technique", "marche": "X", "acheteur": "", "entreprise": "Solo SARL",
            "date": "Octobre 2026", "sections": [{"id": 1, "numero": "1", "intitule": "Méthodologie", "parents": [],
                                                  "texte": texte}],
            "police_rendu": ax.police(None), "ia": True}


def test_export_propre_sans_marque_interne_avec_tableau_et_alerte():
    from app.services import atelier_export as ax
    from app.services.memoire_export import placeholders
    ctx = _ctx_rendu(MD_SALE)
    assert len(placeholders(MD_SALE)) == 2                       # compté : l'export de l'atelier se bloque
    pdf = _pdf_text(ax.to_pdf(ctx)[0])
    assert "[S1" not in pdf and "S2]" not in pdf and "site occupé." in pdf and "```" not in pdf
    doc = _docx(ax.to_docx(ctx))
    assert doc.tables and [c.text for c in doc.tables[0].rows[1].cells] == ["Dépose", "5 j"]
    word = " ".join(p.text for p in doc.paragraphs)
    assert "[S" not in word and "[À compléter par l'entreprise] : références" in word and "```" not in word
    # Le passage à compléter est surligné dans Word.
    hl = [r for p in doc.paragraphs for r in p.runs if r.font.highlight_color is not None]
    assert any("À compléter" in r.text for r in hl)


def test_panne_ia_aucun_message_technique_dans_le_livrable(client, auth, fake_llm, monkeypatch):
    h = auth["headers"]
    pid = _project(client, h, "Marché panne")
    _set_analysis(pid, {"details": {"intitule_marche": "Marché panne"}})

    def boom(**kw):
        raise RuntimeError("Error code: 529 overloaded_error — request_id=req_123")
    fake = fake_llm(_memoire_answer("x"))
    monkeypatch.setattr(fake, "create", boom)
    _atelier(client, h, pid)
    j = _rediger_tout(client, h, pid)
    assert j["status"] == "error" and "529" not in j["error"] and "req_123" not in j["error"]
    out = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid}).json()["dossier"]
    zf = _zip(out)
    assert not any("moire" in n.lower() for n in zf.namelist()), zf.namelist()
    everything = " ".join(_pdf_text(zf.read(n)) for n in zf.namelist() if n.endswith(".pdf"))
    assert "indisponible :" not in everything and "529" not in everything and "req_123" not in everything
    assert out["memoire_fichiers"] == []
    assert any(w.startswith("BLOQUANT") and "Dossier de présentation" in w for w in out["warnings"])
    assert not any("529" in w or "req_123" in w for w in out["warnings"])


def test_texte_coupe_signale_et_jamais_livre_comme_complet(client, auth, fake_llm):
    from app.services.memoire_export import TRUNCATED_NOTE
    h = auth["headers"]
    pid = _project(client, h, "Marché coupé")
    _set_analysis(pid, {"details": {"intitule_marche": "Marché coupé"}})
    fake_llm(_memoire_answer("Nous commençons par la dépose des", "max_tokens"))
    st = _atelier(client, h, pid)
    sid = st["sections"][0]["id"]
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
    assert _attendre_tache(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    st = client.get(f"/api/atelier/{pid}", headers=h).json()
    v = next(s for s in st["sections"] if s["id"] == sid)["version_courante"]
    assert v["coupe"] is True and v["texte"].rstrip().endswith(TRUNCATED_NOTE)
    # Jamais livré comme complet : la mention compte comme un passage « À compléter », l'export bloque.
    assert any(x["code"] == "a_completer" and any(y["section_id"] == sid for y in x["sections"])
               for x in st["blocage"]["raisons"])


def test_section_et_reponse_coupees_marquees(fake_llm):
    from app.services.agents import memoire, questionnaire
    fake_llm(lambda kw: ("Texte interrompu [S1", "max_tokens"))
    chunks = [{"doc_name": "a.pdf", "chunk_id": 1, "text": "x" * 50}]
    s = memoire.write_section({"titre": "Méthodo"}, chunks)
    assert s["coupe"] is True and "interrompue" in s["content"]
    a = questionnaire._answer_from_chunks("Q ?", chunks)
    assert a["coupe"] is True and "interrompue" in a["answer"]


# ═══ 5. Archivage Word + PDF, reconnu par la liste des pièces ═════════════════════════════════════

def test_archivage_word_pdf_reconnu_par_la_liste_des_pieces(client, auth):
    h = auth["headers"]
    _company(client, h, team=[])
    pid = _project(client, h, "Marché archivage")
    _set_analysis(pid, ANALYSIS)
    # Ancienne archive brute (jamais reconnue) : elle est retirée à la génération suivante.
    from app.services.dossier import save_to_dossier
    db = _db()
    try:
        save_to_dossier(db, _user_id(auth["email"]), pid, "Mémoire technique", "memoire_technique.md", b"# vieux [S1]")
    finally:
        db.close()
    ck = client.get(f"/api/checklist/{pid}", headers=h).json()
    item = next(i for i in ck["checklist"] if i["id"] == "memoire_technique")
    assert item["status"] == "manquant"            # l'ancienne archive .md ne compte pas
    _atelier(client, h, pid)
    _ecrire_tout(client, h, pid, "Texte complet et relu.", valider=True)
    r = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid, "cotraitants": []})
    assert r.status_code == 200, r.text
    d = r.json()["dossier"]
    assert "_memoire_files" not in d and d["memoire_nom"] == "Mémoire justificatif"
    assert not any(w.startswith("BLOQUANT — Dossier de présentation") for w in d["warnings"])
    folders = client.get(f"/api/projects/{pid}/documents", headers=h).json()["folders"]
    names = {doc["name"] for f in folders if f["folder"] == "Mémoire technique" for doc in f["documents"]}
    assert names == {"Mémoire justificatif.docx", "Mémoire justificatif.pdf"}, names
    ck = client.get(f"/api/checklist/{pid}", headers=h).json()
    item = next(i for i in ck["checklist"] if i["id"] == "memoire_technique")
    assert item["status"] == "ok" and item["document_name"].startswith("Mémoire justificatif")
    assert item["name"] == "Dossier de présentation (mémoire justificatif)"
    # Pas de doublon : le « Mémoire justificatif » exigé par le règlement EST cette pièce.
    assert not any(i["from_ia"] and "moire" in i["name"].lower() for i in ck["checklist"])
    # Aperçu du dossier commun : mêmes noms que le ZIP.
    prev = client.get(f"/api/projects/{pid}/dossier-preview", headers=h).json()
    assert {"Mémoire justificatif.docx", "Mémoire justificatif.pdf"} <= {f["name"] for f in prev["files"]}


def test_dossier_avec_trous_ou_coupe_pas_disponible_dans_la_liste(client, auth):
    h = auth["headers"]
    _company(client, h, team=[])
    pid = _project(client, h, "Marché à compléter")
    _set_analysis(pid, {"details": {"intitule_marche": "Marché à compléter"}})
    _atelier(client, h, pid)
    st = _ecrire_tout(client, h, pid, "Texte relu.")
    sec = st["sections"][0]
    r = client.put(f"/api/atelier/{pid}/sections/{sec['id']}/texte", headers=h,
                   json={"texte": "Références : [À compléter par l'entreprise]",
                         "base_version_id": sec["version_courante"]["id"]})
    assert r.status_code == 200
    # Export bloqué : rien n'est joint tant qu'une dérogation n'est pas enregistrée.
    d = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid}).json()["dossier"]
    assert d["memoire_fichiers"] == []
    gate = client.get(f"/api/atelier/{pid}", headers=h).json()["blocage"]
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "completerai_dans_word", "texte": "Références à joindre",
                          "etat_hash": gate["etat_hash"]})
    assert r.status_code == 200, r.text
    assert client.post("/api/sourcing/documents", headers=h,
                       json={"project_id": pid, "cotraitants": []}).status_code == 200
    ck = client.get(f"/api/checklist/{pid}", headers=h).json()
    item = next(i for i in ck["checklist"] if i["id"] == "memoire_technique")
    assert item["status"] == "a_completer", item
    assert "1 passage « À compléter »" in item["context_note"]
    assert ck["stats"]["a_completer"] == 1
    assert 'a_completer: "À terminer"' in _static("app.js")


def test_nom_du_reglement_seulement_s_il_figure_dans_le_dossier():
    from app.services.memoire_export import memoire_label
    a = {"details": {"pieces_requises": [{"nom": "Mémoire de synthèse"}]}, "dce_excerpt": "rien de tel"}
    assert memoire_label(a) == "Mémoire technique"               # proposé par l'IA, absent du texte
    a = {"pieces_requises": ["Note méthodologique (5 p.)"], "dce_excerpt": "une NOTE  méthodologique"}
    assert memoire_label(a) == "Note méthodologique"
    assert memoire_label(None) == "Mémoire technique"


# ═══ 7. Promesses ramenées au réel ════════════════════════════════════════════════════════════════

def test_promesses_retirees_de_l_ecran():
    html, js = _static("app.html"), _static("app.js")
    assert "Aucune invention" not in html
    assert "(sourcé)" not in html and "Basé sur votre savoir-faire" not in html
    assert "kb_used" not in html and "100% sourcés" not in js and "IA sourcée" not in html
    from app.routers.checklist import PIECES_STANDARD
    refs = next(p for p in PIECES_STANDARD if p["id"] == "references")
    assert "alimentent le mémoire" not in refs["help"] and "DC2" not in refs["help"]


def test_seuls_les_extraits_reellement_cites_sont_montres(fake_llm):
    from app.services.agents import memoire, questionnaire
    chunks = [{"doc_name": f"d{i}.pdf", "chunk_id": i, "text": "savoir-faire " * 5} for i in (1, 2, 3)]
    fake_llm(lambda kw: ("Nous sommes certifiés ISO 14001.", "end_turn"))   # aucune citation
    s = memoire.write_section({"titre": "Qualité"}, chunks)
    assert s["sources"] == [] and "grounded" not in s and s["extraits_fournis"] == 3
    a = questionnaire._answer_from_chunks("Certifications ?", chunks)
    assert a["sources"] == []
    fake_llm(lambda kw: ("Démarche qualité [S2]. Hors borne [S9].", "end_turn"))
    s = memoire.write_section({"titre": "Qualité"}, chunks)
    assert [x["ref"] for x in s["sources"]] == ["S2"]
    a = questionnaire._answer_from_chunks("Qualité ?", chunks)
    assert [x["ref"] for x in a["sources"]] == ["S2"]


def test_moteur_savoir_faire_noms_masques_dans_les_fils(client, auth, fake_llm):
    """Rédaction de plusieurs parties (fils parallèles) : la portée « noms masqués » suit chaque fil."""
    h = auth["headers"]
    _company(client, h)
    client.post("/api/knowledge/text", headers=h, json={
        "name": "CV Yann Tanguy", "kind": "reference",
        "text": "Yann Tanguy dirige les chantiers de couverture : sécurité, méthodologie, encadrement."})
    pid = _project(client, h, "Marché fils")
    _set_analysis(pid, {"details": {"intitule_marche": "Marché fils"}})
    fake = fake_llm(lambda kw: ("[Chef de chantier 1] encadre [S1].", "end_turn"))
    _atelier(client, h, pid)
    assert _rediger_tout(client, h, pid)["status"] == "done"
    assert len(fake.calls) >= 4
    for nom in NOMS:
        assert nom not in fake.sent(), f"« {nom} » est parti à l'IA"
    st = client.get(f"/api/atelier/{pid}", headers=h).json()
    assert all("Yann Tanguy encadre" in s["version_courante"]["texte"] for s in st["sections"])


def test_compte_rendu_document_interne_au_chef_de_file():
    from app.services.consortium_report import generate_consortium_report_pdf
    data = {"readiness": {"pct": 50, "submitted": 1, "invited": 1, "missing": []},
            "partners": [{"company": "Alpha", "status": "submitted", "completeness": 50, "pieces_count": 0}],
            "lots": [], "state": {}}
    txt = " ".join(_pdf_text(generate_consortium_report_pdf("AO", "Chef", data, [])).split())
    assert "Document interne au chef de file" in txt
    html = _static("app.html")
    assert "Compte-rendu interne (PDF)" in html and "Document interne au chef de file" in html


# ═══ 9. Un seul vocabulaire ═══════════════════════════════════════════════════════════════════════

def test_vocabulaire_dossier_de_presentation_a_l_ecran():
    html, inv, js = _static("app.html"), _static("invite.html"), _static("app.js")
    assert "Paragraphe de mémoire technique" not in inv and "Paragraphe pour le dossier de présentation" in inv
    assert "parts soumises dans le mémoire" not in html
    # Rédaction, vague 2 : le dossier se rédige dans l'atelier du marché, qui s'appelle ainsi à l'écran.
    assert "Mémoire généré" not in js and 'atelier: "Dossier de présentation"' in js
    assert 'folderLabel(f) { return f === "Mémoire technique" ? "Dossier de présentation"' in js
    assert "{{ g.folder }}" not in html
    # Bulle (i) « mémoire technique » à côté du dossier de présentation livré (page du marché ; l'étape 4
    # de la recherche, qui en portait une seconde, a disparu : décision 4 de l'étude organisation).
    assert html.count("« mémoire technique »</b> demandé par l'acheteur") >= 1


# ═══ 10. Limite d'appels IA simultanés + coût d'un dossier ════════════════════════════════════════

def _slow_client(monkeypatch, delay=0.15, nested=False):
    state = {"now": 0, "peak": 0}
    lock = threading.Lock()

    class C:
        class messages:
            @staticmethod
            def create(**kw):
                with lock:
                    state["now"] += 1
                    state["peak"] = max(state["peak"], state["now"])
                try:
                    time.sleep(delay)
                    return _Resp("ok")
                finally:
                    with lock:
                        state["now"] -= 1
    monkeypatch.setattr(llm, "client", lambda: C)
    monkeypatch.setattr(llm, "_use_mistral", lambda: False)
    llm._CB.update(fails=0, open_until=0.0)
    return state


def test_limite_globale_d_appels_simultanes(monkeypatch):
    monkeypatch.setattr(llm, "_SLOTS", threading.BoundedSemaphore(2))
    monkeypatch.setattr(llm, "_MAX_CONCURRENT", 2)
    state = _slow_client(monkeypatch)
    errors = []

    def call():
        try:
            llm.messages_create(model=llm.MODEL_FAST, max_tokens=5, messages=[{"role": "user", "content": "x"}])
        except Exception as e:      # pragma: no cover - échec = test rouge
            errors.append(e)
    ts = [threading.Thread(target=call) for _ in range(8)]
    for t in ts:
        t.start()
    for t in ts:
        t.join(10)
    assert not errors and state["peak"] <= 2, state
    assert llm.concurrency_snapshot()["max"] == 2


def test_place_attendue_trop_longtemps_echec_propre_sans_ouvrir_le_disjoncteur(monkeypatch):
    sem = threading.BoundedSemaphore(1)
    monkeypatch.setattr(llm, "_SLOTS", sem)
    monkeypatch.setattr(llm, "_SLOT_WAIT", 0.05)
    _slow_client(monkeypatch)
    assert sem.acquire(timeout=1)          # une place prise ailleurs, jamais rendue pendant le test
    try:
        with pytest.raises(llm.LLMUnavailable):
            llm.messages_create(model=llm.MODEL, max_tokens=5, messages=[])
        assert llm._CB["fails"] == 0       # attendre une place n'est pas une panne du fournisseur
    finally:
        sem.release()


def test_appel_imbrique_n_attend_pas_sa_propre_place(monkeypatch):
    """Réentrance : un appel passé PENDANT un appel (même fil) ne redemande pas de place → pas
    d'interblocage même avec une seule place."""
    monkeypatch.setattr(llm, "_SLOTS", threading.BoundedSemaphore(1))
    monkeypatch.setattr(llm, "_SLOT_WAIT", 2)
    monkeypatch.setattr(llm, "_use_mistral", lambda: False)
    llm._CB.update(fails=0, open_until=0.0)
    depth = {"n": 0}

    class C:
        class messages:
            @staticmethod
            def create(**kw):
                depth["n"] += 1
                if depth["n"] == 1:
                    llm.messages_create(model=llm.MODEL, max_tokens=5, messages=[])
                return _Resp("ok")
    monkeypatch.setattr(llm, "client", lambda: C)
    done = []
    t = threading.Thread(target=lambda: done.append(llm.messages_create(model=llm.MODEL, max_tokens=5, messages=[])))
    t.start()
    t.join(5)
    assert done and depth["n"] == 2
    assert llm._SLOTS.acquire(blocking=False)    # la place a bien été rendue
    llm._SLOTS.release()


def test_cout_ia_d_un_dossier_complet_journalise(monkeypatch, caplog):
    from concurrent.futures import ThreadPoolExecutor
    usages = iter([(1000, 500), (2000, 1000)])

    class C:
        class messages:
            @staticmethod
            def create(**kw):
                i, o = next(usages)
                return _Resp("ok", usage=_Usage(i, o))
    monkeypatch.setattr(llm, "client", lambda: C)
    monkeypatch.setattr(llm, "_use_mistral", lambda: False)
    monkeypatch.setattr(llm, "_resolve_provider", lambda: "anthropic")
    llm._CB.update(fails=0, open_until=0.0)
    caplog.set_level(logging.INFO, logger="adjugo")
    with llm.cost_meter("dossier de test", 42) as m:
        llm.messages_create(model=llm.MODEL, max_tokens=5, messages=[])
        with ThreadPoolExecutor(1) as ex:       # un appel dans un fil compte aussi (map_in_context)
            llm.map_in_context(ex, lambda _x: llm.messages_create(model=llm.MODEL_FAST, max_tokens=5,
                                                                  messages=[]), [0])
    s = m.summary()
    assert s["appels"] == 2 and s["tokens_entree"] == 3000 and s["tokens_sortie"] == 1500
    # Sonnet 3 $/15 $ et Haiku 1 $/5 $ par million : 0,003 + 0,0075 + 0,002 + 0,005.
    assert s["cout_estime_usd"] == pytest.approx(0.0175)
    assert any("Coût IA « dossier de test »" in r.getMessage() and "0.0175" in r.getMessage()
               for r in caplog.records)


def test_atelier_journalise_son_cout_et_le_cumule(client, auth, fake_llm, caplog):
    h = auth["headers"]
    _company(client, h, team=[])
    pid = _project(client, h, "Marché coût")
    _set_analysis(pid, {"details": {"intitule_marche": "Marché coût"}})
    fake_llm(_memoire_answer("Texte."))
    caplog.set_level(logging.INFO, logger="adjugo")
    st = _atelier(client, h, pid)
    r = client.post(f"/api/atelier/{pid}/sections/{st['sections'][0]['id']}/rediger", headers=h, json={})
    assert _attendre_tache(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    assert any("Coût IA « atelier : rédiger »" in rec.getMessage() and "1 appel" in rec.getMessage()
               for rec in caplog.records)
    from app.models import DossierPresentation
    db = _db()
    try:
        d = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).first()
        assert d.cout_ia["appels"] == 1 and d.cout_ia["tokens_entree"] == 100
    finally:
        db.close()
    assert "cout_ia" not in json.dumps(client.get(f"/api/atelier/{pid}", headers=h).json())   # jamais montré
