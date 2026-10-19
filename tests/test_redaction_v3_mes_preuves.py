"""
Rédaction, vague 3 — E1 « Mes preuves » côté serveur (étude § 6.7 E1, doctrine § 7, décision 8a).

 • REPRISE de l'existant (profil de l'entreprise : collaborateurs, références, qualifications,
   effectif et chiffres d'affaires ; profil du réseau ; justificatifs du coffre-fort) : aperçu,
   l'utilisateur coche, rien n'est perdu (les données d'origine ne changent pas), aucun doublon
   (même après modification ou retrait de la fiche reprise, deux origines qui disent la même chose
   n'en font qu'une, double envoi sans effet), contenu relu côté serveur (jamais pris du client) ;
 • ATELIER : une référence ou un CV choisi s'insère TEL QUEL dans une partie (nouvelle version,
   aucune IA), seulement s'il est utilisable ; rien d'interne dans le texte ni dans le Word ;
 • AUCUN CV VERS L'IA : « Raccourcir » refuse une partie qui contient un CV de « Mes preuves »
   AVANT tout appel et sans débit ; « Rédiger » n'envoie ni le nom ni le CV d'une fiche « Équipe » ;
 • isolation par organisation (404, rien d'un autre compte dans la reprise ni dans l'atelier),
   membres d'une organisation : même reprise ; démo en lecture seule ; RGPD (export, suppression,
   effacement du compte) ; jamais « vérifié ».
Aucun appel réseau ; IA simulée (et un faux client qui échoue s'il est appelé là où il ne doit pas).
"""
import io
import json
import uuid
from datetime import date, timedelta

import pytest

from app.services import llm
from app.services import preuves_atelier as pa
from tests.test_redaction_socle0 import _db, fake_llm  # noqa: F401
from tests.test_redaction_v2_atelier import _attendre, _creer, _docx_xml, _ecrire, _etat, _projet

TEAM = [{"nom": "Gwendal Kerhervé", "fonction": "Chef de chantier", "qualifications": "Brevet de compagnon couvreur",
         "references": "Chapelle Sainte-Barbe du Faouët"},
        {"nom": "Soizic Le Goff", "fonction": "", "qualifications": "CAP", "references": ""}]
REFS = [{"intitule": "Réfection de la toiture de l'école Jules Ferry", "client": "Mairie de Quimper",
         "montant": 85000, "annee": 2024},
        {"intitule": "Couverture du gymnase", "client": "Ville de Brest", "montant": "pas un nombre", "annee": 1700}]
QUALS = ["Qualibat 3194", {"name": "RGE Qualibat", "detail": "Couverture", "expiration": "2027-06-30"}]


def _h(tok):
    return {"Authorization": f"Bearer {tok}"}


def _register(client, prefix="e1"):
    email = f"{prefix}_{uuid.uuid4().hex[:8]}@test.fr"
    r = client.post("/api/auth/register", json={"email": email, "password": "motdepasse123",
                                                "full_name": "Dirigeant", "company_name": "Breizh Couverture"})
    assert r.status_code == 201, r.text
    from app.models import User
    db = _db()
    try:
        u = db.query(User).filter(User.email == email).first()
        u.email_verified = True
        db.commit()
        uid = u.id
    finally:
        db.close()
    return email, _h(r.json()["access_token"]), uid


def _ok(r, code=200):
    assert r.status_code == code, r.text
    return r.json()


def _company(client, h, team=TEAM, refs=REFS, quals=QUALS):
    _ok(client.put("/api/company/", headers=h, json={
        "name": "Breizh Couverture", "siret": "12345678900011", "city": "Quimper", "effectif": 14,
        "ca_n1": 1250000, "team": team, "references": refs, "qualifications": quals}))


def _reseau(uid, refs):
    from app.models import PartnerProfile
    db = _db()
    try:
        db.add(PartnerProfile(public_id=uuid.uuid4().hex, user_id=uid, references_publiques=refs))
        db.commit()
    finally:
        db.close()


def _coffre(uid, name="Certificat_Qualibat-3194.pdf", expiration=None, category="qualifications", deleted=False):
    from datetime import datetime, timezone
    from app.models import DocCategory, Document
    db = _db()
    try:
        d = Document(user_id=uid, name=name, category=DocCategory(category), file_key=f"k/{uuid.uuid4().hex}",
                     expiration_date=expiration, deleted_at=datetime.now(timezone.utc) if deleted else None)
        db.add(d)
        db.commit()
        return d.id
    finally:
        db.close()


def _preuve(client, h, **body):
    body.setdefault("type_source", "declare")
    return _ok(client.post("/api/preuves", headers=h, json=body), 201)


def _sans_ia(monkeypatch):
    """Faux client IA qui échoue au premier appel : la fonction testée ne doit JAMAIS l'appeler."""
    calls = []

    class _Boom:
        def __init__(self):
            self.messages = self

        def create(self, **kw):
            calls.append(kw)
            raise AssertionError("appel à l'IA interdit ici")
    monkeypatch.setattr(llm, "client", lambda: _Boom())
    monkeypatch.setattr(llm, "_use_mistral", lambda: False)
    return calls


# ═══ 1. Reprise de l'existant : sans doublon ni perte ═══════════════════════════════════

def test_reprise_profil_reseau_et_coffre_fort_apercu_puis_coche(client, monkeypatch):
    calls = _sans_ia(monkeypatch)
    email, h, uid = _register(client)
    _company(client, h)
    # Le profil du réseau reprend une référence du profil (autre casse, sans accents) + une nouvelle.
    _reseau(uid, [{"intitule": "REFECTION DE LA TOITURE DE L'ECOLE JULES FERRY", "client": "Mairie de Quimper",
                   "annee": 2024},
                  {"intitule": "Charpente de la salle des fêtes", "client": "Commune de Rosporden", "annee": 2023}])
    fin = date.today() + timedelta(days=200)
    doc_id = _coffre(uid, expiration=fin)
    _coffre(uid, name="Kbis.pdf", category="administratif")              # pas une qualification
    _coffre(uid, name="Ancien_certificat.pdf", deleted=True)             # à la corbeille
    ap = _ok(client.get("/api/preuves/reprise", headers=h))
    els = ap["elements"]
    par = {(e["nature"], e["valeur"].get("objet") or e["valeur"].get("nom") or e["valeur"].get("libelle")): e
           for e in els}
    assert ap["deja"] == 0
    # Équipe : le CV minimal (qualifications, chantiers) ; sans fonction, la fiche ne peut pas être reprise.
    g = par[("salarie", "Gwendal Kerhervé")]
    assert g["valeur"] == {"nom": "Gwendal Kerhervé", "fonction": "Chef de chantier",
                           "qualifications": "Brevet de compagnon couvreur",
                           "chantiers": "Chapelle Sainte-Barbe du Faouët"}
    assert g["provenance_libelle"] == "Mon entreprise (Identité)" and g["source_libelle"] == "déclaré par vous"
    s = par[("salarie", "Soizic Le Goff")]
    assert s["manque"] == "Indiquez sa fonction dans « Mon entreprise » pour la reprendre."
    # Références : la même référence du profil et du réseau n'en fait qu'UNE ; les valeurs illisibles
    # (montant, année hors bornes) sont écartées, la fiche est gardée.
    refs = [e for e in els if e["nature"] == "reference"]
    assert [e["valeur"]["objet"] for e in refs] == ["Réfection de la toiture de l'école Jules Ferry",
                                                    "Couverture du gymnase", "Charpente de la salle des fêtes"]
    assert refs[0]["valeur"] == {"objet": "Réfection de la toiture de l'école Jules Ferry",
                                 "acheteur": "Mairie de Quimper", "annee": 2024, "montant_ht": 85000}
    assert refs[1]["valeur"] == {"objet": "Couverture du gymnase", "acheteur": "Ville de Brest"}
    assert refs[2]["provenance"] == "reseau"
    # Qualifications : texte simple, objet avec date de fin, justificatif du coffre-fort.
    assert par[("qualification", "RGE Qualibat")]["date_fin_validite"] == "2027-06-30"
    j = par[("qualification", "Certificat Qualibat 3194")]
    assert j["type_source"] == "documents" and j["source_libelle"] == "justificatif déposé"
    assert j["date_fin_validite"] == fin.isoformat() and j["provenance_libelle"] == "Mes documents"
    assert not any("Kbis" in e["texte"] or "Ancien" in e["texte"] for e in els)
    # Chiffres clés du profil.
    assert par[("chiffre_cle", "Effectif")]["valeur"] == {"libelle": "Effectif", "valeur": 14, "unite": "personnes"}
    assert ("chiffre_cle", "Chiffre d'affaires N-1") in par
    assert ap["a_reprendre"] == len(els) - 1
    # L'aperçu ne crée rien.
    assert _ok(client.get("/api/preuves", headers=h))["preuves"] == []
    # L'utilisateur coche : tout sauf la qualification « Qualibat 3194 » (+ une clé inconnue, ignorée).
    cles = [e["cle"] for e in els if e["valeur"].get("nom") != "Qualibat 3194"] + ["0" * 24]
    r = _ok(client.post("/api/preuves/reprise", headers=h, json={"cles": cles}))
    assert len(r["creees"]) == len(els) - 2 and r["deja"] == 0
    assert r["ignorees"] == [{"texte": "Équipe : Soizic Le Goff ; qualifications : CAP",
                              "message": "Indiquez sa fonction dans « Mon entreprise » pour la reprendre."}]
    ps = _ok(client.get("/api/preuves", headers=h))["preuves"]
    assert {p["statut"] for p in ps} == {"acceptee"} and {p["origine"] for p in ps} == {"profil"}
    assert all(p["utilisable"] for p in ps)
    pj = next(p for p in ps if p["type_source"] == "documents")
    assert pj["source"] == {"document_id": doc_id, "document_nom": "Certificat_Qualibat-3194.pdf",
                            "lien_justificatif": f"/api/documents/{doc_id}/download"}
    # Ce qui n'a pas été coché reste proposé ; le reste est « déjà dans vos preuves ».
    ap2 = _ok(client.get("/api/preuves/reprise", headers=h))
    assert [e["valeur"]["nom"] for e in ap2["elements"]] == ["Soizic Le Goff", "Qualibat 3194"]
    assert ap2["deja"] == len(els) - 2
    assert calls == []                                                   # aucune IA


def test_reprise_sans_perte_ni_doublon_meme_apres_modification(client):
    email, h, uid = _register(client)
    _company(client, h)
    _coffre(uid)
    avant = _ok(client.get("/api/company/", headers=h))
    els = _ok(client.get("/api/preuves/reprise", headers=h))["elements"]
    cles = [e["cle"] for e in els]
    n = len(_ok(client.post("/api/preuves/reprise", headers=h, json={"cles": cles}))["creees"])
    # Aucune perte : le profil n'a pas changé.
    apres = _ok(client.get("/api/company/", headers=h))
    for k in ("team", "references", "qualifications", "effectif", "ca_n1"):
        assert apres[k] == avant[k], k
    # Double envoi : rien de plus.
    r = _ok(client.post("/api/preuves/reprise", headers=h, json={"cles": cles}))
    assert r["creees"] == [] and r["deja"] == n
    ps = _ok(client.get("/api/preuves", headers=h))["preuves"]
    assert len(ps) == n
    # Fiche modifiée (autre acheteur) puis retirée : jamais reprise une seconde fois.
    ref = next(p for p in ps if p["nature"] == "reference" and p["valeur"]["objet"].startswith("Réfection"))
    _ok(client.patch(f"/api/preuves/{ref['id']}", headers=h,
                     json={"valeur": dict(ref["valeur"], acheteur="Quimper Bretagne Occidentale"),
                           "version": ref["version"]}))
    q = next(p for p in ps if p["nature"] == "qualification" and p["type_source"] == "documents")
    _ok(client.patch(f"/api/preuves/{q['id']}", headers=h, json={"valeur": {"nom": "Qualibat 3194 (couverture)"}}))
    eff = next(p for p in ps if p["nature"] == "chiffre_cle" and p["valeur"]["libelle"] == "Effectif")
    _ok(client.post(f"/api/preuves/{eff['id']}/retirer", headers=h, json={}))
    ap = _ok(client.get("/api/preuves/reprise", headers=h))
    assert [e["valeur"].get("nom") for e in ap["elements"]] == ["Soizic Le Goff"]    # seule celle sans fonction
    # Un nouveau collaborateur dans le profil : seul lui est proposé.
    team = TEAM + [{"nom": "Erwan Le Bris", "fonction": "Couvreur", "qualifications": "", "references": ""}]
    _company(client, h, team=team)
    ap = _ok(client.get("/api/preuves/reprise", headers=h))
    assert [e["valeur"].get("nom") for e in ap["elements"]] == ["Soizic Le Goff", "Erwan Le Bris"]


def test_reprise_isolee_par_organisation_partagee_entre_membres(client):
    from app.models import PlanType, User
    email, h, uid = _register(client)
    _company(client, h)
    _coffre(uid)
    cles = [e["cle"] for e in _ok(client.get("/api/preuves/reprise", headers=h))["elements"]]
    # Un autre compte : rien de l'organisation A, et les clés de A ne créent rien chez lui.
    _e2, h2, _u2 = _register(client)
    assert _ok(client.get("/api/preuves/reprise", headers=h2))["elements"] == []
    r = _ok(client.post("/api/preuves/reprise", headers=h2, json={"cles": cles}))
    assert r["creees"] == [] and r["deja"] == 0
    assert _ok(client.get("/api/preuves", headers=h2))["preuves"] == []
    # Un membre de l'organisation A voit la même reprise et reprend pour l'organisation.
    db = _db()
    try:
        db.query(User).filter(User.id == uid).first().plan = PlanType.pro
        db.commit()
    finally:
        db.close()
    inv = _ok(client.post("/api/org/invite", headers=h, json={"email": f"membre_{uuid.uuid4().hex[:6]}@test.fr"}), 201)
    login = _ok(client.post("/api/auth/login", json={"email": inv["email"], "password": inv["temp_password"]}))
    hm = _h(login["access_token"])
    assert [e["cle"] for e in _ok(client.get("/api/preuves/reprise", headers=hm))["elements"]] == cles
    els = _ok(client.get("/api/preuves/reprise", headers=hm))["elements"]
    deux = [e["cle"] for e in els if not e["manque"]][:2]
    _ok(client.post("/api/preuves/reprise", headers=hm, json={"cles": deux}))
    assert len(_ok(client.get("/api/preuves", headers=h))["preuves"]) == 2
    assert len(_ok(client.get("/api/preuves/reprise", headers=h))["elements"]) == len(cles) - 2


# ═══ 2. L'atelier insère une fiche choisie, telle quelle, sans IA ═══════════════════════

def _fiches_de_test(client, h, autre_pid=None):
    ref = _preuve(client, h, nature="reference", valeur={
        "objet": "Réfection de la toiture de l'école Jules Ferry", "acheteur": "Mairie de Quimper", "annee": 2024,
        "montant_ht": 85000, "role": "seul"})
    cv = _preuve(client, h, nature="salarie", valeur={
        "nom": "Gwendal Kerhervé", "fonction": "Chef de chantier", "experience_annees": 15,
        "diplomes": "Brevet de compagnon couvreur zingueur", "qualifications": "CACES R486",
        "chantiers": "Chapelle Sainte-Barbe du Faouët"})
    hier = (date.today() - timedelta(days=1)).isoformat()
    vieille = _preuve(client, h, nature="qualification", valeur={"nom": "Qualibat 3194"}, date_fin_validite=hier)
    prop = _preuve(client, h, nature="materiel", valeur={"designation": "Nacelle 18 m"}, statut="proposee")
    est = None
    if autre_pid:
        est = _preuve(client, h, nature="chiffre_cle", valeur={"libelle": "Délai", "valeur": 6, "unite": "semaines"},
                      type_source="estimation", source={"project_id": autre_pid})
    return ref, cv, vieille, prop, est


def test_atelier_insere_une_reference_et_un_cv_tels_quels_sans_ia(client, monkeypatch):
    email, h, uid = _register(client)
    pid = _projet(client, h)
    autre = _projet(client, h, with_dce=False, name="Autre marché")
    ref, cv, vieille, prop, est = _fiches_de_test(client, h, autre_pid=autre)
    st = _creer(client, h, pid)
    calls = _sans_ia(monkeypatch)
    # Ce que l'atelier peut insérer : fiches utilisables pour CE marché, avec le texte inséré.
    f = _ok(client.get(f"/api/atelier/{pid}/preuves", headers=h))
    assert [x["id"] for x in f["fiches"]] == [ref["id"], cv["id"]]
    assert [(x["id"], x["message"]) for x in f["ecartees"]] == [
        (vieille["id"], "Date de fin de validité dépassée : Adjugo ne s'en sert plus.")]
    fr = f["fiches"][0]
    assert fr["titre"] == "Réfection de la toiture de l'école Jules Ferry" and fr["source_libelle"] == "déclaré par vous"
    assert fr["bloc"] == ("**Réfection de la toiture de l'école Jules Ferry**\n- Acheteur : Mairie de Quimper\n"
                          "- Année : 2024\n- Montant : 85 000 € HT\n- Réalisé seul")
    assert f["fiches"][1]["bloc"] == ("**Gwendal Kerhervé, Chef de chantier**\n- Expérience : 15 ans\n"
                                      "- Diplômes : Brevet de compagnon couvreur zingueur\n"
                                      "- Qualifications et habilitations : CACES R486\n"
                                      "- Chantiers : Chapelle Sainte-Barbe du Faouët")
    # L'estimation vaut pour son marché.
    assert [x["id"] for x in _ok(client.get(f"/api/atelier/{autre}/preuves", headers=h))["fiches"]] == [
        ref["id"], cv["id"], est["id"]]
    sec = st["sections"][0]
    url = f"/api/atelier/{pid}/sections/{sec['id']}/inserer-preuve"
    s1 = _ok(client.post(url, headers=h, json={"preuve_id": ref["id"], "base_version_id": None}))
    v1 = s1["version_courante"]
    assert v1["texte"] == fr["bloc"] and v1["origine"] == "utilisateur" and v1["ia"] is False
    assert v1["faits_a_relire"] == 0 and v1["a_completer"] == 0
    s2 = _ok(client.post(url, headers=h, json={"preuve_id": cv["id"], "base_version_id": v1["id"]}))
    v2 = s2["version_courante"]
    assert v2["texte"] == fr["bloc"] + "\n\n" + f["fiches"][1]["bloc"] and v2["numero"] == v1["numero"] + 1
    # Rien d'interne dans le texte : ni numéro « F », ni libellé de source.
    for interdit in (ref["ref"], cv["ref"], "déclaré par vous", "trouvé dans vos documents", "[F"):
        assert interdit not in v2["texte"]
    # Version de départ périmée, fiche non utilisable, fiche d'un autre compte.
    r = client.post(url, headers=h, json={"preuve_id": ref["id"], "base_version_id": v1["id"]})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "texte_modifie_entre_temps"
    for p, raison in ((vieille, "expiree"), (prop, "proposee"), (est, "autre_marche")):
        r = client.post(url, headers=h, json={"preuve_id": p["id"], "base_version_id": v2["id"]})
        assert r.status_code == 409, r.text
        assert r.json()["detail"]["code"] == "preuve_non_utilisable" and r.json()["detail"]["raison"] == raison
    _e2, h2, _u2 = _register(client)
    autre_ref = _preuve(client, h2, nature="reference", valeur={"objet": "Gymnase de Brest"})
    r = client.post(url, headers=h, json={"preuve_id": autre_ref["id"], "base_version_id": v2["id"]})
    assert r.status_code == 404 and r.json()["detail"]["code"] == "preuve_introuvable"
    # Un autre compte ne voit ni ne touche l'atelier du marché.
    assert client.get(f"/api/atelier/{pid}/preuves", headers=h2).status_code == 404
    r = client.post(url, headers=h2, json={"preuve_id": autre_ref["id"], "base_version_id": v2["id"]})
    assert r.status_code == 404
    assert calls == []                                                   # aucune IA


def test_cv_insere_tel_quel_dans_le_word_sans_rien_d_interne(client, monkeypatch):
    import docx
    email, h, uid = _register(client)
    pid = _projet(client, h)
    _ref, cv, *_ = _fiches_de_test(client, h)
    st = _creer(client, h, pid)
    _sans_ia(monkeypatch)
    for s in st["sections"]:
        s2 = _ecrire(client, h, pid, s, "Nous intervenons par zones, en site occupé.")
        if s["id"] == st["sections"][1]["id"]:
            s2 = _ok(client.post(f"/api/atelier/{pid}/sections/{s['id']}/inserer-preuve", headers=h,
                                 json={"preuve_id": cv["id"], "base_version_id": s2["version_courante"]["id"]}))
        _ok(client.post(f"/api/atelier/{pid}/sections/{s['id']}/valider", headers=h,
                        json={"version_id": s2["version_courante"]["id"]}))
    st = _etat(client, h, pid)
    assert st["blocage"]["peut_exporter"] is True and st["export"]["ia_utilisee"] is False
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx"})
    assert r.status_code == 200, r.text
    full = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    for attendu in ("Gwendal Kerhervé, Chef de chantier", "Expérience : 15 ans",
                    "Diplômes : Brevet de compagnon couvreur zingueur", "Chantiers : Chapelle Sainte-Barbe du Faouët"):
        assert attendu in full, attendu
    xml = _docx_xml(r.content)
    for interdit in ("[F", "déclaré par vous", "Mes preuves", "**"):
        assert interdit not in xml, interdit


# ═══ 3. Aucun CV vers l'IA ═══════════════════════════════════════════════════════════════

def _usage(uid):
    from app.models import User
    db = _db()
    try:
        return db.get(User, uid).analyses_used_this_month or 0
    finally:
        db.close()


def test_raccourcir_refuse_une_partie_qui_contient_un_cv_avant_tout_appel(client, fake_llm):  # noqa: F811
    email, h, uid = _register(client)
    pid = _projet(client, h)
    _ref, cv, *_ = _fiches_de_test(client, h)
    st = _creer(client, h, pid)
    sec = st["sections"][0]
    long = " ".join(["Notre encadrement suit le chantier au quotidien."] * 120)
    sec = _ecrire(client, h, pid, sec, long)
    sec = _ok(client.post(f"/api/atelier/{pid}/sections/{sec['id']}/inserer-preuve", headers=h,
                          json={"preuve_id": cv["id"], "base_version_id": sec["version_courante"]["id"]}))
    fake = fake_llm(lambda kw: ("Texte raccourci.", "end_turn"))
    avant = _usage(uid)
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/raccourcir", headers=h, json={"cible_mots": 300})
    assert r.status_code == 409, r.text
    d = r.json()["detail"]
    assert d["code"] == "cv_dans_la_partie" and "Gwendal Kerhervé" in d["message"]
    assert "n'envoie jamais un CV à l'IA" in d["message"]
    assert fake.calls == [] and _usage(uid) == avant
    assert _etat(client, h, pid)["sections"][0]["redaction_en_cours"] is None
    # Le CV retouché à la main (une ligne gardée) est toujours repéré.
    retouche = long + "\n\n**Gwendal Kerhervé**\n- Diplômes : Brevet de compagnon couvreur zingueur"
    sec = _ecrire(client, h, pid, sec, retouche)
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/raccourcir", headers=h, json={"cible_mots": 300})
    assert r.status_code == 409 and fake.calls == []
    # Sans le CV, « Raccourcir » part à l'IA comme avant.
    sec = _ecrire(client, h, pid, sec, long + "\n\nGwendal Kerhervé encadre le chantier.")
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/raccourcir", headers=h, json={"cible_mots": 300})
    assert r.status_code == 202, r.text
    assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    assert len(fake.calls) == 1
    envoye = fake.sent()
    assert "Kerhervé" not in envoye and "Gwendal" not in envoye        # le nom reste masqué


def test_rediger_n_envoie_ni_le_nom_ni_le_cv_des_fiches_equipe(client, fake_llm):  # noqa: F811
    email, h, uid = _register(client)
    pid = _projet(client, h)
    _fiches_de_test(client, h)
    st = _creer(client, h, pid)
    fake = fake_llm(lambda kw: ("Notre équipe intervient en site occupé.", "end_turn"))
    sec = st["sections"][0]
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/rediger", headers=h, json={})
    assert r.status_code == 202, r.text
    assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    envoye = fake.sent()
    assert envoye
    for interdit in ("Kerhervé", "Kerherve", "Gwendal", "Brevet de compagnon couvreur zingueur", "CACES R486",
                     "Chapelle Sainte-Barbe"):
        assert interdit not in envoye, interdit


def test_reperage_d_un_cv_sur_le_texte_seul():
    from app.services import preuves as pv
    v = {"nom": "Yann Tanguy", "fonction": "Couvreur", "experience_annees": 1, "chantiers": "- École A\n- École B"}
    assert pa.lignes_cv(v) == ["Expérience : 1 an", "Chantiers : École A ; École B"]
    b = pa.bloc(pv.SALARIE, v)
    assert b == "**Yann Tanguy, Couvreur**\n- Expérience : 1 an\n- Chantiers : École A ; École B"
    # Repérage aux accents, à la casse, aux espaces et au gras près.
    assert pa._plat("Chantiers : école a ;  ÉCOLE B") in pa._plat("Texte.\n- **Chantiers : École A ; École B**")


# ═══ 4. Démo, RGPD, vocabulaire ══════════════════════════════════════════════════════════

def test_demo_ne_peut_ni_reprendre_ni_inserer(client):
    from app.core.demo_guard import blocked_action
    assert blocked_action("POST", "/api/preuves/reprise") == "renseigner vos preuves"
    assert blocked_action("POST", "/api/atelier/3/sections/4/inserer-preuve") == "préparer votre dossier de présentation"
    assert blocked_action("GET", "/api/preuves/reprise") is None
    assert blocked_action("GET", "/api/atelier/3/preuves") is None
    r = client.post("/api/auth/demo")
    assert r.status_code == 200, r.text
    hd = _h(r.json()["access_token"])
    assert client.get("/api/preuves/reprise", headers=hd).status_code == 200
    r = client.post("/api/preuves/reprise", headers=hd, json={"cles": []})
    assert r.status_code == 403 and "renseigner vos preuves" in r.json()["detail"]


def test_rgpd_les_fiches_reprises_s_exportent_et_s_effacent(client):
    from app.models import Preuve
    email, h, uid = _register(client)
    _company(client, h)
    cles = [e["cle"] for e in _ok(client.get("/api/preuves/reprise", headers=h))["elements"]]
    _ok(client.post("/api/preuves/reprise", headers=h, json={"cles": cles}))
    data = _ok(client.get("/api/rgpd/export", headers=h))
    g = next(e for e in data["mes_preuves"] if e["nature"] == "salarie")
    assert g["valeur"]["chantiers"] == "Chapelle Sainte-Barbe du Faouët" and g["origine"] == "profil"
    # Supprimer la fiche efface son CV (le profil, lui, reste tel que l'utilisateur l'a saisi).
    pid = next(p["id"] for p in _ok(client.get("/api/preuves", headers=h))["preuves"] if p["nature"] == "salarie")
    assert client.delete(f"/api/preuves/{pid}", headers=h).status_code == 204
    data = _ok(client.get("/api/rgpd/export", headers=h))
    assert "Chapelle Sainte-Barbe" not in json.dumps([e for e in data["mes_preuves"]], ensure_ascii=False)
    db = _db()
    try:
        row = db.get(Preuve, pid)
        assert row.valeur is None and row.statut == "supprimee"
    finally:
        db.close()
    r = client.post("/api/rgpd/delete-account", headers=h, json={"password": "motdepasse123", "confirm": "SUPPRIMER"})
    assert r.status_code == 200, r.text
    db = _db()
    try:
        assert db.query(Preuve).filter(Preuve.owner_id == uid).count() == 0
    finally:
        db.close()


def test_jamais_verifie_dans_la_reprise_ni_dans_l_atelier(client):
    email, h, uid = _register(client)
    _company(client, h)
    _coffre(uid)
    pid = _projet(client, h, with_dce=False)
    _fiches_de_test(client, h)
    txt = json.dumps([client.get("/api/preuves/reprise", headers=h).json(),
                      client.get(f"/api/atelier/{pid}/preuves", headers=h).json()], ensure_ascii=False).lower()
    assert "vérifié" not in txt and "verifie" not in txt
    import inspect
    from app.routers import preuves_ecrans
    from app.services import preuves_reprise
    for mod in (pa, preuves_reprise, preuves_ecrans):
        src = inspect.getsource(mod).lower()
        assert "vérifié" not in src and "verifie" not in src, mod.__name__


@pytest.mark.parametrize("nature,valeur,attendu", [
    ("qualification", {"nom": "Qualibat 3194", "organisme": "Qualibat", "numero": "12345", "detail": "Couverture"},
     "**Qualibat 3194**\n- Délivrée par : Qualibat\n- Numéro : 12345\n- Valable jusqu'au 31/12/2027\n\nCouverture"),
    ("materiel", {"designation": "Nacelle 18 m", "quantite": 2, "caracteristiques": "Ciseaux", "possession": "location"},
     "**Nacelle 18 m**\n- Quantité : 2\n- Caractéristiques : Ciseaux\n- En location"),
    ("chiffre_cle", {"libelle": "Effectif", "valeur": 14, "unite": "personnes", "annee": 2025},
     "Effectif : 14 personnes (2025)"),
    ("reference", {"objet": "Toiture **école**", "description": "Ligne 1\n## Ligne 2"},
     "**Toiture école**\n\nLigne 1\n## Ligne 2"),
])
def test_texte_insere_par_nature(nature, valeur, attendu):
    assert pa.bloc(nature, valeur, "2027-12-31") == attendu
