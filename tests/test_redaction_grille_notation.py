"""
Rédaction, vague 1 — chantier C : « Comment vous serez noté » (étude rédaction §6.3, C0 + C1 +
C2 minimal) et la grille enfin affichée au chef de file (socle 0, point 3).

Rejoue la sonde de l'étude (§2, constat 3) : un cadre de réponse .docx rangé derrière un gros
CCTP disparaissait de la lecture, et avec CCTP.pdf + RC.pdf + Cadre_reponse.docx c'était le
règlement lui-même qui sautait à 120 000 caractères. Ici, le règlement et le cadre sont
TOUJOURS lus, en premier et en entier.

DCE synthétiques réalistes générés dans les tests (RC en PDF texte, cadre .docx numéroté,
CCTP volumineux, ZIP). Aucun réseau, aucune IA réelle (le client IA est simulé).
"""
import io
import json
import re
import zipfile

import pytest

from app.services import dce_pieces as dp
from app.services.analysis import citation_guard, extract_dce_text, select_dce_content
from app.services import grille_notation as gn


# ── Fabriques ────────────────────────────────────────────────────────────────────────

def _pdf(paragraphs) -> bytes:
    """PDF texte mis en page par reportlab (retours à la ligne au milieu des phrases)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    st = getSampleStyleSheet()["Normal"]
    doc.build([Paragraph(t, st) for t in paragraphs])
    return buf.getvalue()


def _numbering(doc, num_id: str, lvl0_text: str):
    """Ajoute une liste Word à numérotation AUTOMATIQUE (absente de paragraph.text)."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    root = doc.part.numbering_part.element
    root.append(parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{num_id}">'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="{lvl0_text}"/></w:lvl>'
        f'<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/></w:lvl>'
        f'</w:abstractNum>'))
    root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{num_id}"/></w:num>'))


def _numbered(doc, text: str, num_id: str, ilvl: int = 0):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    p = doc.add_paragraph(text)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'))
    return p


def _cadre_docx() -> bytes:
    """Cadre de réponse technique réaliste : questions numérotées par Word, tableau, limites."""
    import docx
    d = docx.Document()
    d.add_heading("CADRE DE REPONSE TECHNIQUE", level=1)
    d.add_paragraph("Ce cadre de réponse doit obligatoirement être utilisé par le candidat, sans en modifier la structure.")
    _numbering(d, "77", "Question %1 :")
    _numbered(d, "Décrivez votre organisation de chantier (2 pages maximum).", "77")
    _numbered(d, "Effectif affecté au chantier.", "77", ilvl=1)
    _numbered(d, "Présentez les moyens matériels mis à disposition.", "77")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Thème"
    t.cell(0, 1).text = "Réponse du candidat"
    t.cell(1, 0).text = "Gestion des déchets de chantier"
    _numbered(d, "Détaillez votre démarche environnementale.", "77")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


RC_PARAS = [
    "REGLEMENT DE LA CONSULTATION",
    "Commune de Plougastel-Daoulas - Réfection de la couverture du groupe scolaire.",
    "Article 8 - Critères de jugement des offres : Valeur technique 60 % appréciée au vu du mémoire technique ; Prix 40 %.",
    "Sous-critere 1 : methodologie 30 points ; Sous-critere 2 : moyens 20 points ; Sous-critere 3 : environnement 10 points.",
    "Article 9 - Le mémoire technique ne devra pas excéder 10 pages, police Arial 11.",
    "La visite des lieux est obligatoire ; une attestation de visite sera remise au candidat.",
]


def _big_cctp(n: int = 900) -> bytes:
    return _pdf(["Article %d. Les travaux de couverture comprennent la depose des ardoises existantes et la mise "
                 "en oeuvre d'un ecran de sous-toiture conforme au DTU 40.11 sur l'ensemble des versants." % i
                 for i in range(1, n)])


def _zip(members: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buf.getvalue()


@pytest.fixture(scope="module")
def cctp():
    return _big_cctp()


@pytest.fixture(scope="module")
def rc():
    return _pdf(RC_PARAS)


@pytest.fixture(scope="module")
def cadre():
    return _cadre_docx()


# ── C0 : pièces reconnues sans IA ────────────────────────────────────────────────────

@pytest.mark.parametrize("name,expected", [
    ("DCE/01_RC.pdf", dp.REGLEMENT), ("Reglement_de_consultation_2026.pdf", dp.REGLEMENT),
    ("04_Cadre_de_reponse_technique.docx", dp.CADRE), ("CRT_lot2.docx", dp.CADRE),
    ("Cadre_memoire_technique.docx", dp.CADRE), ("02_CCAP.pdf", dp.CCAP),
    ("03_CCTP_Lot2_Couverture.pdf", dp.CCTP), ("DPGF_lot1.csv", dp.BORDEREAU), ("BPU.pdf", dp.BORDEREAU),
    ("Acte_d_engagement.pdf", dp.ACTE_ENGAGEMENT), ("AE_lot3.pdf", dp.ACTE_ENGAGEMENT),
    ("plans_facades.pdf", dp.AUTRE),
])
def test_piece_reconnue_par_son_nom(name, expected):
    assert dp.classify_piece(name, "") == expected


def test_piece_reconnue_par_ses_premieres_lignes():
    # Nom muet : le titre en tête du document parle.
    assert dp.classify_piece("Consultation_2026-15.pdf",
                             "MAIRIE DE QUIMPER\nREGLEMENT DE LA CONSULTATION\nObjet : ...") == dp.REGLEMENT
    assert dp.classify_piece("pieces.pdf", "ACTE D'ENGAGEMENT\nLot 1") == dp.ACTE_ENGAGEMENT
    # Une abréviation cède devant un titre contraire (« RC » = responsabilité civile ici).
    assert dp.classify_piece("RC_annexe.pdf",
                             "CAHIER DES CLAUSES ADMINISTRATIVES PARTICULIERES\nLe présent CCAP...") == dp.CCAP
    # Le titre qui arrive EN PREMIER gagne : un CCTP qui cite le règlement reste un CCTP.
    assert dp.classify_piece("doc.pdf", "CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n"
                                        "Conformément au règlement de la consultation...") == dp.CCTP
    assert dp.classify_piece("annexe.pdf", "Liste des plans du bâtiment") == dp.AUTRE


def test_cadre_docx_lu_avec_ses_numeros_de_questions_dans_l_ordre(cadre):
    t = extract_dce_text("Cadre_reponse.docx", cadre)
    assert "Question 1 : Décrivez votre organisation de chantier" in t
    assert "a) Effectif affecté au chantier." in t
    assert "Question 2 : Présentez les moyens matériels" in t
    assert "Question 3 : Détaillez votre démarche environnementale." in t
    # Le tableau est lu à sa place (entre la question 2 et la question 3), pas rejeté en fin.
    assert t.index("Question 2") < t.index("Gestion des déchets de chantier") < t.index("Question 3")


# ── C0 : la sonde de l'étude, rejouée ────────────────────────────────────────────────

def _sonde_zip(rc, cadre, cctp):
    return _zip({
        "DCE/01_RC.pdf": rc,
        "DCE/02_CCAP.pdf": _pdf(["CCAP article %d penalites de retard 1/1000." % i for i in range(1, 200)]),
        "DCE/03_CCTP_Lot2_Couverture.pdf": cctp,
        "DCE/04_Cadre_de_reponse_technique.docx": cadre,
    })


@pytest.mark.parametrize("max_chars", [40000, 120000])
def test_sonde_cadre_derriere_un_gros_cctp_toujours_lu(rc, cadre, cctp, max_chars):
    t, meta = extract_dce_text("dce.zip", _sonde_zip(rc, cadre, cctp), max_chars=max_chars, with_pieces=True)
    assert len(t) <= max_chars
    assert "Critères de jugement des offres" in t and "Sous-critere 3" in t, "règlement absent"
    assert "CADRE DE REPONSE TECHNIQUE" in t and "Question 3 : Détaillez" in t, "cadre de réponse absent"
    # Ordre de lecture : règlement, cadre, puis les autres pièces.
    assert t.index("REGLEMENT DE LA CONSULTATION") < t.index("CADRE DE REPONSE") < t.index("Article 1. Les travaux")
    kinds = {m["nom"]: m for m in meta}
    assert kinds["01_RC.pdf"]["type"] == dp.REGLEMENT and kinds["01_RC.pdf"]["lu_en_entier"] is True
    assert kinds["04_Cadre_de_reponse_technique.docx"]["type"] == dp.CADRE
    assert kinds["04_Cadre_de_reponse_technique.docx"]["lu_en_entier"] is True
    assert kinds["03_CCTP_Lot2_Couverture.pdf"]["lu_en_entier"] is False   # le gros CCTP, lui, est coupé
    # Les autres pièces ne disparaissent plus : le CCAP garde sa part.
    assert "penalites de retard" in t
    # Ce que reçoit l'IA : règlement et cadre EN ENTIER, puis des fenêtres ciblées.
    sel = select_dce_content(t)
    flat = re.sub(r"\s+", " ", sel)            # le PDF coupe ses lignes au milieu des phrases
    assert "Sous-critere 3 : environnement 10 points" in flat
    assert "Question 3 : Détaillez votre démarche environnementale." in flat
    assert "La visite des lieux est obligatoire" in flat
    assert "règlement de la consultation, lu en entier" in sel and "cadre de réponse, lu en entier" in sel
    assert len(sel) < len(t) or len(t) <= 30000


def test_sonde_noms_courts_le_reglement_ne_disparait_plus(rc, cadre, cctp):
    """Variante de la sonde : CCTP.pdf + RC.pdf + Cadre_reponse.docx à 120 000 caractères."""
    z = _zip({"CCTP.pdf": cctp, "RC.pdf": rc, "Cadre_reponse.docx": cadre})
    t = extract_dce_text("dce.zip", z, max_chars=120000)
    assert "Critères de jugement des offres" in t, "le règlement a encore disparu"
    assert "CADRE DE REPONSE TECHNIQUE" in t
    assert len(t) <= 120000
    assert t.startswith("=== RC.pdf ===")


def test_compose_partage_equitable_et_fichier_seul_inchange():
    pieces = [{"nom": "CCTP.pdf", "texte": "x" * 50000}, {"nom": "AE.pdf", "texte": "acte " * 400},
              {"nom": "RC.pdf", "texte": "REGLEMENT DE LA CONSULTATION " + "r" * 3000}]
    t, meta = dp.compose(pieces, 20000, headers=True)
    assert len(t) <= 20000
    m = {x["nom"]: x for x in meta}
    assert m["RC.pdf"]["lu_en_entier"] and m["AE.pdf"]["lu_en_entier"] and not m["CCTP.pdf"]["lu_en_entier"]
    # Fichier seul : texte brut, sans en-tête (comportement historique).
    t1, _ = dp.compose([{"nom": "RC.pdf", "texte": "abc " * 10}], 1000, headers=False)
    assert t1 == "abc " * 10


# ── C1 : critères retrouvés dans le texte (IA simulée) ───────────────────────────────

def _fake_llm(payload: dict, seen: list = None):
    class _U:
        input_tokens = 10
        output_tokens = 20

    class _Resp:
        usage = _U()
        content = [type("B", (), {"text": json.dumps(payload, ensure_ascii=False)})()]

    def _create(**kw):
        if seen is not None:
            seen.append(kw)
        return _Resp()
    return type("C", (), {"messages": type("M", (), {"create": staticmethod(_create)})})


def _payload(criteres):
    return {"match_score": 60, "go_decision": "a_etudier", "summary": "Analyse simulée.",
            "details": {"intitule_marche": "Réfection de la couverture du groupe scolaire",
                        "criteres_attribution": criteres}}


def _analyze(text, criteres, monkeypatch, pieces=None, seen=None):
    from app.services import llm
    from app.services.analysis import analyze_dce_text
    llm._CB["fails"] = 0
    llm._CB["open_until"] = 0.0
    llm._TENANT_CB.clear()
    monkeypatch.setattr(llm, "_use_mistral", lambda: False)
    monkeypatch.setattr(llm, "client", lambda: _fake_llm(_payload(criteres), seen))
    return analyze_dce_text(text, {"name": "Toitures de l'Ouest"}, {}, lang_name="français", pieces=pieces)


def _crit(g, name):
    return next(x for x in g["criteres"] if x["critere"] == name)


def test_criteres_phrase_retrouvee_ou_a_verifier(rc, cadre, cctp, monkeypatch):
    t, meta = extract_dce_text("dce.zip", _sonde_zip(rc, cadre, cctp), max_chars=120000, with_pieces=True)
    seen = []
    r = _analyze(t, [
        {"critere": "Valeur technique", "ponderation": "60 %",
         "extrait": "Valeur technique 60 % appréciée au vu du mémoire technique"},          # recopiée
        {"critere": "Prix", "ponderation": "40%", "extrait": "le prix compte pour 40 %"},  # reformulée
        {"critere": "Délai d'exécution", "ponderation": "10 %",
         "extrait": "Délai d'exécution : 10 %"},                                          # inventée
    ], monkeypatch, pieces=meta, seen=seen)
    g = r["details"]["grille_notation"]
    vt, prix, delai = _crit(g, "Valeur technique"), _crit(g, "Prix"), _crit(g, "Délai d'exécution")
    # Recopiée : phrase d'origine conservée, poids TEL QU'ÉCRIT.
    assert vt["statut"] == gn.RETROUVE and vt["poids"] == "60 %" and vt["dans_le_plan"]
    assert "Valeur technique 60 %" in vt["phrase"] and citation_guard(t, vt["phrase"])
    assert vt["source"].startswith("Règlement de la consultation")
    # Reformulée par l'IA : la vraie phrase est retrouvée SANS IA (nom + même poids écrit).
    assert prix["statut"] == gn.RETROUVE and prix["poids"] == "40 %"
    assert "Prix 40 %" in prix["phrase"] and citation_guard(t, prix["phrase"], min_len=4)
    # Inventée : « à vérifier dans le règlement », pas de poids affiché, hors du plan.
    assert delai["statut"] == gn.A_VERIFIER and delai["poids"] == "" and not delai["dans_le_plan"]
    assert delai["phrase"] == ""
    assert g["compteurs"]["a_verifier"] == 1
    assert gn.criteres_plan(g) == [{"critere": "Valeur technique", "poids": "60 %", "source": "reglement"},
                                   {"critere": "Prix", "poids": "40 %", "source": "reglement"}]
    # Contrôle sans IA : 60 + 40 = 100 %.
    assert g["controle_total"]["statut"] == "ok" and g["controle_total"]["total"] == 100
    # Ordre du règlement conservé (Valeur technique est citée avant Prix).
    assert [x["critere"] for x in g["criteres"]][:2] == ["Valeur technique", "Prix"]
    # Ce que l'IA a reçu : le règlement et le cadre en entier.
    prompt = re.sub(r"\s+", " ", seen[0]["messages"][0]["content"])
    assert "Sous-critere 3 : environnement 10 points" in prompt and "Question 3 : Détaillez" in prompt
    # Lecture pièce par pièce gardée dans l'analyse.
    assert g["lecture"]["reglement"] is True and g["lecture"]["cadre"] is True


def test_poids_invente_par_l_ia_jamais_affiche(monkeypatch):
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 6 - Critères d'attribution\n"
          "Valeur technique : 60 %\nPrix : 40 %\n")
    g = _analyze(rc, [{"critere": "Valeur technique", "ponderation": "70 %", "extrait": "Valeur technique : 60 %"},
                      {"critere": "Prix", "ponderation": "40 %", "extrait": "Prix : 40 %"}],
                 monkeypatch)["details"]["grille_notation"]
    vt = _crit(g, "Valeur technique")
    assert vt["statut"] == gn.A_VERIFIER and vt["poids"] == "" and not vt["dans_le_plan"]
    assert vt["phrase"] == "Valeur technique : 60 %"        # la vraie phrase, sans le poids inventé
    assert "70" not in json.dumps(gn.public_grille(g))
    assert _crit(g, "Prix")["poids"] == "40 %"


def test_poids_non_annonce_jamais_ajoute(monkeypatch):
    """L'IA ne donne pas de poids : on n'en prend jamais un « à côté du nom » (ici celui d'un
    sous-critère), et la pondération reste à vérifier, sans prétendre « non pondérés »."""
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Critères de jugement des offres\n"
          "Valeur technique : sous-critère méthodologie 30 points\nPrix\n")
    g = _analyze(rc, [{"critere": "Valeur technique", "ponderation": "",
                       "extrait": "Valeur technique : sous-critère méthodologie 30 points"}],
                 monkeypatch)["details"]["grille_notation"]
    vt = _crit(g, "Valeur technique")
    assert vt["statut"] == gn.RETROUVE and vt["poids"] == "" and vt["poids_retenu"] == ""
    assert g["ponderation"]["statut"] == "incomplet"
    # L'utilisateur peut alors donner le poids qu'il a lu (ligne retrouvée sans poids).
    g2 = gn.apply_choice(g, vt["id"], "confirme", poids="60 %")
    assert _crit(g2, "Valeur technique")["poids_retenu"] == "60 %"


def test_total_incoherent_avertit(monkeypatch):
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Critères de jugement des offres\n"
          "Valeur technique : 60 %\nPrix : 30 %\n")
    g = _analyze(rc, [{"critere": "Valeur technique", "ponderation": "60 %", "extrait": "Valeur technique : 60 %"},
                      {"critere": "Prix", "ponderation": "30 %", "extrait": "Prix : 30 %"}],
                 monkeypatch)["details"]["grille_notation"]
    ct = g["controle_total"]
    assert ct["statut"] == "ecart" and ct["total"] == 90 and ct["attendu"] == 100
    assert "90 %" in ct["message"] and "100 %" in ct["message"]


def test_points_et_total_annonce(monkeypatch):
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Critères de jugement des offres\n"
          "Les offres sont notées sur un total de 100 points.\n"
          "Valeur technique : 60 points\nPrix : 30 points\n")
    g = _analyze(rc, [{"critere": "Valeur technique", "ponderation": "60 points", "extrait": "Valeur technique : 60 points"},
                      {"critere": "Prix", "ponderation": "30 points", "extrait": "Prix : 30 points"}],
                 monkeypatch)["details"]["grille_notation"]
    assert g["total_annonce"]["valeur"] == 100
    assert g["controle_total"]["statut"] == "ecart"
    assert "90 points au lieu de 100 points" in g["controle_total"]["message"]


def test_criteres_non_ponderes_sans_ordre_invente(monkeypatch):
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Critères de jugement des offres\n"
          "Le jugement des offres sera effectué selon les critères suivants : la valeur technique de l'offre, "
          "le prix des prestations, le délai d'exécution. Ces critères ne sont pas hiérarchisés.\n")
    g = _analyze(rc, [{"critere": "Valeur technique", "ponderation": "", "extrait": ""},
                      {"critere": "Prix", "ponderation": "", "extrait": ""},
                      {"critere": "Délai d'exécution", "ponderation": "", "extrait": ""}],
                 monkeypatch)["details"]["grille_notation"]
    assert g["ponderation"]["statut"] == "non_ponderes"
    assert g["ponderation"]["message"] == gn.MSG_NON_PONDERES == "Critères non pondérés, sans ordre indiqué."
    assert g["ordre"]["indique"] is False and "ne sont pas hiérarchisés" in g["ordre"]["phrase"]
    assert all(x["poids"] == "" for x in g["criteres"])
    assert all(x["statut"] == gn.RETROUVE for x in g["criteres"]), g["criteres"]
    assert g["controle_total"]["statut"] == "sans_objet"
    assert gn.plan_text(g).startswith("critères non pondérés, sans ordre indiqué : ")


def test_non_ponderes_mais_ordre_annonce_par_le_reglement(monkeypatch):
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Critères d'attribution\n"
          "Les critères sont classés par ordre décroissant d'importance : 1. Valeur technique ; 2. Prix.\n")
    g = _analyze(rc, [{"critere": "Valeur technique", "ponderation": "", "extrait": ""},
                      {"critere": "Prix", "ponderation": "", "extrait": ""}], monkeypatch)["details"]["grille_notation"]
    assert g["ponderation"]["statut"] == "non_ponderes"
    assert g["ponderation"]["message"] == gn.MSG_NON_PONDERES_ORDRE
    assert g["ordre"]["indique"] is True and "ordre décroissant" in g["ordre"]["phrase"]


def test_poids_ecrits_mais_non_rapportes_par_l_ia_jamais_non_ponderes(monkeypatch):
    """L'IA oublie les poids alors que le règlement en donne : on n'affirme pas « non pondérés »."""
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Critères de jugement des offres\n"
          "Pondération : technique 55 %, financier 45 % (voir tableau).\n")
    g = _analyze(rc, [{"critere": "Valeur technique", "ponderation": "", "extrait": ""},
                      {"critere": "Prix", "ponderation": "", "extrait": ""}], monkeypatch)["details"]["grille_notation"]
    assert g["ponderation"]["statut"] != "non_ponderes"


# ── C2 minimal : ce que le dossier impose ────────────────────────────────────────────

RC_LIMITES = (
    "REGLEMENT DE LA CONSULTATION\n"
    "Article 9 - Présentation des offres\n"
    "Le mémoire technique fera 10 pages maximum, annexes comprises.\n"
    "La réponse à chaque question est limitée à 15 000 caractères espaces compris.\n"
    "Le document sera rédigé en police Arial 11.\n"
    "La visite des lieux est obligatoire avant la remise des offres.\n"
    "Le candidat devra impérativement compléter le cadre de réponse technique joint au dossier.\n"
)


def test_limites_cadre_impose_et_visite_reperes():
    g = gn.build_grille({}, RC_LIMITES)
    c = g["contraintes"]
    lim = {(x["type"], x.get("valeur") or x.get("police")) for x in c["limites"]}
    assert ("pages", 10) in lim and ("caracteres", 15000) in lim and ("police", "Arial") in lim
    pol = next(x for x in c["limites"] if x["type"] == "police")
    assert pol["taille"] == 11
    assert c["visite"] and "visite des lieux est obligatoire" in c["visite"]["phrase"]
    assert c["fichier_impose"] is True
    assert c["message_fichier"] == "Ce marché impose son propre fichier de réponse : vous devrez y recopier votre texte."
    assert "impérativement compléter le cadre de réponse" in c["cadre"]["obligatoire"]["phrase"]
    assert c["cadre"]["fourni"] is False and c["cadre"]["mentionne"] is True   # cité, pas dans les fichiers
    # Chaque repérage garde une phrase VRAIMENT présente dans le texte.
    for x in c["limites"] + [c["visite"], c["cadre"]["obligatoire"]]:
        assert x["phrase"] and citation_guard(RC_LIMITES, x["phrase"]), x


def test_limites_pas_de_faux_positifs_evidents():
    rc = ("REGLEMENT DE LA CONSULTATION\n"
          "Le CCTP comporte 45 pages et les pénalités sont plafonnées à 10 %.\n"
          "Le candidat fournit une attestation de police d'assurance responsabilité civile.\n"
          "La visite des lieux n'est pas obligatoire.\n"
          "Une visite facultative pourra être organisée.\n"
          "Le délai d'exécution est de 10 jours maximum.\n"
          "Le mémoire comportera au minimum 2 pages.\n"
          "Le candidat pourra utiliser le cadre de réponse proposé, à titre indicatif.\n"
          "La police municipale sera informée du chantier.\n")
    c = gn.build_grille({}, rc)["contraintes"]
    assert c["limites"] == [], c["limites"]
    assert c["visite"] is None
    assert c["fichier_impose"] is False and c["message_fichier"] == ""
    assert c["cadre"]["mentionne"] is True     # mentionné, mais facultatif


def test_limites_du_cctp_ignorees_quand_le_reglement_est_la():
    """« Rapport mensuel de 3 pages maximum » (exécution du marché, dans le CCTP) n'est pas une
    limite de l'offre ; une visite de maintenance obligatoire non plus."""
    t, _ = dp.compose([
        {"nom": "RC.pdf", "texte": "REGLEMENT DE LA CONSULTATION\nLe mémoire technique ne devra pas excéder 12 pages."},
        {"nom": "CCTP.pdf", "texte": "CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\nLe rapport mensuel fera 3 pages maximum. "
                                     "Une visite obligatoire de maintenance est due chaque année."},
    ], 50000, headers=True)
    c = gn.build_grille({}, t)["contraintes"]
    assert [(x["type"], x["valeur"]) for x in c["limites"]] == [("pages", 12)]
    assert c["visite"] is None


def test_cadre_fourni_dans_le_zip_et_limite_de_question(rc, cadre, cctp):
    t, meta = extract_dce_text("dce.zip", _sonde_zip(rc, cadre, cctp), max_chars=120000, with_pieces=True)
    g = gn.build_grille({}, t, pieces=meta)
    c = g["contraintes"]
    assert c["cadre"]["fourni"] is True and "04_Cadre_de_reponse_technique.docx" in c["cadre"]["pieces"]
    assert c["fichier_impose"] is True        # « doit obligatoirement être utilisé »
    pages = sorted(x["valeur"] for x in c["limites"] if x["type"] == "pages")
    assert pages == [2, 10]                   # 2 pages (question 1 du cadre), 10 pages (règlement)
    assert c["visite"] is not None


# ── Choix de l'utilisateur, reprise, vue partenaire ─────────────────────────────────

def _grille_simple():
    rc = ("REGLEMENT DE LA CONSULTATION\nArticle 7 - Critères de jugement des offres\n"
          "Valeur technique : 60 %\nPrix : 30 %\nLe mémoire technique ne devra pas excéder 8 pages.\n")
    return gn.build_grille({"criteres_attribution": [
        {"critere": "Valeur technique", "ponderation": "60 %", "extrait": "Valeur technique : 60 %"},
        {"critere": "Prix", "ponderation": "30 %", "extrait": "Prix : 30 %"},
        {"critere": "Délai", "ponderation": "10 %", "extrait": "Délai : 10 %"}]}, rc)


def test_ligne_a_verifier_entre_dans_le_plan_seulement_apres_confirmation():
    g = _grille_simple()
    d = _crit(g, "Délai")
    assert d["statut"] == gn.A_VERIFIER and not d["dans_le_plan"]
    assert g["controle_total"]["statut"] == "ecart"           # 60 + 30 = 90
    g2 = gn.apply_choice(g, d["id"], "confirme", poids="10 %", when="2026-10-12T10:00:00")
    d2 = _crit(g2, "Délai")
    assert d2["dans_le_plan"] and d2["poids_retenu"] == "10 %" and d2["poids_source"] == "utilisateur"
    assert g2["controle_total"]["statut"] == "ok"
    assert {"critere": "Délai", "poids": "10 %", "source": "utilisateur"} in gn.criteres_plan(g2)
    assert _crit(g, "Délai")["dans_le_plan"] is False        # la grille d'origine n'est pas modifiée
    g3 = gn.apply_choice(g2, d["id"], "annule")
    assert not _crit(g3, "Délai")["dans_le_plan"]
    g4 = gn.apply_choice(g, _crit(g, "Prix")["id"], "ecarte")
    assert not _crit(g4, "Prix")["dans_le_plan"]
    with pytest.raises(ValueError):
        gn.apply_choice(g, d["id"], "confirme", poids="dix pour")     # poids illisible
    with pytest.raises(ValueError):
        gn.apply_choice(g, "c99", "confirme")


def test_reprise_des_choix_a_la_reanalyse():
    g = gn.apply_choice(_grille_simple(), "c3", "confirme", poids="10 %")
    fresh = _grille_simple()
    out = gn.carry_over(g, fresh)
    assert _crit(out, "Délai")["dans_le_plan"] and _crit(out, "Délai")["poids_retenu"] == "10 %"


def test_vue_partenaire_bridee_sans_fuite():
    from types import SimpleNamespace
    from app.routers import invites
    g = gn.apply_choice(_grille_simple(), "c3", "confirme", poids="10 %", when="2026-10-12T10:00:00")
    g["lecture"]["pieces"] = [{"nom": "RC_interne_v3.pdf", "type": "reglement", "libelle": "x", "lu_en_entier": True}]
    analysis = {"match_score": 72, "details": {"intitule_marche": "Couverture", "acheteur": "Commune",
                                               "budget_estime": "180 000 EUR HT", "grille_notation": g}}
    pub = invites._safe_analysis(analysis)["grille_notation"]
    blob = json.dumps(pub, ensure_ascii=False)
    assert "RC_interne_v3.pdf" not in blob and "2026-10-12" not in blob and "180 000" not in blob
    assert "lecture" not in pub and "choix_utilisateur" not in blob and "source" not in json.dumps(pub["criteres"])
    assert [x["critere"] for x in pub["criteres"]] == ["Valeur technique", "Prix", "Délai"]
    assert pub["criteres"][2]["confirme"] is True and pub["criteres"][2]["poids"] == "10 %"
    assert pub["contraintes"]["limites"][0]["valeur"] == 8
    # Aperçu minimal (invitation réseau non acceptée) : ni acheteur, ni grille.
    inv = SimpleNamespace(recipient="a@b.fr", company_name="Partenaire", role="cotraitant", can_view_docs=True,
                          can_contribute=True, created_at=None)
    project = SimpleNamespace(name="Couverture", client="Commune", deadline=None, source_url=None, ai_analysis=analysis)
    orig = invites._network_pending
    try:
        invites._network_pending = lambda i, db: True
        out = invites._invite_public(inv, project, None, [], db=None)
    finally:
        invites._network_pending = orig
    assert "grille_notation" not in out["project"]["analysis"] and "acheteur" not in out["project"]["analysis"]
    # Les phrases d'état s'adressent au partenaire, pas au chef de file.
    msg = gn.public_grille(_grille_simple())["ponderation"]["message"]
    assert "le chef de file doit le confirmer" in msg and "votre plan" not in msg
    assert "votre plan" in _grille_simple()["ponderation"]["message"]


def test_annonce_seule_sans_liste_de_pieces():
    g = gn.build_grille({"criteres_attribution": [{"critere": "Prix", "ponderation": "100 %",
                                                   "extrait": "Critères d'attribution : prix 100 %"}]},
                        "Avis de marché. Critères d'attribution : prix 100 %. Date limite : 12/12/2026.")
    assert g["lecture"]["pieces"] == [] and g["lecture"]["reglement"] is False
    assert _crit(g, "Prix")["poids"] == "100 %" and g["controle_total"]["statut"] == "ok"


def test_redaction_ne_recoit_que_le_plan():
    from app.services.agents.redaction import _criteres_du_plan
    g = _grille_simple()
    txt = _criteres_du_plan({"grille_notation": g, "criteres_attribution": [{"critere": "Délai"}]})
    assert "Valeur technique (60 %)" in txt and "Prix (30 %)" in txt and "Délai" not in txt
    assert _criteres_du_plan({"criteres_attribution": [{"critere": "X"}]}) == "[{'critere': 'X'}]"


# ── Parcours complet : import du DCE, carte, confirmation (IA simulée) ───────────────

def test_parcours_import_dce_puis_confirmation(client, auth, monkeypatch, rc, cadre, cctp):
    h = auth["headers"]
    pr = client.post("/api/projects/", headers=h, json={"name": "Couverture école", "budget": 180000})
    assert pr.status_code == 201, pr.text
    pid = pr.json()["id"]
    from app.services import llm
    llm._CB["fails"] = 0
    llm._CB["open_until"] = 0.0
    llm._TENANT_CB.clear()
    monkeypatch.setattr(llm, "_use_mistral", lambda: False)
    monkeypatch.setattr(llm, "client", lambda: _fake_llm(_payload([
        {"critere": "Valeur technique", "ponderation": "60 %", "extrait": "Valeur technique 60 %"},
        {"critere": "Prix", "ponderation": "40 %", "extrait": "Prix 40 %."},
        {"critere": "Délai", "ponderation": "", "extrait": "Délai de réalisation"}])))
    z = _sonde_zip(rc, cadre, cctp)
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE.zip", z, "application/zip")})
    assert r.status_code == 200, r.text
    g = r.json()["details"]["grille_notation"]
    assert g["lecture"]["reglement"] and g["lecture"]["cadre"]
    assert _crit(g, "Délai")["statut"] == gn.A_VERIFIER
    assert g["contraintes"]["fichier_impose"] is True

    bad = client.post("/api/sourcing/grille/choix", headers=h,
                      json={"project_id": pid, "critere_id": _crit(g, "Délai")["id"], "choix": "confirme", "poids": "beaucoup"})
    assert bad.status_code == 422
    ok = client.post("/api/sourcing/grille/choix", headers=h,
                     json={"project_id": pid, "critere_id": _crit(g, "Délai")["id"], "choix": "confirme", "poids": ""})
    assert ok.status_code == 200, ok.text
    assert _crit(ok.json()["grille_notation"], "Délai")["dans_le_plan"] is True
    p = client.get(f"/api/projects/{pid}", headers=h).json()
    assert _crit(p["ai_analysis"]["details"]["grille_notation"], "Délai")["dans_le_plan"] is True

    # Un autre compte ne peut rien changer (404, pas de fuite d'existence).
    import uuid
    other = client.post("/api/auth/register", json={"email": f"o_{uuid.uuid4().hex[:8]}@test.fr",
                                                     "password": "motdepasse123", "full_name": "O", "company_name": "O"})
    oh = {"Authorization": f"Bearer {other.json()['access_token']}"}
    r404 = client.post("/api/sourcing/grille/choix", headers=oh,
                       json={"project_id": pid, "critere_id": "c1", "choix": "ecarte"})
    assert r404.status_code == 404

    # Ré-analyse du dossier archivé : le choix de l'utilisateur est repris.
    rr = client.post("/api/sourcing/reanalyze", headers=h, data={"project_id": str(pid)})
    assert rr.status_code == 200, rr.text
    assert _crit(rr.json()["details"]["grille_notation"], "Délai")["dans_le_plan"] is True
    llm._CB["fails"] = 0
    llm._CB["open_until"] = 0.0


# ── Écrans : carte compacte, lecture seule côté partenaire, zéro émoji ──────────────

def _static(name):
    import os
    return open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app", "static", name),
                encoding="utf-8").read()


def test_ecrans_carte_comment_vous_serez_note():
    js, html, inv = _static("app.js"), _static("app.html"), _static("invite.html")
    assert '__adjApp.component("grille-carte"' in js
    assert "<grille-carte" in html and ':edit="true"' in html and ':edit="false"' in html
    assert "Comment vous serez noté" in html and "Comment vous serez noté" in inv
    assert "function renderGrille" in inv
    # Détails repliés par défaut : <details> sans attribut open.
    comp = js[js.index('__adjApp.component("grille-carte"'):js.index("// ── Réseau de partenaires (vague 3) : petits composants")]
    assert "<details" in comp and not re.search(r"<details[^>]*\bopen\b", comp) and ":open" not in comp
    assert "à vérifier dans le règlement" in comp
    # On ne dit jamais « vérifié » ; jamais de v-html.
    assert "vérifié" not in comp and "v-html" not in comp
    emoji = re.compile("[\U0001F300-\U0001FAFF☀-➿⭐✅]")
    for src in (comp, inv[inv.index("function renderGrille"):inv.index("function renderMonAO")]):
        assert not emoji.search(src)
