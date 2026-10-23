"""
Rédaction, vague 4 — le PARCOURS COMPLET du dossier de groupement cloisonné, de bout en bout, par l'API (étude
« rédaction » § 6.8 B ; contrat docs/REDACTION_V4_GROUPEMENT.md § 4, § 5, § 15 à § 17). IA simulée, aucun réseau.

Le chef de file répartit → demande les parts → le partenaire A (lien magique, SANS compte) rédige une partie avec
l'aide de l'IA et une autre lui-même → il envoie → le chef de file reçoit le texte et des COMPTEURS de sources, jamais
un extrait → il modifie une part → A revalide → le partenaire B (AVEC un compte Adjugo) rédige à partir de SES preuves
et de SON savoir-faire → assemblage (ordre du cadre, parties tirées du groupement) → « Qui fait quoi » validée par
chacun → « Sections validées 6 sur 6 » → dossier assemblé validé par chacun → export Word sans marque interne.
Plus : la tâche d'un partenaire rend 404 à un membre de l'organisation du chef de file ; un marché sans groupement
garde l'atelier d'une seule entreprise ; tout est inscrit au journal chaîné.
"""
import io
import re

from app.services import atelier_groupement as ag
from app.services import llm
from tests.test_convention_groupement import _register
from tests.test_redaction_socle0 import _company, _db, fake_llm  # noqa: F401
from tests.test_redaction_v2_atelier import _creer, _etat, _projet
from tests.test_redaction_v4_assemblage import _convention_validee, _generee, _valider_chef
from tests.test_redaction_v4_groupement import (P, PC, SECRET_BROUILLON, SECRET_CHEF, SECRET_DOC, SECRET_EXTRAIT,
                                                _aide, _audit, _groupement, _part_a, _section_par, _vues_du_chef)

TEXTE_IA_A = ("Nous trions les déchets amiantés sur place, dans des bennes étanches fermées. "
              "[À compléter : planning des rotations]")
TEXTE_IA_B = ("Nos électriciens habilités interviennent en site occupé [S1], avec une nacelle élévatrice de 18 m [F1] "
              "pour les travaux en hauteur.")
METHODO_A = "Nous intervenons en site occupé, par zones fermées et balisées, avec un confinement à chaque phase."


def _email_a(invite_id, adresse):
    from app.models import ProjectInvite
    db = _db()
    try:
        db.get(ProjectInvite, invite_id).recipient = adresse
        db.commit()
    finally:
        db.close()


def _relire_tout(p, sid):
    """Le partenaire relit lui-même chaque fait encore « à relire » de son brouillon (« Je l'ai relu »)."""
    b = p.get(f"/sections/{sid}").json()["brouillon"]
    for f in b["faits"]:
        if f["statut"] == "a_relire" and f["type"] != "reste":
            r = p.post(f"/sections/{sid}/faits/{f['id']}", {"statut": "relu_par_vous"})
            assert r.status_code == 200, r.text
    return p.get(f"/sections/{sid}").json()["brouillon"]


def _envoyer_lu(p, sid):
    b = p.get(f"/sections/{sid}").json()["brouillon"]
    r = p.post(f"/sections/{sid}/envoyer", {"brouillon_id": b["id"]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "avertissement_requis"      # décision 3a
    r = p.post(f"/sections/{sid}/envoyer", {"brouillon_id": b["id"], "avertissement_lu": True})
    assert r.status_code == 200, r.text
    assert r.json()["statut"] == "envoyee"
    return r.json()


def _valider_part(p, sid):
    cur = p.get(f"/sections/{sid}").json()["texte_commun"]["version_id"]
    r = p.post(f"/sections/{sid}/valider", {"version_id": cur})
    assert r.status_code == 200, r.text
    assert r.json()["statut"] == "validee"


def _docx_texte(contenu: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(contenu))
    lignes = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lignes.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lignes)


def test_parcours_complet_du_dossier_de_groupement(client, auth, fake_llm, monkeypatch):
    h = auth["headers"]
    envois = []
    from app.services import email as _email
    monkeypatch.setattr(_email, "send_email", lambda to, subj, text, **k: envois.append((to, subj, text)) or True)

    # ── 0. Le groupement est monté : A (désamiantage, lien sans compte), B (électricité) ; parts soumises, convention
    #       générée et validée par les trois membres. Le chef de file a son propre savoir-faire (jamais lu par eux).
    pid, st, a, b, vue = _groupement(client, auth)
    assert client.post("/api/knowledge/text", headers=h, json={
        "name": "Mémoire du chef", "kind": "memoire",
        "text": f"Équipe d'exécution et gestion des déchets du chantier par le chef de file {SECRET_CHEF}."}).status_code == 200
    _convention_validee(client, auth, pid, [a, b])
    # B : un compte Adjugo vérifié, SES preuves (matériel) et SON savoir-faire ; son lien vise son adresse.
    hb, eb = _register(client, company="Le Gall Électricité")
    r = client.post("/api/preuves", headers=hb, json={"nature": "materiel", "type_source": "declare",
                                                       "valeur": {"designation": "Nacelle élévatrice 18 m"}})
    assert r.status_code == 201, r.text
    assert client.post("/api/knowledge/text", headers=hb, json={
        "name": SECRET_DOC, "kind": "memoire",
        "text": (f"Équipe d'exécution : six électriciens habilités, compagnons affectés au chantier en site occupé "
                 f"{SECRET_EXTRAIT}.")}).status_code == 200
    _email_a(b["id"], eb)
    assert client.post(f"/api/invite/{b['token']}/claim", headers=hb).status_code == 200

    # ── 1. Le chef de file répartit (pré-remplissage sans IA, puis ses choix) ──
    env, met = _part_a(vue), _section_par(vue, "méthodologie")
    equ, enc = _section_par(vue, "équipe"), _section_par(vue, "encadrement")
    assert env["redacteur"]["invite_id"] == a["id"] and env["choix"] == "propose"
    for ligne, inv in ((met, a), (equ, b)):
        r = client.put(f"/api/atelier/{pid}/repartition/{ligne['section_id']}", headers=h,
                       json={"redacteur": "partenaire", "invite_id": inv["id"], "lock_version": ligne["lock_version"]})
        assert r.status_code == 200, r.text
    r = client.put(f"/api/atelier/{pid}/repartition/{env['section_id']}", headers=h,
                   json={"longueur_mots": 150, "date_limite": "2030-11-02"})
    assert r.status_code == 200 and r.json()["longueur_mots"] == 150
    vue = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()
    cadre = [l for l in vue["sections"] if not l.get("genere")]
    assert [l["redacteur"]["nom"] for l in cadre] == ["Chef de file", "Le Gall Électricité", "Breizh Désamiantage",
                                                      "Breizh Désamiantage"]

    # ── 2. « Demander les parts » : B est prévenu par email (son adresse) ; A n'a pas d'adresse (lien transmis) ──
    del envois[:]                                   # les emails de la convention, envoyés avant
    r = client.post(f"/api/atelier/{pid}/repartition/demander", headers=h, json={})
    assert r.status_code == 200, r.text
    d = r.json()
    assert [(x["nom"], x["parties"], x["email_parti"]) for x in d["demandees"]] == [("Le Gall Électricité", 1, True)]
    assert [x["nom"] for x in d["sans_adresse"]] == ["Breizh Désamiantage"]
    assert len(envois) == 1 and envois[0][0] == eb and "Équipe d'exécution" in envois[0][2]

    # ── 3. A, par son lien, SANS compte : l'aide de l'IA sur une partie, une autre écrite lui-même ──
    vus = []

    def repondre(kw):
        t = llm._CURRENT_TENANT.get()
        vus.append((t, llm._flatten_content(kw.get("system")) + "\n".join(
            llm._flatten_content(m.get("content")) for m in kw.get("messages") or [])))
        return (TEXTE_IA_B if t == f"inv:{b['id']}" else TEXTE_IA_A), "end_turn"
    fake = fake_llm(repondre)
    pa = P(client, a["token"])
    sid_env, sid_met, sid_equ = env["section_id"], met["section_id"], equ["section_id"]
    job_a, jid_a = _aide(pa, sid_env)
    assert job_a["status"] == "done" and "texte" not in str(job_a["result"])
    envoye_a = "\n".join(x for t, x in vus if t == f"inv:{a['id']}")
    assert vus and all(t == f"inv:{a['id']}" for t, _x in vus)                       # sa portée de facturation
    assert "Qualibat 1552" in envoye_a and "REGLEMENT DE LA CONSULTATION" in envoye_a  # D5 : sa part, l'acheteur
    for interdit in (SECRET_CHEF, SECRET_EXTRAIT, "Nacelle élévatrice", "Électricité courants forts"):
        assert interdit not in envoye_a, interdit
    b1 = pa.get(f"/sections/{sid_env}").json()["brouillon"]
    assert b1["origine"] == "ia" and b1["a_completer"] == 1
    r = pa.put(f"/sections/{sid_env}/brouillon", {"texte": b1["texte"].replace(
        "[À compléter : planning des rotations]", "Le planning des rotations est joint au dossier."),
        "base_brouillon_id": b1["id"]})
    assert r.status_code == 200 and r.json()["brouillon"]["a_completer"] == 0
    _relire_tout(pa, sid_env)
    # La 2e partie, écrite par lui (un brouillon privé d'abord, jamais lu du chef de file).
    assert pa.put(f"/sections/{sid_met}/brouillon", {"texte": f"Premier jet {SECRET_BROUILLON}.",
                                                     "base_brouillon_id": None}).status_code == 200
    bm = pa.get(f"/sections/{sid_met}").json()["brouillon"]
    assert pa.put(f"/sections/{sid_met}/brouillon", {"texte": METHODO_A, "base_brouillon_id": bm["id"]}).status_code == 200
    _relire_tout(pa, sid_met)
    for sid in (sid_env, sid_met):
        _envoyer_lu(pa, sid)
    n_ia = len(fake.calls)

    # ── 4. Le chef de file reçoit le texte et des compteurs par type, jamais une source, un extrait, un brouillon ──
    for sid in (sid_env, sid_met):
        ligne = client.get(f"/api/atelier/{pid}/repartition/{sid}", headers=h).json()
        assert ligne["statut"] == "envoyee" and ligne["envoyee_le"]
        tc = ligne["texte_commun"]
        assert set(tc["compteurs"]) == {"par_type", "avec_source", "a_relire", "engagements"}
        assert "sources" not in tc and "brouillon" not in ligne
    assert "Le planning des rotations est joint au dossier." in client.get(
        f"/api/atelier/{pid}/repartition/{sid_env}", headers=h).json()["texte_commun"]["texte"]
    tout = _vues_du_chef(client, h, pid, sid_met, token=a["token"]) + _vues_du_chef(client, h, pid, sid_env)
    assert METHODO_A in tout and SECRET_BROUILLON not in tout
    assert not re.search(r'"extrait":\s*"[^"]', tout)

    # ── 5. Il modifie une part : « à revalider par » A, l'avant et l'après ; A valide la modification ──
    cur = client.get(f"/api/atelier/{pid}/repartition/{sid_env}", headers=h).json()["texte_commun"]
    r = client.put(f"/api/atelier/{pid}/repartition/{sid_env}/texte", headers=h, json={
        "texte": cur["texte"].replace("bennes étanches fermées", "bennes étanches fermées et étiquetées"),
        "base_version_id": cur["version_id"]})
    assert r.status_code == 200, r.text
    p = pa.get(f"/sections/{sid_env}").json()
    assert p["statut"] == "a_revalider" and p["mention"] == "Modifié par le chef de file, à revalider par Breizh Désamiantage"
    assert "étiquetées" in p["comparaison"]["apres"]["texte"] and "étiquetées" not in p["comparaison"]["avant"]["texte"]
    _valider_part(pa, sid_env)
    # Le contrôle des faits a tourné chez lui, sur SES faits : ce que le texte du chef de file ajoute et qu'il ne peut
    # pas sourcer reste à relire (rien n'est blanchi), et bloque l'export ; sa relecture vaut pour le texte commun,
    # sans toucher au texte ni à sa validation.
    assert client.get(f"/api/atelier/{pid}/repartition/{sid_env}", headers=h).json()[
        "texte_commun"]["compteurs"]["a_relire"] >= 1
    _relire_tout(pa, sid_env)
    ligne = client.get(f"/api/atelier/{pid}/repartition/{sid_env}", headers=h).json()
    assert ligne["texte_commun"]["compteurs"]["a_relire"] == 0 and ligne["statut"] == "validee"
    _valider_part(pa, sid_met)
    assert client.get(f"/api/atelier/{pid}/repartition/{sid_env}", headers=h).json()["mention"] == \
        "Modifié par le chef de file, validé par Breizh Désamiantage"

    # ── 6. B, par son COMPTE : l'aide de l'IA à partir de SES preuves et de SON savoir-faire, avec SES droits ──
    pb = PC(client, pid, hb)
    vb = pb.get().json()
    assert vb["mode"] == "compte" and vb["message_mode"] == ag.MSG_MODE_COMPTE and vb["identite"]["a_prouver"] is False
    assert [x["section_id"] for x in vb["parties"]] == [sid_equ]
    job_b, jid_b = _aide(pb, sid_equ)
    assert job_b["status"] == "done", job_b
    envoye_b = "\n".join(x for t, x in vus[n_ia:] if t == f"inv:{b['id']}")
    assert all(t == f"inv:{b['id']}" for t, _x in vus[n_ia:])
    assert "Nacelle élévatrice 18 m" in envoye_b and SECRET_EXTRAIT in envoye_b       # ses preuves, son savoir-faire
    for interdit in (SECRET_CHEF, "Qualibat 1552", METHODO_A, "Désamiantage et gestion"):
        assert interdit not in envoye_b, interdit
    bb = pb.get(f"/sections/{sid_equ}").json()["brouillon"]
    assert bb["origine"] == "ia" and any(SECRET_EXTRAIT in (x.get("extrait") or "") for x in bb["sources"])
    _relire_tout(pb, sid_equ)
    _envoyer_lu(pb, sid_equ)
    ligne = client.get(f"/api/atelier/{pid}/repartition/{sid_equ}", headers=h).json()
    assert ligne["texte_commun"]["compteurs"]["avec_source"] >= 1
    tout = _vues_du_chef(client, h, pid, sid_equ, token=b["token"])
    assert "nacelle élévatrice de 18 m" in tout
    for secret in (SECRET_EXTRAIT, SECRET_DOC):
        assert secret not in tout, secret
    _valider_part(pb, sid_equ)

    # ── 7. La tâche d'un partenaire n'existe pas pour l'organisation du chef de file (404) ──
    from app.models import Job, User
    hc, ec = _register(client, company="Collègue du chef")
    db = _db()
    try:
        org_id = db.query(User).filter(User.email == auth["email"]).first().org_id
        db.query(User).filter(User.email == ec).first().org_id = org_id
        db.commit()
        for jid, inv in ((jid_a, a), (jid_b, b)):
            j = db.get(Job, jid)
            assert j.user_id is None and j.portee == "partenaire" and j.invite_id == inv["id"]
    finally:
        db.close()
    for hh in (h, hc):
        for jid in (jid_a, jid_b):
            assert client.get(f"/api/jobs/{jid}", headers=hh).status_code == 404
            assert client.get(f"/api/atelier/{pid}/taches/{jid}", headers=hh).status_code == 404
    assert pa.get(f"/taches/{jid_b}").status_code == 404 and pb.get(f"/taches/{jid_a}").status_code == 404

    # ── 8. Le chef de file écrit et valide SA partie dans l'atelier (aucune IA) ──
    s_enc = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == enc["section_id"])
    r = client.put(f"/api/atelier/{pid}/sections/{s_enc['id']}/texte", headers=h, json={
        "texte": "Un conducteur de travaux encadre le chantier du début à la fin.", "base_version_id": None})
    assert r.status_code == 200, r.text
    assert client.post(f"/api/atelier/{pid}/sections/{s_enc['id']}/valider", headers=h,
                       json={"version_id": r.json()["version_courante"]["id"]}).status_code == 200

    # ── 9. Assemblage : l'ordre du cadre, les deux parties tirées du groupement en tête ──
    dos = client.get(f"/api/atelier/{pid}/repartition/dossier", headers=h).json()
    titres = [x["intitule"] for x in dos["parties"]]
    assert titres[:2] == ["Organisation du groupement", "Qui fait quoi dans le groupement"]
    assert [x["section_id"] for x in dos["parties"][2:]] == [enc["section_id"], sid_equ, sid_met, sid_env]
    assert dos["note_interne"]["lignes"]                                     # au chef de file seul
    lu_a = pa.get("/dossier").json()
    assert [x["intitule"] for x in lu_a["parties"]] == titres and "note_interne" not in lu_a
    qfq_txt = next(x for x in lu_a["parties"] if x.get("genere") == "grp_qfq")["texte"]
    assert "Breizh Désamiantage" in qfq_txt and "Le Gall Électricité" in qfq_txt
    for interdit in ("aucune qualification", "manque", "justifiée", "déclarée", "€"):
        assert interdit not in qfq_txt, interdit

    # ── 10. « Qui fait quoi » validée par chacun ; « Organisation » par le chef de file (convention validée par tous) ──
    pret = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["dossier"]["pret"]
    assert pret["message"] == "Sections validées 4 sur 6" and pret["complet"] is False
    for part in (pa, pb):
        q = part.get().json()["qui_fait_quoi"]
        r = part.post(f"/sections/{q['section_id']}/valider", {"version_id": q["texte_commun"]["version_id"]})
        assert r.status_code == 200, r.text
    st = _etat(client, h, pid)
    for kind in ("grp_orga", "grp_qfq"):
        assert _valider_chef(client, h, pid, _generee(st, kind)["id"]).status_code == 200
    vue = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()
    assert vue["dossier"]["pret"]["message"] == "Sections validées 6 sur 6"
    assert [m["message"] for m in vue["dossier"]["pret"]["manques"]] == [
        "Dossier assemblé : à valider par le chef de file", "Dossier assemblé : à valider par Breizh Désamiantage",
        "Dossier assemblé : à valider par Le Gall Électricité"]

    # ── 11. Le dossier assemblé validé par chaque membre, sur ce qu'il a lu : « Prêt à déposer » ──
    r = client.post(f"/api/atelier/{pid}/repartition/valider-dossier", headers=h,
                    json={"empreinte": vue["dossier"]["empreinte"]})
    assert r.status_code == 200, r.text
    for part in (pa, pb):
        assert part.get().json()["dossier"]["peut_valider"] is True
        r = part.post("/dossier/valider", {"empreinte": part.get("/dossier").json()["empreinte"]})
        assert r.status_code == 200 and r.json()["validation"]["statut"] == "valide"
    pret = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["dossier"]["pret"]
    assert pret["message"] == "Sections validées 6 sur 6" and pret["complet"] is True and pret["manques"] == []

    # ── 12. Export Word : le dossier assemblé complet, sans aucune marque interne ──
    st = _etat(client, h, pid)
    assert st["blocage"]["peut_exporter"] is True and st["blocage"]["derogation"] is None, st["blocage"]["raisons"]
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx", "etat_hash": st["blocage"]["etat_hash"]})
    assert r.status_code == 200, r.text
    plat = re.sub(r"\s+", " ", _docx_texte(r.content))
    ordre = ["Organisation du groupement", "Qui fait quoi dans le groupement", "Encadrement du chantier",
             "Équipe d'exécution", "Méthodologie", "Démarche environnementale"]
    assert all(x in plat for x in ordre) and [plat.index(x) for x in ordre] == sorted(plat.index(x) for x in ordre)
    for x in ("Un conducteur de travaux encadre", "nacelle élévatrice de 18 m", METHODO_A, "étiquetées",
              "Le planning des rotations est joint au dossier.", "Breizh Désamiantage", "Le Gall Électricité"):
        assert x in plat, x
    for interdit in ("[F1]", "[F", "[S", "[P", "justifiée", "déclarée", "à revalider", "à valider", "Rédigé par le chef",
                     "Modifié par le chef", "validé par", "Note interne", "note interne", "À compléter", "Adjugo",
                     "€", SECRET_EXTRAIT, SECRET_DOC, SECRET_BROUILLON, SECRET_CHEF, "source du partenaire"):
        assert interdit not in plat, interdit

    # ── 13. Journal chaîné : chaque étape y est, chaque ligne porte son empreinte ──
    journal = _audit(project_id=pid)
    actions = [x[0] for x in journal]
    for act in ("redaction.repartition", "redaction.partie_confiee", "redaction.parts_demandees",
                "redaction.part_envoyee", "redaction.part_modifiee", "redaction.part_validee",
                "redaction.partie_groupement_validee", "redaction.dossier_valide", "atelier.export"):
        assert any(x.startswith(act) for x in actions), act
    assert all(x[4] for x in journal if x[0].startswith("redaction."))


def test_marche_sans_groupement_inchange(client, auth, fake_llm):
    """Un marché sans groupement : l'atelier d'une seule entreprise, exactement comme avant la vague 4 — aucun onglet,
    aucune partie tirée du groupement, aucune répartition, rien à assembler ni à harmoniser."""
    fake = fake_llm(lambda kw: (_ for _ in ()).throw(AssertionError("aucune IA pour lire l'atelier")))
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h, name="Marché seul")
    st = _creer(client, h, pid)
    st = _etat(client, h, pid)
    assert st["groupement"] == {"onglet": False, "actif": False, "membres": 0}
    assert [s.get("genere") for s in st["sections"]] == [None] * len(st["sections"]) and len(st["sections"]) == 4
    assert all(s["peut_ecrire"] and s["prive_visible"] and "part" not in s for s in st["sections"])
    assert all(s["auteur"]["type"] == "organisation" for s in st["sections"])
    assert client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["actif"] is False
    r = client.get(f"/api/atelier/{pid}/repartition/dossier", headers=h)
    assert r.status_code == 404 and r.json()["detail"]["code"] == "groupement_non_prepare"
    r = client.post(f"/api/atelier/{pid}/repartition/harmoniser", headers=h, json={})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "groupement_non_prepare"
    assert fake.calls == []
    from app.models import DossierRepartition, DossierValidation
    db = _db()
    try:
        assert db.query(DossierRepartition).filter(DossierRepartition.project_id == pid).count() == 0
        assert db.query(DossierValidation).filter(DossierValidation.project_id == pid).count() == 0
    finally:
        db.close()
