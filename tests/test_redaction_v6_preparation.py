"""
Rédaction, vague 6 — H : « Comment ce dossier a été préparé » (étude § 6.12, décision 9 a). Contrat :
docs/REDACTION_V6.md, section H ; code : app/services/atelier_preparation.py.

Invariants vérifiés :
 • la case est DÉCOCHÉE par défaut, jamais enregistrée dans les options : sans elle, ni le Word ni le PDF n'ont la
   page ; avec elle, la page est la DERNIÈRE du fichier ;
 • contenu exact, partie par partie : l'entreprise qui l'a rédigée, l'aide de l'IA (oui / non), le nombre de faits
   par type de source, qui l'a validée et quand ; une source perdue compte « sans source retrouvée », comme l'export
   la lit ;
 • aucune fuite : ni texte de source, ni nom de document, ni phrase d'un fait, ni nom de personne, ni adresse, ni
   nom d'outil, jamais « vérifié » ;
 • groupement : chaque membre et sa contribution réelle (partie modifiée par le chef de file, partie tirée de la
   convention ou des parts des membres, « Qui fait quoi » validée par chacun), jamais un brouillon d'un partenaire
   ni ce qu'il a relu lui-même ; rien dans ce que lit le partenaire ;
 • sans IA : aucun appel d'IA pour la page ; la page et la mention « Préparé avec assistance IA » des propriétés du
   fichier disent toujours la même chose ; le mode « sans IA » de la vague 5 est dit ;
 • démo : la page se lit (GET, sans écriture), l'export reste refusé.
Aucun appel réseau : le client IA est simulé (et échoue s'il est appelé).
"""
import io
import json
import re
import uuid

import pytest

from app.services import atelier as at
from app.services import atelier_preparation as prep
from tests.test_redaction_socle0 import NOMS, _company, _db, fake_llm  # noqa: F401
from tests.test_redaction_v2_atelier import _creer, _ecrire, _etat, _projet

TITRE = "Comment ce dossier a été préparé"


# ── Aides ───────────────────────────────────────────────────────────────────────────

def _docx(contenu: bytes):
    import docx
    return docx.Document(io.BytesIO(contenu))


def _docx_texte(contenu: bytes) -> str:
    d = _docx(contenu)
    lignes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lignes.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lignes)


def _pdf_pages(contenu: bytes) -> list:
    from pypdf import PdfReader
    return [re.sub(r"\s+", " ", p.extract_text() or "") for p in PdfReader(io.BytesIO(contenu)).pages]


def _plat(s: str) -> str:
    return re.sub(r"\s+", " ", s or "")


def _export(client, h, pid, fmt="docx", **extra):
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json=dict({"format": fmt}, **extra))
    assert r.status_code == 200, r.text
    return r.content


def _apercu(client, h, pid):
    r = client.get(f"/api/atelier/{pid}/preparation", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _valider(client, h, pid, sid):
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sid)
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/valider", headers=h,
                    json={"version_id": s["version_courante"]["id"]})
    assert r.status_code == 200, r.text


def _valide_le(sid) -> str:
    from app.models import DossierSection
    db = _db()
    try:
        return prep.date_fr(db.get(DossierSection, sid).validated_at)
    finally:
        db.close()


def _fait(fid, statut, type_source=None, kind="chiffre", **src):
    """Un fait relevé (forme de A). Ce qui est privé est marqué pour prouver qu'il ne sort jamais."""
    x = {"id": fid, "fait": f"FAIT-PRIVE-{fid}", "type": kind, "phrase": f"PHRASE-PRIVEE-{fid}",
         "message": "Adjugo ne trouve pas d'où vient ce fait : relisez-le.", "statut": statut, "a_propos": "entreprise",
         "source": None}
    if type_source:
        x["source"] = dict({"type": type_source, "ref": src.pop("ref", ""), "libelle": f"LIBELLE-PRIVE-{fid}",
                            "doc_name": "Registre-interne-XYZ.pdf", "extrait": "EXTRAIT-PRIVE-XYZ"}, **src)
    return x


def _injecter(pid, sid, texte, faits, ia=True, sources=None):
    """Une version « Rédigé avec Adjugo » contrôlée (ses faits) : comme après une rédaction de l'IA."""
    from app.models import DossierSection
    db = _db()
    try:
        s = db.get(DossierSection, sid)
        at.ajouter_version(db, s, texte=texte, origine=at.IA, ia=ia, faits=faits,
                           sources=sources if sources is not None else
                           [{"ref": "S1", "doc_name": "Registre-interne-XYZ.pdf", "chunk_id": None,
                             "extrait": "EXTRAIT-PRIVE-XYZ"}])
        db.commit()
    finally:
        db.close()


@pytest.fixture
def ia_interdite(fake_llm):
    """Le client d'IA échoue s'il est appelé : la page se construit sans IA."""
    def ans(_kw):
        raise AssertionError("aucun appel d'IA n'est permis pour la page « Comment ce dossier a été préparé »")
    return fake_llm(ans)


INTERDITS = ("FAIT-PRIVE", "PHRASE-PRIVEE", "LIBELLE-PRIVE", "Registre-interne-XYZ", "EXTRAIT-PRIVE", "Adjugo",
             "adjugo", "vérifié", "Vérifié", "Eliot Viegas", "1 rue du Port", "[S1]", "[F", "À compléter", "relisez")


# ═══ 1. La case : décochée par défaut, jamais enregistrée, la page à la fin ════════════

def test_case_decochee_par_defaut_et_page_seulement_sur_demande(client, auth, ia_interdite):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    _creer(client, h, pid)
    st = _etat(client, h, pid)
    for s in st["sections"]:
        v = _ecrire(client, h, pid, s, "Nous intervenons par zones, en site occupé.")
        r = client.post(f"/api/atelier/{pid}/sections/{s['id']}/valider", headers=h,
                        json={"version_id": v["version_courante"]["id"]})
        assert r.status_code == 200, r.text
    st = _etat(client, h, pid)
    assert st["export"]["page_preparation"] == {"libelle": "Joindre une page « Comment ce dossier a été préparé »",
                                                "bulle": prep.BULLE, "par_defaut": False}
    assert "votre décision" in prep.BULLE and "« Rédiger sans IA »" in prep.BULLE
    assert st["export"]["options"] == {"page_garde": True, "sommaire": True}     # contrat de la vague 2 inchangé
    # Sans la case (absente ou fausse) : ni le Word ni le PDF n'ont la page.
    for corps in ({}, {"page_preparation": False}):
        assert TITRE not in _docx_texte(_export(client, h, pid, "docx", **corps))
        assert all(TITRE not in p for p in _pdf_pages(_export(client, h, pid, "pdf", **corps)))
    # Les options ne l'enregistrent jamais : l'export suivant est sans la page.
    r = client.put(f"/api/atelier/{pid}/options", headers=h, json={"page_preparation": True})
    assert r.status_code == 200 and r.json()["options"] == {"page_garde": True, "sommaire": True}
    assert TITRE not in _docx_texte(_export(client, h, pid, "docx"))
    from app.models import DossierPresentation
    db = _db()
    try:
        d = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).first()
        assert d.dernier_export["page_preparation"] is False and "page_preparation" not in (d.options_export or {})
    finally:
        db.close()
    # Avec la case : la page est la DERNIÈRE du fichier, avec son tableau.
    d = _docx(_export(client, h, pid, "docx", page_preparation=True))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads[-1] == TITRE and heads[-2] == "3. Démarche environnementale"
    t = d.tables[-1]
    assert [c.text for c in t.rows[0].cells] == ["Partie", "Rédigée par", "Aide de l'IA", "Faits et sources", "Validée"]
    assert len(t.rows) == 1 + len(st["sections"])
    pages = _pdf_pages(_export(client, h, pid, "pdf", page_preparation=True))
    assert TITRE in pages[-1] and all(TITRE not in p for p in pages[:-1])
    assert "Démarche environnementale" in pages[-1] and "Breizh Couverture" in pages[-1]
    # Inscrit au dernier export et au journal chaîné.
    from app.models import AuditLog
    db = _db()
    try:
        d = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).first()
        assert d.dernier_export["page_preparation"] is True and d.dernier_export["format"] == "pdf"
        e = (db.query(AuditLog).filter(AuditLog.action == "atelier.export", AuditLog.project_id == pid)
             .order_by(AuditLog.id.desc()).first())
        assert e.meta["page_preparation"] is True
    finally:
        db.close()
    assert ia_interdite.calls == []


# ═══ 2. Contenu exact, sans fuite (une seule entreprise) ══════════════════════════════

def test_contenu_exact_par_partie_sans_fuite(client, auth, ia_interdite):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    _creer(client, h, pid)
    secs = {s["intitule"]: s for s in _etat(client, h, pid)["sections"]}
    enc, equ = secs["Encadrement du chantier"], secs["Équipe d'exécution"]
    meth, env = secs["Méthodologie"], secs["Démarche environnementale"]
    # 1.1 : écrite par l'entreprise elle-même, sans IA, sans fait.
    _ecrire(client, h, pid, enc, "Un conducteur de travaux suit le chantier chaque jour.")
    _valider(client, h, pid, enc["id"])
    # 1.2 : préparée avec l'IA ; des faits de chaque type de source, un relu par l'entreprise, un engagement
    # (pas un fait) ; validée.
    _injecter(pid, equ["id"], "Notre équipe est présentée ci-dessous.", [
        _fait("f1", "avec_source", "documents", ref="S1"),
        _fait("f2", "avec_source", "documents", ref="S2"),
        _fait("f3", "avec_source", "justificatif", ref="F1"),
        _fait("f4", "avec_source", "declare", ref="P1"),
        _fait("f5", "avec_source", "vous"),
        _fait("f6", "avec_source", "partenaire", ref="G1"),
        _fait("f7", "avec_source", "public", ref="F2"),
        _fait("f8", "avec_source", "estimation", ref="C1"),
        _fait("f9", "avec_source", "calcul"),
        _fait("f10", "avec_source", "acheteur"),
        _fait("f11", "relu_par_vous"),
        _fait("f12", "relu_par_vous", kind="engagement"),
    ])
    _valider(client, h, pid, equ["id"])
    # 2 : préparée avec l'IA, un fait relié à un extrait qui n'existe plus (source perdue : sans source, comme à
    # l'export) ; validée.
    _injecter(pid, meth["id"], "Nous phasons le chantier.", [
        _fait("g1", "avec_source", "documents", ref="S1", chunk_id=987654321),
        _fait("g2", "avec_source", "declare", ref="P1"),
    ])
    _valider(client, h, pid, meth["id"])
    # 3 : écrite par l'entreprise, jamais validée.
    _ecrire(client, h, pid, env, "Nous trions les déchets sur place.")
    ap = _apercu(client, h, pid)
    page = ap["page"]
    assert page["titre"] == TITRE and page["introduction"] == prep.INTRODUCTION and page["legende"] == prep.LEGENDE
    assert page["colonnes"] == ["Partie", "Rédigée par", "Aide de l'IA", "Faits et sources", "Validée"]
    lignes = [{k: x[k] for k in ("partie", "redigee_par", "ia_libelle", "faits_libelle", "validee_libelle")}
              for x in page["lignes"]]
    assert lignes == [
        {"partie": "1.1 Encadrement du chantier", "redigee_par": "Breizh Couverture", "ia_libelle": "Non",
         "faits_libelle": "Aucun fait relevé",
         "validee_libelle": f"Par Breizh Couverture, le {_valide_le(enc['id'])}"},
        {"partie": "1.2 Équipe d'exécution", "redigee_par": "Breizh Couverture", "ia_libelle": "Oui",
         "faits_libelle": ("10 faits reliés à une source : 3 tirés de ses documents, 2 déclarés par l'entreprise, "
                           "1 déclaré par un membre du groupement, 1 de source publique, 1 selon son estimation, "
                           "1 calculé à partir de ces faits, 1 tiré du dossier de l'acheteur ; 1 fait relu par "
                           "l'entreprise, sans autre source"),
         "validee_libelle": f"Par Breizh Couverture, le {_valide_le(equ['id'])}"},
        {"partie": "2. Méthodologie", "redigee_par": "Breizh Couverture", "ia_libelle": "Oui",
         "faits_libelle": "1 fait relié à une source : 1 déclaré par l'entreprise ; 1 fait sans source retrouvée",
         "validee_libelle": f"Par Breizh Couverture, le {_valide_le(meth['id'])}"},
        {"partie": "3. Démarche environnementale", "redigee_par": "Breizh Couverture", "ia_libelle": "Non",
         "faits_libelle": "Aucun fait relevé", "validee_libelle": "Non validée"},
    ]
    assert page["lignes"][1]["faits"] == {"avec_source": 10, "sans_source": 0, "relus": 1, "par_type": {
        "documents": 3, "declare": 2, "partenaire": 1, "public": 1, "estimation": 1, "calcul": 1, "acheteur": 1}}
    assert page["resume"] == ["Parties préparées avec l'aide de l'IA : 2 sur 4.",
                              "Le fichier porte dans ses propriétés la mention « Préparé avec assistance IA ».",
                              "Parties validées : 3 sur 4."]
    assert [w["code"] for w in ap["avertissements"]] == ["non_validees", "sans_source", "limite_pages"]
    assert ap["avertissements"][2]["message"] == ("Le règlement limite le dossier à 8 pages : cette page s'ajoute à "
                                                   "la fin du fichier et n'entre pas dans le compteur.")
    # Rien de privé, aucun nom de personne, aucun nom d'outil, jamais « vérifié ».
    brut = json.dumps(page, ensure_ascii=False)
    for x in INTERDITS + tuple(NOMS):
        assert x not in brut, x
    assert auth["email"] not in brut
    # L'export (dérogation pour le fait sans source et la partie non validée) écrit EXACTEMENT ce que l'écran montre.
    gate = _etat(client, h, pid)["blocage"]
    assert [r_["code"] for r_ in gate["raisons"]] == ["faits_a_relire"]
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "relu_moi_meme", "texte": "Relu avec le conducteur de travaux.",
                          "etat_hash": gate["etat_hash"]})
    assert r.status_code == 200, r.text
    data = _export(client, h, pid, "docx", page_preparation=True)
    d = _docx(data)
    assert d.core_properties.subject == "Préparé avec assistance IA"         # même chose que la page
    t = d.tables[-1]
    rows = [[c.text for c in row.cells] for row in t.rows[1:]]
    assert rows == [[x["partie"], x["redigee_par"], x["ia_libelle"], x["faits_libelle"], x["validee_libelle"]]
                    for x in lignes]
    texte = _docx_texte(data)
    page_txt = texte[texte.index(TITRE):]
    for x in page["resume"] + [prep.INTRODUCTION, prep.LEGENDE]:
        assert x in page_txt, x
    for x in INTERDITS + tuple(NOMS):
        assert x not in texte, x
    pages = _pdf_pages(_export(client, h, pid, "pdf", page_preparation=True))
    for x in ("1 fait sans source retrouvée", "Non validée", "Parties préparées avec l'aide de l'IA : 2 sur 4."):
        assert x in pages[-1], x
    for x in INTERDITS:
        assert all(x not in p for p in pages), x
    assert ia_interdite.calls == []


def test_page_et_mention_ia_toujours_d_accord(client, auth, ia_interdite):
    """Aucune partie préparée avec l'IA : la page le dit et le fichier ne porte aucune mention ; une seule partie de
    l'IA : la page dit « Oui » pour elle et le fichier porte la mention."""
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h, with_dce=False)
    _creer(client, h, pid)
    st = _etat(client, h, pid)
    for s in st["sections"]:
        _ecrire(client, h, pid, s, "Nous intervenons par zones.")
        _valider(client, h, pid, s["id"])
    page = _apercu(client, h, pid)["page"]
    assert page["parties_ia"] == 0 and page["mention_ia"] == ""
    assert page["resume"][0] == "Aucune partie n'a été préparée avec l'aide de l'IA."
    assert not any("propriétés" in x for x in page["resume"]) and {x["ia_libelle"] for x in page["lignes"]} == {"Non"}
    d = _docx(_export(client, h, pid, "docx", page_preparation=True))
    assert d.core_properties.subject == "" and d.core_properties.keywords == ""
    # Une partie préparée avec l'IA.
    s0 = st["sections"][0]
    _injecter(pid, s0["id"], "Nous intervenons par zones, avec un balisage quotidien.", [])
    _valider(client, h, pid, s0["id"])
    page = _apercu(client, h, pid)["page"]
    assert page["parties_ia"] == 1 and page["mention_ia"] == "Préparé avec assistance IA"
    assert [x["ia_libelle"] for x in page["lignes"]][0] == "Oui"
    assert page["resume"][:2] == [f"Parties préparées avec l'aide de l'IA : 1 sur {len(st['sections'])}.",
                                  "Le fichier porte dans ses propriétés la mention « Préparé avec assistance IA »."]
    assert _etat(client, h, pid)["export"]["mention_ia"] == page["mention_ia"]
    d = _docx(_export(client, h, pid, "docx", page_preparation=True))
    assert d.core_properties.subject == "Préparé avec assistance IA"
    assert ia_interdite.calls == []


def test_mode_sans_ia_dit_sur_la_page(client, auth, ia_interdite):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    _creer(client, h, pid)
    st = _etat(client, h, pid)
    s0 = st["sections"][0]
    _injecter(pid, s0["id"], "Texte préparé avec l'IA avant le choix du mode sans IA.", [])
    r = client.put(f"/api/atelier/{pid}/options", headers=h, json={"sans_ia": True})
    assert r.status_code == 200 and r.json()["sans_ia"]["actif"] is True
    le = r.json()["sans_ia"]["le"]
    from datetime import datetime
    page = _apercu(client, h, pid)["page"]
    assert (f"L'entreprise a choisi de rédiger ce dossier sans IA à partir du "
            f"{prep.date_fr(datetime.fromisoformat(le))}.") in page["resume"]
    # La partie préparée avant reste « Oui » : la page et la mention du fichier ne mentent jamais.
    assert page["lignes"][0]["ia_libelle"] == "Oui" and page["mention_ia"] == "Préparé avec assistance IA"
    assert ia_interdite.calls == []


# ═══ 3. Groupement : la contribution réelle de chaque membre, rien de privé ═══════════

def test_groupement_contribution_de_chaque_membre_sans_rien_de_prive(client, auth, mails, ia_interdite):
    from tests.test_redaction_v4_assemblage import (_generee, _groupement_convention, _valider_chef, _valider_part,
                                                    _valider_qfq, _ecrire_chef)
    from tests.test_redaction_v4_groupement import P, _ecrire as _ecrire_p, _envoyer, _part_a
    h = auth["headers"]
    pid, a, _b, vue, _conv = _groupement_convention(client, auth)
    pa = P(client, a["token"])
    sid = _part_a(vue)["section_id"]
    _ecrire_p(pa, sid, "Nous trions les déchets amiantés sur place. Notre équipe compte 9 salariés.")
    _envoyer(pa, sid)
    _valider_part(pa, sid)
    assert _valider_qfq(pa).status_code == 200
    st = _etat(client, h, pid)
    for s in st["sections"]:
        if s["auteur"]["type"] == "organisation":
            _ecrire_chef(client, h, pid, s["id"], f"Notre réponse pour {s['intitule']}.", valider=True)
    for kind in ("grp_orga", "grp_qfq"):
        assert _valider_chef(client, h, pid, _generee(st, kind)["id"]).status_code == 200
    # Un brouillon PRIVÉ du partenaire, jamais envoyé : il n'entre nulle part.
    _ecrire_p(pa, sid, "BROUILLON-PRIVE-MARQUEUR : 47 salariés et une nacelle de 40 m.")
    ap = _apercu(client, h, pid)
    page = ap["page"]
    par = {x["partie"]: x for x in page["lignes"]}
    orga, qfq = par["Organisation du groupement"], par["Qui fait quoi dans le groupement"]
    assert (orga["redigee_par"], orga["ia_libelle"], orga["faits_libelle"]) == (
        "Le groupement, d'après sa convention", "Non", "Repris de la convention de groupement")
    assert (qfq["redigee_par"], qfq["faits_libelle"]) == ("Le groupement, d'après les parts des membres",
                                                         "Repris des parts déclarées par les membres")
    assert orga["validee_libelle"].startswith("Par Breizh Couverture, le ")
    assert qfq["validee_libelle"].startswith("Par Breizh Couverture et Breizh Désamiantage, le ")
    env = next(x for x in page["lignes"] if x["section_id"] == sid)
    assert env["redigee_par"] == "Breizh Désamiantage" and env["ia_libelle"] == "Non"
    assert env["validee_libelle"] == f"Par Breizh Désamiantage, le {_valide_le(sid)}"
    assert env["faits"]["relus"] is None          # ce qu'il a relu lui-même ne sort pas de chez lui
    assert env["faits_libelle"] == "1 fait relié à une source : 1 déclaré par l'entreprise"   # « 9 salariés », sa part
    autres = [x for x in page["lignes"] if x["section_id"] != sid and x["partie"] not in par.keys() & {
        "Organisation du groupement", "Qui fait quoi dans le groupement"}]
    assert len(autres) == 3 and {x["redigee_par"] for x in autres} == {"Breizh Couverture"}
    assert all(x["validee_libelle"] == f"Par Breizh Couverture, le {_valide_le(x['section_id'])}" for x in autres)
    assert page["resume"] == [
        "Aucune partie n'a été préparée avec l'aide de l'IA.",
        "Parties rédigées : Breizh Couverture, 3 ; Breizh Désamiantage, 1 ; le groupement (convention et parts des "
        "membres), 2.",
        "Parties validées : 6 sur 6."]
    brut = json.dumps(ap, ensure_ascii=False)
    for x in INTERDITS + ("BROUILLON-PRIVE", "47 salariés", "nacelle", "40 000", "€", "Jean Le Gall", "Note interne"):
        assert x not in brut, x
    # Le chef de file modifie la part du partenaire : la page le dit, et elle n'est plus validée.
    cur = client.get(f"/api/atelier/{pid}/repartition/{sid}", headers=h).json()["texte_commun"]["version_id"]
    r = client.put(f"/api/atelier/{pid}/repartition/{sid}/texte", headers=h,
                   json={"texte": "Nous trions les déchets amiantés sur place, en big-bags.", "base_version_id": cur})
    assert r.status_code == 200, r.text
    env = next(x for x in _apercu(client, h, pid)["page"]["lignes"] if x["section_id"] == sid)
    assert env["redigee_par"] == "Breizh Désamiantage ; modifiée par Breizh Couverture"
    assert env["validee_libelle"] == "Non validée"
    # Rien dans ce que lit le partenaire, et la page n'est lue que par l'organisation du chef de file.
    for chemin in ("", "/dossier", f"/sections/{sid}", "/export"):
        r = pa.get(chemin)
        assert r.status_code == 200, (chemin, r.text)
        assert TITRE not in r.text and "Rédigée par" not in r.text, chemin
    autre = _autre_compte(client)
    assert client.get(f"/api/atelier/{pid}/preparation", headers=autre).status_code == 404
    assert client.get(f"/api/atelier/{pid}/preparation").status_code in (401, 403)
    assert ia_interdite.calls == []


def _autre_compte(client):
    email = f"autre_{uuid.uuid4().hex[:8]}@test.fr"
    r = client.post("/api/auth/register", json={"email": email, "password": "motdepasse123", "full_name": "Autre",
                                                "company_name": "Autre entreprise"})
    assert r.status_code == 201, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_page_impossible_jamais_un_fichier_sans_elle(client, auth, monkeypatch):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h, with_dce=False)
    _creer(client, h, pid)
    for s in _etat(client, h, pid)["sections"]:
        _ecrire(client, h, pid, s, "Nous intervenons par zones.")
        _valider(client, h, pid, s["id"])

    def panne(*_a, **_k):
        raise RuntimeError("panne")
    monkeypatch.setattr(prep, "page", panne)
    for fmt in ("docx", "pdf"):
        r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": fmt, "page_preparation": True})
        assert r.status_code == 503 and r.json()["detail"]["code"] == "page_preparation_indisponible", r.text
    r = client.get(f"/api/atelier/{pid}/preparation", headers=h)
    assert r.status_code == 503 and r.json()["detail"]["message"] == prep.MSG_INDISPONIBLE
    from app.models import Document
    db = _db()
    try:
        assert db.query(Document).filter(Document.project_id == pid, Document.folder == "Mémoire technique").count() == 0
    finally:
        db.close()
    assert TITRE not in _docx_texte(_export(client, h, pid, "docx"))      # sans la case : l'export marche


# ═══ 4. Démo : la page se lit, l'export reste refusé, rien n'est écrit ═══════════════

def test_demo_lit_la_page_sans_rien_ecrire(client, ia_interdite):
    r = client.post("/api/auth/demo")
    assert r.status_code == 200, r.text
    hd = {"Authorization": f"Bearer {r.json()['access_token']}"}
    from app.models import DossierPresentation, Project, User
    from app.services.demo_seed import DEMO_EMAIL
    db = _db()
    try:
        u = db.query(User).filter(User.email == DEMO_EMAIL).one()
        p = next(x for x in db.query(Project).filter(Project.user_id == u.id) if x.name.startswith("Réfection toiture"))
        toiture = {"id": p.id}
        avant = db.query(DossierPresentation).filter(DossierPresentation.project_id == p.id).one()
        export_avant = avant.dernier_export
    finally:
        db.close()
    ap = _apercu(client, hd, toiture["id"])
    page = ap["page"]
    assert page["parties"] >= 4 and page["parties_ia"] >= 3 and page["parties_validees"] == 3
    assert any(x["faits"] and x["faits"]["avec_source"] for x in page["lignes"])
    brut = json.dumps(ap, ensure_ascii=False)
    for x in ("Adjugo", "vérifié", "[S", "[F", "doc_name", "extrait"):
        assert x not in brut, x
    r = client.post(f"/api/atelier/{toiture['id']}/export", headers=hd,
                    json={"format": "docx", "page_preparation": True})
    assert r.status_code == 403 and "lecture guidée" in json.dumps(r.json(), ensure_ascii=False)
    db = _db()
    try:
        apres = db.query(DossierPresentation).filter(DossierPresentation.project_id == toiture["id"]).first()
        assert apres.dernier_export == export_avant
    finally:
        db.close()
    assert ia_interdite.calls == []


def test_ecran_case_bulle_et_page_lue_avant_l_export():
    import pathlib
    root = pathlib.Path(__file__).resolve().parent.parent / "app" / "static"
    html = (root / "app.html").read_text(encoding="utf-8")
    js = (root / "app.js").read_text(encoding="utf-8")
    assert 'v-model="atl.prepJoindre"' in html and "atl.st.export.page_preparation.libelle" in html
    assert "atl.st.export.page_preparation.bulle" in html and "atlPrepVoir()" in html
    assert "prepJoindre: false" in js and "page_preparation: !!a.prepJoindre" in js
    assert '"/preparation"' in js
    # Interface : français, sans émoji.
    from tests.zones import zone
    bloc = zone(html, "atelier-page-preparation-case") + zone(html, "atelier-page-preparation")
    assert not re.search("[\U0001F300-\U0001FAFF☀-➿]", bloc)


@pytest.fixture
def mails(monkeypatch):
    sent = []
    import app.routers.invites as inv
    monkeypatch.setattr(inv, "_send_mails_async", lambda m: sent.extend(m))
    return sent
