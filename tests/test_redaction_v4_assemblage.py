"""
Rédaction, vague 4 — chantier 3 : ASSEMBLAGE et VALIDATION du dossier de groupement (étude § 6.8 B, parcours 4
et 5). Contrat : docs/REDACTION_V4_GROUPEMENT.md § 15 ; code : app/services/atelier_assemblage.py (derrière
app/services/atelier_groupement.py).

Invariants vérifiés :
 • assemblage SANS IA, dans l'ordre du cadre de l'acheteur, avec deux parties TIRÉES DU GROUPEMENT
   (« Organisation du groupement », « Qui fait quoi dans le groupement ») placées en tête (ou après la partie du
   plan qui parle du groupement), jamais écrites à la main, retirables ;
 • « Organisation du groupement » tirée de la convention EN VIGUEUR (forme, chef de file et ses pouvoirs,
   répartition, jamais un prix) ; convention absente, en brouillon ou périmée : dit honnêtement
   (« [À compléter : …] », qui bloque l'export), rien d'inventé ; validée par le chef de file une fois la
   convention validée par tous ;
 • « Qui fait quoi » : ce que chaque membre APPORTE (part, qualifications « justifiée » / « déclarée » à
   l'écran seulement, moyens), JAMAIS ce qui manque au chef de file ; la justification du groupement reste une
   note interne au chef de file ; validée par CHAQUE membre ;
 • « Harmoniser le ton » : des phrases de transition seulement, au moins 20 faits ajoutés de formes variées
   refusés sans IA, le texte des membres jamais réécrit, les parties touchées « à revalider », débit d'une
   refonte rendu si l'IA tombe ou si rien n'est retenu, noms des salariés masqués ;
 • validation : chaque membre valide ses parties et « Qui fait quoi » ; « Sections validées 4 sur 5 » et ce qui
   manque ; toute modification après validation → « à revalider » ; journal chaîné ;
 • export Word et PDF du dossier assemblé complet, sans marque interne ni renvoi [F#], bloqué (dérogation
   journalisée) par les parts des partenaires et par les parties tirées du groupement.
Aucun appel réseau : le client IA est simulé.
"""
import io
import json
import re
import time
import uuid

import pytest

from app.services import atelier_assemblage as asm
from app.services import atelier_groupement as ag
from app.services import llm
from tests.test_convention_groupement import _gen, _invite, _put, _register, _submit, _valider
from tests.test_redaction_socle0 import _company, _db, _user_id, fake_llm  # noqa: F401
from tests.test_redaction_v2_atelier import _creer, _etat, _projet, _usage
from tests.test_redaction_v4_groupement import (LOT_A, LOT_B, P, _audit, _ecrire, _envoyer, _groupement, _part_a,
                                                _section_par)

LOT_CHEF = "Couverture et zinguerie"


# ── Aides ───────────────────────────────────────────────────────────────────────────

def _generee(st, kind):
    return next(s for s in st["sections"] if s.get("genere") == kind)


def _valider_chef(client, h, pid, sid):
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sid)
    return client.post(f"/api/atelier/{pid}/sections/{sid}/valider", headers=h,
                       json={"version_id": s["version_courante"]["id"]})


def _ecrire_chef(client, h, pid, sid, texte, valider=False):
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sid)
    r = client.put(f"/api/atelier/{pid}/sections/{sid}/texte", headers=h,
                   json={"texte": texte, "base_version_id": (s.get("version_courante") or {}).get("id")})
    assert r.status_code == 200, r.text
    if valider:
        r2 = client.post(f"/api/atelier/{pid}/sections/{sid}/valider", headers=h,
                         json={"version_id": r.json()["version_courante"]["id"]})
        assert r2.status_code == 200, r2.text
    return r.json()


def _valider_part(p, sid):
    cur = p.get(f"/sections/{sid}").json()["texte_commun"]["version_id"]
    r = p.post(f"/sections/{sid}/valider", {"version_id": cur})
    assert r.status_code == 200, r.text
    return r.json()


def _valider_qfq(p):
    q = p.get().json()["qui_fait_quoi"]
    return p.post(f"/sections/{q['section_id']}/valider", {"version_id": q["texte_commun"]["version_id"]})


def _parts_soumises(client, invs, prix=(40000.0, 30000.0)):
    for i, inv in enumerate(invs):
        r = _put(client, inv["token"], prix_lot_ht=prix[i % len(prix)], representant_legal="Jean Le Gall",
                 signataire_qualite="Gérant", siret="1" * 9 + str(10000 + i))
        assert r.status_code == 200, r.text
        _submit(client, inv["token"])


def _convention_validee(client, auth, pid, invs, *, forme="conjoint", tous=True, solidaire=False):
    """Parts soumises (avec leur prix), convention conjointe générée, accord du chef de file et (tous=True) de
    chaque partenaire : la convention est alors « validée par tous les membres »."""
    h = auth["headers"]
    _parts_soumises(client, invs)
    r = client.put(f"/api/projects/{pid}/groupement/convention", headers=h, json={
        "forme": forme, "mandataire_solidaire": solidaire,
        "params": {"part_mandataire_ht": 60000, "lot_mandataire": LOT_CHEF, "mandataire_qualite": "Président"}})
    assert r.status_code == 200, r.text
    r = _gen(client, h, pid)
    assert r.status_code == 200, r.text
    conv = r.json()
    assert _valider(client, f"/api/projects/{pid}/groupement/convention/valider", conv, headers=h,
                    nom="Eliot Viegas").status_code == 200
    if tous:
        for inv in invs:
            r = _valider(client, f"/api/invite/{inv['token']}/convention/valider", conv)
            assert r.status_code == 200, r.text
    return conv


def _groupement_convention(client, auth, avec_b=False, tous=True):
    pid, st, a, b, vue = _groupement(client, auth, avec_b=avec_b)
    conv = _convention_validee(client, auth, pid, [x for x in (a, b) if x], tous=tous)
    return pid, a, b, vue, conv


def _attendre_chef(client, h, pid, jid):
    r = None
    for _ in range(300):
        r = client.get(f"/api/atelier/{pid}/taches/{jid}", headers=h)
        assert r.status_code == 200, r.text
        if r.json()["status"] in ("done", "error"):
            return r.json()
        time.sleep(0.05)
    raise AssertionError(f"tâche {jid} jamais terminée : {r.json() if r is not None else None}")


def _docx_texte(contenu: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(contenu))
    lignes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lignes.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lignes)


def _pdf_texte(contenu: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(contenu)).pages)


# ═══ 1. L'assemblage suit l'ordre du cadre, avec les parties tirées du groupement ═════

def test_ordre_du_cadre_et_parties_tirees_du_groupement(client, auth):
    h = auth["headers"]
    pid, st0, a, b, vue = _groupement(client, auth)
    st = _etat(client, h, pid)
    assert [(s["numero_affiche"], s["intitule"]) for s in st["sections"]] == [
        ("", "Organisation du groupement"), ("", "Qui fait quoi dans le groupement"),
        ("1.1", "Encadrement du chantier"), ("1.2", "Équipe d'exécution"), ("2", "Méthodologie"),
        ("3", "Démarche environnementale")]
    orga, qfq = _generee(st, "grp_orga"), _generee(st, "grp_qfq")
    for s in (orga, qfq):
        assert s["auteur"]["type"] == "groupement" and s["peut_ecrire"] is False and s["prive_visible"] is False
        assert s["version_courante"]["texte"] and s["version_courante"]["ia"] is False
        assert s["version_courante"]["origine"] == "genere"
    assert orga["auteur"]["libelle"] == "Tirée de la convention de groupement"
    # « Qui écrit quoi », le dossier assemblé du partenaire et celui du chef de file : le même ordre.
    rep = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()
    assert [x["intitule"] for x in rep["sections"]] == [s["intitule"] for s in st["sections"]]
    assert rep["sections"][0]["redacteur"]["type"] == "groupement" and rep["sections"][0]["choix"] == "genere"
    assert rep["statuts"]["a_valider"] == "À valider"
    pa = P(client, a["token"])
    assert [x["intitule"] for x in pa.get("/dossier").json()["parties"]] == [s["intitule"] for s in st["sections"]]
    chef = client.get(f"/api/atelier/{pid}/repartition/dossier", headers=h).json()
    assert [x["intitule"] for x in chef["parties"]] == [s["intitule"] for s in st["sections"]]
    # Préparer de nouveau : idempotent, aucune partie en double.
    assert client.post(f"/api/atelier/{pid}/repartition", headers=h, json={}).json()["preparation"][
        "parties_du_groupement"] == 0
    assert len(_etat(client, h, pid)["sections"]) == 6
    # Personne ne les écrit à la main ; elles n'ont pas de rédacteur.
    r = client.put(f"/api/atelier/{pid}/sections/{orga['id']}/texte", headers=h,
                   json={"texte": "Texte à la main.", "base_version_id": orga["version_courante"]["id"]})
    assert r.status_code == 403 and r.json()["detail"]["code"] == "partie_du_groupement"
    r = client.post(f"/api/atelier/{pid}/sections/{qfq['id']}/rediger", headers=h, json={})
    assert r.status_code == 403 and r.json()["detail"]["code"] == "partie_du_groupement"
    r = client.patch(f"/api/atelier/{pid}/sections/{qfq['id']}", headers=h, json={"intitule": "Autre"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "intitule_impose"
    r = client.put(f"/api/atelier/{pid}/repartition/{orga['id']}", headers=h,
                   json={"redacteur": "partenaire", "invite_id": a["id"]})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "partie_du_groupement"
    r = client.put(f"/api/atelier/{pid}/repartition/{orga['id']}/texte", headers=h,
                   json={"texte": "x", "base_version_id": None})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "partie_du_groupement"
    # Le chef de file peut la retirer (elle n'est jamais recréée), puis la remettre à SA place.
    assert client.delete(f"/api/atelier/{pid}/sections/{qfq['id']}", headers=h).status_code == 200
    client.post(f"/api/atelier/{pid}/repartition", headers=h, json={})
    st = _etat(client, h, pid)
    assert [s["intitule"] for s in st["sections"]][:2] == ["Organisation du groupement", "Encadrement du chantier"]
    assert any(x["id"] == qfq["id"] for x in st["sections_retirees"])
    assert client.post(f"/api/atelier/{pid}/sections/{qfq['id']}/remettre", headers=h).status_code == 200
    assert [s["intitule"] for s in _etat(client, h, pid)["sections"]][:3] == [
        "Organisation du groupement", "Qui fait quoi dans le groupement", "Encadrement du chantier"]
    # Reconstruire le plan : elles gardent leur place, en tête.
    r = client.post(f"/api/atelier/{pid}/plan", headers=h, json={"base": "cadre", "remplacer": True})
    assert r.status_code == 200, r.text
    assert [s.get("genere") for s in r.json()["sections"][:2]] == ["grp_orga", "grp_qfq"]
    # Le quota de l'IA ne les compte pas : elles ne sont jamais rédigées par l'IA.
    assert r.json()["quota"]["premieres_restantes"] == len([x for x in r.json()["sections"] if not x["genere"]])


def test_plan_simple_parties_apres_la_presentation_du_groupement(client, auth):
    """Plan simple d'un marché en groupement : « Présentation du groupement » d'abord (écrite par le chef de
    file), puis les deux parties tirées du groupement, puis le reste du plan, dans son ordre."""
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h, with_dce=False)
    a = _invite(client, h, pid, "Breizh Désamiantage")
    assert _put(client, a["token"], company_name="Breizh Désamiantage", lot=LOT_A).status_code == 200
    _creer(client, h, pid, base="libre")
    assert client.post(f"/api/atelier/{pid}/repartition", headers=h, json={}).status_code == 200
    titres = [s["intitule"] for s in _etat(client, h, pid)["sections"]]
    assert titres[:4] == ["Présentation du groupement", "Organisation du groupement", "Qui fait quoi dans le groupement",
                          "Compréhension du besoin"]
    assert len(titres) == 9


# ═══ 2. « Organisation du groupement » : tirée de la convention en vigueur ════════════

def test_organisation_tiree_de_la_convention(client, auth, mails):
    h = auth["headers"]
    pid, a, b, vue, conv = _groupement_convention(client, auth, tous=False)
    # Convention générée, accord du partenaire en attente : le texte la reprend, la validation attend.
    st = _etat(client, h, pid)
    orga = _generee(st, "grp_orga")
    t = orga["version_courante"]["texte"]
    assert "**Forme du groupement.** Groupement conjoint : chacun est responsable de sa part." in t
    assert "**Chef de file (mandataire).** Breizh Couverture est le mandataire du groupement." in t
    assert "- représenter le groupement auprès de l'acheteur et coordonner le travail de tous les membres ;" in t
    assert f"- Breizh Couverture, chef de file (mandataire) : {LOT_CHEF}." in t
    assert f"- Breizh Désamiantage, co-traitant : {LOT_A}." in t
    # Jamais un prix (le dossier de présentation n'est pas l'offre de prix), rien d'inventé, rien d'interne.
    assert "€" not in t and "40 000" not in t and "[À compléter" not in t and "Adjugo" not in t
    r = _valider_chef(client, h, pid, orga["id"])
    assert r.status_code == 409 and r.json()["detail"]["code"] == "convention_a_valider"
    assert "Breizh Désamiantage" in r.json()["detail"]["message"]
    d = client.get(f"/api/atelier/{pid}/repartition/{orga['id']}", headers=h).json()
    assert d["convention"]["statut"] == "a_signer" and d["convention"]["accords_en_attente"] == ["Breizh Désamiantage"]
    assert d["statut"] == "a_valider" and d["valide_par"] == "chef_de_file"
    # Accord du partenaire : convention validée par tous → le chef de file valide la partie.
    assert _valider(client, f"/api/invite/{a['token']}/convention/valider", conv).status_code == 200
    r = _valider_chef(client, h, pid, orga["id"])
    assert r.status_code == 200, r.text
    assert r.json()["statut"] == "valide"
    d = client.get(f"/api/atelier/{pid}/repartition/{orga['id']}", headers=h).json()
    assert d["statut"] == "validee" and d["convention"]["toutes_validees"] is True
    assert "redaction.partie_groupement_validee" in [x[0] for x in _audit(project_id=pid)]
    # Le partenaire la lit dans le dossier assemblé, mais ne la valide pas (elle reprend la convention
    # qu'il a déjà validée) ; ni n'y écrit.
    pa = P(client, a["token"])
    part = _section_par(pa.get("/dossier").json()["parties"], "organisation du groupement")
    assert part["genere"] == "grp_orga" and part["statut"] == "validee" and LOT_A in part["texte"]
    assert pa.post(f"/sections/{orga['id']}/valider", {"version_id": orga["version_courante"]["id"]}).status_code == 404
    # Groupement solidaire : la forme suit la convention régénérée.
    r = client.put(f"/api/projects/{pid}/groupement/convention", headers=h, json={"forme": "solidaire"})
    assert r.status_code in (200, 409)


def test_organisation_convention_absente_brouillon_ou_perimee(client, auth, mails):
    """Pas de convention, convention en brouillon ou périmée : la partie le DIT (« [À compléter : …] », qui bloque
    l'export) au lieu d'inventer une forme ou une répartition ; elle repasse « à revalider » quand la convention
    devient caduque."""
    h = auth["headers"]
    pid, st0, a, b, vue = _groupement(client, auth, avec_b=False)
    st = _etat(client, h, pid)
    orga = _generee(st, "grp_orga")
    t = orga["version_courante"]["texte"]
    assert t.startswith("[À compléter : la convention de groupement n'est pas encore préparée.")
    assert "conjoint" not in t and "solidaire" not in t and "mandataire" not in t
    codes = {r["code"]: r for r in st["blocage"]["raisons"]}
    assert any(x["section_id"] == orga["id"] for x in codes["a_completer"]["sections"])
    r = _valider_chef(client, h, pid, orga["id"])
    assert r.status_code == 409 and r.json()["detail"]["code"] == "partie_incomplete"
    # Brouillon de convention : toujours pas de texte inventé.
    assert client.put(f"/api/projects/{pid}/groupement/convention", headers=h,
                      json={"forme": "conjoint"}).status_code == 200
    t = _generee(_etat(client, h, pid), "grp_orga")["version_courante"]["texte"]
    assert t.startswith("[À compléter : la convention de groupement n'est pas encore générée.")
    # Générée et validée par tous : le texte la reprend, le chef de file valide.
    _convention_validee(client, auth, pid, [a])
    orga = _generee(_etat(client, h, pid), "grp_orga")
    assert "Groupement conjoint" in orga["version_courante"]["texte"]
    assert _valider_chef(client, h, pid, orga["id"]).status_code == 200
    # Le partenaire change son prix : la convention est caduque, la partie le dit et repasse « à revalider ».
    assert _put(client, a["token"], prix_lot_ht=41000).status_code == 200
    orga = _generee(_etat(client, h, pid), "grp_orga")
    assert orga["version_courante"]["texte"].startswith("[À compléter : la convention de groupement n'est plus à jour")
    assert orga["statut"] == "a_revalider"
    rep = client.get(f"/api/atelier/{pid}/repartition/{orga['id']}", headers=h).json()
    assert rep["statut"] == "a_revalider" and rep["convention"]["statut"] == "perimee"
    assert "redaction.partie_groupement_a_jour" in [x[0] for x in _audit(project_id=pid)]
    # L'historique reste entier (aucune version effacée).
    from app.models import DossierSectionVersion
    db = _db()
    try:
        assert db.query(DossierSectionVersion).filter(DossierSectionVersion.section_id == orga["id"]).count() >= 4
    finally:
        db.close()


# ═══ 3. « Qui fait quoi » : ce que chacun apporte, jamais ce qui manque ══════════════

_GAP_EXIGE = "Qualification désamiantage SS3 exigée"
_GAP_A = "Non renseigné dans le profil du chef de file"


def _preparer_qfq(client, auth):
    """Chef de file : une qualification du profil (déclarée), une périmée, une preuve publique acceptée
    (justifiée), du matériel ; partenaire A : deux qualifications dont une avec justificatif joint ; B : sa part
    seule ; un sous-traitant ; une composition retenue dont les manques du chef de file sont connus."""
    from app.models import Company, ContributionPiece, Project, ProjectContribution, User
    from app.services import preuves as pv
    h = auth["headers"]
    pid, st0, a, b, vue = _groupement(client, auth)
    st_inv = client.post(f"/api/projects/{pid}/invites", headers=h,
                         json={"recipient": "", "company_name": "Sous-Traitant Échafaudages", "role": "sous_traitant"})
    assert st_inv.status_code == 200, st_inv.text
    assert _put(client, a["token"], qualifications=["Qualibat 1552", "Certification amiante SS3"],
                capacites={"ca_n1": 900000, "effectif": 9}).status_code == 200
    _submit(client, a["token"])
    _submit(client, b["token"])
    assert client.put(f"/api/projects/{pid}/groupement/convention", headers=h,
                      json={"params": {"lot_mandataire": LOT_CHEF}}).status_code == 200
    db = _db()
    try:
        u = db.query(User).filter(User.email == auth["email"]).first()
        c = db.query(Company).filter(Company.user_id == u.id).first()
        c.qualifications = [{"name": "Qualibat 3112", "expiration": "2099-01-01"},
                            {"name": "Ancienne qualification 9999", "expiration": "2001-01-01"}]
        pv.creer(db, u.id, nature=pv.QUALIFICATION, valeur={"nom": "ISO 14001", "organisme": "AFNOR"},
                 type_source=pv.PUBLIC, source={"nom": "Annuaire public", "url": "https://exemple.invalid/iso",
                                                "consulte_le": "2026-10-01"})
        pv.creer(db, u.id, nature=pv.MATERIEL, valeur={"designation": "Nacelle articulée"})
        contrib = db.query(ProjectContribution).filter(ProjectContribution.invite_id == a["id"]).first()
        db.add(ContributionPiece(contribution_id=contrib.id, project_id=pid, owner_id=u.id,
                                 name="Qualibat_1552_2026.pdf", file_key="x/qualibat.pdf", piece_type="qualification"))
        p = db.get(Project, pid)
        p.groupement_plan = {"auto": False, "score_seul": 43, "score_groupement": 81, "members": [],
                             "gaps_combles": [{"key": "qualifs", "label": "Qualifications exigées",
                                               "required": _GAP_EXIGE, "have": _GAP_A,
                                               "apporte_par": ["Breizh Désamiantage"]}]}
        db.commit()
    finally:
        db.close()
    return pid, a, b


def test_qui_fait_quoi_complementarite_sans_jamais_les_manques(client, auth):
    h = auth["headers"]
    pid, a, b = _preparer_qfq(client, auth)
    st = _etat(client, h, pid)
    qfq = _generee(st, "grp_qfq")
    t = qfq["version_courante"]["texte"]
    assert t.startswith("Le groupement réunit 3 entreprises.")
    assert "**Breizh Couverture**, chef de file (mandataire)" in t
    assert f"- Part : {LOT_CHEF}." in t
    assert "- Qualifications : ISO 14001, Qualibat 3112." in t
    assert "- Moyens : 12 salariés, Nacelle articulée." in t
    assert "**Breizh Désamiantage**, co-traitant" in t and f"- Part : {LOT_A}." in t
    assert "- Qualifications : Qualibat 1552, Certification amiante SS3." in t and "- Moyens : 9 salariés." in t
    assert "**Le Gall Électricité**, co-traitant" in t and f"- Part : {LOT_B}." in t
    # Le sous-traitant n'est pas membre ; la qualification périmée n'est pas reprise.
    assert "Sous-Traitant" not in t and "9999" not in t
    # Ni libellé interne, ni manque, ni justification : seulement ce que chacun apporte.
    assert "justifiée" not in t and "déclarée" not in t and "[À compléter" not in t
    assert not re.search(r"manqu|n'a pas|ne dispose|aucune? |insuffis|faute de|lacune|pas de qualification", t,
                         re.IGNORECASE)
    assert _GAP_EXIGE not in t and _GAP_A not in t and "43 %" not in t
    # À l'écran : les libellés « justifiée » / « déclarée », et la note interne au chef de file SEUL.
    d = client.get(f"/api/atelier/{pid}/repartition/{qfq['id']}", headers=h).json()
    chef = next(m for m in d["membres"] if m["role"] == "chef_de_file")
    assert {q["libelle"]: q["statut_libelle"] for q in chef["qualifications"]} == {
        "ISO 14001": "justifiée", "Qualibat 3112": "déclarée"}
    pa_m = next(m for m in d["membres"] if m["nom"] == "Breizh Désamiantage")
    assert {q["libelle"]: q["statut_libelle"] for q in pa_m["qualifications"]} == {
        "Qualibat 1552": "justifiée", "Certification amiante SS3": "déclarée"}
    note = " ".join(d["note_interne"]["lignes"])
    assert "43 % seul, 81 %" in note and _GAP_EXIGE in note and "apporté par Breizh Désamiantage" in note
    assert "jamais exportée" in d["note_interne"]["message"]
    assert "note_interne" in client.get(f"/api/atelier/{pid}/repartition/dossier", headers=h).json()
    # Les partenaires ne voient jamais la note interne (elle dit ce que le chef de file ne pouvait pas faire seul).
    pa = P(client, a["token"])
    vues_partenaire = "\n".join([pa.get().text, pa.get("/dossier").text, pa.get(f"/sections/{qfq['id']}").text])
    assert "note_interne" not in vues_partenaire and _GAP_EXIGE not in vues_partenaire and _GAP_A not in vues_partenaire
    assert "43 %" not in vues_partenaire
    # Validée par CHAQUE membre (chef de file, puis chaque partenaire qui rédige).
    assert _valider_chef(client, h, pid, qfq["id"]).status_code == 200
    assert _generee(_etat(client, h, pid), "grp_qfq")["statut"] != "valide"
    pb = P(client, b["token"])
    for p in (pa, pb):
        r = _valider_qfq(p)
        assert r.status_code == 200, r.text
    q = pa.get().json()["qui_fait_quoi"]
    assert q["statut"] == "validee" and {x["statut"] for x in q["validations"]} == {"valide"}
    assert q["ma_validation"]["statut"] == "valide"
    assert _generee(_etat(client, h, pid), "grp_qfq")["statut"] == "valide"
    # Une part modifiée (puis renvoyée) : le texte suit, « à revalider » pour tous.
    assert _put(client, a["token"], capacites={"ca_n1": 900000, "effectif": 11}).status_code == 200
    _submit(client, a["token"])
    q = pa.get().json()["qui_fait_quoi"]
    assert "11 salariés" in q["texte_commun"]["texte"] and q["statut"] == "a_revalider"
    assert {x["statut"] for x in q["validations"]} == {"a_revalider"}
    assert [x[0] for x in _audit(project_id=pid)].count("redaction.partie_groupement_validee") == 3


def test_qui_fait_quoi_part_non_envoyee_et_chef_sans_qualification(client, auth):
    """Une part pas encore envoyée : « à compléter » (honnête), jamais un brouillon du partenaire ; un chef de file
    sans qualification ni moyen déclaré : aucune ligne « aucune qualification » (jamais un manque)."""
    h = auth["headers"]
    _company(client, h, team=[])
    from tests.test_convention_groupement import _set_company
    _set_company(auth["email"], effectif=0, qualifications=[])
    pid = _projet(client, h)
    _creer(client, h, pid)
    a = _invite(client, h, pid, "Breizh Désamiantage")
    assert _put(client, a["token"], company_name="Breizh Désamiantage", lot=LOT_A,
                qualifications=["BROUILLON-NON-ENVOYE"]).status_code == 200
    assert client.post(f"/api/atelier/{pid}/repartition", headers=h, json={}).status_code == 200
    t = _generee(_etat(client, h, pid), "grp_qfq")["version_courante"]["texte"]
    assert "[À compléter : la part de Breizh Désamiantage n'est pas encore envoyée" in t
    assert "BROUILLON-NON-ENVOYE" not in t
    assert "Part : [À compléter : votre part (chef de file)" in t
    chef = t.split("**Breizh Couverture**, chef de file (mandataire)")[1].split("**Breizh Désamiantage**")[0]
    assert "Qualifications" not in chef and "Moyens" not in chef
    r = _valider_chef(client, h, pid, _generee(_etat(client, h, pid), "grp_qfq")["id"])
    assert r.status_code == 409 and r.json()["detail"]["code"] == "partie_incomplete"


# ═══ 4. « Harmoniser le ton » : des transitions, et aucun fait ajouté ════════════════

SOURCES_AVANT = ["Nous trions les déchets amiantés sur place, avec 9 salariés formés.",
                 "Notre conducteur de travaux encadre le chantier.", "Breizh Désamiantage", "Le Gall Électricité",
                 "Démarche environnementale"]

FAITS_AJOUTES = [
    ("Cette partie s'appuie sur 12 compagnons.", "chiffre"),
    ("Le groupement mobilise douze compagnons sur ce chantier.", "nombre_en_lettres"),
    ("Cette partie couvre 15 % des prestations.", "chiffre"),
    ("Le groupement intervient sur ce type de chantier depuis 2019.", "chiffre"),
    ("La démarche suit le calendrier fixé au 15 mars.", "chiffre"),
    ("Le chantier démarre en septembre, comme le précise cette partie.", "date"),
    ("Cette partie se déroule au printemps.", "date"),
    ("Cette partie représente 25 000 € de travaux.", "chiffre"),
    ("Cette partie traite 3,5 tonnes de déchets.", "chiffre"),
    ("Le groupement est certifié ISO 9001 pour cette partie.", "certification"),
    ("Cette partie s'appuie sur la qualification Qualibat 7131.", "certification"),
    ("Les équipes de cette partie sont labellisées RGE.", "certification"),
    ("Le groupement applique la démarche MASE à cette partie.", "certification"),
    ("Les entreprises sont certifiées pour ces travaux.", "certification"),
    ("La Mairie de Quimper suit cette partie de près.", "nom_propre"),
    ("Cette partie est pilotée par Jean Dupont.", "nom_propre"),
    ("Kerbrat coordonne cette partie.", "nom_propre"),
    ("Cette partie mobilise une nacelle articulée.", "materiel"),
    ("Les travaux de cette partie utilisent une mini-pelle.", "materiel"),
    ("Le groupement monte un échafaudage pour cette partie.", "materiel"),
    ("Le groupement intervient sans retard sur cette partie.", "engagement"),
    ("Nous nous engageons sur cette partie.", "engagement"),
    ("Une centaine de chantiers illustrent cette partie.", "nombre_en_lettres"),
    ("La moitié des équipes intervient sur cette partie.", "quantite"),
    ("Le groupement intervient 24h/24 sur cette partie.", "engagement"),
    ("Cette partie se termine sous huit jours.", "nombre_en_lettres"),
]


def test_controle_des_transitions_refuse_tout_fait_ajoute():
    """Au moins 20 faits ajoutés, de formes variées, sont refusés sans IA ; une vraie transition passe, comme une
    phrase qui ne reprend qu'un fait DÉJÀ dans le dossier."""
    assert len(FAITS_AJOUTES) >= 20
    for phrase, genre in FAITS_AJOUTES:
        r = asm.controle_transition(phrase, SOURCES_AVANT)
        assert r["ok"] is False, phrase
        assert r["faits"], phrase
        assert genre in {f["type"] for f in r["faits"]}, (phrase, r["faits"])
        assert "absent" in r["raison"]
    assert len({g for _p, g in FAITS_AJOUTES}) >= 7
    for forme in ("[À compléter : une transition]", "### Titre", "- Une liste", "Une. Deux phrases. Trois phrases.",
                  "Cette partie " + "relie les parts " * 15 + ".", "Cette partie [S1] suit.", ""):
        assert asm.controle_transition(forme, SOURCES_AVANT)["ok"] is False, forme
    for bonne in ("Cette partie présente ensuite la démarche environnementale du groupement.",
                  "Après l'organisation du chantier, cette partie décrit la gestion des déchets.",
                  "Le groupement s'appuie sur 9 salariés formés pour cette partie.",
                  "Breizh Désamiantage décrit ensuite le tri des déchets.",
                  "S'agissant de l'environnement, la suite précise la méthode retenue."):
        r = asm.controle_transition(bonne, SOURCES_AVANT)
        assert r["ok"] is True, (bonne, r)


def _faux_harmoniseur(transitions, vus=None, erreur=False):
    """Faux fournisseur IA : pour l'harmonisation, rend une phrase par partie demandée (dans l'ordre), ou rien."""
    def ans(kw):
        sys_ = llm._flatten_content(kw.get("system"))
        if "Tu harmonises le ton" not in sys_:
            return "Texte.", "end_turn"
        if vus is not None:
            vus.append(llm._CURRENT_TENANT.get())
        if erreur:
            raise RuntimeError("panne simulée")
        user = llm._flatten_content(kw["messages"][0]["content"])
        ids = re.findall(r"PARTIE (\d+) :", user)
        out = [{"id": i, "texte": t} for i, t in zip(ids, transitions) if t]
        return json.dumps({"transitions": out}, ensure_ascii=False), "end_turn"
    return ans


def _groupement_trois_jonctions(client, auth):
    """1.1 et 2 au chef de file (validées), 1.2 à B (envoyée), 3 à A (envoyée et validée) : trois jonctions."""
    h = auth["headers"]
    pid, st0, a, b, vue = _groupement(client, auth)
    st = _etat(client, h, pid)
    s11, s12 = _section_par(st, "encadrement"), _section_par(st, "équipe d'exécution")
    s2, s3 = _section_par(st, "méthodologie"), _section_par(st, "environnement")
    r = client.put(f"/api/atelier/{pid}/repartition/{s12['id']}", headers=h,
                   json={"redacteur": "partenaire", "invite_id": b["id"]})
    assert r.status_code == 200, r.text
    _ecrire_chef(client, h, pid, s11["id"], "Notre conducteur de travaux Yann Tanguy encadre le chantier au quotidien.",
                 valider=True)
    _ecrire_chef(client, h, pid, s2["id"], "Nous intervenons par zones, en site occupé.", valider=True)
    pa, pb = P(client, a["token"]), P(client, b["token"])
    _ecrire(pb, s12["id"], "Nos électriciens interviennent en binôme sur chaque zone.")
    _envoyer(pb, s12["id"])
    _ecrire(pa, s3["id"], "Nous trions les déchets amiantés sur place, avec 9 salariés formés.")
    _envoyer(pa, s3["id"])
    _valider_part(pa, s3["id"])
    return pid, a, b, pa, pb, {"1.1": s11["id"], "1.2": s12["id"], "2": s2["id"], "3": s3["id"]}


def _texte(client, h, pid, sid):
    return client.get(f"/api/atelier/{pid}/repartition/{sid}", headers=h).json()


def _harmoniser(client, h, pid, **body):
    r = client.post(f"/api/atelier/{pid}/repartition/harmoniser", headers=h, json=body)
    assert r.status_code == 202, r.text
    return _attendre_chef(client, h, pid, r.json()["tache"]["id"]), r.json()["tache"]["id"]


def test_harmoniser_ajoute_des_transitions_et_refuse_les_faits(client, auth, fake_llm):
    h = auth["headers"]
    pid, a, b, pa, pb, s = _groupement_trois_jonctions(client, auth)
    avant = {k: _texte(client, h, pid, v)["texte_commun"] for k, v in s.items() if k in ("1.2", "3")}
    st0 = _etat(client, h, pid)
    t2_avant = _section_par(st0, "méthodologie")["version_courante"]["texte"]
    d0 = client.get(f"/api/atelier/{pid}/repartition/dossier", headers=h).json()
    assert d0["harmoniser"]["jonctions"] == 3 and "à revalider" in d0["harmoniser"]["message"]
    usage0 = _usage(auth["email"])
    vus = []
    fake = fake_llm(_faux_harmoniseur([
        "Cette partie présente ensuite l'équipe d'exécution qui intervient sur chaque zone.",
        "La méthodologie qui suit s'appuie sur 12 compagnons.",
        "Le groupement décrit enfin sa démarche environnementale."], vus))
    job, jid = _harmoniser(client, h, pid)
    assert job["status"] == "done", job
    res = job["result"]
    assert res["ajoutees"] == 2 and res["refusees"] == 1 and "_debit" not in res
    refus = next(x for x in res["transitions"] if not x["ok"])
    assert refus["section_id"] == s["2"] and "12 compagnons" in refus["raison"]
    # L'IA n'a reçu ni nom de salarié, ni nom d'entreprise, ni matière privée ; espace facturé du chef de file.
    envoye = fake.sent()
    assert "Tanguy" not in envoye and "Yann" not in envoye and "Breizh" not in envoye
    assert vus and all(v == _user_id(auth["email"]) for v in vus)
    # Le texte des membres n'est jamais réécrit : la phrase est ajoutée devant, en paragraphe séparé.
    t12 = _texte(client, h, pid, s["1.2"])
    assert t12["texte_commun"]["texte"] == ("Cette partie présente ensuite l'équipe d'exécution qui intervient sur "
                                            "chaque zone.\n\n" + avant["1.2"]["texte"])
    t3 = _texte(client, h, pid, s["3"])
    assert t3["texte_commun"]["texte"] == ("Le groupement décrit enfin sa démarche environnementale.\n\n"
                                           + avant["3"]["texte"])
    # La partie dont la phrase a été refusée ne change pas, et reste validée.
    st = _etat(client, h, pid)
    m = _section_par(st, "méthodologie")
    assert m["version_courante"]["texte"] == t2_avant and m["statut"] == "valide"
    # Les parties touchées repassent « à revalider » par leur auteur, avec l'avant et l'après.
    assert t3["statut"] == "a_revalider" and t3["mention"] == "Modifié par le chef de file, à revalider par Breizh Désamiantage"
    assert t3["comparaison"]["avant"]["texte"] == avant["3"]["texte"]
    assert t12["statut"] == "a_revalider"
    vpa = pa.get().json()
    assert vpa["parties"][0]["statut"] == "a_revalider"
    assert "redaction.transition_ajoutee" in [x[0] for x in _audit(project_id=pid, invite_id=a["id"])]
    assert "redaction.harmonisation" in [x[0] for x in _audit(project_id=pid)]
    j = [e["label"] for e in client.get(f"/api/invite/{a['token']}/journal").json()["entries"]]
    assert any("transition" in x for x in j)
    # Débit : la 1re action de l'IA du dossier (1 analyse), comptée comme une refonte.
    assert _usage(auth["email"]) == usage0 + 1
    assert st["quota"]["regenerations_utilisees"] == 1
    # Le partenaire valide ou rétablit sa version (rien n'est imposé).
    _valider_part(pa, s["3"])
    cur = pb.get(f"/sections/{s['1.2']}").json()["texte_commun"]["version_id"]
    r = pb.post(f"/sections/{s['1.2']}/retablir", {"base_version_id": cur})
    assert r.status_code == 200 and r.json()["texte_commun"]["texte"] == avant["1.2"]["texte"]
    # Harmoniser de nouveau : la transition précédente est remplacée, jamais empilée.
    fake_llm(_faux_harmoniseur([None, None, "Pour finir, cette partie traite de l'environnement."]))
    job, _ = _harmoniser(client, h, pid)
    assert job["status"] == "done" and job["result"]["ajoutees"] == 1
    t3 = _texte(client, h, pid, s["3"])["texte_commun"]["texte"]
    assert t3 == "Pour finir, cette partie traite de l'environnement.\n\n" + avant["3"]["texte"]
    assert "Le groupement décrit enfin" not in t3


def test_harmoniser_panne_ou_rien_retenu_rien_n_est_decompte(client, auth, fake_llm):
    h = auth["headers"]
    pid, a, b, pa, pb, s = _groupement_trois_jonctions(client, auth)
    textes = {k: _texte(client, h, pid, v)["texte_commun"]["texte"] for k, v in s.items() if k in ("1.2", "3")}
    usage0 = _usage(auth["email"])
    fake_llm(_faux_harmoniseur([], erreur=True))
    job, _ = _harmoniser(client, h, pid)
    assert job["status"] == "error"
    st = _etat(client, h, pid)
    assert st["quota"]["regenerations_utilisees"] == 0 and _usage(auth["email"]) == usage0
    assert st["quota"]["analyse_debitee"] is False
    # Toutes les phrases refusées (un fait ajouté chacune) : rien ne change, rien n'est décompté.
    fake_llm(_faux_harmoniseur(["Cette partie mobilise 3 nacelles.", "Le chantier démarre en mars.",
                                "La Mairie de Quimper suit cette partie."]))
    job, _ = _harmoniser(client, h, pid)
    assert job["status"] == "done" and job["result"]["ajoutees"] == 0 and job["result"]["refusees"] == 3
    assert "rien n'est décompté" in job["result"]["message"]
    st = _etat(client, h, pid)
    assert st["quota"]["regenerations_utilisees"] == 0 and _usage(auth["email"]) == usage0
    for k in ("1.2", "3"):
        assert _texte(client, h, pid, s[k])["texte_commun"]["texte"] == textes[k]


def test_harmoniser_rien_a_relier_et_partenaire_jamais(client, auth, fake_llm):
    h = auth["headers"]
    pid, st0, a, b, vue = _groupement(client, auth, avec_b=False)
    r = client.post(f"/api/atelier/{pid}/repartition/harmoniser", headers=h, json={})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "rien_a_harmoniser"
    # Un autre compte : le marché n'existe pas pour lui.
    h2, _e2 = _register(client, company="Autre SARL")
    assert client.post(f"/api/atelier/{pid}/repartition/harmoniser", headers=h2, json={}).status_code == 404
    assert client.get(f"/api/atelier/{pid}/repartition/dossier", headers=h2).status_code == 404


# ═══ 5. Validation : chaque membre, compteur, « à revalider », ce qui manque ═════════

def test_validation_compteur_revalidation_et_ce_qui_manque(client, auth, mails):
    h = auth["headers"]
    pid, a, b, vue, conv = _groupement_convention(client, auth)
    pa = P(client, a["token"])
    sid = _part_a(vue)["section_id"]
    pret = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["dossier"]["pret"]
    assert pret["message"] == "Sections validées 0 sur 6" and pret["complet"] is False
    msgs = [m["message"] for m in pret["manques"]]
    assert "« 3. Démarche environnementale » : à rédiger par Breizh Désamiantage" in msgs
    assert "« Organisation du groupement » : à valider par le chef de file" in msgs
    assert "« Qui fait quoi dans le groupement » : à valider par le chef de file et Breizh Désamiantage" in msgs
    assert "« 1.1 Encadrement du chantier » : à rédiger par le chef de file" in msgs
    assert "Dossier assemblé : à valider par le chef de file" in msgs
    assert "Dossier assemblé : à valider par Breizh Désamiantage" in msgs
    # Le partenaire : sa partie, « Qui fait quoi », puis le dossier assemblé.
    _ecrire(pa, sid, "Nous trions les déchets amiantés sur place.")
    _envoyer(pa, sid)
    pret = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["dossier"]["pret"]
    assert "« 3. Démarche environnementale » : à valider par Breizh Désamiantage" in [m["message"] for m in pret["manques"]]
    _valider_part(pa, sid)
    r = pa.post("/dossier/valider", {})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "parties_a_valider"
    assert [x["titre"] for x in r.json()["detail"]["sections"]] == ["Qui fait quoi dans le groupement"]
    assert _valider_qfq(pa).status_code == 200
    assert pa.get().json()["dossier"]["peut_valider"] is True
    emp = pa.get("/dossier").json()["empreinte"]
    assert pa.post("/dossier/valider", {"empreinte": emp}).status_code == 200
    # Le chef de file : ses parties, « Organisation » (convention validée par tous), « Qui fait quoi ».
    st = _etat(client, h, pid)
    for s in st["sections"]:
        if s["auteur"]["type"] == "organisation":
            _ecrire_chef(client, h, pid, s["id"], f"Texte du chef pour {s['intitule']}.", valider=True)
    r = client.post(f"/api/atelier/{pid}/repartition/valider-dossier", headers=h, json={})
    assert r.status_code == 409 and {x["titre"] for x in r.json()["detail"]["sections"]} == {
        "Organisation du groupement", "Qui fait quoi dans le groupement"}
    assert _valider_chef(client, h, pid, _generee(st, "grp_orga")["id"]).status_code == 200
    assert _valider_chef(client, h, pid, _generee(st, "grp_qfq")["id"]).status_code == 200
    vue2 = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()
    assert vue2["dossier"]["pret"]["message"] == "Sections validées 6 sur 6"
    r = client.post(f"/api/atelier/{pid}/repartition/valider-dossier", headers=h,
                    json={"empreinte": vue2["dossier"]["empreinte"]})
    assert r.status_code == 200, r.text
    # Les parties validées du chef de file ont changé le dossier : le partenaire le revalide.
    assert pa.get().json()["dossier"]["ma_validation"]["statut"] == "a_revalider"
    assert pa.post("/dossier/valider", {"empreinte": pa.get("/dossier").json()["empreinte"]}).status_code == 200
    pret = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["dossier"]["pret"]
    assert pret["complet"] is True and pret["manques"] == [] and pret["message"] == "Sections validées 6 sur 6"
    # Toute modification après validation → « à revalider » (même empreinte que la convention périmée).
    cur = client.get(f"/api/atelier/{pid}/repartition/{sid}", headers=h).json()["texte_commun"]["version_id"]
    assert client.put(f"/api/atelier/{pid}/repartition/{sid}/texte", headers=h,
                      json={"texte": "Nous trions les déchets amiantés.", "base_version_id": cur}).status_code == 200
    pret = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["dossier"]["pret"]
    assert pret["message"] == "Sections validées 5 sur 6" and pret["complet"] is False
    msgs = [m["message"] for m in pret["manques"]]
    assert "« 3. Démarche environnementale » : à revalider par Breizh Désamiantage" in msgs
    assert "Dossier assemblé : à revalider par le chef de file" in msgs
    assert "Dossier assemblé : à revalider par Breizh Désamiantage" in msgs
    # Une part de partenaire modifiée (et renvoyée) : « Qui fait quoi » repasse « à revalider » pour tous.
    assert _put(client, a["token"], capacites={"ca_n1": 900000, "effectif": 10}).status_code == 200
    _submit(client, a["token"])
    pret = client.get(f"/api/atelier/{pid}/repartition", headers=h).json()["dossier"]["pret"]
    assert any(m["message"].startswith("« Qui fait quoi dans le groupement » : à revalider par") for m in pret["manques"])
    # Tout est inscrit au journal chaîné.
    actions = [x[0] for x in _audit(project_id=pid)]
    for act in ("redaction.partie_groupement_validee", "redaction.dossier_valide", "redaction.dossier_a_revalider",
                "redaction.partie_groupement_a_jour"):
        assert act in actions, act
    assert all(x[4] for x in _audit(project_id=pid) if x[0].startswith("redaction."))


# ═══ 6. Export : le dossier assemblé complet, propre, bloqué tant qu'il manque quelque chose ══

def test_export_du_dossier_assemble_propre_et_bloque(client, auth, mails, fake_llm):
    h = auth["headers"]
    pid, a, b, vue, conv = _groupement_convention(client, auth)
    pa = P(client, a["token"])
    sid = _part_a(vue)["section_id"]
    # La part du partenaire cite SES sources ([F#], [S#]) : renvois retirés à l'export, jamais montrés.
    _ecrire(pa, sid, "Nous trions les déchets amiantés sur place [S1]. Notre équipe compte 9 salariés [P1].")
    _envoyer(pa, sid)
    # Tant que la part du partenaire a un « [À compléter] », l'export est bloqué (dérogation journalisée).
    st = _etat(client, h, pid)
    for s in st["sections"]:
        if s["auteur"]["type"] == "organisation":
            _ecrire_chef(client, h, pid, s["id"], f"Notre réponse pour {s['intitule']}.", valider=True)
    cur = pa.get(f"/sections/{sid}").json()["brouillon"]["id"]
    r = pa.put(f"/sections/{sid}/brouillon", {"texte": "Nous trions les déchets [S1]. [À compléter : planning]",
                                               "base_brouillon_id": cur})
    assert r.status_code == 200
    _envoyer(pa, sid)
    gate = _etat(client, h, pid)["blocage"]
    codes = {r_["code"]: r_ for r_ in gate["raisons"]}
    assert any(x["section_id"] == sid for x in codes["a_completer"]["sections"])
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "export_bloque"
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "depot_en_l_etat", "texte": "Le partenaire complètera son planning.",
                          "etat_hash": gate["etat_hash"]})
    assert r.status_code == 200, r.text
    assert "atelier.derogation_export" in [x[0] for x in _audit(project_id=pid)]
    assert client.post(f"/api/atelier/{pid}/export", headers=h,
                       json={"format": "docx", "etat_hash": gate["etat_hash"]}).status_code == 200
    # Le partenaire complète ; tout est validé : l'export passe sans dérogation.
    cur = pa.get(f"/sections/{sid}").json()["brouillon"]["id"]
    pa.put(f"/sections/{sid}/brouillon", {"texte": "Nous trions les déchets amiantés sur place [S1]. "
                                                   "Notre équipe compte 9 salariés [P1].", "base_brouillon_id": cur})
    _envoyer(pa, sid)
    _valider_part(pa, sid)
    assert _valider_qfq(pa).status_code == 200
    st = _etat(client, h, pid)
    for kind in ("grp_orga", "grp_qfq"):
        assert _valider_chef(client, h, pid, _generee(st, kind)["id"]).status_code == 200
    st = _etat(client, h, pid)
    faits = [r_["code"] for r_ in st["blocage"]["raisons"]]
    if "faits_a_relire" in faits:          # « 9 salariés » écrit par le partenaire : sa déclaration
        pytest.fail(f"faits à relire inattendus : {st['blocage']['raisons']}")
    assert st["blocage"]["peut_exporter"] is True and st["blocage"]["derogation"] is None, st["blocage"]
    for fmt in ("docx", "pdf"):
        r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": fmt,
                                                                        "etat_hash": st["blocage"]["etat_hash"]})
        assert r.status_code == 200, r.text
        txt = _docx_texte(r.content) if fmt == "docx" else _pdf_texte(r.content)
        plat = re.sub(r"\s+", " ", txt)
        # Le dossier assemblé COMPLET, dans l'ordre du cadre, parties tirées du groupement en tête.
        for x in ("Organisation du groupement", "Qui fait quoi dans le groupement", "Encadrement du chantier",
                  "Démarche environnementale", "Groupement conjoint", LOT_A, "Nous trions les déchets amiantés"):
            assert x in plat, (fmt, x)
        assert plat.index("Organisation du groupement") < plat.index("Qui fait quoi dans le groupement") \
            < plat.index("Encadrement du chantier") < plat.index("Démarche environnementale")
        # Rien d'interne : ni renvoi, ni libellé de source, ni mention de rédaction, ni note interne, ni prix.
        for interdit in ("[S1]", "[P1]", "[F", "justifiée", "déclarée", "à revalider", "Rédigé par le chef",
                         "Note interne", "note interne", "À compléter", "Adjugo", "40 000", "€"):
            assert interdit not in plat, (fmt, interdit)


# ═══ 7. Démo : l'harmonisation n'est jamais lancée par la vitrine ═══════════════════

def test_harmoniser_refuse_a_la_demo(client):
    r = client.post("/api/auth/demo")
    assert r.status_code == 200
    hd = {"Authorization": f"Bearer {r.json()['access_token']}"}
    from app.core.demo_guard import blocked_action
    assert blocked_action("POST", "/api/atelier/7/repartition/harmoniser")
    r = client.post("/api/atelier/1/repartition/harmoniser", headers=hd, json={})
    assert r.status_code == 403 and "lecture guidée" in json.dumps(r.json(), ensure_ascii=False)


@pytest.fixture
def mails(monkeypatch):
    sent = []
    import app.routers.invites as inv
    monkeypatch.setattr(inv, "_send_mails_async", lambda m: sent.extend(m))
    return sent
