"""
Rédaction, vague 2 — l'atelier du dossier de présentation (étude §6.4 G + A0), côté serveur.

 • modèle : un dossier par marché et par entreprise, parties ordonnées, versions EN AJOUT SEUL,
   validation (qui, quand, empreinte), « à revalider » après modification ;
 • plan : cadre imposé (ordre, numéros, intitulés), grille retrouvée (poids tel qu'écrit), plan
   simple sinon ; jamais de critère ni de poids inventés ;
 • moteur : tâche de fond, espace facturé reposé dans le fil, noms des salariés masqués, faits
   autorisés seuls, règlement et cadre lus en entier, remboursement si l'IA tombe ;
 • A0 « à relire » : phrases piégées, le dossier de l'acheteur n'est jamais une source ;
 • export Word et PDF : aucune marque interne, mention IA seulement si l'IA a servi, numéros du
   cadre ; blocage et dérogation inscrite au journal chaîné ; compteur « environ » ;
 • quota (1 analyse par dossier, refaire gratuit jusqu'au plafond), mesure « fichier imposé » ;
 • isolation entre comptes (404), matière privée d'un auteur partenaire (vague 4) jamais exposée.
Aucun appel réseau : le client IA est simulé.
"""
import io
import json
import re
import time
import zipfile

import pytest

from app.services import atelier as at
from app.services import atelier_plan as ap
from app.services import faits_a_relire as far
from app.services import grille_notation as gn
from app.services import llm
from tests.test_redaction_socle0 import NOMS, TEAM, _company, _db, _user_id, fake_llm  # noqa: F401

RC = ("REGLEMENT DE LA CONSULTATION\n"
      "Article 1 - Objet\nRéfection de la toiture du groupe scolaire Jean Moulin.\n"
      "Article 7 - Critères de jugement des offres\nValeur technique : 60 %\nPrix : 40 %\n"
      "Article 8 - Présentation des offres\n"
      "Le candidat remettra un mémoire justificatif. Le mémoire justificatif ne devra pas excéder 8 pages.\n"
      "Le cadre de réponse technique joint devra obligatoirement être complété.\n"
      "Les entreprises devront être certifiées ISO 14001.\n"
      "Les documents seront rédigés en police Arial 11.\n")
CADRE = ("CADRE DE REPONSE TECHNIQUE\n"
         "1. Moyens humains\n"
         "1.1 Encadrement du chantier\nPrésentez l'encadrement du chantier (2 pages maximum).\n"
         "1.2 Équipe d'exécution\nPrésentez les compagnons affectés au chantier.\n"
         "2. Méthodologie\nDécrivez la méthodologie d'intervention en site occupé.\n"
         "3. Démarche environnementale\nDécrivez la gestion des déchets du chantier.\n")
PROPOSES = [{"critere": "Valeur technique", "ponderation": "60 %", "extrait": "Valeur technique : 60 %"},
            {"critere": "Prix", "ponderation": "40 %", "extrait": "Prix : 40 %"}]
CCTP_MARK = "CCTP-MARQUEUR-FIN-DE-DOSSIER"


def _zip(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, text in files.items():
            zf.writestr(name, text)
    return buf.getvalue()


def _analysis(text: str, pieces=None, impose=True):
    det = {"intitule_marche": "Réfection de la toiture du groupe scolaire", "acheteur": "Ville de Lorient",
           "criteres_attribution": PROPOSES,
           "pieces_requises": [{"nom": "Mémoire justificatif", "disponible": False}]}
    det["grille_notation"] = gn.build_grille(det, text, pieces=pieces)
    if not impose:
        det["grille_notation"]["contraintes"]["fichier_impose"] = False
    return {"summary": "Toiture", "dce_available": True, "dce_excerpt": text[:30000], "details": det}


def _projet(client, h, with_dce=True, impose=True, name="Toiture du groupe scolaire"):
    r = client.post("/api/projects/", headers=h, json={"name": name, "client": "Ville de Lorient"})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    from app.models import Project
    from app.services.analysis import extract_dce_text
    from app.services.dossier import save_to_dossier
    db = _db()
    try:
        p = db.get(Project, pid)
        if with_dce:
            content = _zip({"01_Reglement_de_la_consultation.txt": RC, "02_Cadre_de_reponse_technique.txt": CADRE,
                            "03_CCTP.txt": "CAHIER DES CLAUSES TECHNIQUES PARTICULIERES\n" + ("Couverture zinc. " * 50)
                            + CCTP_MARK})
            text, pieces = extract_dce_text("DCE.zip", content, max_chars=120000, with_pieces=True)
            p.ai_analysis = _analysis(text, pieces, impose=impose)
            save_to_dossier(db, p.user_id, pid, "DCE", "DCE.zip", content, "application/zip")
        else:
            p.ai_analysis = {"summary": "x", "details": {"intitule_marche": name, "acheteur": "Ville de Lorient"}}
        db.commit()
    finally:
        db.close()
    return pid


def _creer(client, h, pid, base="auto"):
    r = client.post(f"/api/atelier/{pid}", headers=h, json={"base": base})
    assert r.status_code == 201, r.text
    return r.json()


def _etat(client, h, pid):
    r = client.get(f"/api/atelier/{pid}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _ecrire(client, h, pid, sec, texte):
    base = (sec.get("version_courante") or {}).get("id")
    r = client.put(f"/api/atelier/{pid}/sections/{sec['id']}/texte", headers=h,
                   json={"texte": texte, "base_version_id": base})
    assert r.status_code == 200, r.text
    return r.json()


def _attendre(client, h, pid, jid):
    r = None
    for _ in range(200):
        r = client.get(f"/api/atelier/{pid}/taches/{jid}", headers=h)
        assert r.status_code == 200, r.text
        if r.json()["status"] in ("done", "error"):
            return r.json()
        time.sleep(0.05)
    raise AssertionError(f"tâche {jid} jamais terminée : {r.json() if r is not None else None}")


def _usage(email):
    from app.models import User
    db = _db()
    try:
        return db.query(User).filter(User.email == email).first().analyses_used_this_month or 0
    finally:
        db.close()


def _redige(texte="Notre équipe intervient en site occupé [S1]."):
    """Faux fournisseur IA : rédiger rend `texte`, raccourcir rend un texte court ; chaque appel
    note l'espace facturé (tenant) vu dans le fil."""
    seen = []

    def ans(kw):
        seen.append(llm._CURRENT_TENANT.get())
        system = llm._flatten_content(kw.get("system"))
        if "Tu raccourcis" in system:
            return "Texte raccourci de la partie.", "end_turn"
        return (texte(kw) if callable(texte) else texte), "end_turn"
    return ans, seen


# ═══ Plan : le cadre de l'acheteur commande ══════════════════════════════════════════

def test_plan_cadre_impose_ordre_numeros_intitules_et_limites():
    from app.services.analysis import extract_dce_text
    content = _zip({"RC.txt": RC, "Cadre de reponse technique.txt": CADRE})
    text, pieces = extract_dce_text("DCE.zip", content, max_chars=120000, with_pieces=True)
    a = _analysis(text, pieces)
    assert a["details"]["grille_notation"]["contraintes"]["fichier_impose"] is True
    plan = ap.build_plan(a, CADRE)
    assert plan["source"] == ap.CADRE and plan["fichier_impose"] is True
    assert [(s["numero"], s["intitule"]) for s in plan["sections"]] == [
        ("1.1", "Encadrement du chantier"), ("1.2", "Équipe d'exécution"), ("2", "Méthodologie"),
        ("3", "Démarche environnementale")]
    assert plan["sections"][0]["parents"] == [{"numero": "1", "intitule": "Moyens humains"}]
    assert [(x["type"], x["valeur"]) for x in plan["sections"][0]["limite"]] == [("pages", 2)]
    assert all(s["poids"] == "" and s["critere"] == "" for s in plan["sections"])     # rien d'inventé
    assert [(x["type"], x["valeur"]) for x in plan["limites"]["document"]] == [("pages", 8)]
    assert plan["limites"]["police"]["police"].lower() == "arial"


def test_plan_grille_poids_tel_qu_ecrit_prix_hors_dossier_et_rien_d_invente():
    g = {"criteres": [
        {"id": "c1", "critere": "Valeur technique", "dans_le_plan": True, "poids_retenu": "60 %", "phrase": "Valeur technique : 60 %"},
        {"id": "c2", "critere": "Délai d'exécution", "dans_le_plan": True, "poids_retenu": "", "phrase": "Délai d'exécution"},
        {"id": "c3", "critere": "Prix des prestations", "dans_le_plan": True, "poids_retenu": "40 %", "phrase": "Prix : 40 %"},
        {"id": "c4", "critere": "Critère à vérifier", "dans_le_plan": False, "poids_retenu": "", "phrase": ""}],
        "contraintes": {"fichier_impose": False, "limites": []}}
    plan = ap.build_plan({"details": {"grille_notation": g}})
    assert plan["source"] == ap.GRILLE
    assert [(s["intitule"], s["poids"]) for s in plan["sections"]] == [("Valeur technique", "60 %"),
                                                                       ("Délai d'exécution", "")]
    assert plan["criteres_hors_dossier"][0]["critere"] == "Prix des prestations"
    assert "Critère à vérifier" not in json.dumps(plan, ensure_ascii=False)
    # Un seul critère : découpage proposé par Adjugo, dit comme tel, poids repris tel quel.
    g1 = dict(g, criteres=[g["criteres"][0], g["criteres"][2]])
    p1 = ap.build_plan({"details": {"grille_notation": g1}})
    assert all(s["partie_de_critere"] and s["poids"] == "60 %" for s in p1["sections"])
    assert "pas imposé par l'acheteur" in p1["message"]


def test_plan_simple_sans_ia_sans_critere_et_cadre_non_decoupe():
    plan = ap.build_plan({"details": {}})
    assert plan["source"] == ap.LIBRE and len(plan["sections"]) == len(ap.PARTIES_LIBRES)
    assert all(s["poids"] == "" and s["critere"] == "" for s in plan["sections"])
    assert "sans IA" in plan["message"]
    g = {"criteres": [], "contraintes": {"fichier_impose": True, "limites": []}}
    p2 = ap.build_plan({"details": {"grille_notation": g}}, cadre_text="Texte libre sans numérotation.")
    assert p2["cadre_non_decoupe"] is True and p2["message"].startswith(ap.MSG_CADRE_NON_DECOUPE)
    # Le sommaire d'un cadre et une phrase qui commence par un nombre ne sont pas des questions.
    qs = ap.parse_cadre("1. Moyens ........ 3\n2. Délais ....... 4\n1. Moyens\nDécrivez.\n"
                        "10 jours ouvrés suffisent.\n2. Délais\nPrécisez.\n")
    assert [(q["numero"], q["intitule"]) for q in qs] == [("1", "Moyens"), ("2", "Délais")]


# ═══ Le dossier : création, isolation, versions ════════════════════════════════════════

def test_creation_proposition_sans_rien_enregistrer_puis_dossier_unique(client, auth):
    h = auth["headers"]
    pid = _projet(client, h)
    r = client.get(f"/api/atelier/{pid}", headers=h)
    assert r.status_code == 200 and r.json()["dossier"] is None
    prop = r.json()["proposition"]
    assert prop["source"] == "cadre" and prop["sections"][0]["numero"] == "1.1"
    from app.models import DossierPresentation
    db = _db()
    try:
        assert db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).count() == 0
    finally:
        db.close()
    st = _creer(client, h, pid)
    assert [s["titre"] for s in st["sections"]] == ["1.1 Encadrement du chantier", "1.2 Équipe d'exécution",
                                                    "2. Méthodologie", "3. Démarche environnementale"]
    assert all(s["statut"] == "a_ecrire" for s in st["sections"])
    assert st["dossier_acheteur"]["source"] == "fichier"
    assert st["export"]["fichier_impose"] is True and st["export"]["plan_du_cadre"] is True
    av = next(a for a in st["blocage"]["avertissements"] if a["code"] == "fichier_impose")
    # Une seule consigne, qui désigne le bon fichier (décision 11) : jamais « recopiez-y » vers le Word d'Adjugo.
    assert av["message"] == ("Ce marché impose son propre fichier de réponse. Le fichier Word d'Adjugo suit l'ordre, "
                             "les numéros et les intitulés de ce cadre : recopiez-en le texte dans le fichier de "
                             "l'acheteur.")
    assert st["dossier"]["plan"]["avertissement_fichier"] == av["message"]
    r = client.post(f"/api/atelier/{pid}", headers=h, json={})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "dossier_existe"


def test_isolation_un_autre_compte_recoit_404_partout(client, auth):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][0]["id"]
    r = client.post("/api/auth/register", json={"email": f"autre_{pid}@test.fr", "password": "motdepasse123",
                                                "full_name": "Autre", "company_name": "Autre SARL"})
    h2 = {"Authorization": f"Bearer {r.json()['access_token']}"}
    calls = [("get", f"/api/atelier/{pid}", None), ("post", f"/api/atelier/{pid}", {}),
             ("put", f"/api/atelier/{pid}/sections/{sid}/texte", {"texte": "x"}),
             ("get", f"/api/atelier/{pid}/sections/{sid}", None),
             ("post", f"/api/atelier/{pid}/sections/{sid}/rediger", {}),
             ("post", f"/api/atelier/{pid}/sections/{sid}/valider", {"version_id": 1}),
             ("post", f"/api/atelier/{pid}/export", {"format": "docx"}),
             ("post", f"/api/atelier/{pid}/derogation", {"raison": "autre", "texte": "abc", "etat_hash": "x"})]
    for m, url, body in calls:
        r = getattr(client, m)(url, headers=h2, **({"json": body} if body is not None else {}))
        assert r.status_code == 404, (url, r.status_code, r.text)
    # Une partie d'un autre dossier, appelée via son propre marché : 404 aussi.
    pid2 = _projet(client, h2)
    _creer(client, h2, pid2)
    r = client.put(f"/api/atelier/{pid2}/sections/{sid}/texte", headers=h2, json={"texte": "x"})
    assert r.status_code == 404 and r.json()["detail"]["code"] == "partie_introuvable"


def test_versions_en_ajout_seul_validation_et_a_revalider(client, auth):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sec = st["sections"][2]
    s1 = _ecrire(client, h, pid, sec, "Nous intervenons en site occupé, par zones.")
    v1 = s1["version_courante"]
    assert s1["statut"] == "brouillon" and v1["numero"] == 1 and v1["origine"] == "utilisateur" and v1["ia"] is False
    # Base périmée : refus explicite, rien n'est écrasé.
    r = client.put(f"/api/atelier/{pid}/sections/{sec['id']}/texte", headers=h,
                   json={"texte": "Autre texte.", "base_version_id": None})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "texte_modifie_entre_temps"
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/relu", headers=h, json={"version_id": v1["id"]})
    assert r.json()["statut"] == "relu"
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/valider", headers=h, json={"version_id": v1["id"]})
    val = r.json()
    assert val["statut"] == "valide" and val["validation"]["version_id"] == v1["id"]
    assert val["validation"]["empreinte"] == v1["empreinte"] and val["validation"]["par"] == _user_id(auth["email"])
    s2 = _ecrire(client, h, pid, val, "Nous intervenons en site occupé, par zones, hors temps scolaire.")
    assert s2["statut"] == "a_revalider" and s2["version_courante"]["numero"] == 2
    # L'historique est entier ; remettre le texte validé redonne « validé ».
    r = client.get(f"/api/atelier/{pid}/sections/{sec['id']}", headers=h)
    assert [v["numero"] for v in r.json()["versions"]] == [2, 1] and "texte" not in r.json()["versions"][0]
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/restaurer", headers=h,
                    json={"version_id": v1["id"], "base_version_id": s2["version_courante"]["id"]})
    assert r.status_code == 200 and r.json()["statut"] == "valide" and r.json()["version_courante"]["numero"] == 3
    assert r.json()["version_courante"]["restaure_de"] == v1["id"]
    # Validation journalisée (journal chaîné).
    from app.models import AuditLog
    db = _db()
    try:
        assert db.query(AuditLog).filter(AuditLog.action == "atelier.partie_validee",
                                         AuditLog.target_id == sec["id"]).count() == 1
    finally:
        db.close()


def test_plan_modifiable_seulement_hors_cadre_et_reconstruction_confirmee(client, auth):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    cadre_sec = st["sections"][0]
    r = client.patch(f"/api/atelier/{pid}/sections/{cadre_sec['id']}", headers=h, json={"intitule": "Autre"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "intitule_impose"
    r = client.patch(f"/api/atelier/{pid}/sections/{cadre_sec['id']}", headers=h, json={"position": 3})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "ordre_impose"
    r = client.post(f"/api/atelier/{pid}/sections", headers=h, json={"intitule": "Annexes"})
    assert r.status_code == 201 and r.json()["numero_affiche"] == "" and r.json()["intitule_modifiable"] is True
    _ecrire(client, h, pid, cadre_sec, "Notre encadrement.")
    r = client.post(f"/api/atelier/{pid}/plan", headers=h, json={"base": "libre"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "plan_remplacement_a_confirmer"
    r = client.post(f"/api/atelier/{pid}/plan", headers=h, json={"base": "libre", "remplacer": True})
    st2 = r.json()
    assert r.status_code == 200 and st2["dossier"]["plan_source"] == "libre"
    assert [s["numero_affiche"] for s in st2["sections"]][:3] == ["1", "2", "3"]
    ret = {x["intitule"]: x for x in st2["sections_retirees"]}
    assert ret["Encadrement du chantier"]["nb_versions"] == 1          # rien n'est perdu
    r = client.post(f"/api/atelier/{pid}/sections/{cadre_sec['id']}/remettre", headers=h)
    assert r.status_code == 200 and r.json()["version_courante"]["texte"] == "Notre encadrement."
    # Plan libre : intitulé et ordre modifiables.
    libre = st2["sections"][1]
    r = client.patch(f"/api/atelier/{pid}/sections/{libre['id']}", headers=h,
                     json={"intitule": "Compréhension du site", "position": 0})
    assert r.status_code == 200 and r.json()["intitule"] == "Compréhension du site" and r.json()["ordre"] == 0
    r = client.delete(f"/api/atelier/{pid}/sections/{libre['id']}", headers=h)
    assert r.status_code == 200 and libre["id"] not in [s["id"] for s in _etat(client, h, pid)["sections"]]


# ═══ Le moteur : tâche de fond, faits autorisés, noms masqués, quota ═══════════════════

def test_rediger_tache_de_fond_faits_autorises_noms_masques_et_tenant(client, auth, fake_llm):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sec = st["sections"][0]
    texte = ("Le chantier est encadré par [Chef de chantier 1], présent chaque jour [S1]. "
             "Nous sommes certifiés ISO 14001. Notre équipe compte 12 salariés. "
             "[À compléter : nombre de compagnons affectés]")
    ans, seen = _redige(texte)
    fake = fake_llm(ans)
    before = _usage(auth["email"])
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/rediger", headers=h, json={})
    assert r.status_code == 202, r.text
    assert r.json()["quota"]["analyse_debitee"] is True and "\"_debit\"" not in json.dumps(r.json())
    job = _attendre(client, h, pid, r.json()["tache"]["id"])
    assert job["status"] == "done" and "\"_debit\"" not in json.dumps(job)
    assert _usage(auth["email"]) == before + 1                     # le dossier compte pour 1 analyse
    owner = _user_id(auth["email"])
    assert seen and all(t == owner for t in seen)                  # espace facturé reposé dans le fil
    sent = fake.sent()
    for nom in NOMS:
        assert nom not in sent, f"« {nom} » est parti à l'IA"
    assert "FAITS AUTORISÉS" in sent and "Effectif : 12" in sent and "Breizh Couverture" in sent
    assert "Encadrement du chantier" in sent and "QUESTION DU CADRE DE RÉPONSE (n° 1.1)" in sent
    assert "Les entreprises devront être certifiées ISO 14001" in sent      # règlement lu en entier
    assert "Présentez les compagnons affectés" in sent                      # cadre lu en entier
    assert CCTP_MARK in sent                                                # le reste, borné
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sec["id"])
    v = s["version_courante"]
    assert s["statut"] == "brouillon" and v["origine"] == "ia" and v["ia"] is True
    assert "Yann Tanguy" in v["texte"]                                      # remis en clair sans IA
    faits = {(f["fait"], f["type"]) for f in v["faits"]}
    assert ("ISO 14001", "certification") in faits                          # exigé par le RC seulement
    assert not any(f["fait"].startswith("12") for f in v["faits"])          # 12 salariés : dans le profil
    assert all(f["message"] == far.MESSAGE and f["statut"] == far.A_RELIRE for f in v["faits"])
    assert s["a_completer"] == 1 and s["redaction_en_cours"] is None


def test_refaire_une_partie_gratuit_jusqu_au_plafond_puis_accord_exprès(client, auth, fake_llm, monkeypatch):
    monkeypatch.setenv("ATELIER_REGEN_GRATUITES", "1")
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][2]["id"]
    ans, _seen = _redige("Méthodologie en site occupé.")
    fake_llm(ans)
    u0 = _usage(auth["email"])
    for attendu_usage, attendu_regen in ((u0 + 1, 0), (u0 + 1, 1)):
        r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
        assert r.status_code == 202, r.text
        assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
        assert _usage(auth["email"]) == attendu_usage
        assert r.json()["quota"]["regenerations_utilisees"] == attendu_regen
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "debit_a_confirmer" and r.json()["detail"]["cout"] == 1
    assert _usage(auth["email"]) == u0 + 1
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={"accepte_debit": True})
    assert r.status_code == 202
    assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    assert _usage(auth["email"]) == u0 + 2
    q = _etat(client, h, pid)["quota"]
    assert q["regenerations_payantes"] == 1 and q["prochaine_refonte_coute"] == 1
    # Aucune rédaction n'a effacé de version : 3 versions de l'IA.
    r = client.get(f"/api/atelier/{pid}/sections/{sid}", headers=h)
    assert [v["origine"] for v in r.json()["versions"]] == ["ia", "ia", "ia"]


def test_panne_ia_rembourse_et_ne_cree_rien(client, auth, fake_llm):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][0]["id"]

    def boom(kw):
        raise RuntimeError("fournisseur indisponible")
    fake_llm(boom)
    u0 = _usage(auth["email"])
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
    assert r.status_code == 202
    job = _attendre(client, h, pid, r.json()["tache"]["id"])
    assert job["status"] == "error" and job["error"] == at.MSG_IA_PANNE      # aucun détail technique
    assert _usage(auth["email"]) == u0                                        # remboursé
    st = _etat(client, h, pid)
    s = next(x for x in st["sections"] if x["id"] == sid)
    assert s["version_courante"] is None and s["statut"] == "a_ecrire" and s["redaction_en_cours"] is None
    assert st["quota"]["analyse_debitee"] is False                            # la prochaine débitera


def test_echec_partiel_seules_les_parties_en_echec_sont_rendues(client, auth, fake_llm, monkeypatch):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    ids = [s["id"] for s in st["sections"]]

    def ans(kw):
        user = llm._flatten_content(kw["messages"][0]["content"])
        if "PARTIE À RÉDIGER : Méthodologie" in user:
            raise RuntimeError("panne sur une partie")
        return "Texte de la partie.", "end_turn"
    fake_llm(ans)
    u0 = _usage(auth["email"])
    r = client.post(f"/api/atelier/{pid}/rediger", headers=h, json={})
    assert r.status_code == 202 and sorted(r.json()["section_ids"]) == sorted(ids)
    job = _attendre(client, h, pid, r.json()["tache"]["id"])
    assert job["status"] == "done" and job["result"]["echecs"] == 1 and "rien n'a été débité" in job["result"]["message"]
    assert _usage(auth["email"]) == u0 + 1
    st = _etat(client, h, pid)
    by = {s["intitule"]: s for s in st["sections"]}
    assert by["Méthodologie"]["statut"] == "a_ecrire" and by["Encadrement du chantier"]["statut"] == "brouillon"
    # Re-rédiger la partie en échec n'est pas une « refonte » (elle n'a jamais été rédigée).
    fake_llm(lambda kw: ("Méthodologie rédigée.", "end_turn"))
    r = client.post(f"/api/atelier/{pid}/sections/{by['Méthodologie']['id']}/rediger", headers=h, json={})
    assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    assert r.json()["quota"]["regenerations_utilisees"] == 0 and _usage(auth["email"]) == u0 + 1


def test_generation_n_ecrase_jamais_un_texte_valide(client, auth, fake_llm):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sec = _ecrire(client, h, pid, st["sections"][2], "Notre méthodologie, relue et validée.")
    vid = sec["version_courante"]["id"]
    client.post(f"/api/atelier/{pid}/sections/{sec['id']}/valider", headers=h, json={"version_id": vid})
    ans, _seen = _redige("Une autre méthodologie proposée.")
    fake_llm(ans)
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/rediger", headers=h, json={})
    assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sec["id"])
    assert s["statut"] == "a_revalider" and s["validation"]["version_id"] == vid
    r = client.get(f"/api/atelier/{pid}/sections/{sec['id']}/versions/{vid}", headers=h)
    assert r.json()["texte"] == "Notre méthodologie, relue et validée."


def test_pas_d_ecriture_pendant_une_redaction(client, auth, fake_llm):
    import threading
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][0]["id"]
    gate = threading.Event()

    def slow(kw):
        gate.wait(5)
        return "Texte.", "end_turn"
    fake_llm(slow)
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
    jid = r.json()["tache"]["id"]
    try:
        r2 = client.put(f"/api/atelier/{pid}/sections/{sid}/texte", headers=h, json={"texte": "x"})
        assert r2.status_code == 409 and r2.json()["detail"]["code"] == "redaction_en_cours"
        r3 = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
        assert r3.status_code == 409
        s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sid)
        assert s["redaction_en_cours"]["job_id"] == jid
    finally:
        gate.set()
    assert _attendre(client, h, pid, jid)["status"] == "done"


def test_raccourcir_vers_la_limite_et_fait_ajoute_a_relire(client, auth, fake_llm):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sec = st["sections"][0]                                  # limite : 2 pages
    long = " ".join(["Notre encadrement suit le chantier au quotidien."] * 200)
    sec = _ecrire(client, h, pid, sec, long)
    calls = []

    def ans(kw):
        calls.append(llm._flatten_content(kw["messages"][0]["content"]))
        return "Notre encadrement suit le chantier. Nous disposons de 3 nacelles.", "end_turn"
    fake_llm(ans)
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/raccourcir", headers=h, json={})
    assert r.status_code == 202, r.text
    assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    cible = int(re.search(r"au plus (\d+) mots", calls[0]).group(1))
    assert 300 < cible < 1600 and long[:60] in calls[0]
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sec["id"])
    v = s["version_courante"]
    assert v["origine"] == "raccourci" and v["ia"] is True
    assert [f["fait"] for f in v["faits"]] == ["3 nacelles"]               # ajouté par l'IA : à relire
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/raccourcir", headers=h, json={"cible_mots": 500})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "deja_assez_court"


def test_tache_orpheline_au_redemarrage_rend_exactement_le_debit(client, auth):
    from app.main import _recover_orphan_jobs
    from app.models import DossierPresentation, DossierSection, Job
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][0]["id"]
    uid = _user_id(auth["email"])
    db = _db()
    try:
        from app.core.quota import consume_analysis
        from app.models import User
        u = db.get(User, uid)
        consume_analysis(u, db)
        d = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).first()
        d.analyse_debitee = True
        j = Job(user_id=uid, kind=at.JOB_REDIGER, status="running", label="x",
                result={"_debit": {"analyse": True, "sections": {str(sid): "incluse"}, "dossier_id": d.id}})
        db.add(j)
        db.flush()
        db.get(DossierSection, sid).job_id = j.id
        db.commit()
        jid = j.id
    finally:
        db.close()
    u1 = _usage(auth["email"])
    _recover_orphan_jobs()
    assert _usage(auth["email"]) == u1 - 1
    db = _db()
    try:
        assert db.get(Job, jid).status == "error" and db.get(DossierSection, sid).job_id is None
        assert db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).first().analyse_debitee is False
    finally:
        db.close()
    _recover_orphan_jobs()                                          # idempotent
    assert _usage(auth["email"]) == u1 - 1


# ═══ A0 « à relire » : phrases piégées ═════════════════════════════════════════════════

SOURCES_ENT = ["Entreprise : Breizh Couverture (SARL), Lorient.\n"
               "Effectif : 12 (salariés, personnes, collaborateurs, compagnons, employés).\n"
               "Qualifications et certifications : Qualibat 3111.",
               "- Yann Tanguy, Chef de chantier ; qualifications : BTS Bâtiment, 15 ans"]


@pytest.mark.parametrize("phrase,attendu", [
    ("Nous sommes certifiés ISO 14001.", {("ISO 14001", far.CERTIFICATION)}),
    ("33 % de l'effectif est formé au travail en hauteur.", {("33 %", far.CHIFFRE)}),
    ("Nous mobilisons quinze compagnons sur le chantier.", {("quinze", far.EN_LETTRES)}),
    ("Nous garantissons 100 % de satisfaction.", {("100 %", far.CHIFFRE)}),
    ("Nous avons couvert la Mairie de Quimper en 2019.", {("Mairie de Quimper", far.NOM_PROPRE), ("2019", far.CHIFFRE)}),
    ("Notre équipe dispose d'une nacelle.", {("nacelle", far.MATERIEL)}),
    ("Ces moyens garantissent zéro accident.", {("zéro", far.EN_LETTRES)}),
])
def test_a0_phrases_piegees_a_relire(phrase, attendu):
    got = {(x["fait"], x["type"]) for x in far.a_relire(phrase, SOURCES_ENT, noms_entreprise=["Breizh Couverture"])}
    assert attendu <= got
    for x in far.a_relire(phrase, SOURCES_ENT):
        assert x["message"] == "Adjugo ne trouve pas d'où vient ce fait : relisez-le." and x["statut"] == far.A_RELIRE


@pytest.mark.parametrize("phrase", [
    "Notre équipe compte 12 salariés.",
    "Nous mobilisons douze compagnons.",
    "Breizh Couverture est titulaire de la qualification Qualibat 3111.",
    "Le chantier est encadré par Yann Tanguy.",
    "Le CCTP prévoit une couverture en zinc de 450 m2.",
    "Le maître d'ouvrage impose la certification ISO 14001.",
])
def test_a0_faits_sources_ou_phrases_sur_le_marche_non_signales(phrase):
    assert far.a_relire(phrase, SOURCES_ENT, noms_entreprise=["Breizh Couverture"]) == []


def test_a0_le_dossier_de_l_acheteur_n_est_jamais_une_source(client, auth, fake_llm):
    """Le règlement exige l'ISO 14001 : « Nous sommes certifiés ISO 14001 » reste à relire, même
    si l'IA a lu le règlement."""
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][3]["id"]
    ans, _seen = _redige("Nous sommes certifiés ISO 14001 et nous triions 100 % des déchets.")
    fake_llm(ans)
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
    assert _attendre(client, h, pid, r.json()["tache"]["id"])["status"] == "done"
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sid)
    faits = {f["fait"] for f in s["version_courante"]["faits"]}
    assert "ISO 14001" in faits and "100 %" in faits
    ctx_src = " ".join(at.Contexte(_db(), _get_project(pid), _user_id(auth["email"]), [_user_id(auth["email"])])
                       .sources_entreprise)
    assert "ISO 14001" not in ctx_src and "REGLEMENT" not in ctx_src


def _get_project(pid):
    from app.models import Project
    db = _db()
    try:
        p = db.get(Project, pid)
        db.expunge(p)
        return p
    finally:
        db.close()


def test_a0_declarations_de_l_utilisateur_et_faits_a_relire_qui_restent(client, auth, fake_llm):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][2]["id"]
    ans, _seen = _redige("Nous sommes certifiés ISO 14001. Nous intervenons par zones.")
    fake_llm(ans)
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/rediger", headers=h, json={})
    _attendre(client, h, pid, r.json()["tache"]["id"])
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sid)
    assert s["faits_a_relire"] == 1
    # L'utilisateur ajoute SON fait (une déclaration) et garde celui de l'IA : seul celui-ci reste.
    s = _ecrire(client, h, pid, s, s["version_courante"]["texte"] + " Nous avons 4 camions.")
    assert [f["fait"] for f in s["version_courante"]["faits"]] == ["ISO 14001"]
    assert s["version_courante"]["ia"] is True                        # texte de l'IA repris
    fid = s["version_courante"]["faits"][0]["id"]
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/faits/{fid}", headers=h, json={"statut": "relu_par_vous"})
    assert r.status_code == 200 and r.json()["fait"]["statut"] == "relu_par_vous"
    assert r.json()["section"]["faits_a_relire"] == 0
    r = client.post(f"/api/atelier/{pid}/sections/{sid}/faits/{fid}", headers=h, json={"statut": "bon"})
    assert r.status_code == 422                                          # aucun libellé positif
    # Il réécrit tout lui-même : plus rien de l'IA, plus de mention IA.
    s = _ecrire(client, h, pid, r.json()["section"] if r.status_code == 200 else s,
                "Nous intervenons par demi-journées, hors temps scolaire, avec un balisage quotidien.")
    assert s["version_courante"]["ia"] is False and s["version_courante"]["faits"] == []


# ═══ Compteur, blocage, dérogation, export ═════════════════════════════════════════════

def _valider_tout(client, h, pid, texte="Nous intervenons par zones, en site occupé."):
    st = _etat(client, h, pid)
    for s in st["sections"]:
        s2 = _ecrire(client, h, pid, s, texte)
        r = client.post(f"/api/atelier/{pid}/sections/{s['id']}/valider", headers=h,
                        json={"version_id": s2["version_courante"]["id"]})
        assert r.status_code == 200, r.text
    return _etat(client, h, pid)


def test_export_bloque_puis_derogation_au_journal_et_export(client, auth, fake_llm):
    h = auth["headers"]
    pid = _projet(client, h)
    _creer(client, h, pid)
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx"})
    assert r.status_code == 409
    d = r.json()["detail"]
    assert d["code"] == "export_bloque" and [x["code"] for x in d["raisons"]] == ["parties_vides"]
    # Seules les raisons qui répondent à ce qui bloque sont proposées.
    assert [x["code"] for x in d["derogations_possibles"]] == ["redigerai_dans_word", "depot_en_l_etat", "autre"]
    st = _valider_tout(client, h, pid, "Nous intervenons par zones. [À compléter : planning détaillé]")
    gate = st["blocage"]
    assert gate["bloque"] and [x["code"] for x in gate["raisons"]] == ["a_completer"] and gate["raisons"][0]["nombre"] == 4
    assert [x["code"] for x in st["derogations_possibles"]] == ["completerai_dans_word", "depot_en_l_etat", "autre"]
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "relu_moi_meme", "texte": "Planning", "etat_hash": gate["etat_hash"]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "raison_invalide"     # aucun fait à relire ici
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "completerai_dans_word", "texte": "Planning", "etat_hash": "faux"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "etat_modifie"
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "inconnue", "texte": "Planning", "etat_hash": gate["etat_hash"]})
    assert r.status_code == 422
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "completerai_dans_word", "texte": "Planning envoyé lundi",
                          "etat_hash": gate["etat_hash"]})
    assert r.status_code == 200, r.text
    assert r.json()["blocage"]["peut_exporter"] is True and r.json()["derogation"]["raisons"] == ["a_completer"]
    from app.models import AuditLog
    from app.services.audit import verify_chain
    db = _db()
    try:
        e = db.query(AuditLog).filter(AuditLog.action == "atelier.derogation_export").order_by(AuditLog.id.desc()).first()
        assert e.meta["texte"] == "Planning envoyé lundi" and e.meta["raisons"] == ["a_completer"]
        assert verify_chain(db, e.owner_id)["intact"] is True
    finally:
        db.close()
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    # Le texte change : la dérogation ne vaut plus.
    s = _etat(client, h, pid)["sections"][0]
    _ecrire(client, h, pid, s, "Nous intervenons. [À compléter : autre point]")
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "pdf"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "export_bloque"


def test_limite_de_pages_depassee_derogation_avec_avertissement_d_irregularite(client, auth):
    h = auth["headers"]
    pid = _projet(client, h)
    _creer(client, h, pid)
    long = "\n\n".join(["Nous organisons le chantier par zones successives, avec un balisage quotidien "
                        "et une information des usagers du groupe scolaire."] * 60)
    st = _valider_tout(client, h, pid, long)
    lims = [x for x in st["limites"] if x["depassee"]]
    assert {x["portee"] for x in lims} == {"document", "section"}
    doc = next(x for x in lims if x["portee"] == "document")
    assert doc["type"] == "pages" and doc["valeur"] == 8 and doc["mesure"] > 8 and doc["environ"] is True
    assert doc["message"].startswith("Dossier de présentation : environ") and "page de garde et sommaire compris" in doc["message"]
    assert doc["avertissement"] == "L'acheteur peut juger votre offre irrégulière."
    assert "environ" in st["compteur"]["message"].lower() or "approximatif" in st["compteur"]["message"]
    assert st["compteur"]["police"]["famille"].lower() == "arial" and st["compteur"]["police"]["connue"] is True
    gate = st["blocage"]
    assert [x["code"] for x in gate["raisons"]] == ["limite_depassee"]
    r = client.post(f"/api/atelier/{pid}/derogation", headers=h,
                    json={"raison": "depot_en_l_etat", "texte": "Annexes comprises", "etat_hash": gate["etat_hash"]})
    assert r.json()["derogation"]["avertissement"] == "L'acheteur peut juger votre offre irrégulière."
    assert client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "pdf"}).status_code == 200


def _docx_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return "\n".join(zf.read(n).decode("utf-8", "replace") for n in zf.namelist() if n.endswith(".xml")
                         or n.endswith(".rels"))


def test_export_word_propre_numeros_du_cadre_et_mention_ia(client, auth, fake_llm):
    import docx
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    _creer(client, h, pid)
    st = _valider_tout(client, h, pid, "Nous intervenons par zones.")
    assert st["export"]["ia_utilisee"] is False and st["export"]["mention_ia"] == ""
    assert st["export"]["nom_fichier_docx"] == "Mémoire justificatif.docx"
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx"})
    assert r.status_code == 200, r.text
    assert "filename*=UTF-8''M%C3%A9moire%20justificatif.docx" in r.headers["content-disposition"]
    d = docx.Document(io.BytesIO(r.content))
    cp = d.core_properties
    assert cp.subject == "" and cp.keywords == "" and cp.author == "" and cp.title == "Mémoire justificatif"
    # IA sur une partie : la mention apparaît, annoncée avant l'export.
    ans, _seen = _redige("Notre méthodologie [S1] en site occupé.")
    fake_llm(ans)
    sec = st["sections"][2]
    r = client.post(f"/api/atelier/{pid}/sections/{sec['id']}/rediger", headers=h, json={})
    _attendre(client, h, pid, r.json()["tache"]["id"])
    s = next(x for x in _etat(client, h, pid)["sections"] if x["id"] == sec["id"])
    client.post(f"/api/atelier/{pid}/sections/{sec['id']}/valider", headers=h,
                json={"version_id": s["version_courante"]["id"]})
    st = _etat(client, h, pid)
    assert st["export"]["ia_utilisee"] is True and st["export"]["mention_ia"] == "Préparé avec assistance IA"
    assert st["export"]["proprietes"]["sujet"] == "Préparé avec assistance IA"
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx", "etat_hash": st["blocage"]["etat_hash"]})
    assert r.status_code == 200, r.text
    d = docx.Document(io.BytesIO(r.content))
    assert d.core_properties.subject == "Préparé avec assistance IA"
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads[:1] == ["Sommaire"]
    body = heads[heads.index("Sommaire") + 1:]
    assert body == ["1. Moyens humains", "1.1 Encadrement du chantier", "1.2 Équipe d'exécution",
                    "2. Méthodologie", "3. Démarche environnementale"]
    full = "\n".join(p.text for p in d.paragraphs)
    assert "Notre méthodologie en site occupé." in full
    xml = _docx_xml(r.content)
    for interdit in ("[S1]", "Adjugo", "adjugo", "python-docx", "À compléter", "BLOQUANT"):
        assert interdit not in xml, interdit
    assert "Breizh Couverture" in full and "Ville de Lorient" in full
    # Un état périmé est refusé.
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx", "etat_hash": "0" * 64})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "etat_modifie"


def test_export_pdf_sans_marque_interne_et_archive_dans_le_dossier(client, auth):
    from pypdf import PdfReader
    h = auth["headers"]
    pid = _projet(client, h)
    _creer(client, h, pid)
    _valider_tout(client, h, pid, "Nous intervenons par zones.\n\n| Phase | Durée |\n|---|---|\n| Dépose | 2 jours |")
    r = client.put(f"/api/atelier/{pid}/options", headers=h, json={"page_garde": False})
    assert r.json()["options"] == {"page_garde": False, "sommaire": True}
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "pdf"})
    assert r.status_code == 200 and r.content[:5] == b"%PDF-"
    reader = PdfReader(io.BytesIO(r.content))
    meta = {k: str(v) for k, v in (reader.metadata or {}).items()}
    assert not any("adjugo" in v.lower() or "reportlab" in v.lower() for v in meta.values()), meta
    text = " ".join(p.extract_text() or "" for p in reader.pages)
    assert "Dépose" in text and "1.1 Encadrement du chantier" in text and "[S" not in text
    from app.models import Document, Project
    db = _db()
    try:
        docs = db.query(Document).filter(Document.project_id == pid, Document.folder == "Mémoire technique").all()
        assert [d.name for d in docs] == ["Mémoire justificatif.pdf"]
        wf = db.get(Project, pid).workflow["dossier_presentation"]
        assert wf["fichiers"] == ["Mémoire justificatif.pdf"] and wf["source"] == "atelier"
    finally:
        db.close()


def test_compteur_mots_caracteres_pages_par_partie_et_total(client, auth):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    _ecrire(client, h, pid, st["sections"][2], "Un deux trois quatre cinq.")
    r = client.get(f"/api/atelier/{pid}/compteur", headers=h)
    assert r.status_code == 200
    c = r.json()
    sec = next(x for x in c["sections"] if x["titre"] == "2. Méthodologie")
    assert sec["compteur"]["mots"] == 5 and sec["compteur"]["caracteres"] == len("Un deux trois quatre cinq.")
    assert c["compteur"]["total"]["mots"] == 5 and c["compteur"]["total"]["pages"] >= 3
    assert c["compteur"]["environ"] is True


# ═══ Dossier commun (ZIP) : le texte de l'atelier, plus aucun moteur A ═════════════════

def test_dossier_commun_prend_le_texte_de_l_atelier_et_n_appelle_pas_l_ia(client, auth, fake_llm):
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    fake = fake_llm(lambda kw: ("NE DOIT PAS SERVIR", "end_turn"))
    r = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid})
    assert r.status_code == 200, r.text
    d = r.json()["dossier"]
    assert any(w.startswith("BLOQUANT — Dossier de présentation non rédigé") for w in d["warnings"])
    assert d["memoire_fichiers"] == [] and fake.calls == []
    _creer(client, h, pid)
    st = _etat(client, h, pid)
    for s in st["sections"]:
        _ecrire(client, h, pid, s, "Nous intervenons par zones.")
    d = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid}).json()["dossier"]
    assert any("Dossier de présentation non validé : 4 parties à valider" in w for w in d["warnings"])
    assert d["memoire_fichiers"] == ["Mémoire justificatif.docx", "Mémoire justificatif.pdf"]
    zf = zipfile.ZipFile(io.BytesIO(__import__("base64").b64decode(d["zip_b64"])))
    assert "Mémoire justificatif.docx" in zf.namelist()
    assert fake.calls == []
    s = _etat(client, h, pid)["sections"][0]
    _ecrire(client, h, pid, s, "[À compléter : encadrement]")
    d = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid}).json()["dossier"]
    assert d["memoire_fichiers"] == [] and any(w.startswith("BLOQUANT — Dossier de présentation non joint")
                                               for w in d["warnings"])


# ═══ Quota, mesure « fichier imposé », savoir-faire, démo ═══════════════════════════════

def test_mesure_fichier_impose_a_l_import_et_metrics(client, auth, monkeypatch):
    from app.services import analysis as an

    def fake_analyze(dce_text, company, gonogo, lang_name=None, pieces=None, **_lecture):   # lecture intégrale
        return _analysis(dce_text, pieces)
    monkeypatch.setattr(an, "analyze_dce_text", fake_analyze)
    h = auth["headers"]
    r = client.post("/api/projects/", headers=h, json={"name": "Import", "client": "Commune"})
    pid = r.json()["id"]
    content = _zip({"RC.txt": RC, "Cadre de reponse technique.txt": CADRE})
    r = client.post("/api/sourcing/analyze-upload", headers=h, data={"project_id": str(pid)},
                    files={"file": ("DCE.zip", content, "application/zip")})
    assert r.status_code == 200, r.text
    from app.models import Project
    db = _db()
    try:
        p = db.get(Project, pid)
        assert p.dce_importe_le is not None and p.dce_fichier_impose is True
        m = at.mesure_fichier_impose(db)
        assert m["dossiers_importes"] >= 1 and m["fichier_impose"] >= 1 and m["seuil_decision"] == 0.2
    finally:
        db.close()
    txt = client.get("/metrics").text
    assert "adjugo_dossiers_fichier_impose " in txt and "adjugo_dossiers_importes " in txt
    assert client.get("/api/admin/mesures/redaction", headers=h).status_code == 403


def test_savoir_faire_ne_redige_plus_le_dossier(client, auth):
    for url in ("/api/knowledge/memoire", "/api/knowledge/memoire-upload"):
        r = client.post(url, headers=auth["headers"], json={"dce_text": "x" * 100})
        assert r.status_code == 410 and "page du marché" in r.json()["detail"]


def test_demo_ne_lance_pas_l_ia_de_l_atelier(client):
    from app.core.demo_guard import blocked_action
    for path in ("/api/atelier/3/sections/4/rediger", "/api/atelier/3/sections/4/raccourcir", "/api/atelier/3/rediger"):
        assert blocked_action("POST", path) == "rédiger avec Adjugo"
    for m, path in (("POST", "/api/atelier/3"), ("PUT", "/api/atelier/3/sections/4/texte"),
                    ("POST", "/api/atelier/3/derogation"), ("POST", "/api/atelier/3/export"),
                    ("POST", "/api/atelier/3/sections/4/faits/abc")):
        assert blocked_action(m, path) == "préparer votre dossier de présentation", path
    assert blocked_action("GET", "/api/atelier/3") is None


# ═══ Vague 4 préparée : la matière privée d'un auteur partenaire n'est jamais exposée ══

def test_partie_d_un_partenaire_matiere_privee_jamais_exposee(client, auth):
    from app.models import DossierPresentation, DossierSection
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    sid = st["sections"][1]["id"]
    db = _db()
    try:
        s = db.get(DossierSection, sid)
        s.author_kind, s.author_invite_id = at.AUTEUR_INVITATION, 999
        db.commit()
        v_brouillon = at.ajouter_version(db, s, texte="Brouillon privé du couvreur", origine="ia", ia=True,
                                         sources=[{"ref": "S1", "doc_name": "secret.pdf", "extrait": "SECRET"}],
                                         faits=[{"id": "f1", "fait": "ISO 9001", "statut": "a_relire"}])
        db.commit()
        vb = v_brouillon.id
        d = db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).first()
        assert at.peut_voir_prive(s, at.Lecteur.organisation(d.owner_id)) is False
        assert at.peut_voir_prive(s, at.Lecteur(at.AUTEUR_INVITATION, d.owner_id, invite_id=999)) is True
    finally:
        db.close()
    st = _etat(client, h, pid)
    raw = json.dumps(st, ensure_ascii=False)
    assert "Brouillon privé" not in raw and "SECRET" not in raw and "secret.pdf" not in raw
    s = next(x for x in st["sections"] if x["id"] == sid)
    assert s["version_courante"] is None and s["peut_ecrire"] is False and s["auteur"]["type"] == "invitation"
    assert client.get(f"/api/atelier/{pid}/sections/{sid}/versions/{vb}", headers=h).status_code == 404
    assert client.get(f"/api/atelier/{pid}/sections/{sid}", headers=h).json()["versions"] == []
    for m, url, body in (("put", f"/api/atelier/{pid}/sections/{sid}/texte", {"texte": "x", "base_version_id": vb}),
                         ("post", f"/api/atelier/{pid}/sections/{sid}/rediger", {}),
                         ("post", f"/api/atelier/{pid}/sections/{sid}/valider", {"version_id": vb})):
        r = getattr(client, m)(url, headers=h, json=body)
        assert r.status_code == 403 and r.json()["detail"]["code"] == "partie_d_un_partenaire", url


# ═══ Migration : montée, descente, remontée (SQLite) ══════════════════════════════════

def test_migration_sqlite_monte_descend_remonte(tmp_path):
    import os
    import subprocess
    import sys
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    db = tmp_path / "mig.db"
    env = dict(os.environ, DATABASE_URL=f"sqlite:///{db}")

    def alembic(*args):
        r = subprocess.run([sys.executable, "-m", "alembic", *args], cwd=root, env=env, capture_output=True, text=True)
        assert r.returncode == 0, r.stdout + r.stderr

    import sqlalchemy as sa
    alembic("upgrade", "head")
    eng = sa.create_engine(f"sqlite:///{db}")
    insp = sa.inspect(eng)
    assert {"dossiers_presentation", "dossier_sections", "dossier_section_versions"} <= set(insp.get_table_names())
    assert {"dce_importe_le", "dce_fichier_impose"} <= {c["name"] for c in insp.get_columns("projects")}
    assert {"options_export", "cout_ia"} <= {c["name"] for c in insp.get_columns("dossiers_presentation")}
    assert {"premieres_incluses", "job_analyse_id"} <= {c["name"] for c in insp.get_columns("dossiers_presentation")}
    assert "poids_source" in {c["name"] for c in insp.get_columns("dossier_sections")}
    alembic("downgrade", "s3t4redactionv2")          # quota par dossier, provenance du poids (révisions suivantes comprises)
    insp = sa.inspect(sa.create_engine(f"sqlite:///{db}"))
    assert "premieres_incluses" not in {c["name"] for c in insp.get_columns("dossiers_presentation")}
    assert "poids_source" not in {c["name"] for c in insp.get_columns("dossier_sections")}
    alembic("downgrade", "r2s3reseauvalid")          # l'atelier lui-même
    insp = sa.inspect(sa.create_engine(f"sqlite:///{db}"))
    assert "dossiers_presentation" not in insp.get_table_names()
    assert "dce_importe_le" not in {c["name"] for c in insp.get_columns("projects")}
    alembic("upgrade", "head")
    alembic("upgrade", "head")                       # rejeu idempotent
    insp = sa.inspect(sa.create_engine(f"sqlite:///{db}"))
    assert "dossier_section_versions" in insp.get_table_names()
    assert "poids_source" in {c["name"] for c in insp.get_columns("dossier_sections")}


# ═══ Compléments : sources citées, tâches génériques, portabilité ══════════════════════

def test_moteur_ne_montre_que_les_extraits_cites_et_aucun_prix(fake_llm):
    from app.services.agents import memoire as moteur
    chunks = [{"doc_name": f"d{i}.pdf", "chunk_id": i, "text": f"savoir-faire {i} " * 5} for i in (1, 2, 3)]
    fake = fake_llm(lambda kw: ("Démarche qualité [S2]. Hors borne [S9].", "end_turn"))
    out = moteur.rediger_partie({"titre": "Qualité", "chunks": chunks, "faits": {"profil": "Breizh"},
                                 "dossier_acheteur": "RC", "autres_parties": ["1. Moyens"]})
    assert [s["ref"] for s in out["sources"]] == ["S2"] and out["erreur"] is False
    sent = fake.sent()
    assert "AUTRES PARTIES DU DOSSIER" in sent and "1. Moyens" in sent
    assert "aucun prix" in sent and "ne prouve RIEN sur l'entreprise" in sent
    fake_llm(lambda kw: ("", "end_turn"))
    assert moteur.rediger_partie({"titre": "Qualité"})["erreur"] is True        # réponse vide = échec


def test_route_generique_des_taches_ne_montre_pas_le_releve_de_debit(client, auth, fake_llm):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    ans, _seen = _redige("Texte.")
    fake_llm(ans)
    r = client.post(f"/api/atelier/{pid}/sections/{st['sections'][0]['id']}/rediger", headers=h, json={})
    jid = r.json()["tache"]["id"]
    _attendre(client, h, pid, jid)
    g = client.get(f"/api/jobs/{jid}", headers=h)
    assert g.status_code == 200 and g.json()["status"] == "done"
    assert '"_debit"' not in g.text and "dossier_id" not in g.text
    # La tâche d'un autre marché n'est pas lisible par la route de l'atelier d'un autre marché.
    pid2 = _projet(client, h)
    _creer(client, h, pid2)
    assert client.get(f"/api/atelier/{pid2}/taches/{jid}", headers=h).status_code == 404


def test_portabilite_et_transfert_de_propriete_suivent_le_dossier(client, auth):
    h = auth["headers"]
    pid = _projet(client, h)
    st = _creer(client, h, pid)
    _ecrire(client, h, pid, st["sections"][0], "Notre encadrement, écrit par nous.")
    r = client.get("/api/rgpd/export", headers=h)
    assert r.status_code == 200
    data = r.json()
    d = next(x for x in data["dossiers_de_presentation"] if x["project_id"] == pid)
    assert d["parties"][0]["texte"] == "Notre encadrement, écrit par nous."
    # Transfert de propriété de l'organisation : le dossier suit le nouveau propriétaire.
    from app.models import DossierPresentation, DossierSection, DossierSectionVersion, PlanType, User
    db = _db()
    try:
        u = db.query(User).filter(User.email == auth["email"]).first()
        u.plan = PlanType.pro
        db.commit()
    finally:
        db.close()
    r = client.post("/api/org/invite", headers=h, json={"email": f"collegue_{pid}@test.fr"})
    assert r.status_code == 201, r.text
    new_id = r.json()["id"]
    r = client.post("/api/org/transfer-ownership", headers=h, json={"new_owner_id": new_id})
    assert r.status_code == 200, r.text
    db = _db()
    try:
        assert db.query(DossierPresentation).filter(DossierPresentation.project_id == pid).one().owner_id == new_id
        assert {s.owner_id for s in db.query(DossierSection).filter(DossierSection.project_id == pid)} == {new_id}
        assert {v.owner_id for v in db.query(DossierSectionVersion).filter(DossierSectionVersion.project_id == pid)} \
            == {new_id}
    finally:
        db.close()
    h2 = {"Authorization": f"Bearer {r.json()['access_token']}"}
    st = _etat(client, h2, pid)                                   # toujours visible par l'organisation
    assert st["sections"][0]["version_courante"]["texte"] == "Notre encadrement, écrit par nous."


def test_dossier_commun_joint_le_texte_valide_de_preference(client, auth):
    import base64
    import docx
    h = auth["headers"]
    _company(client, h)
    pid = _projet(client, h)
    _creer(client, h, pid)
    st = _valider_tout(client, h, pid, "Texte validé par le dirigeant.")
    s = st["sections"][2]
    _ecrire(client, h, pid, s, "Brouillon postérieur, jamais validé.")
    d = client.post("/api/sourcing/documents", headers=h, json={"project_id": pid}).json()["dossier"]
    assert any("2. Méthodologie" in w and "texte VALIDÉ" in w for w in d["warnings"])
    assert not any(w.startswith("BLOQUANT — Dossier de présentation non validé") for w in d["warnings"])
    zf = zipfile.ZipFile(io.BytesIO(base64.b64decode(d["zip_b64"])))
    word = "\n".join(p.text for p in docx.Document(io.BytesIO(zf.read("Mémoire justificatif.docx"))).paragraphs)
    assert "Brouillon postérieur" not in word and word.count("Texte validé par le dirigeant.") == 4


def test_logo_de_l_entreprise_sur_la_page_de_garde(client, auth):
    import docx
    from app.services.dossier import save_to_dossier
    h = auth["headers"]
    pid = _projet(client, h)
    _creer(client, h, pid)
    _valider_tout(client, h, pid, "Nous intervenons par zones.")
    png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
           b"\x00\x00\x00\rIDATx\x9cc\xf8\xcf\xc0\xf0\x1f\x00\x05\x00\x01\xff\x89\x99=\x1d\x00\x00\x00\x00IEND\xaeB`\x82")
    db = _db()
    try:
        save_to_dossier(db, _user_id(auth["email"]), None, "", "logo-entreprise.png", png, "image/png")
    finally:
        db.close()
    st = _etat(client, h, pid)
    assert st["export"]["logo"] is True
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "docx"})
    assert r.status_code == 200 and len(docx.Document(io.BytesIO(r.content)).inline_shapes) == 1
    r = client.post(f"/api/atelier/{pid}/export", headers=h, json={"format": "pdf"})
    assert r.status_code == 200
